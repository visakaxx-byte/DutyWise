from __future__ import annotations

import math
import json
import re
import uuid
import asyncio
import hashlib
import inspect
import base64
import io
import itertools
import time
from copy import copy
from dataclasses import asdict, dataclass, field, replace
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Optional

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from pypdf import PdfReader
import pypdfium2 as pdfium
from PIL import Image

from crawler_client import StrictTaxCrawler
from llm_client import LLMClient
from price_search import PriceEvidence, estimate_declared_unit_price_from_web


APP_DIR = Path(__file__).resolve().parent
LOCAL_TEMPLATE_PATH = APP_DIR / "templates/清关模板.xlsx"
REFERENCE_ROOT = APP_DIR / "reference"
REPLACEMENT_WORKBOOK_PATH = APP_DIR / "docs/海关编码查找.xlsx"
CUSTOMER_CODEBOOK_PATH = REFERENCE_ROOT / "编码库(1).xlsx"
RULES_ROOT = APP_DIR / "rules"
MAX_OUTPUT_ITEMS = 30
BASE_TAX_LIMIT = 0.2
CUSTOMER_CODEBOOK_TAX_LIMIT = 0.30
TAX_TOLERANCE_USD = 1.0
TAX_FINAL_TOLERANCE_USD = 20.0
MIN_ROW_TAX_AMOUNT_USD = 30.0
MAX_ZERO_TAX_ROWS = 4
MAX_TAX_OVER_TARGET_RATIO = 0.1
BILL_PRODUCT_QUERY_RETRIES = 2
UNDETAILED_BILL_TAX_SHARE = 0.20
UNDETAILED_BILL_CARTON_SHARE = 0.06
PRICE_FIT_MIN_REFERENCE_RATIO = 0.67
PRICE_REPAIR_POOL_SIZE = 24
PRICE_REPAIR_MAX_PASSES = 12
PRICE_REPAIR_MAX_WEB_LOOKUPS = 24
PRICE_REPAIR_MIN_SCORE_GAIN = 5.0
PRICE_REPAIR_SEARCH_TIMEOUT_SECONDS = 3.0
PRICE_REPAIR_SEARCH_MAX_PAGES = 2
PRICE_REPAIR_REPLACEMENT_QUERY_LIMIT = 120
PRICE_REPAIR_REPLACEMENT_POOL_SIZE = 48
PRICE_SEARCH_TIMEOUT_SECONDS = 3.0
PRICE_SEARCH_MAX_PAGES = 2
MANUAL_INVOICE_REPLACEMENT_QUERY_LIMIT = 24
MANUAL_INVOICE_REPLACEMENT_POOL_SIZE = 28
MANUAL_INVOICE_ANCHOR_HS = {
    "8471602000",
    "8306290000",
    "3924104000",
    "6912004810",
    "6913105000",
    "7013992000",
    "6303922010",
}
LLM_DRAFT_TIMEOUT_SECONDS = 180.0
LLM_TRANSLATION_TIMEOUT_SECONDS = 60.0
DECLARED_RETAIL_PRICE_RATIO = 0.3
DEFAULT_KG_PER_CTN_MIN = 0.5
DEFAULT_KG_PER_CTN_MAX = 80.0
DEFAULT_KG_PER_PC_MIN = 0.01
DEFAULT_KG_PER_PC_MAX = 50.0
DEFAULT_UNIT_PRICE_MIN = 0.05
DEFAULT_UNIT_PRICE_MAX = 50.0
BILL_TEXT_MIN_CHARS = 80
BILL_PARSER_TEXT_CHARS = 12000
BILL_VISION_MAX_PAGES = 2
BILL_VISION_MAX_SIDE = 1800
BILL_VISION_JPEG_QUALITY = 80
BILL_PARSER_PROMPT_VERSION = "v4"
BILL_PACKAGE_RESOLUTION_PROMPT_VERSION = "v1"
BILL_PACKAGE_RECOGNITION_MAX_ATTEMPTS = 3
BILL_PACKAGE_FINAL_RETRY_MAX_PAGES = 6
BILL_PACKAGE_FINAL_RETRY_TEXT_CHARS = 24000
MANIFEST_SCHEMA_PROMPT_VERSION = "v1"
MANIFEST_SCHEMA_CACHE_REVISION = "net-weight-v1"
MANIFEST_SCHEMA_MIN_CONFIDENCE = 0.80
MANIFEST_WEIGHT_ABSOLUTE_TOLERANCE_KG = 0.5
MANIFEST_WEIGHT_RELATIVE_TOLERANCE = 0.001
LLM_MAX_CONCURRENCY = 2
LLM_PLAUSIBILITY_BATCH_SIZE = 5
LLM_PLAUSIBILITY_PROMPT_VERSION = "v2-batch"
LLM_OUTPUT_REVIEW_PROMPT_VERSION = "v1"
ADAPTIVE_REPLACEMENT_BATCH_SIZE = 4

DEFAULT_REFERENCE_STYLE_ROWS = [
    {"中文品名": "铁制昆虫饰品", "英文品名": "Iron insect ornaments", "商品编码": "8306290000", "材质": "Iron", "用途": "Decoration", "单价": 0.69, "综合税率": 0.10, "箱数": 70, "数量": 910, "毛重": 1648},
    {"中文品名": "天鹅摆件", "英文品名": "Swan figurines", "商品编码": "3926400090", "材质": "Plastic", "用途": "Decoration", "单价": 0.52, "综合税率": 0.153, "箱数": 138, "数量": 1656, "毛重": 2044.21},
    {"中文品名": "花瓶", "英文品名": "Vase", "商品编码": "6913105000", "材质": "Ceramic", "用途": "Decoration", "单价": 0.73, "综合税率": 0.175, "箱数": 70, "数量": 630, "毛重": 1353.15},
    {"中文品名": "厨具收纳架", "英文品名": "Kitchen utensil storage rack", "商品编码": "3924104000", "材质": "ABS", "用途": "Kitchenware", "单价": 1.05, "综合税率": 0.134, "箱数": 157, "数量": 942, "毛重": 2665.9},
    {"中文品名": "陶瓷杯", "英文品名": "Ceramic cup", "商品编码": "6912004810", "材质": "Ceramic", "用途": "HOME", "单价": 0.35, "综合税率": 0.198, "箱数": 55, "数量": 825, "毛重": 763.21},
    {"中文品名": "笔筒", "英文品名": "Pen holder", "商品编码": "3926100000", "材质": "Plastic", "用途": "Office", "单价": 0.33, "综合税率": 0.153, "箱数": 93, "数量": 1209, "毛重": 1044.23},
    {"中文品名": "塑料发夹", "英文品名": "Plastic hairpin", "商品编码": "9615115000", "材质": "Plastic", "用途": "Decoration", "单价": 0.32, "综合税率": 0.10, "箱数": 48, "数量": 1440, "毛重": 913.45},
    {"中文品名": "塑料钥匙扣", "英文品名": "Plastic keychain", "商品编码": "3926909989", "材质": "Acrylic", "用途": "Decoration", "单价": 0.30, "综合税率": 0.153, "箱数": 45, "数量": 1350, "毛重": 1056.8},
    {"中文品名": "玻璃杯", "英文品名": "Glass cup", "商品编码": "7013992000", "材质": "Glass", "用途": "HOME", "单价": 0.35, "综合税率": 0.225, "箱数": 55, "数量": 1100, "毛重": 1066.85},
    {"中文品名": "窗帘", "英文品名": "Curtain", "商品编码": "6303922010", "材质": "Polyester", "用途": "HOME", "单价": 1.20, "综合税率": 0.288, "箱数": 57, "数量": 228, "毛重": 384.65},
    {"中文品名": "键盘", "英文品名": "Keyboard", "商品编码": "8471602000", "材质": "ABS", "用途": "HOME", "单价": 2.90, "综合税率": 0.0, "箱数": 115, "数量": 1380, "毛重": 2465.7},
    {"中文品名": "花园围栏", "英文品名": "Garden fence", "商品编码": "3926400090", "材质": "Plastic", "用途": "HOME", "单价": 1.3},
    {"中文品名": "塑料罩", "英文品名": "Plastic cover", "商品编码": "3924901050", "材质": "Plastic", "用途": "HOME", "单价": 1},
    {"中文品名": "键鼠套装", "英文品名": "Keyboard and mouse set", "商品编码": "8471602000", "材质": "Plastic", "用途": "HOME", "单价": 2},
    {"中文品名": "窗帘", "英文品名": "Curtains", "商品编码": "6303192120", "材质": "Polyester", "用途": "HOME", "单价": 2},
    {"中文品名": "LED灯", "英文品名": "LED lights", "商品编码": "8539520091", "材质": "Iron/Glass", "用途": "HOME", "单价": 2},
    {"中文品名": "杯架", "英文品名": "Cup holder", "商品编码": "3924104000", "材质": "Plastic", "用途": "HOME", "单价": 0.8},
    {"中文品名": "毛巾", "英文品名": "Towel", "商品编码": "6302991520", "材质": "Nylon", "用途": "HOME", "单价": 0.6},
    {"中文品名": "兵乒球拍", "英文品名": "Table tennis paddle", "商品编码": "9506400000", "材质": "Plastic", "用途": "HOME", "单价": 0.7},
    {"中文品名": "背景板", "英文品名": "Backdrop", "商品编码": "3926400090", "材质": "Acrylic", "用途": "HOME", "单价": 1},
    {"中文品名": "花架", "英文品名": "Flower stand", "商品编码": "8306290000", "材质": "Iron", "用途": "HOME", "单价": 2},
    {"中文品名": "手机支架", "英文品名": "Phone holder", "商品编码": "3926100000", "材质": "Plastic", "用途": "HOME", "单价": 0.6},
    {"中文品名": "充气泵", "英文品名": "Inflatable pump", "商品编码": "8414904190", "材质": "Plastic/Iron", "用途": "HOME", "单价": 3},
    {"中文品名": "花园装饰", "英文品名": "Garden decoration", "商品编码": "8306290000", "材质": "Iron", "用途": "HOME", "单价": 1},
    {"中文品名": "戒尺", "英文品名": "Ruler", "商品编码": "9017800000", "材质": "Iron", "用途": "HOME", "单价": 0.7},
    {"中文品名": "卡扣", "英文品名": "Buckle", "商品编码": "8302426000", "材质": "Iron", "用途": "HOME", "单价": 0.5},
    {"中文品名": "装饰雕塑", "英文品名": "Decorative sculpture", "商品编码": "3926400090", "材质": "Plastic", "用途": "HOME", "单价": 1},
    {"中文品名": "塑料花", "英文品名": "Plastic flowers", "商品编码": "6702104000", "材质": "Plastic", "用途": "HOME", "单价": 0.15},
    {"中文品名": "卡片", "英文品名": "Card", "商品编码": "4909004000", "材质": "Paper", "用途": "HOME", "单价": 0.1},
]

HEADERS = (
    "中文品名",
    "英文品名",
    "商品编码",
    "材质",
    "用途",
    "箱数",
    "数量",
    "单位",
    "币制",
    "单价",
    "总价",
    "净重",
    "毛重",
    "原产国",
)
DISPLAY_TAX_RATE_FIELD = "税率"
DISPLAY_TAX_AMOUNT_FIELD = "税金"
WORKBOOK_HEADERS = (*HEADERS, DISPLAY_TAX_RATE_FIELD, DISPLAY_TAX_AMOUNT_FIELD)

FIELD_TRANSLATION_ACRONYMS = {
    "ABS",
    "CBM",
    "DIY",
    "EVA",
    "LED",
    "OPP",
    "PC",
    "PE",
    "PP",
    "PU",
    "PVC",
    "TPE",
}

MATERIAL_TRANSLATIONS = (
    ("不锈钢", "Stainless steel"),
    ("合成皮革", "Synthetic leather"),
    ("聚酯纤维", "Polyester fiber"),
    ("牛津布", "Oxford cloth"),
    ("无纺布", "Non-woven fabric"),
    ("镀锌板", "Galvanized sheet"),
    ("锌合金", "Zinc alloy"),
    ("矿物质粉", "Mineral powder"),
    ("人造纤维", "Synthetic fiber"),
    ("人造丝", "Rayon"),
    ("亚克力", "Acrylic"),
    ("塑料片", "Plastic sheet"),
    ("塑胶", "Plastic"),
    ("塑料", "Plastic"),
    ("涤纶", "Polyester"),
    ("尼龙", "Nylon"),
    ("金属", "Metal"),
    ("钢铁", "Iron and steel"),
    ("黄铜", "Brass"),
    ("橡胶", "Rubber"),
    ("树脂", "Resin"),
    ("玻璃", "Glass"),
    ("皮革", "Leather"),
    ("纤维", "Fiber"),
    ("纸", "Paper"),
    ("铁", "Iron"),
    ("铜", "Copper"),
    ("木制", "Wood"),
    ("木架", "Wood frame"),
    ("木", "Wood"),
    ("布料", "Fabric"),
    ("布", "Cloth"),
    ("海绵", "Sponge"),
    ("陶瓷", "Ceramic"),
    ("硅胶", "Silicone"),
    ("泡沫", "Foam"),
    ("绒面", "Suede"),
    ("棉布", "Cotton cloth"),
    ("棉", "Cotton"),
    ("鞋面", "Upper"),
    ("鞋底", "Sole"),
    ("人造", "Artificial"),
)

USAGE_TRANSLATIONS = (
    ("高尔夫、电力工程、室内装修、户外广告安装等必备测量工具", "Measurement tool"),
    ("用于焊接配件", "Welding accessories"),
    ("家居窗户装饰", "Window treatment"),
    ("汽车行驶", "Automobile driving"),
    ("厨房用品", "Kitchen supplies"),
    ("宠物用品", "Pet supplies"),
    ("庭院用品", "Garden supplies"),
    ("户外用品", "Outdoor supplies"),
    ("生活用品", "Daily necessities"),
    ("日常使用", "Daily use"),
    ("家居用品", "Home supplies"),
    ("装饰用品", "Decoration supplies"),
    ("园艺作业", "Gardening"),
    ("测量戒指大小", "Ring size measurement"),
    ("测量雨量", "Rain gauge"),
    ("装饰照明", "Decorative lighting"),
    ("礼品包装", "Gift packaging"),
    ("居家生活用", "Home use"),
    ("茶具套装", "Tea set"),
    ("餐椅使用", "Dining chair use"),
    ("过滤土", "Soil filtering"),
    ("刮刀工具", "Caulk remover tool"),
    ("切药器", "Pill cutter"),
    ("输入数据", "Data input"),
    ("挂墙展示", "Wall display"),
    ("球存储", "Ball storage"),
    ("写贺卡", "Greeting card writing"),
    ("花园装饰", "Garden decoration"),
    ("生活日用", "Daily use"),
    ("化妆用", "Makeup"),
    ("穿着服饰", "Clothing"),
    ("化妆", "Makeup"),
    ("穿戴", "Wearing"),
    ("穿着", "Wearing"),
    ("家居", "Home use"),
    ("家用", "Home use"),
    ("户外", "Outdoor use"),
    ("办公", "Office use"),
    ("装饰", "Decoration"),
    ("保护", "Protection"),
    ("娱乐", "Entertainment"),
    ("防虫", "Insect protection"),
    ("喂鸟", "Bird feeding"),
    ("打字", "Typing"),
    ("摆件", "Decoration"),
    ("种菜", "Gardening"),
    ("健身", "Fitness"),
    ("凳子", "Stool"),
    ("收纳", "Storage"),
    ("过滤", "Filtering"),
    ("支撑", "Support"),
    ("饮具", "Drinkware"),
    ("粘贴", "Adhesive use"),
    ("眼扣", "Eyelet"),
    ("穿", "Wearing"),
)

ENGLISH_TERM_NORMALIZATIONS = (
    ("home supplies", "Home supplies"),
    ("home supply", "Home supplies"),
    ("household", "Household"),
    ("houshold", "Household"),
    ("daily necessities", "Daily necessities"),
    ("daily use", "Daily use"),
    ("outdoor supplies", "Outdoor supplies"),
    ("outdoor use", "Outdoor use"),
    ("garden supplies", "Garden supplies"),
    ("kitchen supplies", "Kitchen supplies"),
    ("pet supplies", "Pet supplies"),
    ("automobile driving", "Automobile driving"),
    ("decoration supplies", "Decoration supplies"),
    ("stainless steel", "Stainless steel"),
    ("zinc alloy", "Zinc alloy"),
    ("synthetic leather", "Synthetic leather"),
    ("polyester fiber", "Polyester fiber"),
    ("oxford cloth", "Oxford cloth"),
    ("non-woven fabric", "Non-woven fabric"),
    ("cotton cloth", "Cotton cloth"),
    ("iron and steel", "Iron and steel"),
    ("wood frame", "Wood frame"),
    ("plasticr", "Plastic"),
    ("plastics", "Plastic"),
    ("polyster", "Polyester"),
    ("snylon", "Nylon"),
    ("fibre", "Fiber"),
    ("plastic", "Plastic"),
    ("polyester", "Polyester"),
    ("nylon", "Nylon"),
    ("metal", "Metal"),
    ("iron", "Iron"),
    ("steel", "Steel"),
    ("glass", "Glass"),
    ("paper", "Paper"),
    ("wooden", "Wood"),
    ("wood", "Wood"),
    ("fabric", "Fabric"),
    ("cloth", "Cloth"),
    ("sponge", "Sponge"),
    ("foam", "Foam"),
    ("rubber", "Rubber"),
    ("acrylic", "Acrylic"),
    ("cotton", "Cotton"),
    ("canvas", "Canvas"),
    ("fiber", "Fiber"),
    ("velvet", "Velvet"),
    ("leather", "Leather"),
    ("ceramic", "Ceramic"),
    ("silicone", "Silicone"),
    ("brass", "Brass"),
    ("copper", "Copper"),
    ("resin", "Resin"),
    ("upper", "Upper"),
    ("sole", "Sole"),
    ("suede", "Suede"),
    ("wearing", "Wearing"),
    ("dress", "Wearing"),
    ("decoration", "Decoration"),
    ("decorate", "Decoration"),
    ("filtered", "Filtering"),
    ("filtering", "Filtering"),
    ("makeup", "Makeup"),
    ("massage", "Massage"),
    ("coaster", "Coaster"),
    ("tool", "Tool"),
    ("regular", "General use"),
    ("toy", "Toy"),
)

@dataclass
class ManifestItem:
    row: int
    zh: str
    en: str
    hs: str
    material: str
    usage: str
    ctns: Optional[float]
    qty: Optional[float]
    unit_price: Optional[float]
    declared_value: Optional[float]
    real_weight: Optional[float]
    gross_weight: Optional[float]


@dataclass
class ManifestSummary:
    filename: str
    row_count: int
    total_ctns: float
    total_real_weight: float
    total_declared_value: float
    categories: list[str]
    items: list[ManifestItem] = field(default_factory=list)
    total_net_weight: Optional[float] = None
    weight_source: str = ""
    weight_evidence: str = ""
    weight_confidence: float = 0.0
    schema_sheet: str = ""
    schema_header_row: int = 0
    schema_data_start_row: int = 0
    schema_data_end_row: int = 0
    schema_summary_rows: list[int] = field(default_factory=list)
    schema_columns: dict[str, int] = field(default_factory=dict)
    schema_confidence: float = 0.0
    weight_strategy: str = ""
    weight_total_cell: str = ""
    weight_detail_range: str = ""
    weight_explicit_total: Optional[float] = None
    weight_detail_sum: Optional[float] = None
    weight_final: Optional[float] = None
    weight_reconciled: bool = False


@dataclass
class ManifestWeightInfo:
    total_weight_kg: Optional[float]
    source: str = ""
    evidence: str = ""
    confidence: float = 0.0


@dataclass
class ManifestSchema:
    sheet_name: str
    header_row: int
    summary_rows: list[int]
    data_start_row: int
    data_end_row: int
    columns: dict[str, int]
    header_labels: dict[str, str]
    weight_unit: str
    confidence: float
    weight_strategy: str
    weight_total_cell: str = ""
    weight_detail_range: str = ""


@dataclass
class ManifestWeightResolution:
    strategy: str
    total_cell: str
    detail_range: str
    explicit_total: Optional[float]
    detail_sum: Optional[float]
    final_weight: float
    reconciled: bool
    evidence: str


@dataclass
class ManifestHsGroup:
    hs: str
    source_rows: list[int] = field(default_factory=list)
    zh_names: list[str] = field(default_factory=list)
    en_names: list[str] = field(default_factory=list)
    materials: list[str] = field(default_factory=list)
    usages: list[str] = field(default_factory=list)
    total_ctns: float = 0.0
    total_qty: float = 0.0
    total_real_weight: float = 0.0
    total_gross_weight: float = 0.0
    total_declared_value: float = 0.0
    unit_price_values: list[float] = field(default_factory=list)
    canonical_zh: str = ""
    canonical_en: str = ""
    representative_material: str = ""
    representative_usage: str = ""
    warnings: list[str] = field(default_factory=list)


@dataclass
class BillProduct:
    name: str
    hs_code_hint: str = ""
    evidence: str = ""
    confidence: float = 1.0


@dataclass
class BillInfo:
    filename: str
    raw_text: str
    products: list[str]
    shipper: str = ""
    consignee: str = ""
    shipment_no: str = ""
    eta: str = ""
    cartons: Optional[float] = None
    carton_evidence: str = ""
    carton_unit: str = ""
    carton_source: str = ""
    carton_reasoning: str = ""
    carton_manifest_comparison: str = ""
    carton_confidence: float = 0.0
    carton_inferred: bool = False
    carton_recognition_attempts: int = 1
    carton_resolution_history: list[dict[str, Any]] = field(default_factory=list)
    gross_weight: Optional[float] = None
    cbm: Optional[float] = None
    product_entries: list[BillProduct] = field(default_factory=list)
    parse_source: str = "text"
    text_chars: int = 0
    vision_pages: int = 0


@dataclass
class BillLLMFields:
    product_entries: list[BillProduct]
    shipper: str = ""
    consignee: str = ""
    carton_count: Optional[float] = None
    carton_evidence: str = ""
    carton_unit: str = ""
    carton_source: str = ""
    carton_reasoning: str = ""
    carton_manifest_comparison: str = ""
    carton_confidence: float = 0.0
    carton_inferred: bool = False


@dataclass(frozen=True)
class ProcessingOptions:
    target_tax_amount: float
    target_item_count: int
    requested_profile: str = "auto"


@dataclass
class SelectionRules:
    allowed_certifications: list[str] = field(default_factory=lambda: ["Lacey Act", "TSCA"])
    blocked_certifications: list[str] = field(default_factory=list)
    allowed_products: list[str] = field(default_factory=list)
    blocked_products: list[str] = field(default_factory=list)


@dataclass
class ProductCandidate:
    source: str
    source_label: str
    zh: str
    en: str
    hs: str
    material: str
    usage: str
    ctns: float = 0.0
    qty: float = 0.0
    unit_price: float = 0.0
    declared_value: float = 0.0
    real_weight: float = 0.0
    gross_weight: float = 0.0
    source_rows: list[int] = field(default_factory=list)
    tax_data: dict[str, Any] = field(default_factory=dict)
    base_tax_rate: float = 0.0
    effective_tax_rate: float = 0.0
    tax_match_source: str = ""
    certification_texts: list[str] = field(default_factory=list)
    filter_reason: str = ""
    original_hs: str = ""
    query_name: str = ""
    query_material: str = ""
    plausibility_range: Optional["PlausibilityRange"] = None
    plausibility_confidence: float = 0.0
    plausibility_basis: str = ""
    price_evidence: dict[str, Any] = field(default_factory=dict)
    llm_reason: str = ""
    row_warnings: list[str] = field(default_factory=list)
    compliance_review_required: bool = False
    compliance_review_reason: str = ""
    bill_product_name: str = ""
    bill_has_manifest_detail: bool = False

    @property
    def score(self) -> float:
        return self.real_weight + self.gross_weight + self.ctns * 8 + self.declared_value / 100 + self.qty / 20


@dataclass(frozen=True)
class PlausibilityRange:
    kg_per_ctn_min: Optional[float] = None
    kg_per_ctn_max: Optional[float] = None
    kg_per_pc_min: Optional[float] = None
    kg_per_pc_max: Optional[float] = None
    unit_price_min: Optional[float] = None
    unit_price_max: Optional[float] = None
    ctns_min: Optional[float] = None
    ctns_max: Optional[float] = None
    qty_per_ctn_min: Optional[float] = None
    qty_per_ctn_max: Optional[float] = None
    source: str = ""


@dataclass(frozen=True)
class RowPlan:
    candidate: ProductCandidate
    ctns: int
    qty: int
    gross_weight: float
    total_value: float
    unit_price: float
    plausibility: PlausibilityRange
    warnings: tuple[str, ...] = ()
    tax_budget: float = 0.0
    price_reference: float = 0.0
    quantity_basis: str = ""


QueryCache = dict[str, dict[str, Any]]
ProgressCallback = Callable[[dict[str, Any]], Any]


async def parse_input_documents_concurrently(
    manifest_path: str | Path,
    bill_path: str | Path,
    manifest_parser: LLMClient,
    bill_parser: LLMClient,
    query_cache: Optional[QueryCache] = None,
    metrics: Optional[dict[str, Any]] = None,
) -> tuple[ManifestSummary, BillInfo]:
    results = await asyncio.gather(
        parse_manifest(manifest_path, manifest_parser, query_cache),
        parse_bill(bill_path, bill_parser, query_cache, allow_missing_carton_count=True),
        return_exceptions=True,
    )
    fallbacks = 0
    manifest, bill = results
    if isinstance(manifest, BaseException):
        if isinstance(manifest, asyncio.CancelledError):
            raise manifest
        fallbacks += 1
        manifest = await parse_manifest(manifest_path, manifest_parser, query_cache)
    if isinstance(bill, BaseException):
        if isinstance(bill, asyncio.CancelledError):
            raise bill
        fallbacks += 1
        recognition_history: list[dict[str, Any]] = []
        bill_attempt = 1
        while isinstance(bill, BaseException) and bill_attempt < BILL_PACKAGE_RECOGNITION_MAX_ATTEMPTS:
            recognition_history.append(
                {
                    "attempt": bill_attempt,
                    "status": "error",
                    "error": summarize_model_error(bill),
                }
            )
            bill_attempt += 1
            try:
                is_final_attempt = bill_attempt >= BILL_PACKAGE_RECOGNITION_MAX_ATTEMPTS
                bill = await parse_bill(
                    bill_path,
                    bill_parser,
                    query_cache,
                    allow_missing_carton_count=True,
                    text_max_chars=(
                        BILL_PACKAGE_FINAL_RETRY_TEXT_CHARS
                        if is_final_attempt
                        else BILL_PARSER_TEXT_CHARS
                    ),
                    vision_max_pages=(
                        BILL_PACKAGE_FINAL_RETRY_MAX_PAGES
                        if is_final_attempt
                        else BILL_VISION_MAX_PAGES
                    ),
                )
            except asyncio.CancelledError:
                raise
            except Exception as exc:
                bill = exc
        if isinstance(bill, BaseException):
            recognition_history.append(
                {
                    "attempt": bill_attempt,
                    "status": "error",
                    "error": summarize_model_error(bill),
                }
            )
            save_bill_package_resolution_audit(
                query_cache,
                bill_path,
                manifest,
                recognition_history,
                status="failed",
            )
            raise RuntimeError(
                "提单解析失败，"
                f"模型已识别 {bill_attempt} 次，任务已中断；"
                f"最后错误：{summarize_model_error(bill)}"
            ) from None
        shifted_history = []
        for entry in bill.carton_resolution_history:
            shifted = dict(entry)
            shifted["attempt"] = bill_attempt
            shifted_history.append(shifted)
        if not shifted_history:
            shifted_history.append(
                {
                    "attempt": bill_attempt,
                    "status": "recognized" if bill.cartons and bill.cartons > 0 else "missing",
                    "package_count": bill.cartons,
                    "package_unit": bill.carton_unit,
                    "source": bill.carton_source or bill.parse_source,
                    "evidence": bill.carton_evidence,
                    "reasoning": bill.carton_reasoning,
                }
            )
        bill = replace(
            bill,
            carton_recognition_attempts=bill_attempt,
            carton_resolution_history=[*recognition_history, *shifted_history],
        )
    bill = await resolve_bill_carton_count(
        bill_path,
        bill,
        manifest,
        bill_parser,
        query_cache,
    )
    if metrics is not None:
        metrics.update(
            {
                "max_concurrency": 2,
                "sequential_fallbacks": fallbacks,
                "carton_recognition_attempts": bill.carton_recognition_attempts,
                "carton_source": bill.carton_source,
                "carton_inferred": bill.carton_inferred,
            }
        )
    return manifest, bill


def record_stage_timing(stage_timings: dict[str, float], stage: str, started_at: float) -> None:
    stage_timings[stage] = round(time.perf_counter() - started_at, 3)


async def build_clearance(
    manifest_path: str | Path,
    bill_path: str | Path,
    output_dir: str | Path,
    requested_profile: str = "auto",
    target_tax_amount: Optional[float] = None,
    target_item_count: Optional[int] = None,
    query_cache: Optional[QueryCache] = None,
    bill_parser: Optional[LLMClient] = None,
    manifest_parser: Optional[LLMClient] = None,
    llm: Optional[LLMClient] = None,
    progress_callback: Optional[ProgressCallback] = None,
) -> dict[str, Any]:
    build_started_at = time.perf_counter()
    stage_timings: dict[str, float] = {}
    options = validate_processing_options(target_tax_amount, target_item_count, requested_profile)
    query_cache = query_cache if query_cache is not None else {"product": {}, "hs": {}, "bill": {}}
    llm_client = llm or bill_parser or manifest_parser or LLMClient()
    llm_parse_cache_before = bool(
        query_cache.get("manifest_schema") or query_cache.get("manifest") or query_cache.get("bill")
    )
    stage_started_at = time.perf_counter()
    parse_metrics: dict[str, Any] = {}
    manifest, bill = await parse_input_documents_concurrently(
        manifest_path,
        bill_path,
        manifest_parser or llm_client,
        bill_parser or llm_client,
        query_cache,
        metrics=parse_metrics,
    )
    record_stage_timing(stage_timings, "parse_documents", stage_started_at)
    llm_parse_used = bool(
        query_cache.get("manifest_schema") or query_cache.get("manifest") or query_cache.get("bill")
    )
    llm_plausibility_used = False
    llm_generation_used = False
    if manifest.total_real_weight <= 0:
        raise RuntimeError("清单 Excel 未识别到有效总重量，不能继续生成")
    if not bill.cartons or bill.cartons <= 0:
        raise RuntimeError("提单未识别到有效总箱数，不能生成与提单箱数对齐的输出")

    flow: list[dict[str, Any]] = [
        {
            "stage": "parse",
            "status": "ok",
            "manifest_rows": manifest.row_count,
            "bill_products": len(bill.products),
            "manifest_total_weight": manifest.total_real_weight,
            "manifest_total_net_weight": manifest.total_net_weight,
            "manifest_net_to_gross_ratio": manifest_net_to_gross_ratio(manifest),
            "manifest_total_ctns": manifest.total_ctns,
            "bill_package_count": bill.cartons,
            "bill_package_unit": bill.carton_unit,
            "bill_package_source": bill.carton_source,
            "bill_package_inferred": bill.carton_inferred,
            "bill_package_recognition_attempts": bill.carton_recognition_attempts,
            "bill_manifest_package_comparison": bill.carton_manifest_comparison,
            "manifest_schema_sheet": manifest.schema_sheet,
            "manifest_schema_columns": manifest.schema_columns,
            "weight_strategy": manifest.weight_strategy,
            "weight_total_cell": manifest.weight_total_cell,
            "weight_explicit_total": manifest.weight_explicit_total,
            "weight_detail_sum": manifest.weight_detail_sum,
            "weight_final": manifest.weight_final,
            "weight_reconciled": manifest.weight_reconciled,
            "target_tax_amount": options.target_tax_amount,
            "target_item_count": options.target_item_count,
        }
    ]
    await emit_progress(
        progress_callback,
        {
            "stage": "parse",
            "status": "ok",
            "progress": 5,
            "message": "文件解析完成",
            "manifest_rows": manifest.row_count,
            "manifest_total_weight": manifest.total_real_weight,
            "manifest_total_net_weight": manifest.total_net_weight,
            "manifest_net_to_gross_ratio": manifest_net_to_gross_ratio(manifest),
        },
    )

    rules = load_selection_rules()
    crawler = StrictTaxCrawler()
    replacement_pool = load_replacement_candidates()
    plausibility_ranges = load_plausibility_ranges()
    if len(bill.products) > options.target_item_count:
        raise RuntimeError(
            f"提单品类 {len(bill.products)} 个超过目标输出行数 {options.target_item_count}，"
            "请提高最终生成条目数"
        )
    manifest_candidates = build_manifest_candidates(manifest)
    ensure_manifest_codeflag_candidates(manifest, manifest_candidates)
    manifest_hs_candidates = sum(1 for candidate in manifest_candidates if len(normalize_hs(candidate.hs)) == 10)
    manifest_product_only_candidates = len(manifest_candidates) - manifest_hs_candidates
    qualified_manifest: list[ProductCandidate] = []
    manifest_filtered: list[ProductCandidate] = []
    manifest_query_limit = len(manifest_candidates)
    stage_started_at = time.perf_counter()
    qualified_manifest, manifest_filtered = await qualify_candidates(
        crawler,
        manifest_candidates[:manifest_query_limit],
        rules,
        query_cache=query_cache,
        enforce_tax_limit=False,
        progress_callback=progress_callback,
        progress_stage="crawler_manifest_products",
        progress_start=12,
        progress_end=52,
    )
    record_stage_timing(stage_timings, "codeflag_manifest_candidates", stage_started_at)
    flow.append(
        {
            "stage": "crawler_manifest_products",
            "status": "ok" if qualified_manifest else "insufficient",
            "queried_rows": manifest_query_limit,
            "candidate_groups": len(manifest_candidates),
            "hs_candidates": manifest_hs_candidates,
            "product_only_candidates": manifest_product_only_candidates,
            "qualified": len(qualified_manifest),
            "filtered": len(manifest_filtered),
            "message": "已按客户清单 HS 归并池全量优先查询税率；人工发票策略下单品税率不作 20% 硬过滤",
        }
    )
    bill_crawler: StrictTaxCrawler | None = None
    bill_required: list[ProductCandidate] = []
    bill_filtered: list[ProductCandidate] = []
    if bill.products:
        stage_started_at = time.perf_counter()
        await emit_progress(
            progress_callback,
            {
                "stage": "bill_products",
                "status": "running",
                "progress": 54,
                "message": "正在用清单候选覆盖提单品类",
            },
        )
        bill_crawler = StrictTaxCrawler()
        bill_required, bill_filtered = await qualify_bill_product_candidates(
            bill_crawler,
            bill,
            qualified_manifest,
            replacement_pool,
            rules,
            options,
            query_cache=query_cache,
            llm=llm_client,
        )
        unique_bill_count = len(unique_bill_products(bill.products))
        reused_manifest_count = sum(
            1 for candidate in bill_required if candidate.source == "manifest_group"
        )
        flow.append(
            {
                "stage": "bill_products",
                "status": "ok" if len(bill_required) == unique_bill_count else "insufficient",
                "bill_products": unique_bill_count,
                "qualified": len(bill_required),
                "manifest_reused": reused_manifest_count,
                "queried_separately": len(bill_required) - reused_manifest_count,
                "filtered": len(bill_filtered),
                "message": "提单品类优先复用已通过 Codeflag 的清单候选；仅未覆盖品类单独查询",
            }
        )
        if len(bill_required) < unique_bill_count:
            missing = format_missing_bill_products(bill.products, bill_filtered)
            raise RuntimeError(
                "提单品类 Codeflag 查询缺少合格归类结果，不能套用替换表品名；"
                f"请补充提单 HS 或调整品名/材质: {missing}"
            )
        record_stage_timing(stage_timings, "codeflag_bill_products", stage_started_at)
    manifest_filtered.extend(bill_filtered)

    bill_clean_replacement_pool = exclude_bill_product_replacements(replacement_pool, bill.products)
    selected = select_initial_candidates(qualified_manifest, bill_required, rules, options.target_item_count)
    replacement_filtered: list[ProductCandidate] = []
    replacement_used = 0
    codebook_initial_fill_used = 0
    codebook_fill_needed = max(0, options.target_item_count - len(selected))
    if len(selected) < options.target_item_count:
        needed = options.target_item_count - len(selected)
        replacements = sorted(bill_clean_replacement_pool, key=lambda item: selection_score(item, rules), reverse=True)
        selected_keys = {candidate_identity(candidate) for candidate in selected}
        selected_semantic_keys = set().union(
            *(candidate_semantic_keys(candidate) for candidate in selected)
        ) if selected else set()
        replacement_attempts = 0
        for replacement in replacements:
            if len(selected) >= options.target_item_count:
                break
            replacement_semantic_keys = candidate_semantic_keys(replacement)
            if (
                candidate_identity(replacement) in selected_keys
                or replacement_semantic_keys & selected_semantic_keys
            ):
                continue
            if is_customer_codebook_candidate(replacement):
                selected.append(replacement)
                selected_keys.add(candidate_identity(replacement))
                selected_semantic_keys.update(replacement_semantic_keys)
                replacement_used += 1
                codebook_initial_fill_used += 1
                continue
            if replacement_attempts > 0:
                await asyncio.sleep(crawler.settings.delay)
            replacement_attempts += 1
            await emit_progress(
                progress_callback,
                {
                    "stage": "replacement_products",
                    "status": "running",
                    "progress": min(88, 62 + replacement_attempts),
                    "message": f"正在查询替换候选 {replacement_attempts}",
                    "current": replacement_attempts,
                    "selected": len(selected),
                    "target": options.target_item_count,
                },
            )
            qualified, filtered = await qualify_candidates(
                crawler,
                [replacement],
                rules,
                query_cache=query_cache,
                enforce_tax_limit=False,
            )
            replacement_filtered.extend(filtered)
            if not qualified:
                continue
            selected.append(qualified[0])
            selected_keys.add(candidate_identity(qualified[0]))
            selected_semantic_keys.update(candidate_semantic_keys(qualified[0]))
            replacement_used += 1
        flow.append(
            {
                "stage": "replacement_products",
                "status": "ok" if len(selected) >= options.target_item_count else "insufficient",
                "needed": needed,
                "used": replacement_used,
                "codebook_fill_needed": codebook_fill_needed,
                "codebook_fill_used": codebook_initial_fill_used,
                "filtered": len(replacement_filtered),
            }
        )

    shortage_notice = ""
    if len(selected) < options.target_item_count:
        shortage_notice = (
            f"合格品名不足，目标 {options.target_item_count} 行，当前仅 {len(selected)} 行；"
            "按客户确认规则输出合格部分，并列出缺口原因"
        )
        flow.append(
            {
                "stage": "candidate_shortage",
                "status": "partial",
                "target": options.target_item_count,
                "selected": len(selected),
                "missing": options.target_item_count - len(selected),
                "message": shortage_notice,
                "filter_summary": summarize_filter_reasons([*manifest_filtered, *replacement_filtered]),
            }
        )
        await emit_progress(
            progress_callback,
            {
                "stage": "candidate_shortage",
                "status": "partial",
                "progress": 89,
                "message": shortage_notice,
                "selected": len(selected),
                "target": options.target_item_count,
            },
        )
    effective_options = options
    if selected and len(selected) != options.target_item_count:
        effective_options = replace(options, target_item_count=len(selected))

    await emit_progress(
        progress_callback,
        {
            "stage": "optimize_output",
            "status": "running",
            "progress": 90,
            "message": "正在使用规则优化器生成最终草案",
        },
    )
    plausibility_metrics: dict[str, Any] = {}
    stage_started_at = time.perf_counter()
    selected, llm_plausibility_used = await ensure_candidate_plausibility_ranges(
        selected,
        manifest,
        bill,
        plausibility_ranges,
        llm_client,
        query_cache,
        metrics=plausibility_metrics,
    )
    record_stage_timing(stage_timings, "llm_plausibility", stage_started_at)

    stage_started_at = time.perf_counter()
    replacement_repair_attempts = 0
    candidate_library_allowed = codebook_fill_needed > 0
    stable_reference_candidates = (
        load_stable_manual_invoice_candidates(selected, bill.products)
        if candidate_library_allowed
        else []
    )
    manual_invoice_pool = build_manual_invoice_candidate_pool(
        selected,
        qualified_manifest,
        stable_reference_candidates,
        allow_candidate_library=candidate_library_allowed,
    )
    selected, price_repair_summary = optimize_selected_candidates_for_manual_invoice(
        selected,
        manual_invoice_pool,
        manifest,
        bill,
        effective_options,
    )
    price_repair_summary["candidate_library_allowed"] = candidate_library_allowed
    price_repair_summary["candidate_library_reason"] = (
        "清单和提单候选不足以达到目标行数，允许候选库补足"
        if candidate_library_allowed
        else "清单和提单候选已达到目标行数，候选库不参与换品"
    )
    adaptive_replacement_skipped = (
        not candidate_library_allowed
        or manual_invoice_solution_is_acceptable(price_repair_summary, effective_options)
    )
    if adaptive_replacement_skipped:
        skip_reason = (
            "清单和提单候选已满足目标行数，候选库仅用于真实行数缺口"
            if not candidate_library_allowed
            else "现有清单候选和补充候选已有可行解"
        )
        flow.append(
            {
                "stage": "manual_invoice_replacement_pool",
                "status": "skipped",
                "reason": skip_reason,
                "attempts": 0,
            }
        )
    else:
        await emit_progress(
            progress_callback,
            {
                "stage": "manual_invoice_replacement_pool",
                "status": "running",
                "progress": 91,
                "message": "首轮优化无可行解，正在按需查询低税/免税替换候选",
            },
        )
        replacement_seeds = build_manual_invoice_replacement_seeds(
            bill_clean_replacement_pool, selected, bill.products
        )
        replacement_repair_candidates: list[ProductCandidate] = []
        replacement_cursor = 0
        while (
            not manual_invoice_solution_is_acceptable(price_repair_summary, effective_options)
            and replacement_cursor < len(replacement_seeds)
            and replacement_repair_attempts < MANUAL_INVOICE_REPLACEMENT_QUERY_LIMIT
        ):
            batch, batch_filtered, replacement_cursor, batch_attempts = await qualify_manual_invoice_replacement_batch(
                crawler,
                replacement_seeds,
                replacement_cursor,
                [*selected, *replacement_repair_candidates],
                rules,
                prior_network_attempts=replacement_repair_attempts,
                batch_size=ADAPTIVE_REPLACEMENT_BATCH_SIZE,
                query_cache=query_cache,
                progress_callback=progress_callback,
            )
            replacement_repair_attempts += batch_attempts
            replacement_repair_candidates.extend(batch)
            replacement_filtered.extend(batch_filtered)
            if not batch and batch_attempts == 0:
                break
            manual_invoice_pool = dedupe_candidates(
                [
                    *selected,
                    *qualified_manifest,
                    *stable_reference_candidates,
                    *replacement_repair_candidates,
                ]
            )
            selected, price_repair_summary = optimize_selected_candidates_for_manual_invoice(
                selected,
                manual_invoice_pool,
                manifest,
                bill,
                effective_options,
            )
    price_repair_summary["replacement_attempts"] = replacement_repair_attempts
    record_stage_timing(stage_timings, "adaptive_candidate_optimization", stage_started_at)

    price_metrics: dict[str, Any] = {}
    stage_started_at = time.perf_counter()
    selected = await attach_price_evidence_to_candidates_concurrently(
        selected,
        query_cache=query_cache,
        max_concurrency=LLM_MAX_CONCURRENCY,
        metrics=price_metrics,
    )
    record_stage_timing(stage_timings, "price_evidence", stage_started_at)
    draft_attempts = 0
    draft_feedback: list[str] = ["最终草案默认由规则优化器生成，LLM 不参与数值草案生成"]
    if price_repair_summary.get("swaps"):
        draft_feedback.append(
            "已按人工发票策略从低税/免税候选重排 "
            f"{price_repair_summary['swaps']} 行"
        )
    stage_started_at = time.perf_counter()
    best_effort_reason = ""
    try:
        rows = build_output_rows(selected, manifest, bill, effective_options)
    except RuntimeError as exc:
        best_effort_reason = str(exc)
        rows = build_output_rows_best_effort(selected, manifest, bill, effective_options, best_effort_reason)
        draft_feedback.append(f"正式数值求解失败，当前为兜底文件: {best_effort_reason}")
        flow.append(
            {
                "stage": "best_effort_output",
                "status": "partial",
                "message": f"规则优化无完整可行解，已输出最接近方案: {best_effort_reason}",
            }
        )
    record_stage_timing(stage_timings, "build_numeric_rows", stage_started_at)

    llm_review_used = False
    llm_review_reoptimized = False
    llm_review_constraints_applied = 0
    llm_review: dict[str, Any] = {"pass": True, "issues": []}
    blocking_review_issues: list[dict[str, Any]] = []
    await emit_progress(
        progress_callback,
        {
            "stage": "llm_output_review",
            "status": "running",
            "progress": 94,
            "message": "正在进行最终商业合理性审查",
        },
    )
    stage_started_at = time.perf_counter()
    try:
        llm_review = await review_output_rows_with_llm(
            llm_client,
            rows,
            selected,
            manifest,
            bill,
            effective_options,
            query_cache=query_cache,
        )
        llm_review_used = True
        adjusted_selected, llm_review_constraints_applied = apply_llm_review_constraints(
            selected, rows, llm_review
        )
        if llm_review_constraints_applied:
            try:
                reviewed_rows = build_output_rows(
                    adjusted_selected, manifest, bill, effective_options
                )
                selected = adjusted_selected
                rows = reviewed_rows
                llm_review_reoptimized = True
                best_effort_reason = ""
            except RuntimeError as exc:
                llm_review_constraints_applied = 0
                draft_feedback.append(f"LLM 审查约束无法形成可行解，保留原方案: {exc}")
        append_llm_review_warnings(rows, llm_review)
        blocking_review_issues = blocking_llm_review_issues(
            llm_review,
            constraints_applied=llm_review_constraints_applied,
            reoptimized=llm_review_reoptimized,
            candidates=selected,
        )
    except Exception as exc:
        draft_feedback.append(f"最终 LLM 合理性审查跳过: {exc}")
    record_stage_timing(stage_timings, "llm_output_review", stage_started_at)

    stage_started_at = time.perf_counter()
    try:
        await asyncio.wait_for(
            translate_output_chinese_names_with_llm(llm_client, rows),
            timeout=LLM_TRANSLATION_TIMEOUT_SECONDS,
        )
    except Exception as exc:
        draft_feedback.append(f"中文品名 LLM 翻译跳过: {exc}")
    record_stage_timing(stage_timings, "llm_translation", stage_started_at)
    ensure_bill_products_present(rows, bill.products, selected)
    validate_output_rows(rows)
    estimated_tax = round(sum((to_float(row.get("总价")) or 0) * (to_float(row.get("综合税率")) or 0) for row in rows), 2)
    tax_gap = round(estimated_tax - options.target_tax_amount, 2)
    tax_within_tolerance = abs(tax_gap) <= TAX_FINAL_TOLERANCE_USD
    constraint_reasons: list[str] = []
    if best_effort_reason:
        constraint_reasons.append(
            "正式数值求解失败，当前文件为兜底结果，未完成税金优化: "
            f"{best_effort_reason}"
        )
    if not tax_within_tolerance:
        constraint_reasons.append(
            f"税金与目标相差 {tax_gap:+.2f} USD，超过允许偏差 {TAX_FINAL_TOLERANCE_USD:.2f} USD"
        )
    if shortage_notice:
        constraint_reasons.append(shortage_notice)
    if blocking_review_issues:
        constraint_reasons.append(
            "最终商业合理性审查高风险提示（仅供人工复核，不阻止输出）: "
            f"{format_blocking_llm_review_issues(blocking_review_issues)}"
        )
    compliance_review_reasons = list(
        dict.fromkeys(
            candidate.compliance_review_reason
            for candidate in selected
            if candidate.compliance_review_required and candidate.compliance_review_reason
        )
    )
    if compliance_review_reasons:
        constraint_reasons.append("合规条件需人工复核: " + "；".join(compliance_review_reasons))
    constraint_status = "passed" if not constraint_reasons else "needs_review"
    validate_price_evidence(rows, selected)
    flow.append(
        {
            "stage": "optimize_output",
            "status": "ok" if constraint_status == "passed" else "needs_review",
            "rows": len(rows),
            "estimated_tax": estimated_tax,
            "tax_gap": tax_gap,
            "tax_within_tolerance": tax_within_tolerance,
            "constraint_reasons": constraint_reasons,
            "final_draft_generator": "best_effort" if best_effort_reason else "rules",
            "llm_draft_attempts": draft_attempts,
            "draft_feedback": draft_feedback[-1] if draft_feedback else "",
            "price_fit_repair": price_repair_summary,
            "adaptive_replacement_skipped": adaptive_replacement_skipped,
            "plausibility_metrics": plausibility_metrics,
            "price_evidence_metrics": price_metrics,
            "llm_review": {
                "used": llm_review_used,
                "issues": len(llm_review.get("issues") or []),
                "constraints_applied": llm_review_constraints_applied,
                "reoptimized": llm_review_reoptimized,
            },
        }
    )

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"清关_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:6]}.xlsx"
    await emit_progress(
        progress_callback,
        {
            "stage": "write_workbook",
            "status": "running",
            "progress": 96,
            "message": "正在写入 Excel",
        },
    )
    stage_started_at = time.perf_counter()
    write_workbook(bill, rows, output_path, metadata={})
    record_stage_timing(stage_timings, "write_workbook", stage_started_at)
    await emit_progress(
        progress_callback,
        {
            "stage": "done",
            "status": "ok",
            "progress": 100,
            "message": "处理完成",
        },
    )

    total_value = round(sum(to_float(row.get("总价")) or 0 for row in rows), 2)
    final_replacement_rows = sum(1 for candidate in selected if candidate.source == "replacement")
    crawler_auth_expired_retries = crawler.auth_expired_retry_count
    if bill_crawler is not None:
        crawler_auth_expired_retries += bill_crawler.auth_expired_retry_count
    stage_timings["build_clearance_total"] = round(time.perf_counter() - build_started_at, 3)
    stats = {
        "profile_hint": options.requested_profile or "auto",
        "constraint_status": constraint_status,
        "constraint_reasons": constraint_reasons,
        "weight_source": manifest.weight_source or "manifest_total_weight",
        "weight_evidence": manifest.weight_evidence,
        "weight_confidence": manifest.weight_confidence,
        "manifest_schema_sheet": manifest.schema_sheet,
        "manifest_schema_header_row": manifest.schema_header_row,
        "manifest_schema_data_start_row": manifest.schema_data_start_row,
        "manifest_schema_data_end_row": manifest.schema_data_end_row,
        "manifest_schema_summary_rows": manifest.schema_summary_rows,
        "manifest_schema_columns": manifest.schema_columns,
        "manifest_schema_confidence": manifest.schema_confidence,
        "weight_strategy": manifest.weight_strategy,
        "weight_total_cell": manifest.weight_total_cell,
        "weight_detail_range": manifest.weight_detail_range,
        "weight_explicit_total": manifest.weight_explicit_total,
        "weight_detail_sum": manifest.weight_detail_sum,
        "weight_final": manifest.weight_final,
        "weight_reconciled": manifest.weight_reconciled,
        "manifest_rows": manifest.row_count,
        "manifest_product_rows": manifest.row_count,
        "manifest_unique_candidates": len(manifest_candidates),
        "manifest_hs_candidates": manifest_hs_candidates,
        "manifest_product_only_candidates": manifest_product_only_candidates,
        "codeflag_queried": manifest_query_limit,
        "codeflag_qualified": len(qualified_manifest),
        "codebook_fill_needed": codebook_fill_needed,
        "codebook_fill_used": codebook_initial_fill_used,
        "bill_parse_source": bill.parse_source,
        "bill_text_chars": bill.text_chars,
        "bill_vision_pages": bill.vision_pages,
        "bill_package_count": bill.cartons,
        "bill_package_unit": bill.carton_unit,
        "bill_package_source": bill.carton_source,
        "bill_package_inferred": bill.carton_inferred,
        "bill_package_confidence": bill.carton_confidence,
        "bill_package_recognition_attempts": bill.carton_recognition_attempts,
        "bill_manifest_package_comparison": bill.carton_manifest_comparison,
        "input_categories": len(manifest.categories),
        "target_item_count": options.target_item_count,
        "actual_item_count": len(rows),
        "candidate_shortage": shortage_notice,
        "output_rows": len(rows),
        "target_tax_amount": options.target_tax_amount,
        "estimated_tax_amount": estimated_tax,
        "tax_gap": tax_gap,
        "tax_within_tolerance": tax_within_tolerance,
        "input_ctns": manifest.total_ctns,
        "output_ctns": round(sum(to_float(row.get("箱数")) or 0 for row in rows), 2),
        "input_real_weight": manifest.total_real_weight,
        "input_net_weight": manifest.total_net_weight,
        "input_net_to_gross_ratio": manifest_net_to_gross_ratio(manifest),
        "output_net_weight": round(sum(to_float(row.get("净重")) or 0 for row in rows), 2),
        "output_gross_weight": round(sum(to_float(row.get("毛重")) or 0 for row in rows), 2),
        "output_net_to_gross_ratio": round(
            sum(to_float(row.get("净重")) or 0 for row in rows)
            / max(0.01, sum(to_float(row.get("毛重")) or 0 for row in rows)),
            6,
        ),
        "bill_gross_weight_ignored": bill.gross_weight,
        "input_declared_value": manifest.total_declared_value,
        "total_value_usd": total_value,
        "plausibility_warnings": sum(1 for row in rows if clean_text(row.get("约束提示"))),
        "bill_products": len(unique_bill_products(bill.products)),
        "bill_required_locked": sum(1 for candidate in selected if candidate.bill_product_name),
        "bill_manifest_detailed_locked": sum(
            1
            for candidate in selected
            if candidate.bill_product_name and candidate.bill_has_manifest_detail
        ),
        "bill_synthetic_locked": sum(
            1
            for candidate in selected
            if candidate.bill_product_name and not candidate.bill_has_manifest_detail
        ),
        "qualified_manifest_candidates": len(qualified_manifest),
        "replacement_candidates_used": final_replacement_rows,
        "replacement_initial_fill_used": replacement_used,
        "replacement_price_repair_attempts": replacement_repair_attempts,
        "adaptive_replacement_skipped": adaptive_replacement_skipped,
        "candidate_library_allowed": candidate_library_allowed,
        "candidate_library_reason": price_repair_summary["candidate_library_reason"],
        "manifest_origin_rows": sum(1 for row in rows if clean_text(row.get("来源")) == "manifest_group"),
        "replacement_ratio": round(final_replacement_rows / max(1, len(rows)), 4),
        "compliance_review_rows": sum(1 for candidate in selected if candidate.compliance_review_required),
        "compliance_review_reasons": compliance_review_reasons,
        "price_evidence_count": sum(1 for candidate in selected if candidate.price_evidence),
        "filtered_candidates": len(manifest_filtered) + len(replacement_filtered),
        "realism_status": "passed" if not any(clean_text(row.get("约束提示")) for row in rows) else "needs_review",
        "realism_warnings": sum(1 for row in rows if clean_text(row.get("约束提示"))),
        "llm_used": llm_parse_used or llm_plausibility_used or llm_generation_used or llm_review_used,
        "llm_parse_used": llm_parse_used,
        "llm_parse_cache_reused": llm_parse_cache_before,
        "llm_parse_metrics": parse_metrics,
        "llm_plausibility_used": llm_plausibility_used,
        "llm_plausibility_metrics": plausibility_metrics,
        "llm_generation_used": llm_generation_used,
        "llm_review_used": llm_review_used,
        "llm_review_issue_count": len(llm_review.get("issues") or []),
        "llm_review_high_risk_warning_count": len(blocking_review_issues),
        "llm_review_constraints_applied": llm_review_constraints_applied,
        "llm_review_reoptimized": llm_review_reoptimized,
        "price_evidence_metrics": price_metrics,
        "stage_timings_seconds": stage_timings,
        "llm_max_concurrency": LLM_MAX_CONCURRENCY,
        "llm_draft_attempts": draft_attempts,
        "final_draft_generator": "best_effort" if best_effort_reason else "rules",
        "solver_status": "best_effort" if best_effort_reason else "optimized",
        "tax_optimization_applied": not bool(best_effort_reason),
        "best_effort_reason": best_effort_reason,
        "crawler_used": True,
        "crawler_auth_expired_retries": crawler_auth_expired_retries,
    }

    return {
        "stats": stats,
        "flow": flow,
        "manifest": manifest_to_public_dict(manifest),
        "bill": bill_to_public_dict(bill),
        "bill_categories": bill.products,
        "output_rows": rows,
        "input_tax_records": {candidate.hs: candidate.tax_data for candidate in qualified_manifest if candidate.hs},
        "output_tax_records": {normalize_hs(row.get("商品编码")): row.get("爬虫品名") for row in rows},
        "audit": build_audit_summary(rows, selected, manifest),
        "filter_summary": summarize_filter_reasons([*manifest_filtered, *replacement_filtered]),
        "output_file": str(output_path),
    }


def validate_processing_options(
    target_tax_amount: Optional[float],
    target_item_count: Optional[int],
    requested_profile: str,
) -> ProcessingOptions:
    tax_amount = to_float(target_tax_amount)
    if tax_amount is None or tax_amount <= 0:
        raise RuntimeError("期望税金必须大于 0")
    try:
        item_count = int(target_item_count) if target_item_count is not None else 0
    except (TypeError, ValueError) as exc:
        raise RuntimeError("最终生成条目数必须是整数") from exc
    if item_count < 1 or item_count > MAX_OUTPUT_ITEMS:
        raise RuntimeError(f"最终生成条目数必须在 1-{MAX_OUTPUT_ITEMS} 之间")
    return ProcessingOptions(
        target_tax_amount=round(tax_amount, 2),
        target_item_count=item_count,
        requested_profile=requested_profile or "auto",
    )


async def qualify_bill_product_candidates(
    crawler: StrictTaxCrawler,
    bill: BillInfo,
    qualified_manifest: list[ProductCandidate],
    replacement_candidates: list[ProductCandidate],
    rules: SelectionRules,
    options: ProcessingOptions,
    query_cache: Optional[QueryCache] = None,
    llm: Optional[LLMClient] = None,
) -> tuple[list[ProductCandidate], list[ProductCandidate]]:
    _ = replacement_candidates
    required: list[ProductCandidate] = []
    filtered: list[ProductCandidate] = []
    seen: set[tuple[str, str, str]] = set()
    for product in unique_bill_products(bill.products):
        bill_entry = bill_product_entry_for(bill, product)
        bill_material = infer_bill_material_from_entry(bill_entry) if bill_entry else ""
        rejected_matches = [
            candidate
            for candidate in qualified_manifest
            if candidate_matches_single_bill_product(candidate, product)
            and (
                not candidate_has_product_tax_match(candidate)
                or not candidate_material_matches_bill(candidate, bill_material)
            )
        ]
        filtered.extend(
            replace(
                candidate,
                filter_reason=f"提单品类候选未通过品名+材质查询确认，不能用于提单品类: {product}",
            )
            for candidate in rejected_matches
        )
        match = next(
            (
                candidate
                for candidate in qualified_manifest
                if candidate_matches_single_bill_product(candidate, product)
                and candidate_has_product_tax_match(candidate)
                and candidate_material_matches_bill(candidate, bill_material)
            ),
            None,
        )
        matched_manifest_detail = match is not None
        if not match:
            if bill_entry and bill_entry.hs_code_hint:
                result = await qualify_bill_hs_tax_candidate(
                    crawler,
                    build_bill_product_candidate(bill, bill_entry, options),
                    rules,
                    query_cache=query_cache,
                    llm=llm,
                )
                if result.filter_reason:
                    filtered.append(result)
                else:
                    match = result
        if not match:
            result = await qualify_bill_product_query_candidate(
                crawler,
                bill,
                product,
                product,
                bill_material,
                rules,
                options,
                query_cache=query_cache,
                match_source="bill_product",
            )
            if result.filter_reason:
                filtered.append(result)
            else:
                match = result
        if not match and bill_material:
            result = await qualify_bill_product_query_candidate(
                crawler,
                bill,
                product,
                product,
                "",
                rules,
                options,
                query_cache=query_cache,
                match_source="bill_product_no_material",
                llm_reason=f"材质查询失败后改用空材质查询: {bill_material}",
            )
            if result.filter_reason:
                filtered.append(result)
            else:
                match = result
        if not match and llm is not None:
            terms = await generate_bill_product_query_terms(
                llm,
                bill,
                product,
                [candidate.filter_reason for candidate in filtered[-6:] if candidate.filter_reason],
                query_cache,
            )
            attempts = 0
            for term in terms:
                if attempts > 0:
                    await asyncio.sleep(crawler.settings.delay)
                attempts += 1
                result = await qualify_bill_product_query_candidate(
                    crawler,
                    bill,
                    product,
                    term.get("product_name") or product,
                    term.get("material") or "",
                    rules,
                    options,
                    query_cache=query_cache,
                    match_source="llm_query",
                    llm_reason=term.get("reason") or "",
                )
                if result.filter_reason:
                    filtered.append(result)
                    continue
                match = result
                break
        if not match:
            recent_reasons = [
                candidate.filter_reason
                for candidate in filtered[-8:]
                if candidate.source == "bill"
                and normalize_bill_product_text(candidate.zh) == normalize_bill_product_text(product)
                and candidate.filter_reason
            ]
            reason_detail = "；".join(recent_reasons[-3:])
            filtered.append(
                ProductCandidate(
                    source="bill",
                    source_label=bill.filename,
                    zh=product,
                    en=product,
                    hs="",
                    material="",
                    usage="",
                    filter_reason=(
                        f"提单品类 Codeflag 查询未返回合格归类结果: {product}"
                        + (f"；最近失败原因: {reason_detail}" if reason_detail else "")
                    ),
                )
            )
            continue
        match = replace(
            match,
            bill_product_name=product,
            bill_has_manifest_detail=matched_manifest_detail,
        )
        key = candidate_identity(match)
        if key not in seen:
            seen.add(key)
            required.append(match)
    if len(required) > options.target_item_count:
        raise RuntimeError(f"提单品类 {len(required)} 个超过目标输出行数 {options.target_item_count}")
    return required, filtered


def unique_bill_products(products: list[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for product in products:
        normalized = normalize_text(product)
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        result.append(product)
    return result


def format_missing_bill_products(products: list[str], filtered: list[ProductCandidate]) -> str:
    parts: list[str] = []
    for product in unique_bill_products(products):
        product_key = normalize_bill_product_text(product)
        reasons = [
            candidate.filter_reason
            for candidate in filtered
            if candidate.filter_reason
            and candidate.source in {"bill", "llm_query"}
            and normalize_bill_product_text(candidate.zh) == product_key
        ]
        detail = "；".join(reasons[-2:])
        parts.append(f"{product}（{detail}）" if detail else product)
    return ", ".join(parts)


async def qualify_bill_product_query_candidate(
    crawler: StrictTaxCrawler,
    bill: BillInfo,
    product: str,
    query_name: str,
    material: str,
    rules: SelectionRules,
    options: ProcessingOptions,
    *,
    query_cache: Optional[QueryCache] = None,
    match_source: str,
    llm_reason: str = "",
) -> ProductCandidate:
    candidate = build_bill_product_query_candidate(bill, product, query_name, material, options, llm_reason)
    product_reason = product_rule_reason(candidate, rules)
    if product_reason:
        return replace(candidate, filter_reason=product_reason)
    errors: list[str] = []
    for attempt in range(1, BILL_PRODUCT_QUERY_RETRIES + 1):
        if attempt > 1:
            await asyncio.sleep(max(1.0, crawler.settings.delay))
        try:
            product_results = await cached_search_product(crawler, query_name or product, material, query_cache)
            selected = select_qualified_tax_data(
                product_results,
                rules,
                enforce_tax_limit=False,
                ignore_certifications=True,
                semantic_name=query_name or product,
                semantic_material=material,
                semantic_usage=candidate.usage,
            )
            if selected:
                return attach_tax_data(candidate, selected, match_source)
            detail = summarize_tax_candidate_rejections(
                product_results,
                rules,
                enforce_tax_limit=False,
                ignore_certifications=True,
            )
            return replace(
                candidate,
                filter_reason=(
                    f"提单品类查询无合格税率/认证结果: {product} -> {query_name}"
                    + (f"；{detail}" if detail else "")
                ),
            )
        except Exception as exc:
            errors.append(str(exc))
    return replace(
        candidate,
        filter_reason=(
            f"提单品类查询失败: {product} -> {query_name}"
            f"；已重试 {BILL_PRODUCT_QUERY_RETRIES} 次: {'；'.join(errors[-2:])}"
        ),
    )


def build_bill_product_query_candidate(
    bill: BillInfo,
    product: str,
    query_name: str,
    material: str,
    options: ProcessingOptions,
    llm_reason: str = "",
) -> ProductCandidate:
    row_count = max(1, options.target_item_count)
    ctns = max(1.0, (bill.cartons or row_count) / row_count)
    gross_weight = max(1.0, row_count)
    qty = max(1.0, ctns)
    name = clean_text(product)
    query = clean_text(query_name) or name
    return ProductCandidate(
        source="bill" if query == name else "llm_query",
        source_label=bill.filename,
        zh=name,
        en=name.title() if name.isupper() else name,
        hs="",
        material=material,
        usage="HOME",
        ctns=ctns,
        qty=qty,
        unit_price=0.0,
        declared_value=0.0,
        real_weight=gross_weight,
        gross_weight=gross_weight,
        query_name=query,
        query_material=material,
        llm_reason=llm_reason,
    )


def bill_product_entry_for(bill: BillInfo, product: str) -> Optional[BillProduct]:
    product_key = normalize_bill_product_text(product)
    for entry in bill.product_entries:
        if normalize_bill_product_text(entry.name) == product_key:
            return entry
    return None


def build_bill_product_candidate(
    bill: BillInfo,
    entry: BillProduct,
    options: ProcessingOptions,
) -> ProductCandidate:
    row_count = max(1, options.target_item_count)
    ctns = max(1.0, (bill.cartons or row_count) / row_count)
    gross_weight = max(1.0, (bill.gross_weight or row_count) / row_count)
    qty = max(1.0, ctns)
    return ProductCandidate(
        source="bill",
        source_label=bill.filename,
        zh=entry.name,
        en=entry.name.title() if entry.name.isupper() else entry.name,
        hs=entry.hs_code_hint,
        material=infer_bill_material_from_entry(entry),
        usage="HOME",
        ctns=ctns,
        qty=qty,
        unit_price=0.0,
        declared_value=0.0,
        real_weight=max(0.01, gross_weight - ctns),
        gross_weight=gross_weight,
    )


async def qualify_bill_hs_tax_candidate(
    crawler: StrictTaxCrawler,
    candidate: ProductCandidate,
    rules: SelectionRules,
    query_cache: Optional[QueryCache] = None,
    llm: Optional[LLMClient] = None,
) -> ProductCandidate:
    product_reason = product_rule_reason(candidate, rules)
    if product_reason:
        return replace(candidate, filter_reason=product_reason)
    if not candidate.hs:
        return replace(candidate, filter_reason=f"提单品类缺少 HS hint: {candidate.zh or candidate.en}")

    try:
        hs_results = await cached_search(crawler, candidate.hs, query_cache)
        selected = select_qualified_tax_data(hs_results, rules, required_hs=candidate.hs, enforce_tax_limit=False)
        if selected:
            material = candidate.material
            if should_infer_material_from_hs(candidate.material):
                material = await infer_material_from_hs_description(
                    crawler,
                    candidate.hs,
                    selected,
                    llm=llm,
                    query_cache=query_cache,
                )
            return attach_bill_hs_tax_data(candidate, selected, material)
        return replace(candidate, filter_reason="提单 HS hint 查询无合格税率/认证结果")
    except Exception as exc:
        return replace(candidate, filter_reason=f"提单 HS hint 查询失败: {exc}")


def select_initial_candidates(
    qualified_manifest: list[ProductCandidate],
    bill_required: list[ProductCandidate],
    rules: SelectionRules,
    target_item_count: int,
) -> list[ProductCandidate]:
    selected: list[ProductCandidate] = []
    selected_keys: set[tuple[str, str, str]] = set()
    selected_semantic_keys: set[str] = set()
    for candidate in bill_required:
        key = candidate_identity(candidate)
        semantic_keys = candidate_semantic_keys(candidate)
        if key not in selected_keys and not (semantic_keys & selected_semantic_keys):
            selected.append(candidate)
            selected_keys.add(key)
            selected_semantic_keys.update(semantic_keys)
    for candidate in sorted(qualified_manifest, key=lambda item: selection_score(item, rules), reverse=True):
        if len(selected) >= target_item_count:
            break
        key = candidate_identity(candidate)
        semantic_keys = candidate_semantic_keys(candidate)
        if key in selected_keys or semantic_keys & selected_semantic_keys:
            continue
        selected.append(candidate)
        selected_keys.add(key)
        selected_semantic_keys.update(semantic_keys)
    return selected


def select_price_repair_pool(
    qualified_manifest: list[ProductCandidate],
    selected: list[ProductCandidate],
    rules: SelectionRules,
    *,
    limit: int = PRICE_REPAIR_POOL_SIZE,
) -> list[ProductCandidate]:
    selected_keys = {candidate_identity(candidate) for candidate in selected}
    pool = [
        candidate
        for candidate in qualified_manifest
        if candidate_identity(candidate) not in selected_keys
    ]
    return sorted(pool, key=lambda item: selection_score(item, rules), reverse=True)[:limit]


def prepare_price_repair_candidates(
    candidates: list[ProductCandidate],
    known_ranges: dict[tuple[str, str, str], PlausibilityRange],
    *,
    query_cache: Optional[QueryCache] = None,
) -> list[ProductCandidate]:
    prepared: list[ProductCandidate] = []
    for candidate in candidates:
        with_range = attach_deterministic_plausibility_range(candidate, known_ranges)
        if not with_range or not plausibility_range_is_complete(with_range.plausibility_range):
            continue
        prepared.append(with_range)
    if not prepared:
        return []
    return prepared


def selected_price_fit_needs_repair(
    selected: list[ProductCandidate],
    manifest: ManifestSummary,
    bill: BillInfo,
    options: ProcessingOptions,
) -> bool:
    try:
        plans = build_plausible_row_plans(
            selected=selected,
            manifest=manifest,
            bill=bill,
            options=options,
            plausibility_ranges={},
        )
    except RuntimeError:
        return True
    return any(plan_price_fit_is_bad(plan) for plan in plans)


def plan_price_fit_is_bad(plan: RowPlan) -> bool:
    if is_undetailed_bill_candidate(plan.candidate):
        return False
    ratio = row_plan_price_fit_ratio(plan)
    if ratio is not None and ratio < PRICE_FIT_MIN_REFERENCE_RATIO:
        return True
    min_price = plan.plausibility.unit_price_min
    return bool(min_price and min_price > 0 and plan.unit_price < min_price - 0.0001)


def assert_selected_price_fit_resolved(
    selected: list[ProductCandidate],
    manifest: ManifestSummary,
    bill: BillInfo,
    options: ProcessingOptions,
    repair_summary: Optional[dict[str, Any]] = None,
) -> None:
    try:
        plans = build_plausible_row_plans(
            selected=selected,
            manifest=manifest,
            bill=bill,
            options=options,
            plausibility_ranges={},
        )
    except RuntimeError as exc:
        raise RuntimeError(f"价格可行性校验失败：无法生成合理重量/价格草案；{exc}") from exc

    unresolved = [plan for plan in plans if plan_price_fit_is_bad(plan)]
    if not unresolved:
        return

    floor_tax = estimate_price_floor_tax(plans)
    tax_upper = target_tax_upper_bound(options.target_tax_amount)
    details = "; ".join(price_fit_unresolved_detail(plan) for plan in unresolved[:5])
    suffix = ""
    if len(unresolved) > 5:
        suffix = f"; 另有 {len(unresolved) - 5} 行未列出"
    swaps = repair_summary.get("swaps") if repair_summary else 0
    attempts = repair_summary.get("replacement_attempts") if repair_summary else None
    attempt_text = f"，替换表已查询 {attempts} 个候选" if attempts is not None else ""
    raise RuntimeError(
        "优化无解：按公开电商零售价 * 30% 的申报参考范围，仍有 "
        f"{len(unresolved)} 行必须压低单价才能满足目标税金。"
        f"若按价格下限估算，最低税金约 {floor_tax} USD，目标上限 {tax_upper} USD；"
        f"已替换 {swaps} 行{attempt_text}。冲突行：{details}{suffix}"
    )


def price_fit_unresolved_detail(plan: RowPlan) -> str:
    floor = row_plan_price_floor(plan) or 0.0
    ratio = row_plan_price_fit_ratio(plan)
    ratio_text = f", ratio={round(ratio, 3)}" if ratio is not None else ""
    return (
        f"{plan.candidate.zh or plan.candidate.en}"
        f" 单价 {round(plan.unit_price, 4)} < 下限 {round(floor, 4)}"
        f"{ratio_text}"
    )


def estimate_price_floor_tax(plans: list[RowPlan]) -> float:
    total = 0.0
    for plan in plans:
        rate = candidate_tax_rate(plan.candidate)
        if rate <= 0:
            continue
        floor = row_plan_price_floor(plan)
        unit_price = max(plan.unit_price, floor or 0.0)
        total += unit_price * plan.qty * rate
    return round(total, 2)


def row_plan_price_floor(plan: RowPlan) -> Optional[float]:
    if is_undetailed_bill_candidate(plan.candidate):
        return None
    floors: list[float] = []
    min_price = plan.plausibility.unit_price_min
    if min_price and min_price > 0:
        floors.append(min_price)
    if plan.price_reference and plan.price_reference > 0:
        floors.append(plan.price_reference * PRICE_FIT_MIN_REFERENCE_RATIO)
    if not floors:
        return None
    return max(floors)


async def qualify_price_repair_replacements(
    crawler: StrictTaxCrawler,
    replacement_pool: list[ProductCandidate],
    selected: list[ProductCandidate],
    rules: SelectionRules,
    *,
    query_cache: Optional[QueryCache] = None,
    progress_callback: Optional[ProgressCallback] = None,
) -> tuple[list[ProductCandidate], list[ProductCandidate], int]:
    selected_keys = {candidate_identity(candidate) for candidate in selected}
    seeds = [
        candidate
        for candidate in sorted(replacement_pool, key=replacement_price_repair_seed_order)
        if candidate_identity(candidate) not in selected_keys
    ]
    qualified: list[ProductCandidate] = []
    filtered: list[ProductCandidate] = []
    attempts = 0
    for seed in seeds:
        if attempts >= PRICE_REPAIR_REPLACEMENT_QUERY_LIMIT:
            break
        if len(qualified) >= PRICE_REPAIR_REPLACEMENT_POOL_SIZE:
            break
        if attempts > 0:
            await asyncio.sleep(crawler.settings.delay)
        attempts += 1
        result, rejected = await qualify_candidates(
            crawler,
            [seed],
            rules,
            query_cache=query_cache,
        )
        filtered.extend(rejected)
        for candidate in result:
            key = candidate_identity(candidate)
            if key in selected_keys:
                continue
            selected_keys.add(key)
            qualified.append(candidate)
        await emit_progress(
            progress_callback,
            {
                "stage": "replacement_price_repair",
                "status": "running",
                "progress": round(91 + min(1.0, attempts / PRICE_REPAIR_REPLACEMENT_QUERY_LIMIT), 2),
                "message": f"已查询替换表价格修复候选 {attempts}/{PRICE_REPAIR_REPLACEMENT_QUERY_LIMIT}",
                "current": attempts,
                "qualified": len(qualified),
                "filtered": len(filtered),
            },
        )
    return qualified, filtered, attempts


def replacement_price_repair_seed_order(candidate: ProductCandidate) -> tuple[float, float, float]:
    unit_price = candidate.unit_price if candidate.unit_price and candidate.unit_price > 0 else DEFAULT_UNIT_PRICE_MAX
    weight = max(candidate.gross_weight or 0.0, candidate.real_weight or 0.0, 0.1)
    qty = candidate.qty if candidate.qty and candidate.qty > 0 else 1.0
    value_density = unit_price / max(weight / qty, DEFAULT_KG_PER_PC_MIN)
    return (value_density, unit_price, -weight)


async def qualify_manual_invoice_replacements(
    crawler: StrictTaxCrawler,
    replacement_pool: list[ProductCandidate],
    selected: list[ProductCandidate],
    rules: SelectionRules,
    *,
    query_cache: Optional[QueryCache] = None,
    progress_callback: Optional[ProgressCallback] = None,
    bill_products: Optional[list[str]] = None,
) -> tuple[list[ProductCandidate], list[ProductCandidate], int]:
    selected_keys = {candidate_identity(candidate) for candidate in selected}
    eligible_replacement_pool = exclude_bill_product_replacements(replacement_pool, bill_products or [])
    stable_reference_candidates = load_stable_manual_invoice_candidates(selected, bill_products or [])
    for candidate in stable_reference_candidates:
        selected_keys.add(candidate_identity(candidate))
    seeds = [
        candidate
        for candidate in sorted(eligible_replacement_pool, key=manual_invoice_replacement_seed_order)
        if candidate_identity(candidate) not in selected_keys
    ]
    qualified: list[ProductCandidate] = list(stable_reference_candidates)
    filtered: list[ProductCandidate] = []
    attempts = 0
    for seed in seeds:
        if is_customer_codebook_candidate(seed):
            key = candidate_identity(seed)
            if key not in selected_keys:
                selected_keys.add(key)
                qualified.append(seed)
            if len(qualified) >= MANUAL_INVOICE_REPLACEMENT_POOL_SIZE:
                break
            continue
        if attempts >= MANUAL_INVOICE_REPLACEMENT_QUERY_LIMIT:
            break
        if len(qualified) >= MANUAL_INVOICE_REPLACEMENT_POOL_SIZE:
            break
        if attempts > 0:
            await asyncio.sleep(crawler.settings.delay)
        attempts += 1
        result, rejected = await qualify_candidates(
            crawler,
            [seed],
            rules,
            query_cache=query_cache,
            enforce_tax_limit=False,
        )
        filtered.extend(rejected)
        for candidate in result:
            key = candidate_identity(candidate)
            if key in selected_keys:
                continue
            selected_keys.add(key)
            qualified.append(candidate)
        await emit_progress(
            progress_callback,
            {
                "stage": "manual_invoice_replacement_pool",
                "status": "running",
                "progress": round(91 + min(1.0, attempts / MANUAL_INVOICE_REPLACEMENT_QUERY_LIMIT), 2),
                "message": f"已查询人工发票替换候选 {attempts}/{MANUAL_INVOICE_REPLACEMENT_QUERY_LIMIT}",
                "current": attempts,
                "qualified": len(qualified),
                "filtered": len(filtered),
            },
        )
    return qualified, filtered, attempts


def load_stable_manual_invoice_candidates(
    selected: list[ProductCandidate],
    bill_products: list[str],
) -> list[ProductCandidate]:
    selected_keys = {candidate_identity(candidate) for candidate in selected}
    return [
        candidate
        for candidate in exclude_bill_product_replacements(
            load_default_reference_manual_candidates(), bill_products
        )
        if candidate_identity(candidate) not in selected_keys
    ]


def build_manual_invoice_replacement_seeds(
    replacement_pool: list[ProductCandidate],
    selected: list[ProductCandidate],
    bill_products: list[str],
) -> list[ProductCandidate]:
    selected_keys = {candidate_identity(candidate) for candidate in selected}
    return [
        candidate
        for candidate in sorted(
            exclude_bill_product_replacements(replacement_pool, bill_products),
            key=manual_invoice_replacement_seed_order,
        )
        if candidate_identity(candidate) not in selected_keys
    ]


async def qualify_manual_invoice_replacement_batch(
    crawler: StrictTaxCrawler,
    seeds: list[ProductCandidate],
    start_index: int,
    selected: list[ProductCandidate],
    rules: SelectionRules,
    *,
    prior_network_attempts: int = 0,
    batch_size: int = ADAPTIVE_REPLACEMENT_BATCH_SIZE,
    query_cache: Optional[QueryCache] = None,
    progress_callback: Optional[ProgressCallback] = None,
) -> tuple[list[ProductCandidate], list[ProductCandidate], int, int]:
    selected_keys = {candidate_identity(candidate) for candidate in selected}
    qualified: list[ProductCandidate] = []
    filtered: list[ProductCandidate] = []
    network_attempts = 0
    cursor = max(0, start_index)
    while cursor < len(seeds) and len(qualified) < batch_size:
        seed = seeds[cursor]
        if candidate_identity(seed) in selected_keys:
            cursor += 1
            continue
        if is_customer_codebook_candidate(seed):
            qualified.append(seed)
            selected_keys.add(candidate_identity(seed))
            cursor += 1
            continue
        if network_attempts >= batch_size:
            break
        if prior_network_attempts + network_attempts > 0:
            await asyncio.sleep(crawler.settings.delay)
        network_attempts += 1
        cursor += 1
        result, rejected = await qualify_candidates(
            crawler,
            [seed],
            rules,
            query_cache=query_cache,
            enforce_tax_limit=False,
        )
        filtered.extend(rejected)
        for candidate in result:
            key = candidate_identity(candidate)
            if key in selected_keys:
                continue
            selected_keys.add(key)
            qualified.append(candidate)
        await emit_progress(
            progress_callback,
            {
                "stage": "manual_invoice_replacement_pool",
                "status": "running",
                "progress": round(
                    91 + min(1.0, (prior_network_attempts + network_attempts) / MANUAL_INVOICE_REPLACEMENT_QUERY_LIMIT),
                    2,
                ),
                "message": (
                    "首轮无解，已按需查询替换候选 "
                    f"{prior_network_attempts + network_attempts}/{MANUAL_INVOICE_REPLACEMENT_QUERY_LIMIT}"
                ),
                "current": prior_network_attempts + network_attempts,
                "qualified": len(qualified),
                "filtered": len(filtered),
            },
        )
    return qualified, filtered, cursor, network_attempts


def manual_invoice_solution_is_acceptable(
    summary: dict[str, Any],
    options: ProcessingOptions,
) -> bool:
    estimated_tax = to_float(summary.get("estimated_tax"))
    score = to_float(summary.get("score"))
    if estimated_tax is None or score is None:
        return False
    if estimated_tax > target_tax_upper_bound(options.target_tax_amount) + 0.01:
        return False
    return abs(estimated_tax - options.target_tax_amount) <= TAX_FINAL_TOLERANCE_USD


def manual_invoice_replacement_seed_order(candidate: ProductCandidate) -> tuple[float, float, float]:
    unit_price = candidate.unit_price if candidate.unit_price and candidate.unit_price > 0 else DEFAULT_UNIT_PRICE_MAX
    weight = max(candidate.gross_weight or 0.0, candidate.real_weight or 0.0, 0.1)
    ctns = max(candidate.ctns or 0.0, 1.0)
    kg_per_ctn = weight / ctns
    anchor_rank = 0 if normalize_hs(candidate.hs) in MANUAL_INVOICE_ANCHOR_HS else 1
    default_rank = 0 if candidate.source_label == "DEFAULT_REFERENCE_STYLE_ROWS" else 1
    return (anchor_rank, default_rank, unit_price, -kg_per_ctn, -weight)


def attach_deterministic_plausibility_range(
    candidate: ProductCandidate,
    known_ranges: dict[tuple[str, str, str], PlausibilityRange],
) -> Optional[ProductCandidate]:
    known = lookup_plausibility_range(candidate, known_ranges)
    if known and plausibility_range_is_complete(known):
        return replace(candidate, plausibility_range=with_default_plausibility_bounds(known))
    derived = derive_candidate_plausibility_range(candidate)
    if plausibility_range_is_complete(derived):
        return replace(candidate, plausibility_range=with_default_plausibility_bounds(derived))
    return None


def candidate_matches_single_bill_product(candidate: ProductCandidate, product: str) -> bool:
    return row_matches_single_bill_product({"中文品名": candidate.zh, "英文品名": candidate.en}, product)


def exclude_bill_product_replacements(
    replacement_candidates: list[ProductCandidate],
    bill_products: list[str],
) -> list[ProductCandidate]:
    products = unique_bill_products(bill_products)
    if not products:
        return list(replacement_candidates)
    return [
        candidate
        for candidate in replacement_candidates
        if not any(candidate_matches_single_bill_product(candidate, product) for product in products)
    ]


def candidate_has_product_tax_match(candidate: ProductCandidate) -> bool:
    return candidate.tax_match_source in {
        "product",
        "manifest_group_hs",
        "bill_hs",
        "bill_product",
        "bill_product_no_material",
        "llm_query",
    } and bool(normalize_hs(candidate.hs))


def infer_bill_material_from_entry(entry: BillProduct) -> str:
    text = normalize_text(entry.name)
    if "polyester" in text or "涤纶" in text or "化纤" in text or "合纤" in text:
        return "Polyester"
    if "cotton" in text or "棉" in text:
        return "Cotton"
    if "wool" in text or "羊毛" in text:
        return "Wool"
    if "plastic" in text or "塑料" in text or "塑胶" in text or "pvc" in text or "polypropylene" in text or "polyethylene" in text:
        return "Plastic"
    if "iron" in text or "steel" in text or "metal" in text or "metallic" in text or "铁" in text or "金属" in text:
        return "Metal"
    if "wood" in text or "木" in text:
        return "Wood"
    if "glass" in text or "玻璃" in text:
        return "Glass"
    if "paper" in text or "纸" in text:
        return "Paper"
    if "fabric" in text or "cloth" in text or "布" in text or "纺" in text:
        return "Fabric"
    if "rubber" in text or "橡胶" in text:
        return "Rubber"
    if "ceramic" in text or "陶瓷" in text:
        return "Ceramic"
    if "silicone" in text or "硅胶" in text:
        return "Silicone"
    return ""


def material_family(value: Any) -> str:
    text = normalize_text(value)
    families = (
        ("polyester", ("polyester", "涤纶")),
        ("cotton", ("cotton", "棉")),
        ("wool", ("wool", "羊毛")),
        ("plastic", ("plastic", "塑料", "塑胶", "pvc", "polyethylene", "polypropylene", "聚乙烯", "聚丙烯")),
        ("metal", ("iron", "steel", "metal", "铁", "钢", "金属")),
        ("wood", ("wood", "木")),
        ("glass", ("glass", "玻璃")),
        ("paper", ("paper", "纸")),
        ("rubber", ("rubber", "橡胶")),
        ("ceramic", ("ceramic", "陶瓷")),
        ("silicone", ("silicone", "硅胶")),
        ("fabric", ("fabric", "cloth", "textile", "布", "纺织")),
    )
    for family, terms in families:
        if any(term in text for term in terms):
            return family
    return ""


def candidate_material_matches_bill(candidate: ProductCandidate, bill_material: str) -> bool:
    expected = material_family(bill_material)
    actual = material_family(candidate.material)
    return not expected or not actual or expected == actual


def product_family(value: Any) -> str:
    text = normalize_text(value)
    families = (
        ("pillowcase", ("pillowcase", "pillow case", "枕套")),
        ("curtain", ("curtain", "drape", "窗帘", "帘")),
        ("scarf", ("scarf", "shawl", "headscarf", "围巾", "披巾", "头巾")),
        ("keychain", ("keychain", "key chain", "keyring", "key ring", "钥匙扣", "钥匙圈")),
        ("wall_hanging", ("wall hanging", "wall tapestry", "挂画", "壁挂")),
        ("bag", ("bag", "bags", "pouch", "sack", "包", "袋", "箱包", "衣箱", "容器")),
        ("jewelry", ("jewelry", "jewellery", "首饰", "饰品")),
        ("bracket", ("bracket", "holder", "stand", "support", "支架", "托架")),
        ("top", ("women s top", "womens top", "blouse", "shirt", "上衣", "女衬衫", "衬衫")),
        ("cup", ("cup", "tumbler", "杯")),
        ("christmas", ("christmas", "圣诞")),
        ("ornament", ("ornament", "statuette", "sculpture", "装饰品", "雕塑")),
    )
    for family, terms in families:
        if any(term in text for term in terms):
            return family
    return ""


def tax_candidate_semantic_mismatch_reason(
    data: dict[str, Any],
    product_name: str,
    material: str,
) -> str:
    description = clean_text(data.get("description_cn"))
    expected_product = product_family(product_name)
    actual_product = product_family(description)
    if expected_product and actual_product and expected_product != actual_product:
        return f"Codeflag 品类 {actual_product} 与申报品类 {expected_product} 不一致"
    if material_semantic_score(description, material) <= -100:
        return f"Codeflag 描述材质与申报材质 {material_family(material) or clean_text(material)} 不一致"
    return ""


def material_semantic_score(description: str, material: str) -> int:
    expected = material_family(material)
    if not expected:
        return 0
    text = normalize_text(description)
    if expected == "polyester":
        if any(term in text for term in ("cotton", "棉", "wool", "羊毛", "polyethylene", "polypropylene", "聚乙烯", "聚丙烯")):
            return -100
        if "polyester" in text or "涤纶" in text:
            return 6
        if any(
            term in text
            for term in (
                "synthetic fiber",
                "chemical fiber",
                "man made textile",
                "化纤",
                "化学纤维",
                "合纤",
                "人造纤维",
                "人造纺织",
            )
        ):
            return 5
        if any(term in text for term in ("textile", "fabric", "纺织", "布")):
            return 2
        actual = material_family(description)
        return -100 if actual and actual != "fabric" else 0
    actual = material_family(description)
    if actual == expected:
        return 4
    if actual:
        return -100
    return 0


def tax_candidate_semantic_score(
    data: dict[str, Any],
    product_name: str,
    material: str,
    usage: str,
) -> int:
    description = clean_text(data.get("description_cn"))
    expected_product = product_family(product_name)
    actual_product = product_family(description)
    score = 0
    if expected_product:
        score += 4 if expected_product == actual_product else -2
    score += material_semantic_score(description, material)
    usage_text = normalize_text(usage)
    description_text = normalize_text(description)
    if any(term in usage_text for term in ("storage", "packaging", "收纳", "储存", "包装")):
        if any(term in description_text for term in ("包装", "储运", "容器", "袋", "包")):
            score += 1
    if any(term in description_text for term in ("bulk", "散货", "散装")) and not any(
        term in normalize_text(product_name) for term in ("bulk", "散货", "散装")
    ):
        score -= 3
    return score


def should_infer_material_from_hs(material: str) -> bool:
    normalized = normalize_text(material)
    return not normalized or normalized in {"general", "mixed"}


async def infer_material_from_hs_description(
    crawler: StrictTaxCrawler,
    hs_hint: str,
    selected: dict[str, Any],
    llm: Optional[LLMClient] = None,
    query_cache: Optional[QueryCache] = None,
) -> str:
    description = clean_text(selected.get("description_cn"))
    if not description:
        return ""
    prompt = (
        "你是海关商品材质识别专家。请根据 HS 描述判断最合适的英文材质，"
        "只输出一个词或短语，如 Plastic, Metal, Wood, Glass, Paper, Fabric, Rubber, Ceramic, Silicone, Acrylic, Polyester, Nylon, Iron. "
        "如果无法判断，输出 Mixed。\n"
        f"HS描述: {description}"
    )
    try:
        llm_client = llm or LLMClient()
        payload = await llm_client.chat_json(
            [
                {"role": "system", "content": "只输出 JSON object。"},
                {"role": "user", "content": f"{prompt}\nJSON格式：{{\"material\":\"\"}}"},
            ],
            temperature=0.0,
        )
        material = clean_text(payload.get("material"))
        return material or ""
    except Exception:
        return ""


def load_selection_rules(rules_dir: Path = RULES_ROOT) -> SelectionRules:
    return SelectionRules(
        allowed_certifications=read_rule_lines(rules_dir / "allowed_certifications.txt") or ["Lacey Act", "TSCA"],
        blocked_certifications=read_rule_lines(rules_dir / "blocked_certifications.txt"),
        allowed_products=read_rule_lines(rules_dir / "allowed_products.txt"),
        blocked_products=read_rule_lines(rules_dir / "blocked_products.txt"),
    )


def read_rule_lines(path: Path) -> list[str]:
    if not path.exists():
        return []
    lines: list[str] = []
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        lines.append(line)
    return lines


def build_manifest_candidates(manifest: ManifestSummary) -> list[ProductCandidate]:
    groups = build_manifest_hs_groups(manifest)
    candidates: list[ProductCandidate] = []
    for group in groups:
        unit_price = median_or_first(group.unit_price_values)
        qty = group.total_qty or group.total_ctns or 1
        gross_weight = group.total_gross_weight or group.total_real_weight or 1
        candidates.append(
            ProductCandidate(
                source="manifest_group",
                source_label=f"{manifest.filename}/{'HS归并' if group.hs else '品名归并'}",
                zh=group.canonical_zh,
                en=group.canonical_en or group.canonical_zh,
                hs=group.hs,
                material=group.representative_material,
                usage=group.representative_usage or "HOME",
                ctns=group.total_ctns,
                qty=qty,
                unit_price=unit_price,
                declared_value=group.total_declared_value,
                real_weight=group.total_real_weight,
                gross_weight=gross_weight,
                source_rows=group.source_rows,
                original_hs=group.hs,
                llm_reason="; ".join(group.warnings),
            )
        )
    return sorted(candidates, key=lambda item: item.score, reverse=True)


def ensure_manifest_codeflag_candidates(
    manifest: ManifestSummary,
    candidates: list[ProductCandidate],
) -> None:
    if manifest.row_count > 0 and not candidates:
        raise RuntimeError(
            "manifest_schema_unresolved: "
            "清单识别到商品行，但没有生成 Codeflag 查询候选，请检查字段映射和商品解析结果"
        )


def build_manifest_hs_groups(manifest: ManifestSummary) -> list[ManifestHsGroup]:
    groups: dict[tuple[str, str, str, str], ManifestHsGroup] = {}
    for item in manifest.items:
        hs = normalize_hs(item.hs)
        if len(hs) != 10:
            hs = ""
        if not hs and not clean_text(item.zh or item.en):
            continue
        key = manifest_group_identity(item, hs)
        group = groups.setdefault(key, ManifestHsGroup(hs=hs))
        group.source_rows.append(item.row)
        append_unique(group.zh_names, item.zh)
        append_unique(group.en_names, item.en)
        append_unique(group.materials, item.material)
        append_unique(group.usages, item.usage)
        group.total_ctns += item.ctns or 0
        group.total_qty += item.qty or 0
        group.total_real_weight += item.real_weight or 0
        group.total_gross_weight += item.gross_weight or item.real_weight or 0
        group.total_declared_value += item.declared_value or 0
        if item.unit_price and item.unit_price > 0:
            group.unit_price_values.append(item.unit_price)

    result: list[ManifestHsGroup] = []
    for group in groups.values():
        group.canonical_zh = choose_representative_text(group.zh_names)
        group.canonical_en = choose_representative_text(group.en_names) or group.canonical_zh
        group.representative_material = choose_representative_text(group.materials) or "Mixed"
        group.representative_usage = choose_representative_text(group.usages) or "HOME"
        if len({normalize_text(name) for name in group.zh_names if normalize_text(name)}) > 6:
            group.warnings.append("同一 HS 下原始品名较多，代表品名仅用于搜索和申报候选")
        if not group.total_qty:
            group.warnings.append("该 HS group 缺少有效数量")
        if not group.total_gross_weight and not group.total_real_weight:
            group.warnings.append("该 HS group 缺少有效重量")
        if not group.hs:
            group.warnings.append("原清单未提供有效 HS，使用品名和材质查询 Codeflag")
        round_manifest_group_totals(group)
        result.append(group)
    return sorted(result, key=lambda item: item.total_gross_weight + item.total_declared_value / 100, reverse=True)


def manifest_group_identity(item: ManifestItem, hs: str) -> tuple[str, str, str, str]:
    return (
        hs,
        normalize_text(item.zh),
        normalize_text(item.en),
        normalize_text(item.material),
    )


def append_unique(values: list[str], value: Any) -> None:
    text = clean_text(value)
    if not text:
        return
    key = normalize_text(text)
    if key and key not in {normalize_text(item) for item in values}:
        values.append(text)


def choose_representative_text(values: list[str]) -> str:
    cleaned = [clean_text(value) for value in values if clean_text(value)]
    if not cleaned:
        return ""
    return sorted(cleaned, key=lambda item: (len(item), -sum(ch.isascii() and ch.isalpha() for ch in item)))[0]


def median_or_first(values: list[float]) -> float:
    cleaned = sorted(value for value in values if value and value > 0)
    if not cleaned:
        return 0.0
    middle = len(cleaned) // 2
    if len(cleaned) % 2:
        return round(cleaned[middle], 4)
    return round((cleaned[middle - 1] + cleaned[middle]) / 2, 4)


def round_manifest_group_totals(group: ManifestHsGroup) -> None:
    group.total_ctns = round(group.total_ctns, 2)
    group.total_qty = round(group.total_qty, 2)
    group.total_real_weight = round(group.total_real_weight, 2)
    group.total_gross_weight = round(group.total_gross_weight, 2)
    group.total_declared_value = round(group.total_declared_value, 2)


async def qualify_candidates(
    crawler: StrictTaxCrawler,
    candidates: list[ProductCandidate],
    rules: SelectionRules,
    query_cache: Optional[QueryCache] = None,
    enforce_tax_limit: bool = True,
    progress_callback: Optional[ProgressCallback] = None,
    progress_stage: str = "",
    progress_start: float = 0,
    progress_end: float = 100,
) -> tuple[list[ProductCandidate], list[ProductCandidate]]:
    qualified: list[ProductCandidate] = []
    filtered: list[ProductCandidate] = []
    for index, candidate in enumerate(candidates):
        if index > 0:
            await asyncio.sleep(crawler.settings.delay)
        result = await qualify_single_candidate(
            crawler,
            candidate,
            rules,
            query_cache=query_cache,
            enforce_tax_limit=enforce_tax_limit,
        )
        if result.filter_reason:
            filtered.append(result)
        else:
            qualified.append(result)
        if progress_callback and progress_stage:
            progress = progress_start
            if candidates:
                progress = progress_start + (progress_end - progress_start) * ((index + 1) / len(candidates))
            await emit_progress(
                progress_callback,
                {
                    "stage": progress_stage,
                    "status": "running",
                    "progress": round(progress, 2),
                    "message": f"已查询 {index + 1}/{len(candidates)} 个候选",
                    "current": index + 1,
                    "total": len(candidates),
                    "qualified": len(qualified),
                    "filtered": len(filtered),
                },
            )
    return qualified, filtered


async def qualify_single_candidate(
    crawler: StrictTaxCrawler,
    candidate: ProductCandidate,
    rules: SelectionRules,
    query_cache: Optional[QueryCache] = None,
    enforce_tax_limit: bool = True,
) -> ProductCandidate:
    product_reason = product_rule_reason(candidate, rules)
    if product_reason:
        return replace(candidate, filter_reason=product_reason)
    if is_customer_codebook_candidate(candidate):
        if candidate_tax_rate(candidate) >= CUSTOMER_CODEBOOK_TAX_LIMIT:
            return replace(
                candidate,
                filter_reason=f"客户编码库总税率 {format_rate(candidate_tax_rate(candidate))} 不小于 30%",
            )
        return replace(candidate, filter_reason="")

    errors: list[str] = []
    if candidate.source in {"replacement", "manifest_group"} and normalize_hs(candidate.hs):
        try:
            hs_results = await cached_search(crawler, candidate.hs, query_cache)
            selected = select_qualified_tax_data(
                hs_results,
                rules,
                required_hs=candidate.hs,
                enforce_tax_limit=enforce_tax_limit,
                prefer_first=True,
                ignore_certifications=True,
            )
            if selected:
                attached = attach_tax_data(candidate, selected, f"{candidate.source}_hs")
                value_limit_reason = tax_description_unit_value_limit_reason(attached, selected)
                if value_limit_reason:
                    return replace(
                        attached,
                        compliance_review_required=True,
                        compliance_review_reason=value_limit_reason,
                        row_warnings=[*attached.row_warnings, value_limit_reason],
                    )
                return attached
            errors.append("原始 HTS 查询无合格结果")
        except Exception as exc:
            errors.append(f"原始 HTS 查询失败: {exc}")

    try:
        product_results = await cached_search_product(crawler, candidate.zh or candidate.en, candidate.material, query_cache)
        required_hs = candidate.hs if candidate.source == "replacement" else ""
        selected = select_qualified_tax_data(
            product_results,
            rules,
            required_hs=required_hs,
            enforce_tax_limit=enforce_tax_limit,
            ignore_certifications=True,
            semantic_name=candidate.zh or candidate.en,
            semantic_material=candidate.material,
            semantic_usage=candidate.usage,
        )
        if selected:
            return attach_tax_data(candidate, selected, "product")
        errors.append(
            "品名+材质查询无合格税率/认证结果"
            + ("，或返回 HS 与替换清单原始 HS 不一致" if candidate.source == "replacement" and candidate.hs else "")
        )
    except Exception as exc:
        errors.append(f"品名+材质查询失败: {exc}")

    if candidate.hs and candidate.source != "replacement":
        try:
            hs_results = await cached_search(crawler, candidate.hs, query_cache)
            selected = select_qualified_tax_data(
                hs_results,
                rules,
                required_hs=candidate.hs,
                enforce_tax_limit=enforce_tax_limit,
                prefer_first=True,
                ignore_certifications=True,
            )
            if selected:
                return attach_tax_data(candidate, selected, "hs")
            errors.append("原始 HTS 兜底查询无合格结果")
        except Exception as exc:
            errors.append(f"原始 HTS 兜底查询失败: {exc}")

    return replace(candidate, filter_reason="；".join(errors) or "无合格税率/认证结果")


async def cached_search_product(
    crawler: StrictTaxCrawler,
    product_name: str,
    material: str,
    query_cache: Optional[QueryCache] = None,
) -> dict[str, dict[str, Any]]:
    if query_cache is None:
        return await crawler.search_product(product_name, material)
    section = query_cache.setdefault("product", {})
    key = product_cache_key(product_name, material)
    if key not in section:
        section[key] = await crawler.search_product(product_name, material)
    return section[key]


async def cached_search(
    crawler: StrictTaxCrawler,
    hs_code: str,
    query_cache: Optional[QueryCache] = None,
) -> dict[str, dict[str, Any]]:
    if query_cache is None:
        return await crawler.search(hs_code)
    section = query_cache.setdefault("hs", {})
    key = normalize_hs(hs_code)
    if key not in section:
        section[key] = await crawler.search(hs_code)
    return section[key]


async def generate_bill_product_query_terms(
    llm: LLMClient,
    bill: BillInfo,
    product: str,
    failures: list[str],
    query_cache: Optional[QueryCache] = None,
) -> list[dict[str, str]]:
    cache_key = hashlib.sha256(
        json.dumps(
            {
                "kind": "bill_query_terms",
                "product": product,
                "hs_hint": (bill_product_entry_for(bill, product).hs_code_hint if bill_product_entry_for(bill, product) else ""),
                "failures": failures[-6:],
            },
            ensure_ascii=False,
            sort_keys=True,
        ).encode("utf-8", errors="ignore")
    ).hexdigest()
    if query_cache is not None:
        section = query_cache.setdefault("llm_query_terms", {})
        if cache_key in section and isinstance(section[cache_key], dict):
            return normalize_llm_query_terms(section[cache_key], product)

    payload = await llm.chat_json(build_bill_query_term_messages(bill, product, failures), temperature=0.1)
    terms = normalize_llm_query_terms(payload, product)
    if query_cache is not None:
        query_cache.setdefault("llm_query_terms", {})[cache_key] = payload
    return terms


def build_bill_query_term_messages(bill: BillInfo, product: str, failures: list[str]) -> list[dict[str, str]]:
    context = {
        "bill_product": product,
        "bill_products": bill.products,
        "hs_hint": (bill_product_entry_for(bill, product).hs_code_hint if bill_product_entry_for(bill, product) else ""),
        "evidence": (bill_product_entry_for(bill, product).evidence if bill_product_entry_for(bill, product) else ""),
        "failed_reasons": failures[-6:],
    }
    return [
        {
            "role": "system",
            "content": (
                "你是美国进口商品归类查询词专家。你的任务是把提单货物品类转换成 codeflagai 更容易查到的"
                "语义等价或高度相近查询词。不要改变成无关品类。只返回 JSON object。"
            ),
        },
        {
            "role": "user",
            "content": (
                "请给出 3-5 个查询词，优先英文，其次中英混合。每个查询词需包含 product_name、material、reason。\n"
                "要求：必须和提单品类语义一致或商业上可申报为同类；不能为了低税率改成无关商品。\n"
                "JSON格式：{\"queries\":[{\"product_name\":\"\",\"material\":\"\",\"reason\":\"\"}]}\n"
                f"上下文：{json.dumps(context, ensure_ascii=False, separators=(',', ':'))}"
            ),
        },
    ]


def normalize_llm_query_terms(payload: dict[str, Any], product: str) -> list[dict[str, str]]:
    raw_terms = payload.get("queries") or payload.get("terms") or payload.get("items")
    if not isinstance(raw_terms, list):
        raw_terms = []
    result: list[dict[str, str]] = []
    seen: set[str] = set()
    for raw in raw_terms:
        if not isinstance(raw, dict):
            name = clean_text(raw)
            material = ""
            reason = ""
        else:
            name = clean_text(raw.get("product_name") or raw.get("name") or raw.get("query"))
            material = clean_text(raw.get("material"))
            reason = clean_text(raw.get("reason"))
        if not name:
            continue
        key = product_cache_key(name, material)
        if key in seen:
            continue
        seen.add(key)
        result.append({"product_name": name, "material": material, "reason": reason})
        if len(result) >= 5:
            break
    direct = {"product_name": product, "material": "", "reason": "提单原始品类"}
    direct_key = product_cache_key(direct["product_name"], direct["material"])
    return ([direct] if direct_key not in seen else []) + result


def product_cache_key(product_name: str, material: str = "") -> str:
    return f"{normalize_text(product_name)}|{normalize_text(material)}"


async def emit_progress(callback: Optional[ProgressCallback], payload: dict[str, Any]) -> None:
    if not callback:
        return
    result = callback(payload)
    if inspect.isawaitable(result):
        await result


def product_rule_reason(candidate: ProductCandidate, rules: SelectionRules) -> str:
    name = normalize_text(f"{candidate.zh} {candidate.en}")
    for blocked in rules.blocked_products:
        if normalize_text(blocked) and normalize_text(blocked) in name:
            return f"命中禁用品名库: {blocked}"
    return ""


def product_is_allowed(candidate: ProductCandidate, rules: SelectionRules) -> bool:
    name = normalize_text(f"{candidate.zh} {candidate.en}")
    return any(normalize_text(item) and normalize_text(item) in name for item in rules.allowed_products)


def candidate_reference_unit_price(candidate: ProductCandidate) -> float:
    evidence_price = to_float(candidate.price_evidence.get("declared_unit_price"))
    if evidence_price and evidence_price > 0:
        return evidence_price
    if candidate.unit_price and candidate.unit_price > 0:
        return candidate.unit_price
    if candidate.declared_value and candidate.qty:
        derived = candidate.declared_value / candidate.qty
        if derived > 0:
            return derived
    return 0.0


def selection_score(candidate: ProductCandidate, rules: SelectionRules) -> tuple[int, int, float, float]:
    rate = candidate_tax_rate(candidate)
    reference_price = candidate_reference_unit_price(candidate)
    price_tax_pressure = max(0.0, reference_price) * max(0.0, rate)
    penalty = 1.0 + rate * 8.0 + price_tax_pressure * 6.0
    feasibility_score = candidate.score / penalty
    return (
        1 if product_is_allowed(candidate, rules) else 0,
        0 if candidate.compliance_review_required else 1,
        feasibility_score,
        candidate.score,
    )


def select_qualified_tax_data(
    candidates: dict[str, dict[str, Any]],
    rules: SelectionRules,
    required_hs: str = "",
    enforce_tax_limit: bool = True,
    prefer_first: bool = False,
    ignore_certifications: bool = False,
    semantic_name: str = "",
    semantic_material: str = "",
    semantic_usage: str = "",
) -> Optional[dict[str, Any]]:
    ranked: list[tuple[Any, ...]] = []
    for order, data in enumerate(candidates.values()):
        if required_hs and not tax_candidate_matches_required_hs(data, required_hs):
            continue
        reason = tax_filter_reason(
            data,
            rules,
            enforce_tax_limit=enforce_tax_limit,
            ignore_certifications=ignore_certifications,
        )
        if reason:
            continue
        if semantic_name and tax_candidate_semantic_mismatch_reason(
            data,
            semantic_name,
            semantic_material,
        ):
            continue
        anti_dumping_penalty = 1 if data.get("anti_dumping") else 0
        if prefer_first:
            ranked.append((order, 0.0, 0.0, 0.0, 0.0, data))
        elif semantic_name or semantic_material or semantic_usage:
            semantic_score = tax_candidate_semantic_score(
                data,
                semantic_name,
                semantic_material,
                semantic_usage,
            )
            certification_penalty = 1 if (
                data.get("certification_required") or data.get("certification_texts")
            ) else 0
            ranked.append(
                (
                    -semantic_score,
                    certification_penalty,
                    anti_dumping_penalty,
                    effective_tax_rate(data),
                    base_tax_rate(data) or 0,
                    order,
                    data,
                )
            )
        else:
            ranked.append(
                (
                    anti_dumping_penalty,
                    effective_tax_rate(data),
                    base_tax_rate(data) or 0,
                    order,
                    data,
                )
            )
    if not ranked:
        return None
    return sorted(ranked, key=lambda item: item[:-1])[0][-1]


def summarize_tax_candidate_rejections(
    candidates: dict[str, dict[str, Any]],
    rules: SelectionRules,
    required_hs: str = "",
    enforce_tax_limit: bool = True,
    ignore_certifications: bool = False,
    limit: int = 4,
) -> str:
    if not candidates:
        return "Codeflag 未返回候选"
    reasons: list[str] = []
    seen: set[str] = set()
    for data in candidates.values():
        hs = normalize_hs(data.get("hs_code_us")) or clean_text(data.get("hs_code")) or "未知HTS"
        if required_hs and not tax_candidate_matches_required_hs(data, required_hs):
            reason = f"不匹配指定 HTS {normalize_hs(required_hs)}"
        else:
            reason = tax_filter_reason(
                data,
                rules,
                enforce_tax_limit=enforce_tax_limit,
                ignore_certifications=ignore_certifications,
            )
        if not reason:
            continue
        item = f"{hs}: {reason}"
        if item in seen:
            continue
        seen.add(item)
        reasons.append(item)
        if len(reasons) >= limit:
            break
    if not reasons:
        return "Codeflag 返回候选但未选中"
    extra = len(candidates) - len(reasons)
    suffix = f"；另有 {extra} 个候选未展开" if extra > 0 and len(reasons) >= limit else ""
    return "候选被过滤: " + "；".join(reasons) + suffix


def tax_candidate_matches_required_hs(data: dict[str, Any], required_hs: str) -> bool:
    required = normalize_hs(required_hs)
    returned = normalize_hs(data.get("hs_code_us"))
    if not required or not returned:
        return True
    if returned == required:
        return True
    if len(required) < 10 and returned.startswith(required):
        return True
    if len(returned) < 10 and required.startswith(returned):
        return True
    return False


def tax_filter_reason(
    data: dict[str, Any],
    rules: SelectionRules,
    enforce_tax_limit: bool = True,
    ignore_certifications: bool = False,
) -> str:
    hs = normalize_hs(data.get("hs_code_us"))
    if len(hs) != 10:
        return "美国 HTS 不是 10 位"
    base_rate = base_tax_rate(data)
    if base_rate is None:
        return "基础税率为空或无法解析"
    if enforce_tax_limit:
        if base_rate >= BASE_TAX_LIMIT:
            return f"基础税率 {format_rate(base_rate)} 不小于 20%"
        combined_rate = effective_tax_rate(data)
        if combined_rate >= BASE_TAX_LIMIT:
            return f"综合税率 {format_rate(combined_rate)} 不小于 20%"
    if not ignore_certifications:
        cert_reason = certification_filter_reason(data.get("certification_texts") or [], rules)
        if cert_reason:
            return cert_reason
    return ""


def tax_description_unit_value_limit_usd(data: dict[str, Any]) -> Optional[float]:
    text = " ".join(
        clean_text(data.get(key))
        for key in ("description_cn", "taric", "source_description_cn")
        if clean_text(data.get(key))
    ).lower()
    limits: list[float] = []
    patterns = (
        (r"valued\s+not\s+over\s+(\d+(?:\.\d+)?)\s*cents?\s+per\s+(?:piece|item|article)", 0.01),
        (r"valued\s+not\s+over\s+\$?\s*(\d+(?:\.\d+)?)\s*(?:usd|dollars?)?\s+per\s+(?:piece|item|article)", 1.0),
        (r"每件价值不超过\s*(\d+(?:\.\d+)?)\s*美分", 0.01),
        (r"每件(?:价值)?不超过\s*(\d+(?:\.\d+)?)\s*(?:美元|usd)", 1.0),
    )
    for pattern, factor in patterns:
        for match in re.findall(pattern, text, flags=re.IGNORECASE):
            limits.append(float(match) * factor)
    return min(limits) if limits else None


def tax_description_unit_value_limit_reason(
    candidate: ProductCandidate,
    data: dict[str, Any],
) -> str:
    limit = tax_description_unit_value_limit_usd(data)
    if limit is None:
        return ""
    unit_price = to_float(candidate.unit_price)
    if (unit_price is None or unit_price <= 0) and candidate.declared_value and candidate.qty:
        unit_price = candidate.declared_value / candidate.qty
    if unit_price is None or unit_price <= limit + 0.0001:
        return ""
    hs = normalize_hs(data.get("hs_code_us")) or normalize_hs(candidate.hs)
    return (
        f"Codeflag HTS {hs} 限定每件价值不超过 {limit:.4f} USD，"
        f"原始清单单价 {unit_price:.4f} USD；仅在合格候选不足时保留并要求人工复核"
    )


def attach_tax_data(candidate: ProductCandidate, data: dict[str, Any], match_source: str) -> ProductCandidate:
    hs = normalize_hs(data.get("hs_code_us")) or candidate.hs
    return replace(
        candidate,
        hs=hs,
        tax_data=data,
        base_tax_rate=base_tax_rate(data) or 0,
        effective_tax_rate=effective_tax_rate(data),
        tax_match_source=match_source,
        certification_texts=list(data.get("certification_texts") or []),
        filter_reason="",
    )


def attach_bill_hs_tax_data(
    candidate: ProductCandidate,
    data: dict[str, Any],
    material: str,
) -> ProductCandidate:
    returned_hs = normalize_hs(data.get("hs_code_us"))
    output_hs = returned_hs if len(returned_hs) == 10 else normalize_hs(candidate.hs)
    return replace(
        candidate,
        hs=output_hs,
        tax_data=data,
        base_tax_rate=base_tax_rate(data) or 0,
        effective_tax_rate=effective_tax_rate(data),
        tax_match_source="bill_hs",
        certification_texts=list(data.get("certification_texts") or []),
        material=material or candidate.material or "Mixed",
        filter_reason="",
    )


def normalize_generated_candidate(candidate: ProductCandidate) -> ProductCandidate:
    return replace(
        candidate,
        source="generated",
        ctns=1,
        qty=20,
        unit_price=1.0,
        declared_value=20.0,
        real_weight=10.0,
        gross_weight=10.0,
        source_rows=[],
    )


async def ensure_candidate_plausibility_ranges(
    candidates: list[ProductCandidate],
    manifest: ManifestSummary,
    bill: BillInfo,
    known_ranges: dict[tuple[str, str, str], PlausibilityRange],
    llm: LLMClient,
    query_cache: Optional[QueryCache] = None,
    metrics: Optional[dict[str, Any]] = None,
) -> tuple[list[ProductCandidate], bool]:
    result: list[Optional[ProductCandidate]] = [None] * len(candidates)
    unresolved: list[tuple[int, ProductCandidate]] = []
    cache_hits = 0
    for index, candidate in enumerate(candidates):
        known = lookup_plausibility_range(candidate, known_ranges)
        if known and plausibility_range_is_complete(known):
            result[index] = replace(candidate, plausibility_range=with_default_plausibility_bounds(known))
            continue
        if not should_use_llm_for_plausibility(candidate, known):
            derived = derive_candidate_plausibility_range(candidate)
            if plausibility_range_is_complete(derived):
                result[index] = replace(candidate, plausibility_range=with_default_plausibility_bounds(derived))
                continue
        cached = cached_candidate_plausibility(llm, candidate, manifest, bill, query_cache)
        if cached is not None:
            result[index] = attach_candidate_plausibility(candidate, cached)
            cache_hits += 1
            continue
        unresolved.append((index, candidate))

    batch_calls = 0
    individual_calls = 0
    fallback_candidates: list[tuple[int, ProductCandidate]] = []
    if len(unresolved) == 1:
        fallback_candidates = unresolved
    elif unresolved:
        semaphore = asyncio.Semaphore(LLM_MAX_CONCURRENCY)
        chunks = [
            unresolved[offset : offset + LLM_PLAUSIBILITY_BATCH_SIZE]
            for offset in range(0, len(unresolved), LLM_PLAUSIBILITY_BATCH_SIZE)
        ]

        async def run_batch(chunk: list[tuple[int, ProductCandidate]]) -> list[tuple[int, ProductCandidate]]:
            nonlocal batch_calls
            batch_calls += 1
            try:
                async with semaphore:
                    payload = await llm.chat_json(
                        build_batch_plausibility_messages(chunk, manifest, bill),
                        temperature=0.1,
                        max_tokens=8192,
                    )
                parsed = normalize_batch_plausibility_payload(payload, chunk)
            except Exception:
                return chunk
            missing: list[tuple[int, ProductCandidate]] = []
            for index, candidate in chunk:
                estimated = parsed.get(index)
                if estimated is None or not plausibility_range_is_complete(estimated):
                    missing.append((index, candidate))
                    continue
                cache_candidate_plausibility(llm, candidate, manifest, bill, estimated, query_cache)
                result[index] = attach_candidate_plausibility(candidate, estimated)
            return missing

        missing_groups = await asyncio.gather(*(run_batch(chunk) for chunk in chunks))
        fallback_candidates = [item for group in missing_groups for item in group]

    if fallback_candidates:
        semaphore = asyncio.Semaphore(LLM_MAX_CONCURRENCY)

        async def run_individual(index: int, candidate: ProductCandidate) -> None:
            nonlocal individual_calls
            individual_calls += 1
            async with semaphore:
                estimated = await estimate_candidate_plausibility_with_llm(
                    llm, candidate, manifest, bill, query_cache
                )
            if not plausibility_range_is_complete(estimated):
                raise RuntimeError(f"LLM 未能给出完整合理范围，不能生成: {candidate.zh}/{candidate.en}")
            result[index] = attach_candidate_plausibility(candidate, estimated)

        await asyncio.gather(*(run_individual(index, candidate) for index, candidate in fallback_candidates))

    finalized = [candidate for candidate in result if candidate is not None]
    if len(finalized) != len(candidates):
        raise RuntimeError("LLM 合理性估算结果数量不完整")
    if metrics is not None:
        metrics.update(
            {
                "candidate_count": len(candidates),
                "llm_candidate_count": len(unresolved),
                "batch_calls": batch_calls,
                "individual_fallback_calls": individual_calls,
                "cache_hits": cache_hits,
                "max_concurrency": LLM_MAX_CONCURRENCY,
            }
        )
    return finalized, bool(unresolved)


def attach_candidate_plausibility(
    candidate: ProductCandidate,
    estimated: PlausibilityRange,
) -> ProductCandidate:
    return replace(
        candidate,
        plausibility_range=with_default_plausibility_bounds(estimated),
        plausibility_confidence=range_confidence(estimated),
        plausibility_basis=estimated.source,
    )


def build_batch_plausibility_messages(
    candidates: list[tuple[int, ProductCandidate]],
    manifest: ManifestSummary,
    bill: BillInfo,
) -> list[dict[str, str]]:
    context = {
        "manifest_total_weight_kg": manifest.total_real_weight,
        "manifest_total_ctns": manifest.total_ctns,
        "bill_products": bill.products,
        "candidates": [
            {"candidate_id": f"candidate-{index}", **candidate_to_llm_dict(candidate)}
            for index, candidate in candidates
        ],
    }
    return [
        {
            "role": "system",
            "content": (
                "你是美国清关商业发票合理性审核专家。批量估算每个候选的合理申报范围。"
                "不得修改 candidate_id、HS 或税率，不要为了满足目标税金压低价格或重量。只返回 JSON object。"
            ),
        },
        {
            "role": "user",
            "content": (
                "为每个候选返回正数范围，min <= max，单位为 USD、kg、PCS/CTN。"
                "必须逐个覆盖输入 candidate_id。\n"
                "JSON格式：{\"candidates\":[{\"candidate_id\":\"candidate-0\","
                "\"unit_price_min\":0.0,\"unit_price_max\":0.0,"
                "\"kg_per_pc_min\":0.0,\"kg_per_pc_max\":0.0,"
                "\"qty_per_ctn_min\":0.0,\"qty_per_ctn_max\":0.0,"
                "\"kg_per_ctn_min\":0.0,\"kg_per_ctn_max\":0.0,"
                "\"confidence\":0.0,\"basis\":\"\"}]}\n"
                f"上下文：{json.dumps(context, ensure_ascii=False, separators=(',', ':'))}"
            ),
        },
    ]


def normalize_batch_plausibility_payload(
    payload: dict[str, Any],
    candidates: list[tuple[int, ProductCandidate]],
) -> dict[int, PlausibilityRange]:
    raw_items = payload.get("candidates") or payload.get("items") or payload.get("results") or []
    if not isinstance(raw_items, list):
        return {}
    by_id = {
        clean_text(item.get("candidate_id")): item
        for item in raw_items
        if isinstance(item, dict) and clean_text(item.get("candidate_id"))
    }
    result: dict[int, PlausibilityRange] = {}
    for index, candidate in candidates:
        item = by_id.get(f"candidate-{index}")
        if not item:
            continue
        ranges = item.get("ranges") if isinstance(item.get("ranges"), dict) else item
        result[index] = normalize_llm_plausibility_payload(ranges, candidate)
    return result


def should_use_llm_for_plausibility(candidate: ProductCandidate, known: Optional[PlausibilityRange]) -> bool:
    if candidate.source in {"bill", "llm_query"}:
        return True
    return known is not None and not plausibility_range_is_complete(known)


def lookup_plausibility_range(
    candidate: ProductCandidate,
    ranges: dict[tuple[str, str, str], PlausibilityRange],
) -> Optional[PlausibilityRange]:
    direct = ranges.get(candidate_identity(candidate))
    if direct:
        return direct
    hs = normalize_hs(candidate.hs)
    name_keys = {normalize_text(candidate.zh), normalize_text(candidate.en)}
    for (zh_key, en_key, hs_key), value in ranges.items():
        if hs and hs_key == hs and (zh_key in name_keys or en_key in name_keys):
            return value
    return None


def derive_candidate_plausibility_range(candidate: ProductCandidate) -> PlausibilityRange:
    gross = candidate.gross_weight or candidate.real_weight or 0
    kg_per_ctn = gross / candidate.ctns if gross and candidate.ctns else None
    kg_per_pc = gross / candidate.qty if gross and candidate.qty else None
    qty_per_ctn = candidate.qty / candidate.ctns if candidate.qty and candidate.ctns else None
    return build_plausibility_range(
        kg_per_ctn=kg_per_ctn,
        kg_per_pc=kg_per_pc,
        unit_price=candidate.unit_price or None,
        ctns=candidate.ctns or None,
        qty_per_ctn=qty_per_ctn,
        source=f"{candidate.source_label} 派生范围",
    )


def plausibility_range_is_complete(value: Optional[PlausibilityRange]) -> bool:
    if value is None:
        return False
    fields = (
        value.kg_per_ctn_min,
        value.kg_per_ctn_max,
        value.kg_per_pc_min,
        value.kg_per_pc_max,
        value.unit_price_min,
        value.unit_price_max,
        value.qty_per_ctn_min,
        value.qty_per_ctn_max,
    )
    return all(item is not None and item > 0 for item in fields)


def range_confidence(value: PlausibilityRange) -> float:
    match = re.search(r"confidence=([0-9.]+)", value.source)
    if not match:
        return 0.0
    parsed = to_float(match.group(1))
    return max(0.0, min(1.0, parsed or 0.0))


def candidate_plausibility_context(
    candidate: ProductCandidate,
    manifest: ManifestSummary,
    bill: BillInfo,
) -> dict[str, Any]:
    return {
        "candidate": candidate_to_llm_dict(candidate),
        "manifest_total_weight_kg": manifest.total_real_weight,
        "manifest_total_ctns": manifest.total_ctns,
        "bill_products": bill.products,
    }


def candidate_plausibility_cache_key(
    llm: LLMClient,
    candidate: ProductCandidate,
    manifest: ManifestSummary,
    bill: BillInfo,
) -> str:
    model = clean_text(getattr(getattr(llm, "settings", None), "model", "")) or llm.__class__.__name__
    context = candidate_plausibility_context(candidate, manifest, bill)
    return hashlib.sha256(
        json.dumps(
            {
                "kind": "plausibility",
                "prompt_version": LLM_PLAUSIBILITY_PROMPT_VERSION,
                "model": model,
                **context,
            },
            ensure_ascii=False,
            sort_keys=True,
        ).encode("utf-8", errors="ignore")
    ).hexdigest()


def cached_candidate_plausibility(
    llm: LLMClient,
    candidate: ProductCandidate,
    manifest: ManifestSummary,
    bill: BillInfo,
    query_cache: Optional[QueryCache],
) -> Optional[PlausibilityRange]:
    if query_cache is None:
        return None
    cache_key = candidate_plausibility_cache_key(llm, candidate, manifest, bill)
    payload = query_cache.setdefault("llm_plausibility", {}).get(cache_key)
    if not isinstance(payload, dict):
        return None
    if payload.get("_normalized_range") is True:
        values = {key: value for key, value in payload.items() if key != "_normalized_range"}
        try:
            return normalize_plausibility_range_bounds(PlausibilityRange(**values))
        except TypeError:
            return None
    return normalize_llm_plausibility_payload(payload, candidate)


def cache_candidate_plausibility(
    llm: LLMClient,
    candidate: ProductCandidate,
    manifest: ManifestSummary,
    bill: BillInfo,
    plausibility: PlausibilityRange,
    query_cache: Optional[QueryCache],
) -> None:
    if query_cache is None:
        return
    cache_key = candidate_plausibility_cache_key(llm, candidate, manifest, bill)
    query_cache.setdefault("llm_plausibility", {})[cache_key] = {
        "_normalized_range": True,
        **asdict(plausibility),
    }


async def estimate_candidate_plausibility_with_llm(
    llm: LLMClient,
    candidate: ProductCandidate,
    manifest: ManifestSummary,
    bill: BillInfo,
    query_cache: Optional[QueryCache] = None,
) -> PlausibilityRange:
    payload_context = candidate_plausibility_context(candidate, manifest, bill)
    cached = cached_candidate_plausibility(llm, candidate, manifest, bill, query_cache)
    if cached is not None:
        return cached

    payload = await llm.chat_json(build_plausibility_messages(payload_context), temperature=0.1)
    plausibility = normalize_llm_plausibility_payload(payload, candidate)
    cache_candidate_plausibility(llm, candidate, manifest, bill, plausibility, query_cache)
    return plausibility


def build_plausibility_messages(context: dict[str, Any]) -> list[dict[str, str]]:
    return [
        {
            "role": "system",
            "content": (
                "你是美国清关商业发票合理性审核专家。你负责根据品类、材质、HS和常识估算合理的申报范围。"
                "不要为了满足税金而压低单价或单重；只返回 JSON object。"
            ),
        },
        {
            "role": "user",
            "content": (
                "请为该候选估算合理范围，单位为 USD、kg、PCS/CTN。\n"
                "必须返回正数区间，min <= max。电器、机器、家居用品等要符合常识；不能把加湿器、挂烫机估成 1 美元。\n"
                "JSON格式：{\"unit_price_min\":0.0,\"unit_price_max\":0.0,"
                "\"kg_per_pc_min\":0.0,\"kg_per_pc_max\":0.0,"
                "\"qty_per_ctn_min\":0.0,\"qty_per_ctn_max\":0.0,"
                "\"kg_per_ctn_min\":0.0,\"kg_per_ctn_max\":0.0,"
                "\"confidence\":0.0,\"basis\":\"\"}\n"
                f"上下文：{json.dumps(context, ensure_ascii=False, separators=(',', ':'))}"
            ),
        },
    ]


def normalize_llm_plausibility_payload(payload: dict[str, Any], candidate: ProductCandidate) -> PlausibilityRange:
    confidence = to_float(payload.get("confidence"))
    confidence = max(0.0, min(1.0, confidence if confidence is not None else 0.0))
    basis = clean_text(payload.get("basis") or payload.get("reason") or "LLM 合理性估算")
    value = PlausibilityRange(
        kg_per_ctn_min=positive_float(payload.get("kg_per_ctn_min")),
        kg_per_ctn_max=positive_float(payload.get("kg_per_ctn_max")),
        kg_per_pc_min=positive_float(payload.get("kg_per_pc_min")),
        kg_per_pc_max=positive_float(payload.get("kg_per_pc_max")),
        unit_price_min=positive_float(payload.get("unit_price_min")),
        unit_price_max=positive_float(payload.get("unit_price_max")),
        qty_per_ctn_min=positive_float(payload.get("qty_per_ctn_min")),
        qty_per_ctn_max=positive_float(payload.get("qty_per_ctn_max")),
        source=f"LLM合理性估算 confidence={confidence}: {candidate.zh}/{candidate.en}; {basis}",
    )
    return normalize_plausibility_range_bounds(value)


def normalize_plausibility_range_bounds(value: PlausibilityRange) -> PlausibilityRange:
    kg_ctn_min, kg_ctn_max = ordered_optional(value.kg_per_ctn_min, value.kg_per_ctn_max)
    kg_pc_min, kg_pc_max = ordered_optional(value.kg_per_pc_min, value.kg_per_pc_max)
    price_min, price_max = ordered_optional(value.unit_price_min, value.unit_price_max)
    qty_min, qty_max = ordered_optional(value.qty_per_ctn_min, value.qty_per_ctn_max)
    return PlausibilityRange(
        kg_per_ctn_min=kg_ctn_min,
        kg_per_ctn_max=kg_ctn_max,
        kg_per_pc_min=kg_pc_min,
        kg_per_pc_max=kg_pc_max,
        unit_price_min=price_min,
        unit_price_max=price_max,
        ctns_min=value.ctns_min,
        ctns_max=value.ctns_max,
        qty_per_ctn_min=qty_min,
        qty_per_ctn_max=qty_max,
        source=value.source,
    )


def attach_price_evidence_to_candidates(
    candidates: list[ProductCandidate],
    *,
    query_cache: Optional[QueryCache] = None,
    price_timeout: float = PRICE_SEARCH_TIMEOUT_SECONDS,
    price_max_pages: int = PRICE_SEARCH_MAX_PAGES,
) -> list[ProductCandidate]:
    result: list[ProductCandidate] = []
    cache = query_cache.setdefault("price", {}) if query_cache is not None else {}
    cache_dir = APP_DIR / "runtime" / "price_cache"
    for candidate in candidates:
        query = build_price_query(candidate)
        key = product_cache_key(query, candidate.material)
        evidence_payload = cache.get(key) if cache is not None else None
        if isinstance(evidence_payload, dict):
            evidence = evidence_payload
        else:
            try:
                evidence_obj = estimate_declared_unit_price_from_web(
                    query,
                    cache_dir=cache_dir,
                    timeout=price_timeout,
                    max_pages=price_max_pages,
                    declaration_ratio=DECLARED_RETAIL_PRICE_RATIO,
                )
                evidence = evidence_obj.to_dict()
            except Exception as exc:
                evidence = PriceEvidence(query=query, basis=f"web search failed: {exc}", confidence=0.0, source="fallback").to_dict()
            if cache is not None:
                cache[key] = evidence
        evidence = apply_price_fallback(candidate, evidence)
        plausibility = adjust_plausibility_with_price_evidence(candidate.plausibility_range, evidence)
        result.append(replace(candidate, price_evidence=evidence, plausibility_range=plausibility))
    return result


async def attach_price_evidence_to_candidates_concurrently(
    candidates: list[ProductCandidate],
    *,
    query_cache: Optional[QueryCache] = None,
    max_concurrency: int = LLM_MAX_CONCURRENCY,
    metrics: Optional[dict[str, Any]] = None,
) -> list[ProductCandidate]:
    semaphore = asyncio.Semaphore(max(1, max_concurrency))
    results: list[Optional[ProductCandidate]] = [None] * len(candidates)
    network_candidates = 0
    trusted_candidates = 0

    async def enrich(index: int, candidate: ProductCandidate) -> None:
        nonlocal network_candidates, trusted_candidates
        if candidate_has_trusted_price_reference(candidate):
            trusted_candidates += 1
            evidence = apply_price_fallback(
                candidate,
                PriceEvidence(
                    query=build_price_query(candidate),
                    basis="trusted candidate unit price",
                    confidence=0.8,
                    source="candidate",
                ).to_dict(),
            )
            plausibility = adjust_plausibility_with_price_evidence(candidate.plausibility_range, evidence)
            results[index] = replace(candidate, price_evidence=evidence, plausibility_range=plausibility)
            return
        network_candidates += 1
        async with semaphore:
            enriched = await asyncio.to_thread(
                attach_price_evidence_to_candidates,
                [candidate],
                query_cache=None,
            )
        results[index] = enriched[0]

    await asyncio.gather(*(enrich(index, candidate) for index, candidate in enumerate(candidates)))
    finalized = [candidate for candidate in results if candidate is not None]
    if len(finalized) != len(candidates):
        raise RuntimeError("价格证据并发处理结果数量不完整")
    if query_cache is not None:
        cache = query_cache.setdefault("price", {})
        for candidate in finalized:
            cache[product_cache_key(build_price_query(candidate), candidate.material)] = candidate.price_evidence
    if metrics is not None:
        metrics.update(
            {
                "candidate_count": len(candidates),
                "network_candidates": network_candidates,
                "trusted_price_candidates": trusted_candidates,
                "max_concurrency": max(1, max_concurrency),
            }
        )
    return finalized


def candidate_has_trusted_price_reference(candidate: ProductCandidate) -> bool:
    return bool(
        candidate.unit_price
        and candidate.unit_price > 0
        and (
            candidate.source_label == "DEFAULT_REFERENCE_STYLE_ROWS"
            or is_customer_codebook_candidate(candidate)
        )
    )


def build_price_query(candidate: ProductCandidate) -> str:
    parts = [
        candidate.en if has_ascii_alpha(candidate.en) else "",
        candidate.zh if not has_ascii_alpha(candidate.en) else "",
        candidate.material,
        "retail price",
    ]
    return clean_text(" ".join(part for part in parts if clean_text(part)))


def has_ascii_alpha(value: Any) -> bool:
    return any(ch.isascii() and ch.isalpha() for ch in str(value or ""))


def apply_price_fallback(candidate: ProductCandidate, evidence: dict[str, Any]) -> dict[str, Any]:
    declared = to_float(evidence.get("declared_unit_price"))
    if declared and declared > 0:
        return evidence
    fallback = candidate.unit_price or None
    source = "manifest grouped median unit price" if candidate.source == "manifest_group" else "candidate original unit price"
    if fallback is None or fallback <= 0:
        plausibility = candidate.plausibility_range
        fallback = plausibility.unit_price_min if plausibility and plausibility.unit_price_min else DEFAULT_UNIT_PRICE_MIN
        source = "plausibility minimum fallback"
    retail = round(fallback / DECLARED_RETAIL_PRICE_RATIO, 4) if fallback else 0.0
    return {
        **evidence,
        "retail_unit_price": retail,
        "declared_unit_price": round(fallback, 4),
        "basis": f"{source}; no usable public web search sample",
        "confidence": max(0.25, to_float(evidence.get("confidence")) or 0.0),
        "source": "fallback",
        "samples": evidence.get("samples") or [],
    }


def adjust_plausibility_with_price_evidence(
    plausibility: Optional[PlausibilityRange],
    evidence: dict[str, Any],
) -> Optional[PlausibilityRange]:
    if plausibility is None:
        return None
    declared = to_float(evidence.get("declared_unit_price"))
    confidence = to_float(evidence.get("confidence")) or 0.0
    if declared is None or declared <= 0 or confidence < 0.2:
        return plausibility
    min_price = max(0.0001, declared * 0.67)
    max_price = max(min_price, declared * 1.5)
    if plausibility.unit_price_min is not None:
        min_price = max(min_price, plausibility.unit_price_min)
    if plausibility.unit_price_max is not None:
        max_price = min(max_price, plausibility.unit_price_max)
        if max_price < min_price:
            max_price = min_price
    return PlausibilityRange(
        kg_per_ctn_min=plausibility.kg_per_ctn_min,
        kg_per_ctn_max=plausibility.kg_per_ctn_max,
        kg_per_pc_min=plausibility.kg_per_pc_min,
        kg_per_pc_max=plausibility.kg_per_pc_max,
        unit_price_min=round(min_price, 4),
        unit_price_max=round(max_price, 4),
        ctns_min=plausibility.ctns_min,
        ctns_max=plausibility.ctns_max,
        qty_per_ctn_min=plausibility.qty_per_ctn_min,
        qty_per_ctn_max=plausibility.qty_per_ctn_max,
        source=f"{plausibility.source}; price evidence: {clean_text(evidence.get('basis'))}",
    )


def validate_price_evidence(rows: list[dict[str, Any]], selected: list[ProductCandidate]) -> None:
    for idx, (row, candidate) in enumerate(zip(rows, selected), start=1):
        evidence = candidate.price_evidence or {}
        declared = to_float(evidence.get("declared_unit_price"))
        confidence = to_float(evidence.get("confidence")) or 0.0
        unit_price = to_float(row.get("单价")) or 0.0
        if declared and confidence >= 0.2:
            low = declared * 0.67
            high = declared * 1.5
            if unit_price < low - 0.0001 or unit_price > high + 0.0001:
                message = (
                    f"单价超出价格证据范围 {unit_price} not in "
                    f"{round(low, 4)}-{round(high, 4)}"
                )
                append_row_warning(row, message)
                row["价格提示"] = message
        row["价格依据"] = evidence.get("basis", "")
        row["价格置信度"] = evidence.get("confidence", "")
        row["零售参考单价"] = evidence.get("retail_unit_price", "")


def append_row_warning(row: dict[str, Any], message: str) -> None:
    existing = clean_text(row.get("约束提示"))
    if not existing:
        row["约束提示"] = message
        return
    parts = [part.strip() for part in existing.split(";") if part.strip()]
    if message not in parts:
        parts.append(message)
    row["约束提示"] = "; ".join(parts)


def build_audit_summary(
    rows: list[dict[str, Any]],
    selected: list[ProductCandidate],
    manifest: ManifestSummary,
) -> dict[str, Any]:
    line_items: list[dict[str, Any]] = []
    for row, candidate in zip(rows, selected):
        line_items.append(
            {
                "name": row.get("中文品名"),
                "hs": normalize_hs(row.get("商品编码")),
                "source": candidate.source,
                "source_rows": candidate.source_rows,
                "tax_match_source": candidate.tax_match_source,
                "crawler_description": row.get("爬虫品名"),
                "price_evidence": candidate.price_evidence,
                "plausibility_source": row.get("合理性来源") or row.get("重量规则来源"),
                "warnings": row.get("约束提示"),
            }
        )
    return {
        "manifest_total_declared_value": manifest.total_declared_value,
        "manifest_origin_rows": sum(1 for candidate in selected if candidate.source == "manifest_group"),
        "replacement_rows": sum(1 for candidate in selected if candidate.source == "replacement"),
        "line_items": line_items,
    }


def ordered_optional(left: Optional[float], right: Optional[float]) -> tuple[Optional[float], Optional[float]]:
    if left is None or right is None:
        return left, right
    return min(left, right), max(left, right)


def positive_float(value: Any) -> Optional[float]:
    parsed = to_float(value)
    if parsed is None or parsed <= 0:
        return None
    return parsed


def base_tax_rate(data: dict[str, Any]) -> Optional[float]:
    return parse_tax_rate(data.get("tax_rate") or data.get("column_rate_of_duty"))


def effective_tax_rate(data: dict[str, Any]) -> float:
    base = base_tax_rate(data)
    if base is None:
        base = 0
    return base + additional_tax_rate(data)


def additional_tax_rate(data: dict[str, Any]) -> float:
    return parse_non_exempt_additional_tax_rate(
        data.get("additional_tax_rate"),
        data.get("additional_tax_details"),
    )


def parse_non_exempt_additional_tax_rate(value: Any, details: Any = None) -> float:
    if isinstance(details, list):
        total = 0.0
        for item in details:
            if not isinstance(item, dict):
                total += parse_non_exempt_additional_tax_rate(item)
                continue
            label = clean_text(item.get("label") or item.get("name") or item.get("type") or item.get("title"))
            rate_value = item.get("rate") or item.get("value") or item.get("tax_rate") or item.get("additionalTaxRate")
            text = " ".join(clean_text(part) for part in (label, rate_value, item.get("text"), item.get("remark")) if clean_text(part))
            if is_base_tax_detail_text(text):
                continue
            if is_exempt_additional_tax_text(text):
                continue
            total += parse_tax_rate(rate_value if rate_value not in (None, "") else text) or 0
        if total or not clean_text(value):
            return total
        return parse_non_exempt_additional_tax_rate(value)
    if isinstance(details, dict):
        return parse_non_exempt_additional_tax_rate([details])

    text = " ".join(clean_text(part) for part in (value, details) if clean_text(part))
    if not text:
        return 0.0
    total = 0.0
    parts = [part.strip() for part in re.split(r"[;\n,，]+", text) if part.strip()]
    if len(parts) <= 1 and "+" in text:
        parts = [part.strip() for part in text.split("+") if part.strip()]
    for part in parts or [text]:
        if is_base_tax_detail_text(part):
            continue
        if is_exempt_additional_tax_text(part):
            continue
        total += parse_tax_rate(part) or 0
    return total


def is_base_tax_detail_text(value: Any) -> bool:
    text = normalize_text(value)
    compact = text.replace(" ", "")
    if not text:
        return False
    markers = (
        "通用税率",
        "普通税率",
        "基础税率",
        "基本税率",
        "进口税率",
        "generalrate",
        "general",
        "baserate",
        "basicrate",
        "column1",
        "columnone",
        "mfn",
        "normaltrade",
    )
    return any(marker in compact for marker in markers)


def is_exempt_additional_tax_text(value: Any) -> bool:
    text = normalize_text(value)
    compact = text.replace(" ", "")
    if not text:
        return False
    exempt_markers = (
        "条件豁免",
        "可豁免",
        "豁免",
        "排除",
        "exempt",
        "exemption",
        "excluded",
        "exclusion",
        "conditionalexemption",
    )
    non_exempt_markers = (
        "无豁免",
        "不豁免",
        "未豁免",
        "nonexempt",
        "noexemption",
        "without exemption",
        "notexempt",
    )
    if any(marker in compact for marker in non_exempt_markers) or any(marker in text for marker in non_exempt_markers):
        return False
    return any(marker in compact for marker in exempt_markers) or any(marker in text for marker in exempt_markers)


def parse_tax_rate(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        if math.isnan(float(value)):
            return None
        number = float(value)
        return number if abs(number) <= 1 else number / 100

    text = str(value).strip()
    if not text:
        return None
    lowered = text.lower()
    if lowered in {"n/a", "na", "none", "null", "-"}:
        return None
    if lowered in {"free", "免税", "无"}:
        return 0.0

    matches = re.findall(r"-?\d+(?:\.\d+)?", text.replace(",", ""))
    if not matches:
        return None
    numbers = [float(match) for match in matches]
    number = sum(numbers) if "+" in text else numbers[0]
    return number / 100 if "%" in text or abs(number) > 1 else number


def certification_filter_reason(texts: list[str], rules: SelectionRules) -> str:
    cleaned = [clean_text(text) for text in texts if clean_text(text)]
    if not cleaned:
        return ""

    joined = " ".join(cleaned)
    joined_key = normalize_certification_text(joined)
    for blocked in rules.blocked_certifications:
        blocked_key = normalize_certification_text(blocked)
        if blocked_key and blocked_key in joined_key:
            return f"命中禁用认证: {blocked}"

    allowed_keys = [normalize_certification_text(item) for item in rules.allowed_certifications if normalize_certification_text(item)]
    present_allowed = [key for key in allowed_keys if key and key in joined_key]
    regulatory_terms = extract_regulatory_terms(joined)
    unknown_terms = [term for term in regulatory_terms if term not in set(allowed_keys)]
    if unknown_terms:
        return "出现未允许认证/监管: " + ", ".join(sorted(unknown_terms))
    if present_allowed and not unknown_terms:
        return ""
    return "出现未允许认证/监管提示"


def normalize_certification_text(value: Any) -> str:
    text = normalize_text(value)
    text = text.replace("lacey act", "laceyact")
    return text.replace(" ", "")


def extract_regulatory_terms(text: str) -> set[str]:
    normalized = normalize_certification_text(text)
    terms: set[str] = set()
    known = {
        "laceyact": ("laceyact", "lacey"),
        "tsca": ("tsca",),
        "fda": ("fda",),
        "fcc": ("fcc",),
        "cpsc": ("cpsc",),
        "cpsia": ("cpsia",),
        "dot": ("dot",),
        "ul": ("ul",),
        "epa": ("epa",),
        "rohs": ("rohs",),
    }
    for canonical, aliases in known.items():
        if any(alias in normalized for alias in aliases):
            terms.add(canonical)
    for acronym in re.findall(r"\b[A-Z]{2,8}\b", text):
        terms.add(normalize_certification_text(acronym))
    return terms


def format_rate(value: float) -> str:
    return f"{round(value * 100, 4)}%"


def load_replacement_candidates(path: Path = REPLACEMENT_WORKBOOK_PATH) -> list[ProductCandidate]:
    candidates: list[ProductCandidate] = []
    if path.exists():
        workbook = load_workbook(path, data_only=True)
        if "常用1" in workbook.sheetnames:
            candidates.extend(load_common_sheet_candidates(workbook["常用1"]))
        if "20260330" in workbook.sheetnames:
            candidates.extend(load_20260330_candidates(workbook["20260330"]))
    candidates.extend(load_customer_codebook_candidates())
    candidates.extend(load_default_reference_replacement_candidates())
    return dedupe_candidates(candidates)


def load_plausibility_ranges(path: Path = REPLACEMENT_WORKBOOK_PATH) -> dict[tuple[str, str, str], PlausibilityRange]:
    if not path.exists():
        return {}
    workbook = load_workbook(path, data_only=True)
    ranges: dict[tuple[str, str, str], PlausibilityRange] = {}
    if "常用1" in workbook.sheetnames:
        ranges.update(load_common_sheet_plausibility_ranges(workbook["常用1"]))
    if "20260330" in workbook.sheetnames:
        ranges.update(load_20260330_plausibility_ranges(workbook["20260330"]))
    return ranges


def load_customer_codebook_candidates(path: Path = CUSTOMER_CODEBOOK_PATH) -> list[ProductCandidate]:
    if not path.exists():
        return []
    workbook = load_workbook(path, data_only=True)
    candidates: list[ProductCandidate] = []
    for sheet in workbook.worksheets:
        header_row = find_customer_codebook_header_row(sheet)
        if not header_row:
            continue
        columns = customer_codebook_columns(sheet, header_row)
        required = {"zh", "en", "material", "hs", "base_rate", "floating_rate", "fixed_rate"}
        if not required.issubset(columns):
            continue
        for row_idx in range(header_row + 1, sheet.max_row + 1):
            zh = clean_text(sheet.cell(row_idx, columns["zh"]).value)
            en = clean_text(sheet.cell(row_idx, columns["en"]).value) or zh
            material = clean_text(sheet.cell(row_idx, columns["material"]).value)
            hs = normalize_hs(sheet.cell(row_idx, columns["hs"]).value)
            if not zh or not hs:
                continue
            base_raw = sheet.cell(row_idx, columns["base_rate"]).value
            floating_raw = sheet.cell(row_idx, columns["floating_rate"]).value
            fixed_raw = sheet.cell(row_idx, columns["fixed_rate"]).value
            base_rate = parse_percentage_component_tax_rate(base_raw)
            floating_rate = parse_percentage_component_tax_rate(floating_raw) or 0.0
            fixed_rate = parse_percentage_component_tax_rate(fixed_raw) or 0.0
            if base_rate is None:
                continue
            total_rate = base_rate + floating_rate + fixed_rate
            if total_rate >= CUSTOMER_CODEBOOK_TAX_LIMIT:
                continue
            row_warnings: list[str] = []
            if has_specific_duty_component(base_raw):
                row_warnings.append(f"复合税率按百分比部分估算，原始基础税率: {clean_text(base_raw)}")
            plausibility = build_plausibility_range(
                kg_per_ctn=10.0,
                kg_per_pc=0.5,
                unit_price=1.0,
                ctns=1.0,
                qty_per_ctn=10.0,
                source="客户编码库默认合理范围",
            )
            candidates.append(
                ProductCandidate(
                    source="replacement",
                    source_label=f"客户编码库/{sheet.title}",
                    zh=zh,
                    en=en,
                    hs=hs,
                    material=material or "Mixed",
                    usage="HOME",
                    real_weight=10.0,
                    gross_weight=10.0,
                    ctns=1.0,
                    qty=10.0,
                    unit_price=1.0,
                    declared_value=10.0,
                    tax_data=customer_codebook_tax_data(hs, base_raw, floating_raw, fixed_raw, total_rate),
                    base_tax_rate=base_rate,
                    effective_tax_rate=total_rate,
                    tax_match_source="customer_codebook",
                    plausibility_range=plausibility,
                    plausibility_confidence=0.65,
                    plausibility_basis=plausibility.source,
                    row_warnings=row_warnings,
                )
            )
    return candidates


def find_customer_codebook_header_row(sheet) -> Optional[int]:
    for row_idx in range(1, min(sheet.max_row, 10) + 1):
        values = [normalize_text(sheet.cell(row_idx, col_idx).value) for col_idx in range(1, sheet.max_column + 1)]
        joined = " ".join(values).replace(" ", "")
        if "customcode" in joined and ("descriptionofgoods" in joined or "中文品名" in joined):
            return row_idx
    return None


def customer_codebook_columns(sheet, header_row: int) -> dict[str, int]:
    columns: dict[str, int] = {}
    for col_idx in range(1, sheet.max_column + 1):
        key = normalize_text(sheet.cell(header_row, col_idx).value)
        compact = key.replace(" ", "")
        if key in {"中文品名", "chineseproductname"}:
            columns["zh"] = col_idx
        elif "descriptionofgoods" in compact or key in {"英文品名", "englishname"}:
            columns["en"] = col_idx
        elif "material" in compact or "材质" in key:
            columns["material"] = col_idx
        elif "customcode" in compact or "海关编码" in key or key in {"hts", "hscode"}:
            columns["hs"] = col_idx
        elif "基本税率" in key or "baserate" in compact:
            columns["base_rate"] = col_idx
        elif "浮动加增" in key or "floating" in compact:
            columns["floating_rate"] = col_idx
        elif "固定加增" in key or "fixed" in compact:
            columns["fixed_rate"] = col_idx
    return columns


def customer_codebook_tax_data(
    hs: str,
    base_raw: Any,
    floating_raw: Any,
    fixed_raw: Any,
    total_rate: float,
) -> dict[str, Any]:
    return {
        "hs_code_us": normalize_hs(hs),
        "tax_rate": clean_text(base_raw) or "Free",
        "additional_tax_rate": "+".join(
            part
            for part in (format_raw_rate(floating_raw), format_raw_rate(fixed_raw))
            if part
        ),
        "description_cn": "客户编码库确认可报",
        "certification_texts": [],
        "anti_dumping": False,
        "customer_codebook_total_tax_rate": total_rate,
    }


def format_raw_rate(value: Any) -> str:
    parsed = parse_percentage_component_tax_rate(value)
    if parsed is None or parsed == 0:
        return ""
    return format_rate(parsed)


def parse_percentage_component_tax_rate(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        if math.isnan(float(value)):
            return None
        number = float(value)
        return number if abs(number) <= 1 else number / 100
    text = clean_text(value)
    if not text:
        return None
    lowered = text.lower()
    if lowered in {"free", "免税", "无", "n/a", "na", "none", "null", "-"}:
        return 0.0 if lowered in {"free", "免税", "无"} else None
    percent_matches = re.findall(r"(-?\d+(?:\.\d+)?)\s*%", text.replace(",", ""))
    if percent_matches:
        return sum(float(match) for match in percent_matches) / 100
    return parse_tax_rate(text)


def has_specific_duty_component(value: Any) -> bool:
    text = normalize_text(value)
    return "each" in text or "¢" in str(value or "") or "cent" in text


def load_common_sheet_plausibility_ranges(sheet) -> dict[tuple[str, str, str], PlausibilityRange]:
    ranges: dict[tuple[str, str, str], PlausibilityRange] = {}
    for row_idx in range(2, sheet.max_row + 1):
        zh = clean_text(sheet.cell(row_idx, 1).value)
        en = clean_text(sheet.cell(row_idx, 2).value)
        hs = normalize_hs(sheet.cell(row_idx, 4).value)
        if not zh and not en and not hs:
            continue
        total_weight = to_float(sheet.cell(row_idx, 5).value)
        ctns = to_float(sheet.cell(row_idx, 6).value)
        qty = to_float(sheet.cell(row_idx, 7).value)
        unit_price = to_float(sheet.cell(row_idx, 8).value)
        item_weight = to_float(sheet.cell(row_idx, 10).value)
        kg_per_ctn = total_weight / ctns if total_weight and ctns else None
        kg_per_pc = item_weight or (total_weight / qty if total_weight and qty else None)
        qty_per_ctn = qty / ctns if qty and ctns else None
        merge_plausibility_range(
            ranges,
            candidate_key_from_parts(zh, en, hs),
            build_plausibility_range(
                kg_per_ctn=kg_per_ctn,
                kg_per_pc=kg_per_pc,
                unit_price=unit_price,
                ctns=ctns,
                qty_per_ctn=qty_per_ctn,
                source="海关编码查找.xlsx/常用1",
            ),
        )
    return ranges


def load_20260330_plausibility_ranges(sheet) -> dict[tuple[str, str, str], PlausibilityRange]:
    ranges: dict[tuple[str, str, str], PlausibilityRange] = {}
    for row_idx in range(3, sheet.max_row + 1):
        zh = clean_text(sheet.cell(row_idx, 2).value)
        hs = normalize_hs(sheet.cell(row_idx, 5).value)
        if not zh and not hs:
            continue
        unit_price_range = parse_range_bounds(sheet.cell(row_idx, 8).value)
        kg_per_pc_range = parse_range_bounds(sheet.cell(row_idx, 9).value)
        ctns_range = parse_range_bounds(sheet.cell(row_idx, 10).value)
        qty_per_ctn_range = parse_range_bounds(sheet.cell(row_idx, 11).value)
        merge_plausibility_range(
            ranges,
            candidate_key_from_parts(zh, zh, hs),
            PlausibilityRange(
                kg_per_ctn_min=None,
                kg_per_ctn_max=None,
                kg_per_pc_min=kg_per_pc_range[0],
                kg_per_pc_max=kg_per_pc_range[1],
                unit_price_min=unit_price_range[0],
                unit_price_max=unit_price_range[1],
                ctns_min=ctns_range[0],
                ctns_max=ctns_range[1],
                qty_per_ctn_min=qty_per_ctn_range[0],
                qty_per_ctn_max=qty_per_ctn_range[1],
                source="海关编码查找.xlsx/20260330",
            ),
        )
    return ranges


def merge_plausibility_range(
    ranges: dict[tuple[str, str, str], PlausibilityRange],
    key: tuple[str, str, str],
    value: PlausibilityRange,
) -> None:
    existing = ranges.get(key)
    if not existing:
        ranges[key] = value
        return
    ranges[key] = PlausibilityRange(
        kg_per_ctn_min=min_optional(existing.kg_per_ctn_min, value.kg_per_ctn_min),
        kg_per_ctn_max=max_optional(existing.kg_per_ctn_max, value.kg_per_ctn_max),
        kg_per_pc_min=min_optional(existing.kg_per_pc_min, value.kg_per_pc_min),
        kg_per_pc_max=max_optional(existing.kg_per_pc_max, value.kg_per_pc_max),
        unit_price_min=min_optional(existing.unit_price_min, value.unit_price_min),
        unit_price_max=max_optional(existing.unit_price_max, value.unit_price_max),
        ctns_min=min_optional(existing.ctns_min, value.ctns_min),
        ctns_max=max_optional(existing.ctns_max, value.ctns_max),
        qty_per_ctn_min=min_optional(existing.qty_per_ctn_min, value.qty_per_ctn_min),
        qty_per_ctn_max=max_optional(existing.qty_per_ctn_max, value.qty_per_ctn_max),
        source=f"{existing.source}; {value.source}" if value.source not in existing.source else existing.source,
    )


def min_optional(left: Optional[float], right: Optional[float]) -> Optional[float]:
    values = [value for value in (left, right) if value is not None]
    return min(values) if values else None


def max_optional(left: Optional[float], right: Optional[float]) -> Optional[float]:
    values = [value for value in (left, right) if value is not None]
    return max(values) if values else None


def build_plausibility_range(
    *,
    kg_per_ctn: Optional[float],
    kg_per_pc: Optional[float],
    unit_price: Optional[float],
    ctns: Optional[float],
    qty_per_ctn: Optional[float],
    source: str,
) -> PlausibilityRange:
    kg_per_ctn_min, kg_per_ctn_max = spread_range(kg_per_ctn, 0.5, 2.0)
    kg_per_pc_min, kg_per_pc_max = spread_range(kg_per_pc, 0.5, 2.0)
    unit_price_min, unit_price_max = spread_range(unit_price, 0.5, 2.0)
    ctns_min, ctns_max = spread_range(ctns, 0.4, 2.5)
    qty_per_ctn_min, qty_per_ctn_max = spread_range(qty_per_ctn, 0.5, 2.0)
    return PlausibilityRange(
        kg_per_ctn_min=kg_per_ctn_min,
        kg_per_ctn_max=kg_per_ctn_max,
        kg_per_pc_min=kg_per_pc_min,
        kg_per_pc_max=kg_per_pc_max,
        unit_price_min=unit_price_min,
        unit_price_max=unit_price_max,
        ctns_min=ctns_min,
        ctns_max=ctns_max,
        qty_per_ctn_min=qty_per_ctn_min,
        qty_per_ctn_max=qty_per_ctn_max,
        source=source,
    )


def spread_range(value: Optional[float], min_factor: float, max_factor: float) -> tuple[Optional[float], Optional[float]]:
    if value is None or value <= 0:
        return None, None
    return max(0.0001, value * min_factor), max(0.0001, value * max_factor)


def load_common_sheet_candidates(sheet) -> list[ProductCandidate]:
    candidates: list[ProductCandidate] = []
    for row_idx in range(2, sheet.max_row + 1):
        zh = clean_text(sheet.cell(row_idx, 1).value)
        en = clean_text(sheet.cell(row_idx, 2).value)
        material = clean_text(sheet.cell(row_idx, 3).value)
        hs = normalize_hs(sheet.cell(row_idx, 4).value)
        if not zh or not hs:
            continue
        candidates.append(
            ProductCandidate(
                source="replacement",
                source_label="海关编码查找.xlsx/常用1",
                zh=zh,
                en=en or zh,
                hs=hs,
                material=material,
                usage="HOME",
                real_weight=parse_range_mid(sheet.cell(row_idx, 5).value) or parse_range_mid(sheet.cell(row_idx, 10).value) or 1,
                gross_weight=parse_range_mid(sheet.cell(row_idx, 5).value) or parse_range_mid(sheet.cell(row_idx, 10).value) or 1,
                ctns=parse_range_mid(sheet.cell(row_idx, 6).value) or 1,
                qty=parse_range_mid(sheet.cell(row_idx, 7).value) or 1,
                unit_price=parse_range_mid(sheet.cell(row_idx, 8).value) or 1,
                declared_value=parse_range_mid(sheet.cell(row_idx, 9).value) or 0,
            )
        )
    return candidates


def load_20260330_candidates(sheet) -> list[ProductCandidate]:
    candidates: list[ProductCandidate] = []
    for row_idx in range(3, sheet.max_row + 1):
        zh = clean_text(sheet.cell(row_idx, 2).value)
        material = clean_text(sheet.cell(row_idx, 4).value)
        hs = normalize_hs(sheet.cell(row_idx, 5).value)
        if not zh or not hs:
            continue
        unit_price = parse_range_mid(sheet.cell(row_idx, 8).value) or 1
        weight = parse_range_mid(sheet.cell(row_idx, 9).value) or 1
        ctns = parse_range_mid(sheet.cell(row_idx, 10).value) or 1
        qty_per_ctn = parse_range_mid(sheet.cell(row_idx, 11).value) or 10
        candidates.append(
            ProductCandidate(
                source="replacement",
                source_label="海关编码查找.xlsx/20260330",
                zh=zh,
                en=zh,
                hs=hs,
                material=material,
                usage="HOME",
                real_weight=weight,
                gross_weight=weight,
                ctns=ctns,
                qty=max(1, round(ctns * qty_per_ctn)),
                unit_price=unit_price,
                declared_value=round(max(1, round(ctns * qty_per_ctn)) * unit_price, 2),
            )
        )
    return candidates


def load_default_reference_replacement_candidates() -> list[ProductCandidate]:
    candidates: list[ProductCandidate] = []
    for row in DEFAULT_REFERENCE_STYLE_ROWS:
        zh = clean_text(row.get("中文品名"))
        en = clean_text(row.get("英文品名")) or zh
        hs = normalize_hs(row.get("商品编码"))
        if not zh or not hs:
            continue
        unit_price = to_float(row.get("单价")) or 1.0
        ctns = to_float(row.get("箱数")) or 80.0
        qty = to_float(row.get("数量")) or max(1.0, ctns * 10.0)
        gross_weight = to_float(row.get("毛重")) or max(1.0, ctns * 15.0)
        tax_rate = to_float(row.get("综合税率"))
        plausibility = derive_reference_style_plausibility_range(
            unit_price=unit_price,
            ctns=ctns,
            qty=qty,
            gross_weight=gross_weight,
            source="DEFAULT_REFERENCE_STYLE_ROWS 人工发票参考范围",
        )
        candidates.append(
            ProductCandidate(
                source="replacement",
                source_label="DEFAULT_REFERENCE_STYLE_ROWS",
                zh=zh,
                en=en,
                hs=hs,
                material=clean_text(row.get("材质")) or "Plastic",
                usage=clean_text(row.get("用途")) or "HOME",
                ctns=ctns,
                qty=qty,
                unit_price=unit_price,
                declared_value=round(unit_price * qty, 2),
                real_weight=gross_weight,
                gross_weight=gross_weight,
                tax_data=reference_tax_data(hs, tax_rate) if tax_rate is not None else {},
                base_tax_rate=max(0.0, tax_rate or 0.0) if tax_rate is not None else 0.0,
                effective_tax_rate=max(0.0, tax_rate or 0.0) if tax_rate is not None else 0.0,
                tax_match_source="manual_reference" if tax_rate is not None else "",
                plausibility_range=plausibility,
                plausibility_confidence=0.85,
                plausibility_basis=plausibility.source,
            )
        )
    return candidates


def load_default_reference_manual_candidates() -> list[ProductCandidate]:
    return [
        candidate
        for candidate in load_default_reference_replacement_candidates()
        if candidate.tax_match_source == "manual_reference"
    ]


def derive_reference_style_plausibility_range(
    *,
    unit_price: float,
    ctns: float,
    qty: float,
    gross_weight: float,
    source: str,
) -> PlausibilityRange:
    kg_per_ctn = gross_weight / ctns if gross_weight and ctns else DEFAULT_KG_PER_CTN_MAX / 2
    kg_per_pc = gross_weight / qty if gross_weight and qty else DEFAULT_KG_PER_PC_MAX / 2
    qty_per_ctn = qty / ctns if qty and ctns else 10.0
    return PlausibilityRange(
        kg_per_ctn_min=round(max(DEFAULT_KG_PER_CTN_MIN, kg_per_ctn * 0.55), 4),
        kg_per_ctn_max=round(max(DEFAULT_KG_PER_CTN_MIN, kg_per_ctn * 1.8), 4),
        kg_per_pc_min=round(max(DEFAULT_KG_PER_PC_MIN, kg_per_pc * 0.35), 6),
        kg_per_pc_max=round(max(DEFAULT_KG_PER_PC_MIN, kg_per_pc * 2.8), 6),
        unit_price_min=round(max(DEFAULT_UNIT_PRICE_MIN, unit_price * 0.35), 4),
        unit_price_max=round(max(DEFAULT_UNIT_PRICE_MIN, unit_price * 1.8), 4),
        qty_per_ctn_min=round(max(1.0, qty_per_ctn * 0.35), 4),
        qty_per_ctn_max=round(max(1.0, qty_per_ctn * 2.4), 4),
        source=source,
    )


def reference_tax_data(hs: str, tax_rate: Optional[float]) -> dict[str, Any]:
    normalized = normalize_hs(hs)
    rate = max(0.0, tax_rate or 0.0)
    return {
        "hs_code_us": normalized,
        "tax_rate": "Free" if rate <= 0 else format_rate(rate),
        "additional_tax_rate": "",
        "description_cn": "人工发票参考行税率",
        "certification_texts": [],
        "anti_dumping": False,
    }


def parse_range_mid(value: Any) -> Optional[float]:
    direct = to_float(value)
    if direct is not None:
        return direct
    text = str(value or "")
    numbers = [float(item) for item in re.findall(r"\d+(?:\.\d+)?", text)]
    if not numbers:
        return None
    return sum(numbers[:2]) / min(len(numbers), 2)


def parse_range_bounds(value: Any) -> tuple[Optional[float], Optional[float]]:
    direct = to_float(value)
    if direct is not None:
        return direct, direct
    text = str(value or "")
    numbers = [float(item) for item in re.findall(r"\d+(?:\.\d+)?", text)]
    if not numbers:
        return None, None
    if len(numbers) == 1:
        return numbers[0], numbers[0]
    low, high = numbers[0], numbers[1]
    return min(low, high), max(low, high)


def dedupe_candidates(candidates: list[ProductCandidate]) -> list[ProductCandidate]:
    result: list[ProductCandidate] = []
    seen: set[tuple[str, str, str]] = set()
    for candidate in candidates:
        key = candidate_identity(candidate)
        if key in seen:
            continue
        seen.add(key)
        result.append(candidate)
    return result


def candidate_identity(candidate: ProductCandidate) -> tuple[str, str, str]:
    return (normalize_text(candidate.zh), normalize_text(candidate.en), normalize_hs(candidate.hs))


def candidate_semantic_keys(candidate: ProductCandidate) -> set[str]:
    keys: set[str] = set()
    zh = normalize_text(candidate.zh)
    en = normalize_text(candidate.en)
    if zh:
        keys.add(f"zh:{zh}")
    if en:
        keys.add(f"en:{en}")
    return keys


def candidates_have_duplicate_names(candidates: list[ProductCandidate]) -> bool:
    seen: set[str] = set()
    for candidate in candidates:
        keys = candidate_semantic_keys(candidate)
        if keys & seen:
            return True
        seen.update(keys)
    return False


def candidate_key_from_parts(zh: Any, en: Any, hs: Any) -> tuple[str, str, str]:
    return (normalize_text(zh), normalize_text(en), normalize_hs(hs))


def target_tax_upper_bound(target_tax_amount: float) -> float:
    return round(target_tax_amount * (1 + MAX_TAX_OVER_TARGET_RATIO), 2)


def output_package_total(bill: BillInfo, row_count: int) -> int:
    package_count = normalize_carton_count(bill.cartons)
    if package_count is None:
        raise RuntimeError("提单未识别到有效包装数量，不能生成与整票包装数量对齐的输出")
    target = int(package_count)
    if target < row_count:
        unit = f" {bill.carton_unit}" if bill.carton_unit else ""
        raise RuntimeError(
            f"模型识别的整票包装数量 {package_count:g}{unit} 小于输出行数 {row_count}，"
            "无法保证每行至少分配一个包装单位并闭合总数"
        )
    return target


def manifest_net_to_gross_ratio(manifest: ManifestSummary) -> Optional[float]:
    gross_weight = to_float(manifest.total_real_weight)
    net_weight = to_float(manifest.total_net_weight)
    if gross_weight is None or gross_weight <= 0 or net_weight is None or net_weight <= 0:
        return None
    if net_weight > gross_weight + 0.01:
        raise RuntimeError(
            "manifest_weight_conflict: "
            f"清单总净重 {net_weight:.4f} kg 大于总毛重 {gross_weight:.4f} kg"
        )
    return min(1.0, net_weight / gross_weight)


def apply_manifest_net_weights(rows: list[dict[str, Any]], manifest: ManifestSummary) -> None:
    ratio = manifest_net_to_gross_ratio(manifest)
    if ratio is None:
        for row in rows:
            gross = to_float(row.get("毛重")) or 0.0
            ctns = to_float(row.get("箱数")) or 0.0
            row["净重"] = round(max(0.01, gross - ctns), 2)
            row["净重计算依据"] = "清单未提供完整净重，按毛重减箱数兼容计算"
            row["净毛重比例"] = round(row["净重"] / gross, 6) if gross else 0
        return

    gross_values = [max(0.01, to_float(row.get("毛重")) or 0.0) for row in rows]
    target_net_weight = round(sum(gross_values) * ratio, 2)
    if target_net_weight < len(rows) * 0.01:
        raise RuntimeError(
            f"原始清单净重/毛重比例 {ratio:.6f} 无法为每个输出行分配正净重"
        )
    net_values = scale_decimal(gross_values, target_net_weight, 2)
    for row, gross, net in zip(rows, gross_values, net_values):
        if net <= 0 or net > gross + 0.01:
            raise RuntimeError(
                f"按原始清单净重/毛重比例分配失败: 净重 {net} kg, 毛重 {gross} kg"
            )
        row["净重"] = net
        row["净重计算依据"] = f"原始清单总净重/总毛重={ratio:.6f}"
        row["净毛重比例"] = round(net / gross, 6)

    actual_net_weight = round(sum(to_float(row.get("净重")) or 0 for row in rows), 2)
    if abs(actual_net_weight - target_net_weight) > 0.01:
        raise RuntimeError(
            f"总净重未按原始清单比例闭合: 目标 {target_net_weight}, 当前 {actual_net_weight}"
        )


def build_output_rows(
    selected: list[ProductCandidate],
    manifest: ManifestSummary,
    bill: BillInfo,
    options: ProcessingOptions,
) -> list[dict[str, Any]]:
    if len(selected) != options.target_item_count:
        raise RuntimeError(f"输出行数必须等于 {options.target_item_count}")
    if not any(candidate_tax_rate(candidate) > 0 for candidate in selected):
        raise RuntimeError("合格候选综合税率均为 0，无法匹配期望税金")
    if manifest.total_real_weight <= 0:
        raise RuntimeError("清单未提取到有效重量，不能按清单重量生成复核草案")

    plausibility_ranges: dict[tuple[str, str, str], PlausibilityRange] = {}
    row_plans = build_plausible_row_plans(
        selected=selected,
        manifest=manifest,
        bill=bill,
        options=options,
        plausibility_ranges=plausibility_ranges,
    )

    rows: list[dict[str, Any]] = []
    for plan in row_plans:
        candidate = plan.candidate
        tax_rate = candidate.effective_tax_rate or candidate.base_tax_rate
        row = {
            "中文品名": candidate.zh,
            "英文品名": candidate.en or candidate.zh,
            "商品编码": hs_cell_value(candidate.hs),
            "材质": translate_material_to_english(candidate.material),
            "用途": translate_usage_to_english(candidate.usage),
            "箱数": plan.ctns,
            "数量": plan.qty,
            "单位": "PCS",
            "币制": "USD",
            "单价": plan.unit_price,
            "总价": plan.total_value,
            "净重": round(max(0.01, plan.gross_weight - plan.ctns), 2),
            "毛重": plan.gross_weight,
            "原产国": "CN",
            "来源": candidate.source,
            "来源文件": candidate.source_label,
            "基础税率": round(candidate.base_tax_rate, 6),
            "综合税率": round(tax_rate, 6),
            "加征税率": clean_text(candidate.tax_data.get("additional_tax_rate")),
            "预计税金": round(plan.total_value * tax_rate, 2),
            "爬虫匹配HS": candidate.tax_data.get("hs_code_us") or candidate.hs,
            "爬虫品名": candidate.tax_data.get("description_cn", ""),
            "认证信息": "; ".join(candidate.certification_texts),
            "合规需复核": candidate.compliance_review_required,
            "合规复核原因": candidate.compliance_review_reason,
            "重量规则来源": plan.plausibility.source or "默认规则",
            "约束提示": "; ".join(plan.warnings),
            "source_rows": candidate.source_rows,
            "单件重量": round(plan.gross_weight / plan.qty, 6) if plan.qty else 0,
            "每箱数量": round(plan.qty / plan.ctns, 6) if plan.ctns else 0,
            "单箱重量": round(plan.gross_weight / plan.ctns, 6) if plan.ctns else 0,
            "税金预算": plan.tax_budget,
            "参考单价": plan.price_reference,
            "数量推导": plan.quantity_basis,
        }
        for warning in candidate.row_warnings:
            append_row_warning(row, warning)
        update_row_tax_display(row)
        rows.append(row)

    apply_manifest_net_weights(rows, manifest)
    normalize_output_language_fields(rows)
    return rows


def build_output_rows_best_effort(
    selected: list[ProductCandidate],
    manifest: ManifestSummary,
    bill: BillInfo,
    options: ProcessingOptions,
    reason: str,
) -> list[dict[str, Any]]:
    if not selected:
        raise RuntimeError(f"无可输出候选，无法生成最接近方案: {reason}")
    target_ctns = output_package_total(bill, len(selected))
    target_weight = round(manifest.total_real_weight or sum(candidate.gross_weight or candidate.real_weight or 1 for candidate in selected), 2)
    ctn_values = [candidate.ctns or 1 for candidate in selected]
    ctns = scale_positive_integers(ctn_values, target_ctns)
    weight_values = [candidate.gross_weight or candidate.real_weight or 1 for candidate in selected]
    weights = scale_decimal(weight_values, target_weight, 2)
    rows: list[dict[str, Any]] = []
    for candidate, row_ctns, gross_weight in zip(selected, ctns, weights):
        qty = max(row_ctns, int(round(candidate.qty or row_ctns)))
        if qty % row_ctns != 0:
            qty = row_ctns * max(1, math.ceil(qty / row_ctns))
        unit_price = max(0.0001, round(candidate_reference_unit_price(candidate) or candidate.unit_price or DEFAULT_UNIT_PRICE_MIN, 4))
        tax_rate = candidate_tax_rate(candidate)
        row = {
            "中文品名": candidate.zh,
            "英文品名": candidate.en or candidate.zh,
            "商品编码": hs_cell_value(candidate.hs),
            "材质": translate_material_to_english(candidate.material),
            "用途": translate_usage_to_english(candidate.usage),
            "箱数": row_ctns,
            "数量": qty,
            "单位": "PCS",
            "币制": "USD",
            "单价": unit_price,
            "总价": round(unit_price * qty, 2),
            "净重": round(max(0.01, gross_weight - row_ctns), 2),
            "毛重": gross_weight,
            "原产国": "CN",
            "来源": candidate.source,
            "来源文件": candidate.source_label,
            "基础税率": round(candidate.base_tax_rate, 6),
            "综合税率": round(tax_rate, 6),
            "加征税率": clean_text(candidate.tax_data.get("additional_tax_rate")),
            "预计税金": round(unit_price * qty * tax_rate, 2),
            "爬虫匹配HS": candidate.tax_data.get("hs_code_us") or candidate.hs,
            "爬虫品名": candidate.tax_data.get("description_cn", ""),
            "认证信息": "; ".join(candidate.certification_texts),
            "合规需复核": candidate.compliance_review_required,
            "合规复核原因": candidate.compliance_review_reason,
            "重量规则来源": "best_effort",
            "约束提示": f"最接近方案，未完全满足原优化约束: {reason}",
            "source_rows": candidate.source_rows,
            "单件重量": round(gross_weight / qty, 6) if qty else 0,
            "每箱数量": round(qty / row_ctns, 6) if row_ctns else 0,
            "单箱重量": round(gross_weight / row_ctns, 6) if row_ctns else 0,
            "税金预算": 0,
            "参考单价": unit_price,
            "数量推导": "best_effort 原始数量/箱数比例，并闭合总箱数和总重量",
        }
        for warning in candidate.row_warnings:
            append_row_warning(row, warning)
        update_row_tax_display(row)
        rows.append(row)
    apply_manifest_net_weights(rows, manifest)
    normalize_output_language_fields(rows)
    return rows


async def generate_valid_output_rows_with_llm(
    llm: LLMClient,
    selected: list[ProductCandidate],
    manifest: ManifestSummary,
    bill: BillInfo,
    options: ProcessingOptions,
) -> tuple[list[dict[str, Any]], int, list[str]]:
    feedback = ""
    feedback_history: list[str] = []
    last_error = ""
    for attempt in range(1, 4):
        try:
            payload = await llm_generate_output_draft(llm, selected, manifest, bill, options, feedback)
            rows = normalize_llm_output_draft(payload, selected)
            validate_llm_output_rows(rows, selected, manifest, bill, options)
            return rows, attempt, feedback_history
        except RuntimeError as exc:
            last_error = str(exc)
            feedback_history.append(last_error)
            feedback = (
                "上一次草案未通过代码硬校验，请只修正数值和行分配后重新输出 JSON。"
                f"错误：{last_error}。不能放宽税率、认证、总重量、目标税金上浮 10% 上限、行数和合理范围。"
            )
    raise RuntimeError(f"LLM 草案连续不合格: {last_error}")


async def llm_generate_output_draft(
    llm: LLMClient,
    selected: list[ProductCandidate],
    manifest: ManifestSummary,
    bill: BillInfo,
    options: ProcessingOptions,
    feedback: str = "",
) -> dict[str, Any]:
    messages = [
        {
            "role": "system",
            "content": (
                "你是美国清关 Commercial Invoice & Packing List 生成专家。你负责让数量、箱数、毛重、单价看起来像真实清关表。"
                "代码会严格校验税率、认证、总重量、总税金、行数、单价/单重/每箱数量范围。"
                "只返回一个合法 JSON object。禁止 Markdown、代码块、解释文字、前后缀。"
                "回复的第一个字符必须是 {，最后一个字符必须是 }。"
            ),
        },
        {
            "role": "user",
            "content": build_output_draft_prompt(selected, manifest, bill, options, feedback),
        },
    ]
    return await llm.chat_json(messages, temperature=0.0, max_tokens=8192, json_mode=True)


def build_output_draft_prompt(
    selected: list[ProductCandidate],
    manifest: ManifestSummary,
    bill: BillInfo,
    options: ProcessingOptions,
    feedback: str = "",
) -> str:
    candidates = [candidate_to_llm_dict(candidate) for candidate in selected]
    if not bill.cartons or bill.cartons <= 0:
        raise RuntimeError("提单未识别到有效总箱数，不能生成与提单箱数对齐的输出")
    return (
        "请基于给定候选生成最终清关行草案。\n"
        "输出格式硬要求：只输出一个 JSON object，不要 ```json，不要说明文字，不要换成数组顶层。\n"
        "JSON 顶层必须是 {\"rows\":[...]}，rows 内每一行只能包含 candidate_index、箱数、数量、单价、毛重。\n"
        "硬要求：\n"
        f"1. 输出 rows 数量必须等于 {options.target_item_count}，且每个候选必须输出一行，不得新增/删除/改名/改 HS。\n"
        f"2. 毛重请按品类合理分配；代码会按 Excel 总重量 {manifest.total_real_weight} kg 等比例倒推并强制闭合。\n"
        "   净重不由模型生成，代码会优先继承原始清单的总净重/总毛重比例。\n"
        f"3. 总税金尽量贴近目标 {options.target_tax_amount} USD，最终不得高于 {target_tax_upper_bound(options.target_tax_amount)} USD（目标上浮 10%）；税金=总价*综合税率；每行税金只能等于 0 或不低于 {MIN_ROW_TAX_AMOUNT_USD} USD，且税金为 0 的行数最多 {MAX_ZERO_TAX_ROWS} 行。\n"
        f"4. 总箱数必须等于提单总箱数 {bill.cartons}，不得使用清单箱数替代。\n"
        "5. 每行数量必须大于等于箱数，且数量必须是箱数的整数倍；每箱数量必须落入 candidate.plausibility_range；单价优先参考价格证据/合理范围，但不得导致总税金超过上限。\n"
        "6. 不要让所有行数量相同，不要让所有行单件重量相同，不要给电器/机器类低到不合理的单价。\n"
        "7. 单价、数量、毛重、箱数要像真实装箱清单，优先使用候选原始参数或合理范围中位数；毛重最终以 Excel 总重量倒推为准。\n"
        f"{'修正反馈：' + feedback if feedback else ''}\n"
        "JSON格式示例：{\"rows\":[{\"candidate_index\":0,\"箱数\":1,\"数量\":1,\"单价\":1.0,\"毛重\":1.0}]}\n"
        f"提单品类：{json.dumps(bill.products, ensure_ascii=False, separators=(',', ':'))}\n"
        f"候选：{json.dumps(candidates, ensure_ascii=False, separators=(',', ':'))}"
    )


def candidate_to_llm_dict(candidate: ProductCandidate) -> dict[str, Any]:
    plausibility = candidate.plausibility_range
    return {
        "zh": candidate.zh,
        "en": candidate.en,
        "hs": candidate.hs,
        "material": candidate.material,
        "usage": candidate.usage,
        "source": candidate.source,
        "source_label": candidate.source_label,
        "query_name": candidate.query_name,
        "query_material": candidate.query_material,
        "tax_match_source": candidate.tax_match_source,
        "base_tax_rate": candidate.base_tax_rate,
        "effective_tax_rate": candidate_tax_rate(candidate),
        "original_ctns": candidate.ctns,
        "original_qty": candidate.qty,
        "original_unit_price": candidate.unit_price,
        "original_gross_weight": candidate.gross_weight,
        "plausibility_range": asdict(plausibility) if plausibility else None,
        "plausibility_confidence": candidate.plausibility_confidence,
        "plausibility_basis": candidate.plausibility_basis,
        "price_evidence": candidate.price_evidence,
        "llm_reason": candidate.llm_reason,
        "compliance_review_required": candidate.compliance_review_required,
        "compliance_review_reason": candidate.compliance_review_reason,
    }


def normalize_llm_output_draft(payload: dict[str, Any], selected: list[ProductCandidate]) -> list[dict[str, Any]]:
    raw_rows = payload.get("rows")
    if not isinstance(raw_rows, list) or not raw_rows:
        raise RuntimeError("LLM 草案未返回 rows")
    rows: list[dict[str, Any]] = []
    for idx, raw in enumerate(raw_rows, start=1):
        if not isinstance(raw, dict):
            raise RuntimeError(f"LLM 草案第 {idx} 行不是 object")
        candidate_index = int(to_float(raw.get("candidate_index")) if to_float(raw.get("candidate_index")) is not None else idx - 1)
        if candidate_index < 0 or candidate_index >= len(selected):
            raise RuntimeError(f"LLM 草案第 {idx} 行 candidate_index 越界: {candidate_index}")
        candidate = selected[candidate_index]
        qty = to_float(raw.get("数量") or raw.get("qty"))
        ctns = to_float(raw.get("箱数") or raw.get("ctns"))
        unit_price = to_float(raw.get("单价") or raw.get("unit_price"))
        gross_weight = to_float(raw.get("毛重") or raw.get("gross_weight"))
        if None in (qty, ctns, unit_price, gross_weight):
            raise RuntimeError(f"LLM 草案第 {idx} 行数值字段不完整")
        tax_rate = candidate_tax_rate(candidate)
        ctns_int = max(1, int(round(ctns)))
        qty_int = max(1, int(round(qty)))
        unit_price = round(unit_price, 4)
        total_value = round(qty_int * unit_price, 2)
        gross_weight = round(gross_weight, 2)
        row = {
            "中文品名": candidate.zh,
            "英文品名": candidate.en or candidate.zh,
            "商品编码": hs_cell_value(candidate.hs),
            "材质": translate_material_to_english(candidate.material),
            "用途": translate_usage_to_english(candidate.usage),
            "箱数": ctns_int,
            "数量": qty_int,
            "单位": "PCS",
            "币制": "USD",
            "单价": unit_price,
            "总价": total_value,
            "净重": round(max(0.01, gross_weight - ctns_int), 2),
            "毛重": gross_weight,
            "原产国": "CN",
            "来源": candidate.source,
            "来源文件": candidate.source_label,
            "基础税率": round(candidate.base_tax_rate, 6),
            "综合税率": round(tax_rate, 6),
            "加征税率": clean_text(candidate.tax_data.get("additional_tax_rate")),
            "预计税金": round(total_value * tax_rate, 2),
            "爬虫匹配HS": candidate.tax_data.get("hs_code_us") or candidate.hs,
            "爬虫品名": candidate.tax_data.get("description_cn", ""),
            "认证信息": "; ".join(candidate.certification_texts),
            "合规需复核": candidate.compliance_review_required,
            "合规复核原因": candidate.compliance_review_reason,
            "重量规则来源": (candidate.plausibility_range.source if candidate.plausibility_range else ""),
            "约束提示": "",
            "source_rows": candidate.source_rows,
            "candidate_index": candidate_index,
            "LLM草案毛重": gross_weight,
            "单件重量": round(gross_weight / qty_int, 6) if qty_int else 0,
            "每箱数量": round(qty_int / ctns_int, 6) if ctns_int else 0,
            "单箱重量": round(gross_weight / ctns_int, 6) if ctns_int else 0,
            "合理性来源": (candidate.plausibility_range.source if candidate.plausibility_range else ""),
            "税率来源": candidate.tax_match_source,
            "查询词": candidate.query_name or candidate.zh,
            "价格依据": candidate.price_evidence.get("basis", ""),
            "价格置信度": candidate.price_evidence.get("confidence", ""),
            "零售参考单价": candidate.price_evidence.get("retail_unit_price", ""),
        }
        update_row_tax_display(row)
        rows.append(row)
    return rows


def validate_llm_output_rows(
    rows: list[dict[str, Any]],
    selected: list[ProductCandidate],
    manifest: ManifestSummary,
    bill: BillInfo,
    options: ProcessingOptions,
) -> None:
    validate_output_rows(rows)
    if len(rows) != options.target_item_count:
        raise RuntimeError(f"LLM 草案行数 {len(rows)} 不等于目标 {options.target_item_count}")
    seen_indexes = sorted(int(row.get("candidate_index")) for row in rows)
    if seen_indexes != list(range(len(selected))):
        raise RuntimeError("LLM 草案必须一行对应一个候选，不能重复或遗漏候选")
    ensure_bill_products_present(rows, bill.products, selected)

    output_package_total(bill, len(rows))
    reconcile_llm_rows_ctns(rows, bill.cartons)
    reconcile_row_quantities_to_cartons(rows, selected)
    close_llm_rows_gross_weight(rows, manifest.total_real_weight)
    apply_manifest_net_weights(rows, manifest)
    close_llm_rows_tax_gap(rows, selected, options.target_tax_amount)
    validate_qty_ctn_relationship(rows)
    validate_row_counts(rows, bill.cartons)
    tax_total = round(sum((to_float(row.get("总价")) or 0) * (to_float(row.get("综合税率")) or 0) for row in rows), 2)
    tax_upper_bound = target_tax_upper_bound(options.target_tax_amount)
    if tax_total > tax_upper_bound + 0.01:
        feasible = estimate_tax_feasible_range(selected)
        raise RuntimeError(
            "税金超过目标上浮 10% 上限；"
            f"目标 {options.target_tax_amount}, 最高 {tax_upper_bound}, 当前 {tax_total}, "
            f"可行税金区间约 {feasible[0]}-{feasible[1]}"
        )
    validate_row_plausibility(rows, selected, bill.products)
    validate_distribution_realism(rows)
    normalize_output_language_fields(rows)
    for row in rows:
        update_row_tax_display(row)


def reconcile_llm_rows_ctns(rows: list[dict[str, Any]], bill_cartons: Optional[float]) -> None:
    target_ctns = round(float(bill_cartons), 2) if bill_cartons and bill_cartons > 0 else 0.0
    if target_ctns <= 0:
        raise RuntimeError("提单未识别到有效总箱数，不能生成与提单箱数对齐的输出")

    current_values = [max(0.01, to_float(row.get("箱数")) or 0.0) for row in rows]
    scaled = scale_integer(current_values, target_ctns)
    for row, old_ctns, ctns in zip(rows, current_values, scaled):
        row["LLM草案箱数"] = row.get("LLM草案箱数", old_ctns)
        row["箱数"] = ctns
        row["箱数闭合调整"] = round(ctns - old_ctns, 2)

    ctn_total = round(sum(to_float(row.get("箱数")) or 0 for row in rows), 2)
    if abs(ctn_total - target_ctns) > 0.01:
        raise RuntimeError(f"总箱数未闭合: 目标 {target_ctns}, 当前 {ctn_total}")


def reconcile_row_quantities_to_cartons(rows: list[dict[str, Any]], selected: list[ProductCandidate]) -> None:
    for row, candidate in zip(rows, selected):
        ctns = max(1, int(round(to_float(row.get("箱数")) or 1)))
        draft_qty = max(1, int(round(to_float(row.get("数量")) or to_float(candidate.qty) or ctns)))
        desired_per_ctn = max(1, int(math.ceil(draft_qty / ctns)))
        plausibility = candidate.plausibility_range
        if plausibility and plausibility.qty_per_ctn_min is not None:
            desired_per_ctn = max(desired_per_ctn, int(math.ceil(plausibility.qty_per_ctn_min)))
        if plausibility and plausibility.qty_per_ctn_max is not None:
            max_per_ctn = max(1, int(math.floor(plausibility.qty_per_ctn_max)))
            desired_per_ctn = min(desired_per_ctn, max_per_ctn)
        qty = max(ctns, ctns * max(1, desired_per_ctn))
        row["LLM草案数量"] = row.get("LLM草案数量", draft_qty)
        row["数量"] = qty
        row["每箱数量"] = round(qty / ctns, 6)


def validate_qty_ctn_relationship(rows: list[dict[str, Any]]) -> None:
    for idx, row in enumerate(rows, start=1):
        ctns_float = to_float(row.get("箱数"))
        qty_float = to_float(row.get("数量"))
        if ctns_float is None or qty_float is None:
            raise RuntimeError(f"第 {idx} 行箱数/数量缺失")
        ctns = int(round(ctns_float))
        qty = int(round(qty_float))
        if abs(ctns_float - ctns) > 0.0001 or ctns <= 0:
            raise RuntimeError(f"第 {idx} 行箱数必须为正整数: {row.get('箱数')}")
        if abs(qty_float - qty) > 0.0001 or qty <= 0:
            raise RuntimeError(f"第 {idx} 行数量必须为正整数: {row.get('数量')}")
        if qty < ctns:
            raise RuntimeError(f"第 {idx} 行数量不能小于箱数: 数量 {qty}, 箱数 {ctns}")
        if qty % ctns != 0:
            raise RuntimeError(f"第 {idx} 行数量必须是箱数的整数倍: 数量 {qty}, 箱数 {ctns}")


def validate_candidate_tax_rates(selected: list[ProductCandidate]) -> None:
    for idx, candidate in enumerate(selected, start=1):
        rate = candidate_tax_rate(candidate)
        if rate >= BASE_TAX_LIMIT:
            raise RuntimeError(
                f"第 {idx} 行综合税率 {format_rate(rate)} 不小于 20%，禁止输出: "
                f"{candidate.zh} / {candidate.hs}"
            )


def validate_row_counts(rows: list[dict[str, Any]], target_ctns: Optional[float]) -> None:
    if not target_ctns or target_ctns <= 0:
        return
    current = round(sum(to_float(row.get("箱数")) or 0 for row in rows), 2)
    if abs(current - round(float(target_ctns), 2)) > 0.01:
        raise RuntimeError(f"总箱数未闭合: 目标 {round(float(target_ctns), 2)}, 当前 {current}")


def close_llm_rows_gross_weight(rows: list[dict[str, Any]], target_gross: float) -> None:
    current_values = [max(0.01, to_float(row.get("毛重")) or 0.0) for row in rows]
    scaled = scale_decimal(current_values, target_gross, 2)
    for row, old_gross, gross in zip(rows, current_values, scaled):
        qty = to_float(row.get("数量")) or 1
        ctns = to_float(row.get("箱数")) or 1
        row["LLM草案毛重"] = row.get("LLM草案毛重", old_gross)
        row["毛重"] = gross
        row["净重"] = round(max(0.01, gross - ctns), 2)
        row["单件重量"] = round(gross / qty, 6) if qty else 0
        row["单箱重量"] = round(gross / ctns, 6) if ctns else 0
        row["毛重闭合调整"] = round(gross - old_gross, 2)
    gross_total = round(sum(to_float(row.get("毛重")) or 0 for row in rows), 2)
    if abs(gross_total - target_gross) > 0.01:
        raise RuntimeError(f"总毛重未闭合: 目标 {target_gross}, 当前 {gross_total}")


def close_llm_rows_tax_gap(rows: list[dict[str, Any]], selected: list[ProductCandidate], target_tax_amount: float) -> None:
    current_tax = round(sum((to_float(row.get("总价")) or 0.0) * candidate_tax_rate(candidate) for row, candidate in zip(rows, selected)), 2)
    if current_tax <= target_tax_upper_bound(target_tax_amount) + 0.01:
        for row in rows:
            update_row_tax_display(row)
        return
    unit_prices: list[float] = []
    quantities: list[int] = []
    mins: list[float] = []
    maxes: list[float] = []
    for row, candidate in zip(rows, selected):
        qty = max(1, int(round(to_float(row.get("数量")) or to_float(candidate.qty) or 1)))
        unit_price = max(0.0001, to_float(row.get("单价")) or 0.0)
        plausibility = candidate.plausibility_range
        min_price = plausibility.unit_price_min if plausibility and plausibility.unit_price_min is not None else DEFAULT_UNIT_PRICE_MIN
        max_price = plausibility.unit_price_max if plausibility and plausibility.unit_price_max is not None else DEFAULT_UNIT_PRICE_MAX
        unit_prices.append(unit_price)
        quantities.append(qty)
        mins.append(max(0.0001, min_price))
        maxes.append(max(mins[-1], max_price))

    adjust_price_gap(unit_prices, mins, maxes, selected, quantities, target_tax_amount, 0.0)
    relax_price_floors_if_tax_requires(
        unit_prices,
        mins,
        maxes,
        selected,
        quantities,
        target_tax_amount,
        0.0,
    )

    for row, unit_price, qty in zip(rows, unit_prices, quantities):
        row["单价"] = round(unit_price, 4)
        row["总价"] = round(unit_price * qty, 2)
        update_row_tax_display(row)


def validate_row_plausibility(
    rows: list[dict[str, Any]],
    selected: list[ProductCandidate],
    bill_products: Optional[list[str]] = None,
) -> None:
    for idx, (row, candidate) in enumerate(zip(rows, selected), start=1):
        if normalize_hs(row.get("商品编码")) != normalize_hs(candidate.hs):
            raise RuntimeError(f"第 {idx} 行 HS 被 LLM 改写，禁止: {row.get('商品编码')} != {candidate.hs}")
        if not candidate.tax_match_source:
            raise RuntimeError(f"第 {idx} 行缺少 codeflagai 税率来源: {candidate.zh}")
        plausibility = candidate.plausibility_range
        if not plausibility:
            raise RuntimeError(f"第 {idx} 行缺少合理范围: {candidate.zh}")
        qty = to_float(row.get("数量")) or 0
        ctns = to_float(row.get("箱数")) or 0
        gross = to_float(row.get("毛重")) or 0
        unit_price = to_float(row.get("单价")) or 0
        qty_per_ctn_min = max(1.0, plausibility.qty_per_ctn_min or 1.0)
        qty_per_ctn_max = max(qty_per_ctn_min, plausibility.qty_per_ctn_max or qty_per_ctn_min)
        weight_hard_limit = not row_has_bill_weight_basis(row, candidate, bill_products or [])
        checks = [
            ("单价", unit_price, plausibility.unit_price_min, plausibility.unit_price_max, False),
            ("单件重量", gross / qty if qty else 0, plausibility.kg_per_pc_min, plausibility.kg_per_pc_max, weight_hard_limit),
            ("每箱数量", qty / ctns if ctns else 0, qty_per_ctn_min, qty_per_ctn_max, True),
            ("单箱重量", gross / ctns if ctns else 0, plausibility.kg_per_ctn_min, plausibility.kg_per_ctn_max, weight_hard_limit),
        ]
        warnings: list[str] = []
        for label, value, low, high, hard_limit in checks:
            if low is not None and value < low - 0.0001:
                message = f"{label}低于合理下限 {round(value, 4)} < {round(low, 4)}"
                if hard_limit:
                    raise RuntimeError(f"第 {idx} 行 {message} ({candidate.zh})")
                warnings.append(message)
                continue
            if high is not None and value > high + 0.0001:
                message = f"{label}高于合理上限 {round(value, 4)} > {round(high, 4)}"
                if hard_limit:
                    raise RuntimeError(f"第 {idx} 行 {message} ({candidate.zh})")
                warnings.append(message)
                continue
            if near_bound(value, low, high):
                warnings.append(f"{label}接近边界 {round(value, 4)}")
        if abs(to_float(row.get("毛重闭合调整")) or 0) > 0.01:
            warnings.append(f"毛重按Excel总重量倒推调整 {row.get('毛重闭合调整')}kg")
        row["约束提示"] = "; ".join(warnings)


def row_has_bill_weight_basis(row: dict[str, Any], candidate: ProductCandidate, bill_products: list[str]) -> bool:
    if bill_products and row_matches_bill_product(row, bill_products):
        return True
    if candidate.source in {"bill", "bill_product", "llm_query"}:
        return True
    return candidate.tax_match_source in {"bill_hs", "bill_product", "llm_query"}


def validate_distribution_realism(rows: list[dict[str, Any]]) -> None:
    if len(rows) < 4:
        return
    quantities = [to_float(row.get("数量")) or 0 for row in rows]
    kg_per_pc = [round((to_float(row.get("毛重")) or 0) / (to_float(row.get("数量")) or 1), 4) for row in rows]
    if len(set(quantities)) <= 2 and len(rows) >= 8:
        raise RuntimeError("数量分布过于机械，多个品类数量几乎相同")
    if len(set(kg_per_pc)) <= 2 and len(rows) >= 8:
        raise RuntimeError("单件重量分布过于机械，多个品类单体重量几乎相同")


def validate_minimum_row_tax(rows: list[dict[str, Any]]) -> None:
    zero_tax_rows = 0
    for idx, row in enumerate(rows, start=1):
        tax_amount = to_float(row.get(DISPLAY_TAX_AMOUNT_FIELD))
        if tax_amount is None:
            rate = to_float(row.get("综合税率")) or 0.0
            total_value = to_float(row.get("总价")) or 0.0
            tax_amount = round(total_value * rate, 2)
        if abs(tax_amount) <= 0.005:
            zero_tax_rows += 1
            continue
        if tax_amount + 0.005 < MIN_ROW_TAX_AMOUNT_USD:
            raise RuntimeError(
                f"第 {idx} 行税金必须等于 0 或不低于 {MIN_ROW_TAX_AMOUNT_USD} USD: "
                f"{row.get('中文品名')} / {row.get('商品编码')} = {tax_amount}"
            )
    if zero_tax_rows > MAX_ZERO_TAX_ROWS:
        raise RuntimeError(f"税金为 0 的行数不能超过 {MAX_ZERO_TAX_ROWS} 行，当前 {zero_tax_rows} 行")


def estimate_tax_feasible_range(selected: list[ProductCandidate]) -> tuple[float, float]:
    min_tax = 0.0
    max_tax = 0.0
    for candidate in selected:
        plausibility = candidate.plausibility_range
        rate = candidate_tax_rate(candidate)
        if not plausibility or rate <= 0:
            continue
        min_qty = max(1, math.ceil((candidate.ctns or 1) * (plausibility.qty_per_ctn_min or 1)))
        max_qty = max(min_qty, math.ceil((candidate.ctns or 1) * (plausibility.qty_per_ctn_max or min_qty)))
        min_tax += max(MIN_ROW_TAX_AMOUNT_USD, min_qty * (plausibility.unit_price_min or 0) * rate)
        max_tax += max_qty * (plausibility.unit_price_max or 0) * rate
    return round(min_tax, 2), round(max_tax, 2)


def build_plausible_row_plans(
    *,
    selected: list[ProductCandidate],
    manifest: ManifestSummary,
    bill: BillInfo,
    options: ProcessingOptions,
    plausibility_ranges: dict[tuple[str, str, str], PlausibilityRange],
) -> list[RowPlan]:
    target_gross = round(manifest.total_real_weight, 2)
    target_ctns = output_package_total(bill, options.target_item_count)
    ranges = [resolve_plausibility_range(candidate, plausibility_ranges) for candidate in selected]
    ctns = allocate_price_first_cartons(selected, target_ctns)
    weights = allocate_plausible_weights(selected, ctns, ranges, target_gross)
    price_references = [
        price_reference_for_candidate(candidate, plausibility)
        for candidate, plausibility in zip(selected, ranges)
    ]
    tax_budgets = allocate_row_tax_budgets(
        selected,
        options.target_tax_amount,
        ctns=ctns,
        weights=weights,
        ranges=ranges,
        price_references=price_references,
    )
    quantities = [
        choose_plausible_quantity(
            candidate,
            row_ctns,
            gross,
            plausibility,
            tax_budget=tax_budget,
            price_reference=price_reference,
        )
        for candidate, row_ctns, gross, plausibility, tax_budget, price_reference in zip(
            selected,
            ctns,
            weights,
            ranges,
            tax_budgets,
            price_references,
        )
    ]
    prices = allocate_plausible_prices(selected, quantities, weights, ranges, options.target_tax_amount, tax_budgets)

    plans: list[RowPlan] = []
    for candidate, row_ctns, qty, gross, price, plausibility, tax_budget, price_reference in zip(
        selected,
        ctns,
        quantities,
        weights,
        prices,
        ranges,
        tax_budgets,
        price_references,
    ):
        unit_price, total_value = price
        warnings = tuple(build_constraint_warnings(row_ctns, qty, gross, unit_price, plausibility))
        plans.append(
            RowPlan(
                candidate=candidate,
                ctns=row_ctns,
                qty=qty,
                gross_weight=gross,
                total_value=total_value,
                unit_price=unit_price,
                plausibility=plausibility,
                warnings=warnings,
                tax_budget=round(tax_budget, 2),
                price_reference=round(price_reference, 4),
                quantity_basis="税金预算/参考单价反推，并受重量、箱数、每箱数量约束",
            )
        )
    return plans


def optimize_selected_candidates_for_price_fit(
    selected: list[ProductCandidate],
    repair_pool: list[ProductCandidate],
    manifest: ManifestSummary,
    bill: BillInfo,
    options: ProcessingOptions,
    *,
    query_cache: Optional[QueryCache] = None,
) -> tuple[list[ProductCandidate], dict[str, Any]]:
    if not repair_pool:
        return selected, {"swaps": 0, "message": "no repair pool"}

    try:
        best_plans = build_plausible_row_plans(
            selected=selected,
            manifest=manifest,
            bill=bill,
            options=options,
            plausibility_ranges={},
        )
    except RuntimeError as exc:
        return selected, {"swaps": 0, "message": f"initial plan failed: {exc}"}

    best_selected = list(selected)
    best_score = price_fit_score_for_plans(best_plans)
    original_score = best_score
    swaps: list[dict[str, Any]] = []
    priced_repair_candidates: dict[tuple[str, str, str], ProductCandidate] = {}
    web_lookups = 0

    for _ in range(PRICE_REPAIR_MAX_PASSES):
        bad_indexes = [
            idx
            for idx, plan in enumerate(best_plans)
            if row_plan_price_fit_ratio(plan) is not None
            and (row_plan_price_fit_ratio(plan) or 0.0) < PRICE_FIT_MIN_REFERENCE_RATIO
            and not is_bill_required_candidate(plan.candidate)
        ]
        bad_indexes.sort(key=lambda idx: row_plan_price_fit_ratio(best_plans[idx]) or 0.0)
        if not bad_indexes:
            break

        improved = False
        current_keys = {candidate_identity(candidate) for candidate in best_selected}
        for bad_idx in bad_indexes:
            for raw_candidate in sorted(repair_pool, key=price_repair_candidate_order):
                replacement_key = candidate_identity(raw_candidate)
                if replacement_key in current_keys:
                    continue
                replacement_candidate = priced_repair_candidates.get(replacement_key)
                if replacement_candidate is None:
                    evidence_price = to_float(raw_candidate.price_evidence.get("declared_unit_price"))
                    if evidence_price and evidence_price > 0:
                        replacement_candidate = raw_candidate
                    elif web_lookups < PRICE_REPAIR_MAX_WEB_LOOKUPS:
                        replacement_candidate = attach_price_evidence_to_candidates(
                            [raw_candidate],
                            query_cache=query_cache,
                            price_timeout=PRICE_REPAIR_SEARCH_TIMEOUT_SECONDS,
                            price_max_pages=PRICE_REPAIR_SEARCH_MAX_PAGES,
                        )[0]
                        web_lookups += 1
                    else:
                        continue
                    priced_repair_candidates[replacement_key] = replacement_candidate
                trial_selected = list(best_selected)
                old_candidate = trial_selected[bad_idx]
                trial_selected[bad_idx] = replacement_candidate
                try:
                    trial_plans = build_plausible_row_plans(
                        selected=trial_selected,
                        manifest=manifest,
                        bill=bill,
                        options=options,
                        plausibility_ranges={},
                    )
                except RuntimeError:
                    continue
                trial_score = price_fit_score_for_plans(trial_plans)
                score_gain = best_score - trial_score
                if score_gain < max(PRICE_REPAIR_MIN_SCORE_GAIN, best_score * 0.05):
                    continue
                best_selected = trial_selected
                best_plans = trial_plans
                best_score = trial_score
                swaps.append(
                    {
                        "from": old_candidate.zh or old_candidate.en,
                        "to": replacement_candidate.zh or replacement_candidate.en,
                        "score_gain": round(score_gain, 2),
                    }
                )
                improved = True
                break
            if improved:
                break
        if not improved:
            break

    return best_selected, {
        "swaps": len(swaps),
        "initial_score": round(original_score, 2),
        "final_score": round(best_score, 2),
        "web_lookups": web_lookups,
        "details": swaps,
    }


def optimize_selected_candidates_for_manual_invoice(
    selected: list[ProductCandidate],
    candidate_pool: list[ProductCandidate],
    manifest: ManifestSummary,
    bill: BillInfo,
    options: ProcessingOptions,
) -> tuple[list[ProductCandidate], dict[str, Any]]:
    required = [candidate for candidate in selected if is_bill_required_candidate(candidate)]
    required_keys = {candidate_identity(candidate) for candidate in required}
    pool = dedupe_candidates([*required, *candidate_pool, *selected])
    optional = [candidate for candidate in pool if candidate_identity(candidate) not in required_keys]
    ranked_optional = sorted(optional, key=manual_invoice_candidate_order)
    if len(required) > options.target_item_count:
        return selected, {"swaps": 0, "message": "required bill candidates exceed target"}

    target_optional_count = options.target_item_count - len(required)
    best_selected = list(selected)
    best_score = float("inf")
    best_plans: list[RowPlan] = []
    evaluated = 0
    search_pool = manual_invoice_search_pool(ranked_optional, target_optional_count)

    for optional_group in itertools.combinations(search_pool, target_optional_count):
        if evaluated >= 5000:
            break
        trial = [*required, *optional_group]
        if not manual_invoice_trial_allowed(trial):
            continue
        try:
            plans = build_plausible_row_plans(
                selected=trial,
                manifest=manifest,
                bill=bill,
                options=options,
                plausibility_ranges={},
            )
        except RuntimeError:
            continue
        evaluated += 1
        score = manual_invoice_plan_score(plans, options)
        if score < best_score:
            best_score = score
            best_selected = trial
            best_plans = plans

    initial_keys = {candidate_identity(candidate) for candidate in selected}
    final_keys = {candidate_identity(candidate) for candidate in best_selected}
    swaps = len([key for key in final_keys if key not in initial_keys])
    estimated_tax = round(sum(plan.total_value * candidate_tax_rate(plan.candidate) for plan in best_plans), 2) if best_plans else None
    zero_tax_rows = sum(1 for candidate in best_selected if candidate_tax_rate(candidate) <= 0)
    return best_selected, {
        "swaps": swaps,
        "strategy": "manual_invoice",
        "evaluated": evaluated,
        "score": round(best_score, 2) if best_score != float("inf") else None,
        "estimated_tax": estimated_tax,
        "zero_tax_rows": zero_tax_rows,
        "pool_size": len(candidate_pool),
        "search_pool_size": len(search_pool),
    }


def build_manual_invoice_candidate_pool(
    selected: list[ProductCandidate],
    qualified_manifest: list[ProductCandidate],
    candidate_library: list[ProductCandidate],
    *,
    allow_candidate_library: bool,
) -> list[ProductCandidate]:
    candidates = [*selected, *qualified_manifest]
    if allow_candidate_library:
        candidates.extend(candidate_library)
    return dedupe_candidates(candidates)


def manual_invoice_search_pool(candidates: list[ProductCandidate], target_optional_count: int) -> list[ProductCandidate]:
    clean_candidates = [candidate for candidate in candidates if not candidate.compliance_review_required]
    if len(clean_candidates) >= target_optional_count:
        candidates = clean_candidates
    manual_refs = [candidate for candidate in candidates if candidate.source_label == "DEFAULT_REFERENCE_STYLE_ROWS"]
    anchors = [candidate for candidate in candidates if normalize_hs(candidate.hs) in MANUAL_INVOICE_ANCHOR_HS]
    low_tax = [candidate for candidate in candidates if candidate_tax_rate(candidate) <= 0.153]
    high_score = candidates[: max(target_optional_count * 2, 18)]
    return dedupe_candidates([*manual_refs[:18], *anchors[:14], *low_tax[:18], *high_score])[:32]


def manual_invoice_trial_allowed(candidates: list[ProductCandidate]) -> bool:
    if candidates_have_duplicate_names(candidates):
        return False
    anchor_hs_counts: dict[str, int] = {}
    for candidate in candidates:
        hs = normalize_hs(candidate.hs)
        if hs not in MANUAL_INVOICE_ANCHOR_HS:
            continue
        anchor_hs_counts[hs] = anchor_hs_counts.get(hs, 0) + 1
    return all(count <= 1 for count in anchor_hs_counts.values())


def manual_invoice_candidate_order(candidate: ProductCandidate) -> tuple[int, int, float, float, float, float]:
    rate = candidate_tax_rate(candidate)
    source_rank = 0 if candidate.source == "manifest_group" else 1
    if candidate.source_label == "DEFAULT_REFERENCE_STYLE_ROWS":
        source_rank -= 3
    if normalize_hs(candidate.hs) in MANUAL_INVOICE_ANCHOR_HS:
        source_rank -= 2
    if rate <= 0:
        source_rank -= 2
    weight = max(candidate.gross_weight or candidate.real_weight or 0.0, 0.1)
    ctns = max(candidate.ctns or 0.0, 1.0)
    kg_per_ctn = weight / ctns
    return (
        1 if candidate.compliance_review_required else 0,
        source_rank,
        rate,
        -(candidate.declared_value or 0.0),
        -kg_per_ctn,
        -candidate.score,
    )


def manual_invoice_plan_score(plans: list[RowPlan], options: ProcessingOptions) -> float:
    tax_total = sum(plan.total_value * candidate_tax_rate(plan.candidate) for plan in plans)
    tax_upper = target_tax_upper_bound(options.target_tax_amount)
    score = 0.0
    if tax_total > tax_upper:
        score += (tax_total - tax_upper) * 10000.0
    else:
        score += max(0.0, options.target_tax_amount - tax_total) * 0.5
    zero_value = sum(plan.total_value for plan in plans if candidate_tax_rate(plan.candidate) <= 0)
    total_value = sum(plan.total_value for plan in plans) or 1.0
    zero_tax_rows = sum(1 for plan in plans if candidate_tax_rate(plan.candidate) <= 0)
    anchor_rows = sum(1 for plan in plans if normalize_hs(plan.candidate.hs) in MANUAL_INVOICE_ANCHOR_HS)
    manual_reference_rows = sum(1 for plan in plans if plan.candidate.source_label == "DEFAULT_REFERENCE_STYLE_ROWS")
    duplicate_hs_penalty = duplicate_count(normalize_hs(plan.candidate.hs) for plan in plans)
    duplicate_name_penalty = duplicate_count(normalize_text(plan.candidate.zh or plan.candidate.en) for plan in plans)
    score -= min(0.75, zero_value / total_value) * 900.0
    score -= zero_tax_rows * 120.0
    score -= anchor_rows * 35.0
    score -= manual_reference_rows * 45.0
    score += duplicate_hs_penalty * 220.0
    score += duplicate_name_penalty * 280.0
    replacement_count = sum(1 for plan in plans if plan.candidate.source == "replacement")
    score += replacement_count * 8.0
    review_required_count = sum(1 for plan in plans if plan.candidate.compliance_review_required)
    score += review_required_count * 1000000.0
    warning_count = sum(len(plan.warnings) for plan in plans)
    score += warning_count * 1.5
    return round(score, 4)


def duplicate_count(values: Any) -> int:
    counts: dict[str, int] = {}
    for value in values:
        key = clean_text(value)
        if not key:
            continue
        counts[key] = counts.get(key, 0) + 1
    return sum(max(0, count - 1) for count in counts.values())


def price_repair_candidate_order(candidate: ProductCandidate) -> tuple[float, float, float]:
    rate = candidate_tax_rate(candidate)
    reference_price = candidate_reference_unit_price(candidate)
    price_tax_pressure = max(0.0, reference_price) * max(0.0, rate)
    return (price_tax_pressure, -candidate.score, -reference_price)


def price_fit_score_for_plans(plans: list[RowPlan]) -> float:
    score = 0.0
    for plan in plans:
        ratio = row_plan_price_fit_ratio(plan)
        if ratio is None:
            continue
        if ratio <= 0:
            score += 10000.0
        elif ratio < PRICE_FIT_MIN_REFERENCE_RATIO:
            row_tax = plan.total_value * candidate_tax_rate(plan.candidate)
            score += ((PRICE_FIT_MIN_REFERENCE_RATIO / ratio) - 1.0) * max(1.0, row_tax)
        min_price = plan.plausibility.unit_price_min
        if min_price and min_price > 0 and plan.unit_price < min_price:
            row_tax = plan.total_value * candidate_tax_rate(plan.candidate)
            severity = (min_price / max(plan.unit_price, 0.0001)) - 1.0
            score += severity * max(1000.0, row_tax * 10.0)
    return round(score, 4)


def row_plan_price_fit_ratio(plan: RowPlan) -> Optional[float]:
    if plan.price_reference <= 0 or plan.unit_price <= 0 or candidate_tax_rate(plan.candidate) <= 0:
        return None
    return plan.unit_price / plan.price_reference


def resolve_plausibility_range(
    candidate: ProductCandidate,
    ranges: dict[tuple[str, str, str], PlausibilityRange],
) -> PlausibilityRange:
    if candidate.plausibility_range:
        return with_default_plausibility_bounds(candidate.plausibility_range)

    direct = lookup_plausibility_range(candidate, ranges)
    if direct:
        return with_default_plausibility_bounds(direct)
    derived = derive_candidate_plausibility_range(candidate)
    if not plausibility_range_is_complete(derived):
        raise RuntimeError(f"候选缺少完整合理范围，不能生成: {candidate.zh}/{candidate.en}")
    return with_default_plausibility_bounds(derived)


def with_default_plausibility_bounds(value: PlausibilityRange) -> PlausibilityRange:
    if not plausibility_range_is_complete(value):
        raise RuntimeError(f"合理范围不完整，不能使用默认规则兜底: {value.source or 'unknown'}")
    return PlausibilityRange(
        kg_per_ctn_min=value.kg_per_ctn_min,
        kg_per_ctn_max=value.kg_per_ctn_max,
        kg_per_pc_min=value.kg_per_pc_min,
        kg_per_pc_max=value.kg_per_pc_max,
        unit_price_min=value.unit_price_min,
        unit_price_max=value.unit_price_max,
        ctns_min=value.ctns_min,
        ctns_max=value.ctns_max,
        qty_per_ctn_min=value.qty_per_ctn_min,
        qty_per_ctn_max=value.qty_per_ctn_max,
        source=value.source,
    )


def allocate_price_first_cartons(selected: list[ProductCandidate], target_ctns: float) -> list[int]:
    target = int(round(target_ctns))
    if target < len(selected):
        raise RuntimeError(
            f"整票包装数量 {target} 小于输出行数 {len(selected)}，无法逐行分配并闭合总数"
        )
    values = [candidate.ctns or 1 for candidate in selected]
    detailed_bill_indexes = [
        idx
        for idx, candidate in enumerate(selected)
        if is_bill_required_candidate(candidate) and candidate.bill_has_manifest_detail
    ]
    undetailed_bill_indexes = [
        idx
        for idx, candidate in enumerate(selected)
        if is_undetailed_bill_candidate(candidate)
    ]

    locked_detailed: dict[int, int] = {}
    unlocked_indexes = [idx for idx in range(len(selected)) if idx not in set(detailed_bill_indexes)]
    detailed_total = sum(max(1, int(round(values[idx]))) for idx in detailed_bill_indexes)
    if detailed_bill_indexes and detailed_total + len(unlocked_indexes) <= target:
        locked_detailed = {
            idx: max(1, int(round(values[idx])))
            for idx in detailed_bill_indexes
        }
        unlocked_total = target - sum(locked_detailed.values())
        unlocked_ctns = scale_positive_integers(
            [values[idx] for idx in unlocked_indexes],
            unlocked_total,
        )
        initial = [0 for _ in selected]
        for idx, value in locked_detailed.items():
            initial[idx] = value
        for idx, value in zip(unlocked_indexes, unlocked_ctns):
            initial[idx] = value
    else:
        initial = scale_positive_integers(values, target)

    if not undetailed_bill_indexes or len(undetailed_bill_indexes) == len(selected):
        return initial

    bill_total_cap = max(
        len(undetailed_bill_indexes),
        int(round(target * UNDETAILED_BILL_CARTON_SHARE)),
    )
    bill_row_cap = max(1, math.ceil(bill_total_cap / len(undetailed_bill_indexes)))
    if sum(initial[idx] for idx in undetailed_bill_indexes) <= bill_total_cap:
        return initial

    result = [0 for _ in selected]
    used_bill = 0
    for idx in undetailed_bill_indexes:
        result[idx] = min(initial[idx], bill_row_cap)
        used_bill += result[idx]

    fixed_detailed_total = sum(locked_detailed.values())
    for idx, value in locked_detailed.items():
        result[idx] = value
    flexible_indexes = [
        idx
        for idx in range(len(selected))
        if idx not in set(undetailed_bill_indexes) and idx not in locked_detailed
    ]
    remaining = max(len(flexible_indexes), target - used_bill - fixed_detailed_total)
    flexible_values = [values[idx] for idx in flexible_indexes]
    flexible_ctns = scale_positive_integers(flexible_values, remaining)
    for idx, ctn_value in zip(flexible_indexes, flexible_ctns):
        result[idx] = ctn_value

    diff = target - sum(result)
    if diff:
        order = sorted(flexible_indexes, key=lambda idx: values[idx], reverse=True)
        if not order:
            order = [idx for idx in range(len(result)) if idx not in locked_detailed]
        if not order:
            order = list(range(len(result)))
        for step in range(abs(diff)):
            idx = order[step % len(order)]
            if diff > 0:
                result[idx] += 1
            elif result[idx] > 1:
                result[idx] -= 1
    return result


def allocate_row_tax_budgets(
    selected: list[ProductCandidate],
    target_tax_amount: float,
    *,
    ctns: Optional[list[int]] = None,
    weights: Optional[list[float]] = None,
    ranges: Optional[list[PlausibilityRange]] = None,
    price_references: Optional[list[float]] = None,
) -> list[float]:
    positive_indexes = [idx for idx, candidate in enumerate(selected) if candidate_tax_rate(candidate) > 0]
    budgets = [0.0 for _ in selected]
    if not positive_indexes:
        return budgets

    base = {idx: MIN_ROW_TAX_AMOUNT_USD for idx in positive_indexes}
    demand = build_reference_tax_demands(selected, ctns, weights, ranges, price_references, base)
    bill_indexes = [idx for idx in positive_indexes if is_undetailed_bill_candidate(selected[idx])]
    non_bill_indexes = [idx for idx in positive_indexes if idx not in set(bill_indexes)]

    if bill_indexes and non_bill_indexes:
        non_bill_min = sum(base[idx] for idx in non_bill_indexes)
        bill_min = sum(base[idx] for idx in bill_indexes)
        bill_total = min(target_tax_amount * UNDETAILED_BILL_TAX_SHARE, max(bill_min, target_tax_amount - non_bill_min))
        bill_total = max(bill_min, bill_total)
        bill_total = min(bill_total, max(bill_min, target_tax_amount - non_bill_min))
        distribute_tax_budget(budgets, bill_indexes, bill_total, base)
        distribute_tax_budget(budgets, non_bill_indexes, max(non_bill_min, target_tax_amount - bill_total), base, demand)
    else:
        distribute_tax_budget(budgets, positive_indexes, target_tax_amount, base, demand)
    return [round(value, 4) for value in budgets]


def build_reference_tax_demands(
    selected: list[ProductCandidate],
    ctns: Optional[list[int]],
    weights: Optional[list[float]],
    ranges: Optional[list[PlausibilityRange]],
    price_references: Optional[list[float]],
    base: dict[int, float],
) -> dict[int, float]:
    if not ctns or not weights or not ranges or not price_references:
        return {}
    demand: dict[int, float] = {}
    for idx, (candidate, row_ctns, gross, plausibility, reference_price) in enumerate(
        zip(selected, ctns, weights, ranges, price_references)
    ):
        rate = candidate_tax_rate(candidate)
        if rate <= 0 or reference_price <= 0:
            continue
        min_qty, _ = plausible_quantity_bounds(candidate, row_ctns, gross, plausibility)
        target_tax = min_qty * reference_price * rate * PRICE_FIT_MIN_REFERENCE_RATIO
        demand[idx] = max(base.get(idx, 0.0), target_tax)
    return demand


def distribute_tax_budget(
    budgets: list[float],
    indexes: list[int],
    total: float,
    base: dict[int, float],
    demand: Optional[dict[int, float]] = None,
) -> None:
    if not indexes:
        return
    base_total = sum(base.get(idx, 0.0) for idx in indexes)
    if total <= base_total:
        for idx in indexes:
            budgets[idx] = base.get(idx, 0.0)
        return
    extra = total - base_total
    demand = demand or {}
    demand_extra = {
        idx: max(0.0, demand.get(idx, base.get(idx, 0.0)) - base.get(idx, 0.0))
        for idx in indexes
    }
    demand_total = sum(demand_extra.values())
    if demand_total > 0:
        used = min(extra, demand_total)
        for idx in indexes:
            budgets[idx] = base.get(idx, 0.0) + used * demand_extra[idx] / demand_total
        extra -= used
    else:
        for idx in indexes:
            budgets[idx] = base.get(idx, 0.0)
    if extra > 0:
        for idx in indexes:
            budgets[idx] += extra / float(len(indexes))


def is_bill_required_candidate(candidate: ProductCandidate) -> bool:
    return (
        bool(normalize_bill_product_text(candidate.bill_product_name))
        or candidate.source in {"bill", "bill_product"}
        or candidate.tax_match_source in {"bill_hs", "bill_product"}
    )


def is_undetailed_bill_candidate(candidate: ProductCandidate) -> bool:
    return is_bill_required_candidate(candidate) and not candidate.bill_has_manifest_detail


def is_customer_codebook_candidate(candidate: ProductCandidate) -> bool:
    return candidate.tax_match_source == "customer_codebook" or candidate.source_label.startswith("客户编码库/")


def price_reference_for_candidate(candidate: ProductCandidate, plausibility: PlausibilityRange) -> float:
    reference = candidate_reference_unit_price(candidate)
    if reference > 0:
        return reference
    min_price = plausibility.unit_price_min or DEFAULT_UNIT_PRICE_MIN
    max_price = plausibility.unit_price_max or DEFAULT_UNIT_PRICE_MAX
    return max(DEFAULT_UNIT_PRICE_MIN, min(max_price, max(min_price, (min_price + max_price) / 2)))


def scale_positive_integers(values: list[float], target_total: float) -> list[int]:
    if not values:
        return []
    target = max(len(values), int(round(target_total)))
    scaled = scale_integer(values, target)
    result = [max(1, value) for value in scaled]
    diff = target - sum(result)
    if diff > 0:
        order = sorted(range(len(result)), key=lambda idx: values[idx], reverse=True)
        for step in range(diff):
            result[order[step % len(order)]] += 1
    elif diff < 0:
        order = sorted(range(len(result)), key=lambda idx: result[idx], reverse=True)
        remaining = -diff
        for idx in order:
            if remaining <= 0:
                break
            reducible = max(0, result[idx] - 1)
            take = min(reducible, remaining)
            result[idx] -= take
            remaining -= take
    return result


def allocate_plausible_weights(
    selected: list[ProductCandidate],
    ctns: list[int],
    ranges: list[PlausibilityRange],
    target_gross: float,
) -> list[float]:
    base_weights = [candidate.gross_weight or candidate.real_weight or 1 for candidate in selected]
    scaled = scale_decimal(base_weights, target_gross, 4)
    mins = [max(0.01, row_ctns * (plausibility.kg_per_ctn_min or DEFAULT_KG_PER_CTN_MIN)) for row_ctns, plausibility in zip(ctns, ranges)]
    maxes = [max(min_weight, row_ctns * (plausibility.kg_per_ctn_max or DEFAULT_KG_PER_CTN_MAX)) for row_ctns, min_weight, plausibility in zip(ctns, mins, ranges)]
    if sum(mins) - target_gross > 0.01:
        raise RuntimeError(
            "优化无解：目标清单重量低于所选类目的合理单箱重量下限；"
            f"清单重量 {target_gross} kg，最低需要 {round(sum(mins), 2)} kg"
        )
    if target_gross - sum(maxes) > 0.01:
        raise RuntimeError(
            "优化无解：目标清单重量超过所选类目的合理单箱重量上限；"
            f"清单重量 {target_gross} kg，最高可承载 {round(sum(maxes), 2)} kg"
        )

    weights = [min(max(value, min_weight), max_weight) for value, min_weight, max_weight in zip(scaled, mins, maxes)]
    adjust_bounded_values(weights, mins, maxes, target_gross, digits=4, label="重量")
    rounded = [round(value, 2) for value in weights]
    diff = round(target_gross - sum(rounded), 2)
    if rounded and diff:
        idx = max(range(len(rounded)), key=lambda item: maxes[item] - rounded[item] if diff > 0 else rounded[item] - mins[item])
        rounded[idx] = round(rounded[idx] + diff, 2)
    if abs(round(sum(rounded) - target_gross, 2)) > 0.01:
        raise RuntimeError("优化无解：输出重量无法精确闭合到清单重量")
    return rounded


def adjust_bounded_values(
    values: list[float],
    mins: list[float],
    maxes: list[float],
    target_total: float,
    *,
    digits: int,
    label: str,
) -> None:
    diff = round(target_total - sum(values), digits)
    if abs(diff) <= 10 ** (-digits):
        return
    if diff > 0:
        capacities = [max_value - value for value, max_value in zip(values, maxes)]
        capacity_total = sum(max(0.0, value) for value in capacities)
        if capacity_total + 10 ** (-digits) < diff:
            raise RuntimeError(f"优化无解：{label}上调空间不足")
        for idx, capacity in sorted(enumerate(capacities), key=lambda item: item[1], reverse=True):
            if diff <= 10 ** (-digits):
                break
            add = min(max(0.0, capacity), diff)
            values[idx] += add
            diff = round(diff - add, digits)
    else:
        need = abs(diff)
        capacities = [value - min_value for value, min_value in zip(values, mins)]
        capacity_total = sum(max(0.0, value) for value in capacities)
        if capacity_total + 10 ** (-digits) < need:
            raise RuntimeError(f"优化无解：{label}下调空间不足")
        for idx, capacity in sorted(enumerate(capacities), key=lambda item: item[1], reverse=True):
            if need <= 10 ** (-digits):
                break
            take = min(max(0.0, capacity), need)
            values[idx] -= take
            need = round(need - take, digits)


def choose_plausible_quantity(
    candidate: ProductCandidate,
    ctns: int,
    gross_weight: float,
    plausibility: PlausibilityRange,
    *,
    tax_budget: float = 0.0,
    price_reference: float = 0.0,
) -> int:
    min_qty, max_qty = plausible_quantity_bounds(candidate, ctns, gross_weight, plausibility)
    tax_rate = candidate_tax_rate(candidate)

    old_ctns = candidate.ctns or ctns or 1
    old_qty = candidate.qty or old_ctns
    qty_per_ctn = old_qty / old_ctns if old_ctns else 1
    if tax_rate > 0 and tax_budget > 0 and price_reference > 0:
        desired = max(1, round(tax_budget / (price_reference * tax_rate)))
    else:
        desired = max(1, round(ctns * qty_per_ctn))
    return int(min(max(desired, min_qty), max_qty))


def plausible_quantity_bounds(
    candidate: ProductCandidate,
    ctns: int,
    gross_weight: float,
    plausibility: PlausibilityRange,
) -> tuple[int, int]:
    min_pc = plausibility.kg_per_pc_min or DEFAULT_KG_PER_PC_MIN
    max_pc = plausibility.kg_per_pc_max or DEFAULT_KG_PER_PC_MAX
    min_qty = max(1, math.ceil(gross_weight / max_pc))
    max_qty = max(min_qty, math.floor(gross_weight / min_pc))
    if plausibility.qty_per_ctn_min:
        min_qty = max(min_qty, math.ceil(ctns * plausibility.qty_per_ctn_min))
    if plausibility.qty_per_ctn_max:
        max_qty = min(max_qty, max(1, math.floor(ctns * plausibility.qty_per_ctn_max)))
    tax_rate = candidate_tax_rate(candidate)
    max_unit_price = plausibility.unit_price_max or DEFAULT_UNIT_PRICE_MAX
    if tax_rate > 0 and max_unit_price > 0:
        min_qty = max(min_qty, math.ceil(MIN_ROW_TAX_AMOUNT_USD / (max_unit_price * tax_rate)))
    if min_qty > max_qty:
        raise RuntimeError(
            "优化无解：单件重量和单箱件数范围无法同时满足；"
            f"{candidate.zh}/{candidate.en} 毛重 {gross_weight} kg，箱数 {ctns}"
        )
    return min_qty, max_qty


def allocate_plausible_prices(
    selected: list[ProductCandidate],
    quantities: list[int],
    weights: list[float],
    ranges: list[PlausibilityRange],
    target_tax_amount: float,
    tax_budgets: Optional[list[float]] = None,
) -> list[tuple[float, float]]:
    positive_weight = sum(weight for candidate, weight in zip(selected, weights) if candidate_tax_rate(candidate) > 0)
    unit_prices: list[float] = []
    mins: list[float] = []
    maxes: list[float] = []
    row_tax_budgets = tax_budgets or []
    for idx, (candidate, qty, weight, plausibility) in enumerate(zip(selected, quantities, weights, ranges)):
        min_price = plausibility.unit_price_min or DEFAULT_UNIT_PRICE_MIN
        max_price = max(min_price, plausibility.unit_price_max or DEFAULT_UNIT_PRICE_MAX)
        mins.append(min_price)
        maxes.append(max_price)
        tax_rate = candidate_tax_rate(candidate)
        if tax_rate > 0 and idx < len(row_tax_budgets) and row_tax_budgets[idx] > 0:
            target_tax = row_tax_budgets[idx]
            desired = target_tax / tax_rate / qty
        elif tax_rate > 0 and positive_weight > 0:
            target_tax = target_tax_amount * weight / positive_weight
            desired = target_tax / tax_rate / qty
        else:
            desired = candidate.unit_price or min_price
        unit_prices.append(min(max(desired, 0.0001), max_price))

    adjust_price_gap(unit_prices, mins, maxes, selected, quantities, target_tax_amount, 0.0)
    relax_price_floors_if_tax_requires(
        unit_prices,
        mins,
        maxes,
        selected,
        quantities,
        target_tax_amount,
        0.0,
    )
    prices: list[tuple[float, float]] = []
    for unit_price, qty in zip(unit_prices, quantities):
        rounded_unit = round(unit_price, 4)
        total_value = round(rounded_unit * qty, 2)
        prices.append((rounded_unit, total_value))

    estimated_tax = round(sum(total * candidate_tax_rate(candidate) for candidate, (_, total) in zip(selected, prices)), 2)
    tax_upper_bound = target_tax_upper_bound(target_tax_amount)
    if estimated_tax > tax_upper_bound + 0.01:
        raise RuntimeError(
            "优化无解：在重量和单价常理范围内无法压到目标税金上浮 10% 以内；"
            f"目标 {target_tax_amount}，最高 {tax_upper_bound}，可行预计 {estimated_tax}"
        )
    return prices


def relax_price_floors_if_tax_requires(
    unit_prices: list[float],
    mins: list[float],
    maxes: list[float],
    selected: list[ProductCandidate],
    quantities: list[int],
    target_tax_amount: float,
    min_row_tax_amount: float,
) -> None:
    tax_upper_bound = target_tax_upper_bound(target_tax_amount)
    current_tax = estimate_tax_for_unit_prices(unit_prices, selected, quantities)
    if current_tax <= tax_upper_bound + 0.01:
        return
    if min_row_tax_amount <= 0:
        adjust_price_gap(unit_prices, [0.0001 for _ in mins], maxes, selected, quantities, target_tax_amount, 0.0)
        return
    relaxed_mins = relaxed_price_mins_for_tax_floor(selected, quantities, mins, min_row_tax_amount)
    minimum_tax = estimate_tax_for_unit_prices(relaxed_mins, selected, quantities)
    if minimum_tax > tax_upper_bound + 0.01:
        raise RuntimeError(
            "优化无解：最低行税金约束已超过目标税金上浮 10% 上限；"
            f"目标 {target_tax_amount}，最高 {tax_upper_bound}，最低可行 {minimum_tax}"
        )
    adjust_price_gap(unit_prices, relaxed_mins, maxes, selected, quantities, target_tax_amount, min_row_tax_amount)


def relaxed_price_mins_for_tax_floor(
    selected: list[ProductCandidate],
    quantities: list[int],
    mins: list[float],
    min_row_tax_amount: float,
) -> list[float]:
    relaxed: list[float] = []
    for candidate, qty, min_price in zip(selected, quantities, mins):
        rate = candidate_tax_rate(candidate)
        if rate <= 0 or qty <= 0 or min_row_tax_amount <= 0:
            relaxed.append(min_price)
            continue
        relaxed.append(max(0.0001, min_row_tax_amount / (qty * rate)))
    return relaxed


def estimate_tax_for_unit_prices(
    unit_prices: list[float],
    selected: list[ProductCandidate],
    quantities: list[int],
) -> float:
    return round(
        sum(price * qty * candidate_tax_rate(candidate) for price, qty, candidate in zip(unit_prices, quantities, selected)),
        2,
    )


def adjust_price_gap(
    unit_prices: list[float],
    mins: list[float],
    maxes: list[float],
    selected: list[ProductCandidate],
    quantities: list[int],
    target_tax_amount: float,
    min_row_tax_amount: float = 0.0,
) -> None:
    current_tax = sum(price * qty * candidate_tax_rate(candidate) for price, qty, candidate in zip(unit_prices, quantities, selected))
    diff = target_tax_amount - current_tax
    if abs(diff) <= 0.0001:
        return
    if diff > 0:
        order = sorted(range(len(unit_prices)), key=lambda idx: (maxes[idx] - unit_prices[idx]) * quantities[idx] * candidate_tax_rate(selected[idx]), reverse=True)
        for idx in order:
            rate = candidate_tax_rate(selected[idx])
            if rate <= 0:
                continue
            capacity_tax = (maxes[idx] - unit_prices[idx]) * quantities[idx] * rate
            if capacity_tax <= 0:
                continue
            add_tax = min(diff, capacity_tax)
            unit_prices[idx] += add_tax / (quantities[idx] * rate)
            diff -= add_tax
            if abs(diff) <= 0.0001:
                break
    else:
        need = abs(diff)
        order = sorted(range(len(unit_prices)), key=lambda idx: (unit_prices[idx] - mins[idx]) * quantities[idx] * candidate_tax_rate(selected[idx]), reverse=True)
        for idx in order:
            rate = candidate_tax_rate(selected[idx])
            if rate <= 0:
                continue
            min_price = mins[idx]
            if min_row_tax_amount > 0:
                min_price = max(min_price, min_row_tax_amount / (quantities[idx] * rate))
            capacity_tax = (unit_prices[idx] - min_price) * quantities[idx] * rate
            if capacity_tax <= 0:
                continue
            take_tax = min(need, capacity_tax)
            unit_prices[idx] -= take_tax / (quantities[idx] * rate)
            need -= take_tax
            if need <= 0.0001:
                break


def enforce_minimum_row_tax(
    unit_prices: list[float],
    mins: list[float],
    maxes: list[float],
    selected: list[ProductCandidate],
    quantities: list[int],
    min_row_tax_amount: float,
) -> None:
    for idx, candidate in enumerate(selected):
        rate = candidate_tax_rate(candidate)
        qty = quantities[idx]
        if rate <= 0 or qty <= 0:
            continue
        current_tax = unit_prices[idx] * qty * rate
        if current_tax + 0.005 >= min_row_tax_amount:
            continue
        required_unit_price = min_row_tax_amount / (qty * rate)
        if required_unit_price > maxes[idx] + 0.0001:
            raise RuntimeError(
                f"第 {idx + 1} 行税金无法达到最低 {min_row_tax_amount} USD: "
                f"{selected[idx].zh} 最大约 {round(maxes[idx] * qty * rate, 2)}"
            )
        unit_prices[idx] = max(unit_prices[idx], mins[idx], required_unit_price)


def candidate_tax_rate(candidate: ProductCandidate) -> float:
    return candidate.effective_tax_rate or candidate.base_tax_rate or 0.0


def build_constraint_warnings(
    ctns: int,
    qty: int,
    gross_weight: float,
    unit_price: float,
    plausibility: PlausibilityRange,
) -> list[str]:
    warnings: list[str] = []
    kg_per_ctn = gross_weight / ctns if ctns else 0
    kg_per_pc = gross_weight / qty if qty else 0
    if near_bound(kg_per_ctn, plausibility.kg_per_ctn_min, plausibility.kg_per_ctn_max):
        warnings.append(f"单箱重量接近边界 {round(kg_per_ctn, 3)}kg/ctn")
    if near_bound(kg_per_pc, plausibility.kg_per_pc_min, plausibility.kg_per_pc_max):
        warnings.append(f"单件重量接近边界 {round(kg_per_pc, 3)}kg/pc")
    if plausibility.unit_price_min is not None and unit_price < plausibility.unit_price_min - 0.0001:
        warnings.append(f"单价低于合理下限 {round(unit_price, 4)} < {round(plausibility.unit_price_min, 4)}")
    elif plausibility.unit_price_max is not None and unit_price > plausibility.unit_price_max + 0.0001:
        warnings.append(f"单价高于合理上限 {round(unit_price, 4)} > {round(plausibility.unit_price_max, 4)}")
    elif near_bound(unit_price, plausibility.unit_price_min, plausibility.unit_price_max):
        warnings.append(f"单价接近边界 {round(unit_price, 4)}")
    return warnings


def near_bound(value: float, low: Optional[float], high: Optional[float]) -> bool:
    if low is not None and value <= low * 1.05:
        return True
    if high is not None and value >= high * 0.95:
        return True
    return False


def adjust_tax_gap(rows: list[dict[str, Any]], target_tax_amount: float) -> None:
    current = sum((to_float(row.get("总价")) or 0) * (to_float(row.get("综合税率")) or 0) for row in rows)
    diff = round(target_tax_amount - current, 2)
    tolerance = TAX_FINAL_TOLERANCE_USD
    if abs(diff) <= tolerance:
        for row in rows:
            update_row_tax_display(row)
        return
    adjustable = [row for row in rows if (to_float(row.get("综合税率")) or 0) > 0]
    if not adjustable:
        raise RuntimeError("没有可调整税金的正税率行")
    row = adjustable[-1]
    rate = to_float(row.get("综合税率")) or 0
    total_value = max(0.01, (to_float(row.get("总价")) or 0) + diff / rate)
    row["总价"] = round(total_value, 2)
    qty = to_float(row.get("数量")) or 1
    row["单价"] = round(row["总价"] / qty, 4)
    for item in rows:
        update_row_tax_display(item)


def update_row_tax_display(row: dict[str, Any]) -> None:
    rate = to_float(row.get("综合税率")) or to_float(row.get(DISPLAY_TAX_RATE_FIELD)) or 0.0
    total_value = to_float(row.get("总价")) or 0.0
    tax_amount = round(total_value * rate, 2)
    row["综合税率"] = round(rate, 6)
    row["预计税金"] = tax_amount
    row[DISPLAY_TAX_RATE_FIELD] = format_rate(rate)
    row[DISPLAY_TAX_AMOUNT_FIELD] = tax_amount


def summarize_filter_reasons(candidates: list[ProductCandidate]) -> dict[str, int]:
    summary: dict[str, int] = {}
    for candidate in candidates:
        reason = candidate.filter_reason or "unknown"
        first = reason.split("；", 1)[0]
        summary[first] = summary.get(first, 0) + 1
    return summary


async def parse_manifest_total_weight(
    workbook,
    source: Path,
    llm: Optional[LLMClient],
    query_cache: Optional[QueryCache] = None,
) -> ManifestWeightInfo:
    summary = build_manifest_weight_context(workbook)
    if not summary["sheets"]:
        return ManifestWeightInfo(total_weight_kg=None)
    cache_key = hashlib.sha256(
        json.dumps(summary, ensure_ascii=False, sort_keys=True).encode("utf-8", errors="ignore")
    ).hexdigest()
    if query_cache is not None:
        manifest_cache = query_cache.setdefault("manifest", {})
        if cache_key in manifest_cache and isinstance(manifest_cache[cache_key], dict):
            return normalize_manifest_weight_payload(manifest_cache[cache_key])
    if llm is None:
        return infer_manifest_weight_from_context(summary)

    payload = await llm.chat_json(build_manifest_weight_parser_messages(source.name, summary), temperature=0.0)
    weight_info = normalize_manifest_weight_payload(payload)
    if query_cache is not None:
        query_cache.setdefault("manifest", {})[cache_key] = payload
    return weight_info


def build_manifest_weight_context(workbook) -> dict[str, Any]:
    sheets: list[dict[str, Any]] = []
    for sheet in workbook.worksheets[:4]:
        headers = [clean_text(sheet.cell(1, col).value) for col in range(1, sheet.max_column + 1)]
        numeric_columns: list[dict[str, Any]] = []
        for col_idx, header in enumerate(headers, start=1):
            if not header:
                continue
            values: list[float] = []
            sample_values: list[float] = []
            for row_idx in range(2, sheet.max_row + 1):
                value = to_float(sheet.cell(row_idx, col_idx).value)
                if value is None or value <= 0:
                    continue
                values.append(value)
                if len(sample_values) < 8:
                    sample_values.append(value)
            if values:
                numeric_columns.append(
                    {
                        "column": col_idx,
                        "header": header,
                        "count": len(values),
                        "sum": round(sum(values), 4),
                        "max": round(max(values), 4),
                        "sample_values": sample_values,
                    }
                )

        sample_rows: list[list[Any]] = []
        for row_idx in range(1, min(sheet.max_row, 12) + 1):
            sample_rows.append(
                [
                    clean_text(sheet.cell(row_idx, col_idx).value)[:120]
                    for col_idx in range(1, min(sheet.max_column, 12) + 1)
                ]
            )
        tail_rows: list[list[Any]] = []
        start_tail = max(1, sheet.max_row - 5)
        for row_idx in range(start_tail, sheet.max_row + 1):
            tail_rows.append(
                [
                    clean_text(sheet.cell(row_idx, col_idx).value)[:120]
                    for col_idx in range(1, min(sheet.max_column, 12) + 1)
                ]
            )
        sheets.append(
            {
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "headers": headers,
                "numeric_columns": numeric_columns,
                "sample_rows": sample_rows,
                "tail_rows": tail_rows,
            }
        )
    return {"sheets": sheets}


def build_manifest_weight_parser_messages(filename: str, summary: dict[str, Any]) -> list[dict[str, str]]:
    context = json.dumps(summary, ensure_ascii=False, separators=(",", ":"))
    return [
        {
            "role": "system",
            "content": (
                "你是装箱清单 Excel 字段识别专家。你的任务是根据工作簿摘要判断清单的 Excel 总重量。"
                "必须优先识别表示整票或逐行总毛重的列，例如 总毛重KGS、Gross Weight、G.W.、Weight KGS。"
                "不要使用提单重量，也不要使用体积、数量、箱数、单价、金额。只返回 JSON object。"
            ),
        },
        {
            "role": "user",
            "content": (
                f"文件名：{filename}\n"
                "请从以下 Excel 摘要中识别清单总重量，单位统一为 kg。\n"
                "规则：\n"
                "1. 如果有明确总重量合计行，返回该合计。\n"
                "2. 如果没有合计行，但存在逐行总毛重列，返回该列有效正数求和。\n"
                "3. 如果只有单箱重量或单件重量，不要猜总重量，返回 null。\n"
                "4. evidence 要说明使用了哪个 sheet、哪一列、如何计算。\n"
                "JSON格式：{\"total_weight_kg\":0.0,\"source\":\"\",\"evidence\":\"\",\"confidence\":0.0}\n"
                f"Excel摘要：\n{context[:18000]}"
            ),
        },
    ]


def normalize_manifest_weight_payload(payload: dict[str, Any]) -> ManifestWeightInfo:
    weight = to_float(payload.get("total_weight_kg"))
    confidence = to_float(payload.get("confidence"))
    confidence = max(0.0, min(1.0, confidence if confidence is not None else 0.0))
    if weight is not None and weight <= 0:
        weight = None
    return ManifestWeightInfo(
        total_weight_kg=weight,
        source=clean_text(payload.get("source")),
        evidence=clean_text(payload.get("evidence")),
        confidence=confidence,
    )


def infer_manifest_weight_from_context(summary: dict[str, Any]) -> ManifestWeightInfo:
    best: Optional[dict[str, Any]] = None
    for sheet in summary.get("sheets") or []:
        for column in sheet.get("numeric_columns") or []:
            header_key = normalize_text(column.get("header"))
            if any(term in header_key for term in ("毛重", "总毛重", "gross weight", "gw", "g w", "weight kgs")):
                if "体积" in header_key or "cbm" in header_key:
                    continue
                if best is None or (column.get("count") or 0) > (best.get("count") or 0):
                    best = {**column, "sheet": sheet.get("name")}
    if not best:
        return ManifestWeightInfo(total_weight_kg=None)
    return ManifestWeightInfo(
        total_weight_kg=to_float(best.get("sum")),
        source=f"{best.get('sheet')}.{best.get('header')}",
        evidence=f"按 {best.get('sheet')} sheet 的 {best.get('header')} 列正数求和",
        confidence=0.7,
    )


async def detect_manifest_schema(
    workbook,
    source: Path,
    llm: Optional[LLMClient],
    query_cache: Optional[QueryCache] = None,
) -> ManifestSchema:
    if llm is None:
        raise manifest_schema_error("缺少 LLM，无法识别清单结构")
    model = clean_text(getattr(getattr(llm, "settings", None), "model", "")) or llm.__class__.__name__
    file_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    cache_key = (
        f"schema:{MANIFEST_SCHEMA_PROMPT_VERSION}:"
        f"{MANIFEST_SCHEMA_CACHE_REVISION}:{model}:{file_hash}"
    )
    cached = (query_cache or {}).get("manifest_schema", {}).get(cache_key)
    if isinstance(cached, dict):
        return validate_manifest_schema(workbook, normalize_manifest_schema_payload(cached))

    context = build_manifest_schema_context(workbook)
    try:
        payload = await llm.chat_json(
            build_manifest_schema_parser_messages(source.name, context),
            temperature=0.0,
            max_tokens=4096,
            json_mode=True,
        )
    except Exception as exc:
        raise manifest_schema_error(f"LLM 结构识别失败: {exc}") from exc
    schema = validate_manifest_schema(workbook, normalize_manifest_schema_payload(payload))
    if query_cache is not None:
        query_cache.setdefault("manifest_schema", {})[cache_key] = payload
    return schema


def build_manifest_schema_context(workbook) -> dict[str, Any]:
    sheets: list[dict[str, Any]] = []
    for sheet in workbook.worksheets[:6]:
        row_indexes = list(range(1, min(sheet.max_row, 30) + 1))
        row_indexes.extend(range(max(1, sheet.max_row - 4), sheet.max_row + 1))
        rows: list[dict[str, Any]] = []
        for row_idx in dict.fromkeys(row_indexes):
            rows.append(
                {
                    "row": row_idx,
                    "values": [
                        clean_text(sheet.cell(row_idx, col_idx).value)[:100]
                        for col_idx in range(1, min(sheet.max_column, 24) + 1)
                    ],
                }
            )
        sheets.append(
            {
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "merged_ranges": [str(item) for item in list(sheet.merged_cells.ranges)[:20]],
                "rows": rows,
            }
        )
    return {"sheets": sheets}


def build_manifest_schema_parser_messages(filename: str, context: dict[str, Any]) -> list[dict[str, str]]:
    workbook_json = json.dumps(context, ensure_ascii=False, separators=(",", ":"))
    return [
        {
            "role": "system",
            "content": (
                "你是装箱清单 Excel 结构识别器。只识别语义和单元格坐标，不计算或猜测最终重量。"
                "返回一个 JSON object，所有行列坐标从 1 开始。"
            ),
        },
        {
            "role": "user",
            "content": (
                f"文件名：{filename}\n"
                "识别商品清单 sheet、表头行、汇总行、数据起止行和字段列。columns 只可使用 "
                "zh_name,en_name,material,usage,qty,ctns,net_weight,gross_weight,declared_value,unit_price,hs_code。"
                "若存在净重列必须返回 net_weight；不存在则省略。\n"
                "header_labels 必须逐字段原样返回对应表头单元格文字，用于代码核验列坐标。\n"
                "weight_strategy 只能是 explicit_total 或 detail_sum；若存在汇总毛重，返回 total_cell；"
                "若存在毛重明细列，detail_range 必须指向 gross_weight 毛重列，禁止指向 net_weight 净重列。"
                "不要返回模型计算出的总重量。\n"
                "JSON格式："
                "{\"sheet_name\":\"\",\"header_row\":1,\"summary_rows\":[],"
                "\"data_start_row\":2,\"data_end_row\":2,\"columns\":{},\"header_labels\":{},"
                "\"weight_unit\":\"kg\",\"confidence\":0.0,\"weight_strategy\":\"detail_sum\","
                "\"weight_total_cell\":\"\",\"weight_detail_range\":\"\"}\n"
                f"工作簿摘要：\n{workbook_json[:30000]}"
            ),
        },
    ]


def normalize_manifest_schema_payload(payload: dict[str, Any]) -> ManifestSchema:
    if not isinstance(payload, dict):
        raise manifest_schema_error("LLM 返回值不是 JSON object")
    raw_columns = payload.get("columns") if isinstance(payload.get("columns"), dict) else {}
    raw_labels = payload.get("header_labels") if isinstance(payload.get("header_labels"), dict) else {}
    aliases = {
        "chinese_name": "zh_name",
        "english_name": "en_name",
        "name_zh": "zh_name",
        "name_en": "en_name",
        "weight": "gross_weight",
        "grossweight": "gross_weight",
        "netweight": "net_weight",
        "amount": "declared_value",
        "value": "declared_value",
        "price": "unit_price",
        "hs": "hs_code",
        "hscode": "hs_code",
    }
    allowed = {
        "zh_name",
        "en_name",
        "material",
        "usage",
        "qty",
        "ctns",
        "net_weight",
        "gross_weight",
        "declared_value",
        "unit_price",
        "hs_code",
    }
    columns: dict[str, int] = {}
    header_labels: dict[str, str] = {}
    for raw_key, raw_value in raw_columns.items():
        key = aliases.get(normalize_header(raw_key), normalize_header(raw_key))
        if key not in allowed or raw_value in (None, ""):
            continue
        value = to_float(raw_value)
        if value is None or not float(value).is_integer():
            raise manifest_schema_error(f"字段 {raw_key} 的列号无效: {raw_value}")
        columns[key] = int(value)
        label = raw_labels.get(raw_key, raw_labels.get(key, ""))
        header_labels[key] = clean_text(label)
    weight = payload.get("weight") if isinstance(payload.get("weight"), dict) else {}
    summary_rows: list[int] = []
    for value in payload.get("summary_rows") or []:
        parsed = to_float(value)
        if parsed is None or not float(parsed).is_integer():
            raise manifest_schema_error(f"汇总行号无效: {value}")
        summary_rows.append(int(parsed))
    return ManifestSchema(
        sheet_name=clean_text(payload.get("sheet_name")),
        header_row=int(to_float(payload.get("header_row")) or 0),
        summary_rows=sorted(set(summary_rows)),
        data_start_row=int(to_float(payload.get("data_start_row")) or 0),
        data_end_row=int(to_float(payload.get("data_end_row")) or 0),
        columns=columns,
        header_labels=header_labels,
        weight_unit=clean_text(payload.get("weight_unit") or weight.get("unit") or "kg"),
        confidence=float(to_float(payload.get("confidence")) or 0.0),
        weight_strategy=normalize_header(payload.get("weight_strategy") or weight.get("strategy")),
        weight_total_cell=clean_text(payload.get("weight_total_cell") or payload.get("total_cell") or weight.get("total_cell")),
        weight_detail_range=clean_text(payload.get("weight_detail_range") or payload.get("detail_range") or weight.get("detail_range")),
    )


def validate_manifest_schema(workbook, schema: ManifestSchema) -> ManifestSchema:
    if schema.sheet_name not in workbook.sheetnames:
        raise manifest_schema_error(f"sheet 不存在: {schema.sheet_name or '(empty)'}")
    sheet = workbook[schema.sheet_name]
    if schema.confidence < MANIFEST_SCHEMA_MIN_CONFIDENCE:
        raise manifest_schema_error(
            f"结构识别置信度 {schema.confidence:.2f} 低于 {MANIFEST_SCHEMA_MIN_CONFIDENCE:.2f}"
        )
    if not 1 <= schema.header_row <= sheet.max_row:
        raise manifest_schema_error(f"表头行越界: {schema.header_row}")
    if not 1 <= schema.data_start_row <= sheet.max_row or schema.data_start_row <= schema.header_row:
        raise manifest_schema_error(f"数据起始行无效: {schema.data_start_row}")
    if schema.data_end_row <= 0:
        schema.data_end_row = sheet.max_row
    if not schema.data_start_row <= schema.data_end_row <= sheet.max_row:
        raise manifest_schema_error(f"数据结束行无效: {schema.data_end_row}")
    if any(row < 1 or row > sheet.max_row or row == schema.header_row for row in schema.summary_rows):
        raise manifest_schema_error(f"汇总行越界或与表头重叠: {schema.summary_rows}")
    if not schema.columns:
        raise manifest_schema_error("LLM 未返回字段列映射")
    if not any(schema.columns.get(key) for key in ("zh_name", "en_name", "hs_code")):
        raise manifest_schema_error("字段映射缺少商品名称和 HS")
    if not schema.columns.get("gross_weight"):
        raise manifest_schema_error("字段映射缺少重量列")

    for key, col_idx in schema.columns.items():
        if not 1 <= col_idx <= sheet.max_column:
            raise manifest_schema_error(f"字段 {key} 的列号越界: {col_idx}")
        expected = schema.header_labels.get(key, "")
        actual = clean_text(sheet.cell(schema.header_row, col_idx).value)
        if not expected:
            raise manifest_schema_error(f"字段 {key} 缺少 header_labels 核验值")
        if normalize_header(expected) != normalize_header(actual):
            raise manifest_schema_error(
                f"字段 {key} 表头不一致: 模型返回 {expected!r}，单元格实际为 {actual!r}"
            )

    data_rows = [
        row
        for row in range(schema.data_start_row, schema.data_end_row + 1)
        if row not in schema.summary_rows
    ]
    for key in ("qty", "ctns", "net_weight", "gross_weight", "declared_value", "unit_price"):
        col_idx = schema.columns.get(key)
        if not col_idx:
            continue
        values = [sheet.cell(row, col_idx).value for row in data_rows]
        nonempty = [value for value in values if clean_text(value)]
        if nonempty and sum(to_float(value) is not None for value in nonempty) / len(nonempty) < 0.6:
            raise manifest_schema_error(f"字段 {key} 的数据主要不是数值")

    hs_col = schema.columns.get("hs_code")
    if hs_col:
        samples = [clean_text(sheet.cell(row, hs_col).value) for row in data_rows]
        samples = [value for value in samples if value][:20]
        if samples:
            valid = sum(is_manifest_hs_sample(value) for value in samples)
            if valid / len(samples) < 0.8:
                raise manifest_schema_error("HS 列样本不符合编码格式")

    if schema.weight_strategy not in {"explicit_total", "detail_sum"}:
        raise manifest_schema_error(f"重量策略无效: {schema.weight_strategy or '(empty)'}")
    if schema.weight_strategy == "explicit_total" and not schema.weight_total_cell:
        raise manifest_schema_error("explicit_total 策略缺少 weight_total_cell")
    if schema.weight_strategy == "detail_sum" and not schema.weight_detail_range:
        raise manifest_schema_error("detail_sum 策略缺少 weight_detail_range")
    if schema.weight_total_cell:
        total_sheet, total_bounds = validate_manifest_reference(
            schema.weight_total_cell, schema.sheet_name, workbook, single_cell=True
        )
        if total_sheet != schema.sheet_name:
            raise manifest_schema_error("weight_total_cell 不在识别的清单 sheet")
        if schema.summary_rows and total_bounds[1] not in schema.summary_rows:
            raise manifest_schema_error("weight_total_cell 不在模型识别的汇总行")
    if schema.weight_detail_range:
        detail_sheet, bounds = validate_manifest_reference(
            schema.weight_detail_range, schema.sheet_name, workbook, single_column=True
        )
        if detail_sheet == schema.sheet_name and bounds[0] == schema.columns.get("net_weight"):
            gross_col = get_column_letter(schema.columns["gross_weight"])
            schema.weight_detail_range = (
                f"{gross_col}{schema.data_start_row}:{gross_col}{schema.data_end_row}"
            )
            detail_sheet, bounds = validate_manifest_reference(
                schema.weight_detail_range, schema.sheet_name, workbook, single_column=True
            )
        if detail_sheet != schema.sheet_name or bounds[0] != schema.columns["gross_weight"]:
            raise manifest_schema_error("weight_detail_range 与识别的重量列不一致")
        if bounds[1] > schema.data_start_row or bounds[3] < schema.data_end_row:
            raise manifest_schema_error("weight_detail_range 未覆盖模型识别的全部数据行")
    return schema


def manifest_schema_error(message: str) -> RuntimeError:
    return RuntimeError(f"manifest_schema_unresolved: {message}")


def is_manifest_hs_sample(value: Any) -> bool:
    text = clean_text(value)
    if re.search(r"[A-Za-z\u4e00-\u9fff]", text):
        return False
    return 6 <= len(normalize_hs(text)) <= 10


def validate_manifest_reference(
    reference: str,
    default_sheet: str,
    workbook,
    *,
    single_cell: bool = False,
    single_column: bool = False,
) -> tuple[str, tuple[int, int, int, int]]:
    sheet_name, coordinate = split_manifest_reference(reference, default_sheet)
    if sheet_name not in workbook.sheetnames:
        raise manifest_schema_error(f"重量坐标 sheet 不存在: {sheet_name}")
    try:
        bounds = range_boundaries(coordinate.replace("$", ""))
    except (TypeError, ValueError) as exc:
        raise manifest_schema_error(f"重量坐标无效: {reference}") from exc
    min_col, min_row, max_col, max_row = bounds
    sheet = workbook[sheet_name]
    if min_col < 1 or min_row < 1 or max_col > sheet.max_column or max_row > sheet.max_row:
        raise manifest_schema_error(f"重量坐标越界: {reference}")
    if single_cell and (min_col != max_col or min_row != max_row):
        raise manifest_schema_error(f"总重量坐标必须是单个单元格: {reference}")
    if single_column and min_col != max_col:
        raise manifest_schema_error(f"重量明细范围必须是单列: {reference}")
    return sheet_name, bounds


def split_manifest_reference(reference: str, default_sheet: str) -> tuple[str, str]:
    text = clean_text(reference)
    if "!" not in text:
        return default_sheet, text
    sheet_name, coordinate = text.rsplit("!", 1)
    return sheet_name.strip().strip("'"), coordinate.strip()


def manifest_weight_unit_factor(unit: str) -> float:
    normalized = normalize_header(unit).replace(".", "")
    if normalized in {"kg", "kgs", "kilogram", "kilograms", "千克", "公斤"}:
        return 1.0
    if normalized in {"lb", "lbs", "pound", "pounds", "磅"}:
        return 0.45359237
    raise manifest_schema_error(f"不支持的重量单位: {unit}")


def resolve_manifest_weight(
    workbook,
    schema: ManifestSchema,
    items: list[ManifestItem],
) -> ManifestWeightResolution:
    factor = manifest_weight_unit_factor(schema.weight_unit)
    explicit_total: Optional[float] = None
    if schema.weight_total_cell:
        total_sheet, bounds = validate_manifest_reference(
            schema.weight_total_cell, schema.sheet_name, workbook, single_cell=True
        )
        raw_total = to_float(workbook[total_sheet].cell(bounds[1], bounds[0]).value)
        if raw_total is None or raw_total <= 0:
            raise manifest_schema_error(
                f"模型给出的总重量单元格无法从原表重算: {schema.weight_total_cell}"
            )
        explicit_total = raw_total * factor

    detail_values = [item.gross_weight for item in items if item.gross_weight and item.gross_weight > 0]
    detail_sum = sum(detail_values) if detail_values else None
    if schema.weight_strategy == "detail_sum" and detail_sum is None:
        raise manifest_schema_error("模型给出的明细重量范围无法从商品行重算")

    reconciled = False
    if explicit_total is not None and detail_sum is not None:
        tolerance = max(
            MANIFEST_WEIGHT_ABSOLUTE_TOLERANCE_KG,
            max(explicit_total, detail_sum) * MANIFEST_WEIGHT_RELATIVE_TOLERANCE,
        )
        if abs(explicit_total - detail_sum) > tolerance:
            raise RuntimeError(
                "manifest_weight_conflict: "
                f"显式汇总重量 {explicit_total:.4f} kg 与明细合计 {detail_sum:.4f} kg 冲突"
            )
        final_weight = explicit_total
        reconciled = True
        strategy = "explicit_total"
    elif explicit_total is not None:
        final_weight = explicit_total
        strategy = "explicit_total"
    elif detail_sum is not None:
        final_weight = detail_sum
        strategy = "detail_sum"
    else:
        raise manifest_schema_error("清单没有可重算的汇总重量或商品明细重量")

    explicit_display = round(explicit_total, 4) if explicit_total is not None else None
    detail_display = round(detail_sum, 4) if detail_sum is not None else None
    evidence = json.dumps(
        {
            "strategy": strategy,
            "total_cell": schema.weight_total_cell,
            "detail_range": schema.weight_detail_range,
            "explicit_total": explicit_display,
            "detail_sum": detail_display,
            "final": round(final_weight, 4),
            "reconciled": reconciled,
        },
        ensure_ascii=False,
        separators=(",", ":"),
    )
    return ManifestWeightResolution(
        strategy=strategy,
        total_cell=schema.weight_total_cell,
        detail_range=schema.weight_detail_range,
        explicit_total=explicit_display,
        detail_sum=detail_display,
        final_weight=round(final_weight, 4),
        reconciled=reconciled,
        evidence=evidence,
    )


def resolve_manifest_net_weight(items: list[ManifestItem], total_gross_weight: float) -> Optional[float]:
    weighted_items = [item for item in items if item.gross_weight is not None and item.gross_weight > 0]
    if not weighted_items or any(item.real_weight is None or item.real_weight <= 0 for item in weighted_items):
        return None

    for item in weighted_items:
        if (item.real_weight or 0) > (item.gross_weight or 0) + 0.01:
            raise RuntimeError(
                "manifest_weight_conflict: "
                f"第 {item.row} 行净重 {item.real_weight:.4f} kg 大于毛重 {item.gross_weight:.4f} kg"
            )

    detail_gross_weight = sum(item.gross_weight or 0 for item in weighted_items)
    tolerance = max(
        MANIFEST_WEIGHT_ABSOLUTE_TOLERANCE_KG,
        max(detail_gross_weight, total_gross_weight) * MANIFEST_WEIGHT_RELATIVE_TOLERANCE,
    )
    if abs(detail_gross_weight - total_gross_weight) > tolerance:
        return None

    total_net_weight = round(sum(item.real_weight or 0 for item in weighted_items), 4)
    if total_net_weight > total_gross_weight + tolerance:
        raise RuntimeError(
            "manifest_weight_conflict: "
            f"明细净重合计 {total_net_weight:.4f} kg 大于总毛重 {total_gross_weight:.4f} kg"
        )
    return total_net_weight


async def parse_manifest(
    path: str | Path,
    llm: Optional[LLMClient] = None,
    query_cache: Optional[QueryCache] = None,
) -> ManifestSummary:
    source = Path(path)
    workbook = load_workbook(source, data_only=True)
    schema = await detect_manifest_schema(workbook, source, llm, query_cache)
    sheet = workbook[schema.sheet_name]
    columns = schema.columns
    factor = manifest_weight_unit_factor(schema.weight_unit)

    def value(row: int, key: str) -> Any:
        col_idx = columns.get(key)
        return sheet.cell(row, col_idx).value if col_idx else None

    total_ctns = 0.0
    total_declared_value = 0.0
    items: list[ManifestItem] = []
    categories: list[str] = []
    seen_categories: set[str] = set()
    for row_idx in range(schema.data_start_row, schema.data_end_row + 1):
        if row_idx in schema.summary_rows:
            continue
        zh = clean_text(value(row_idx, "zh_name"))
        en = clean_text(value(row_idx, "en_name"))
        raw_hs = clean_text(value(row_idx, "hs_code"))
        hs = normalize_hs(raw_hs) if is_manifest_hs_sample(raw_hs) else ""
        if not (zh or en or hs):
            continue
        ctns = to_float(value(row_idx, "ctns"))
        qty = to_float(value(row_idx, "qty"))
        net_weight = to_float(value(row_idx, "net_weight"))
        if net_weight is not None:
            net_weight *= factor
        gross_weight = to_float(value(row_idx, "gross_weight"))
        if gross_weight is not None:
            gross_weight *= factor
        declared_value = to_float(value(row_idx, "declared_value"))
        unit_price = to_float(value(row_idx, "unit_price"))
        category = zh or en or hs
        category_key = normalize_text(category)
        if category_key and category_key not in seen_categories:
            seen_categories.add(category_key)
            categories.append(category)
        total_ctns += ctns or 0
        total_declared_value += declared_value or 0
        items.append(
            ManifestItem(
                row=row_idx,
                zh=zh,
                en=en,
                hs=hs,
                material=clean_text(value(row_idx, "material")),
                usage=clean_text(value(row_idx, "usage")),
                ctns=ctns,
                qty=qty,
                unit_price=unit_price,
                declared_value=declared_value,
                real_weight=net_weight,
                gross_weight=gross_weight,
            )
        )

    weight = resolve_manifest_weight(workbook, schema, items)
    total_net_weight = resolve_manifest_net_weight(items, weight.final_weight)
    return ManifestSummary(
        filename=source.name,
        row_count=len(items),
        total_ctns=round(total_ctns, 2),
        total_real_weight=round(weight.final_weight, 2),
        total_declared_value=round(total_declared_value, 2),
        categories=categories,
        items=items,
        total_net_weight=round(total_net_weight, 2) if total_net_weight is not None else None,
        weight_source=weight.total_cell or weight.detail_range or weight.strategy,
        weight_evidence=weight.evidence,
        weight_confidence=round(schema.confidence, 4),
        schema_sheet=schema.sheet_name,
        schema_header_row=schema.header_row,
        schema_data_start_row=schema.data_start_row,
        schema_data_end_row=schema.data_end_row,
        schema_summary_rows=schema.summary_rows,
        schema_columns=schema.columns,
        schema_confidence=round(schema.confidence, 4),
        weight_strategy=weight.strategy,
        weight_total_cell=weight.total_cell,
        weight_detail_range=weight.detail_range,
        weight_explicit_total=weight.explicit_total,
        weight_detail_sum=weight.detail_sum,
        weight_final=weight.final_weight,
        weight_reconciled=weight.reconciled,
    )


async def parse_bill(
    path: str | Path,
    llm: LLMClient,
    query_cache: Optional[QueryCache] = None,
    *,
    allow_missing_carton_count: bool = False,
    text_max_chars: int = BILL_PARSER_TEXT_CHARS,
    vision_max_pages: int = BILL_VISION_MAX_PAGES,
) -> BillInfo:
    source = Path(path)
    text = extract_bill_text(source)
    text_chars = len(normalize_text(text))
    parse_source = "text"
    vision_pages = 0
    if text_chars >= BILL_TEXT_MIN_CHARS:
        parsed_fields = await parse_bill_fields_from_text(
            text,
            llm,
            query_cache,
            max_chars=text_max_chars,
        )
    else:
        image_data_urls = render_bill_pdf_pages(source, max_pages=vision_max_pages)
        vision_pages = len(image_data_urls)
        parse_source = "vision"
        try:
            parsed_fields = await parse_bill_fields_from_images(source, image_data_urls, llm, query_cache)
        except RuntimeError as exc:
            if "LLM 未从提单中识别到可靠货物品类" in str(exc):
                raise RuntimeError("提单为扫描件，视觉模型未识别到可靠货物品类") from exc
            raise
    product_entries = parsed_fields.product_entries
    products = [entry.name for entry in product_entries]
    text_carton_match = re.search(
        r"(?<![\d,])([+\-−–—]?\s*\d+(?:,\d{3})*(?:\.\d+)?)\s*CARTONS?\b",
        text,
        flags=re.IGNORECASE,
    )
    text_cartons = normalize_carton_count(
        text_carton_match.group(1) if text_carton_match else None
    )
    cartons = parsed_fields.carton_count or text_cartons
    if (not cartons or cartons <= 0) and not allow_missing_carton_count:
        source_label = "扫描件视觉模型" if parse_source == "vision" else "文本/LLM"
        raise RuntimeError(f"提单未识别到有效总箱数，{source_label}未返回 carton_count")
    carton_source = parsed_fields.carton_source
    carton_unit = parsed_fields.carton_unit
    carton_evidence = parsed_fields.carton_evidence
    if cartons and cartons > 0 and not parsed_fields.carton_count:
        carton_source = "bill_text_cartons_regex"
        carton_unit = carton_unit or "CARTONS"
        carton_evidence = carton_evidence or f"{cartons:g} CARTONS"
    initial_status = "recognized" if cartons and cartons > 0 else "missing"
    initial_history = [
        {
            "attempt": 1,
            "status": initial_status,
            "package_count": cartons,
            "package_unit": carton_unit,
            "source": carton_source or parse_source,
            "evidence": carton_evidence,
            "reasoning": parsed_fields.carton_reasoning,
        }
    ]
    return BillInfo(
        filename=source.name,
        raw_text=text,
        products=products,
        shipper=parsed_fields.shipper or extract_shipper(text),
        consignee=parsed_fields.consignee or extract_consignee(text),
        shipment_no=extract_shipment_no(text),
        eta=extract_eta(text),
        cartons=cartons,
        carton_evidence=carton_evidence,
        carton_unit=carton_unit,
        carton_source=carton_source or parse_source,
        carton_reasoning=parsed_fields.carton_reasoning,
        carton_manifest_comparison=parsed_fields.carton_manifest_comparison,
        carton_confidence=parsed_fields.carton_confidence,
        carton_inferred=parsed_fields.carton_inferred,
        carton_recognition_attempts=1,
        carton_resolution_history=initial_history,
        gross_weight=extract_number(r"(\d+(?:\.\d+)?)\s*KGS?", text),
        cbm=extract_number(r"(\d+(?:\.\d+)?)\s*CBM", text),
        product_entries=product_entries,
        parse_source=parse_source,
        text_chars=text_chars,
        vision_pages=vision_pages,
    )


def bill_and_manifest_carton_comparison(
    package_count: Optional[float],
    manifest: ManifestSummary,
) -> str:
    manifest_count = to_float(manifest.total_ctns)
    if not manifest_count or manifest_count <= 0:
        return "清单未识别到有效总箱数，由模型根据提单选择整票包装数量"
    if package_count and math.isclose(float(package_count), manifest_count, rel_tol=0.0, abs_tol=0.01):
        return f"提单包装数量 {package_count:g} 与清单总箱数 {manifest_count:g} 一致"
    if package_count and package_count > 0:
        return (
            f"提单候选包装数量 {package_count:g} 与清单总箱数 {manifest_count:g} 不一致，"
            "已交给模型结合两份资料复核"
        )
    return f"提单首次未识别到包装数量，清单总箱数候选为 {manifest_count:g}"


def build_bill_package_resolution_context(
    manifest: ManifestSummary,
    bill: BillInfo,
    *,
    attempt: int,
    feedback: str,
) -> dict[str, Any]:
    return {
        "attempt": attempt,
        "max_attempts": BILL_PACKAGE_RECOGNITION_MAX_ATTEMPTS,
        "feedback_from_previous_attempt": feedback,
        "bill": {
            "filename": bill.filename,
            "parse_source": bill.parse_source,
            "previous_package_count": bill.cartons,
            "previous_package_unit": bill.carton_unit,
            "previous_evidence": bill.carton_evidence,
            "previous_source": bill.carton_source,
            "previous_reasoning": bill.carton_reasoning,
            "previous_manifest_comparison": bill.carton_manifest_comparison,
            "previous_confidence": bill.carton_confidence,
            "products": bill.products,
            "gross_weight": bill.gross_weight,
            "cbm": bill.cbm,
        },
        "manifest": {
            "filename": manifest.filename,
            "total_ctns": manifest.total_ctns,
            "row_count": manifest.row_count,
            "categories": manifest.categories[:40],
            "rows": [
                {
                    "row": item.row,
                    "zh_name": item.zh,
                    "en_name": item.en,
                    "ctns": item.ctns,
                    "qty": item.qty,
                }
                for item in manifest.items[:80]
            ],
        },
    }


def build_bill_package_resolution_prompt(
    manifest: ManifestSummary,
    bill: BillInfo,
    *,
    attempt: int,
    feedback: str,
) -> str:
    context = build_bill_package_resolution_context(
        manifest,
        bill,
        attempt=attempt,
        feedback=feedback,
    )
    return (
        "你是国际运输单据识别专家。请重新识别这票货物最终应采用的整票包装数量，并结合清单交叉验证。\n"
        "不要把识别范围限制在 CTNS、CARTONS、PACKAGES 或 NO. OF PKGS；你可以根据单据版式和语义，"
        "识别任何实际代表整票运输包装/交运件数的单位和栏位。航空运单的 No. of Pieces RCP、"
        "No. of Pieces、RCP 等标准栏位可以作为整票包装数量。\n"
        "清单与提单数字一致时直接采用；不一致时请自行判断每个数字代表的层级，并选择你认为正确的最终值。"
        "如果上一轮没有识别出来，请完整重看单据，不要机械重复上一轮答案。\n"
        "必须返回一个 JSON object。识别成功时 package_count 必须是正数；确实无法判断时返回 null。\n"
        "JSON格式：{\"package_count\":null,\"package_unit\":\"\",\"source_document\":\"\","
        "\"field_label\":\"\",\"evidence\":\"\",\"manifest_comparison\":\"\","
        "\"reasoning_summary\":\"\",\"confidence\":0.0,\"inferred\":true}\n"
        f"交叉验证上下文：\n{json.dumps(context, ensure_ascii=False, separators=(',', ':'))}"
    )


def normalize_boolean(value: Any, *, default: bool = False) -> bool:
    if isinstance(value, bool):
        return value
    text = clean_text(value).lower()
    if text in {"1", "true", "yes", "y", "是"}:
        return True
    if text in {"0", "false", "no", "n", "否"}:
        return False
    return default


def normalize_bill_package_resolution(payload: dict[str, Any]) -> BillLLMFields:
    confidence = to_float(payload.get("confidence"))
    if confidence is None:
        confidence = 0.0
    return BillLLMFields(
        product_entries=[],
        carton_count=normalize_carton_count(
            payload.get("package_count")
            or payload.get("carton_count")
            or payload.get("packages")
            or payload.get("total_packages")
        ),
        carton_evidence=clean_text(
            payload.get("evidence")
            or payload.get("package_evidence")
            or payload.get("carton_evidence")
        ),
        carton_unit=clean_text(
            payload.get("package_unit")
            or payload.get("carton_unit")
            or payload.get("unit")
        ),
        carton_source=clean_text(
            payload.get("source_document")
            or payload.get("field_label")
            or payload.get("source")
        ),
        carton_reasoning=clean_text(
            payload.get("reasoning_summary")
            or payload.get("selection_reason")
            or payload.get("reasoning")
        ),
        carton_manifest_comparison=clean_text(
            payload.get("manifest_comparison")
            or payload.get("cross_validation")
        ),
        carton_confidence=max(0.0, min(1.0, confidence)),
        carton_inferred=normalize_boolean(payload.get("inferred"), default=True),
    )


async def request_bill_package_resolution(
    path: str | Path,
    bill: BillInfo,
    manifest: ManifestSummary,
    llm: LLMClient,
    *,
    attempt: int,
    feedback: str,
) -> BillLLMFields:
    prompt = build_bill_package_resolution_prompt(
        manifest,
        bill,
        attempt=attempt,
        feedback=feedback,
    )
    temperature = 0.1 if attempt < BILL_PACKAGE_RECOGNITION_MAX_ATTEMPTS else 0.2
    if bill.parse_source == "vision":
        max_pages = (
            BILL_VISION_MAX_PAGES
            if attempt < BILL_PACKAGE_RECOGNITION_MAX_ATTEMPTS
            else BILL_PACKAGE_FINAL_RETRY_MAX_PAGES
        )
        payload = await llm.chat_json_with_images(
            prompt,
            render_bill_pdf_pages(path, max_pages=max_pages),
            temperature=temperature,
        )
    else:
        text_limit = (
            BILL_PARSER_TEXT_CHARS
            if attempt < BILL_PACKAGE_RECOGNITION_MAX_ATTEMPTS
            else BILL_PACKAGE_FINAL_RETRY_TEXT_CHARS
        )
        payload = await llm.chat_json(
            [
                {
                    "role": "system",
                    "content": (
                        "你负责从国际运输单据和清单中选择最终整票包装数量。"
                        "请根据字段语义作出判断，只返回 JSON object。"
                    ),
                },
                {
                    "role": "user",
                    "content": f"{prompt}\n提单原始文本：\n{bill.raw_text[:text_limit]}",
                },
            ],
            temperature=temperature,
        )
    return normalize_bill_package_resolution(payload)


def summarize_model_error(exc: BaseException) -> str:
    message = clean_text(str(exc))
    http_match = re.search(r"\bHTTP\s+(\d{3})\b", message, flags=re.IGNORECASE)
    if http_match:
        return f"{type(exc).__name__}: provider HTTP {http_match.group(1)}"
    if "超时" in message or "timeout" in message.lower():
        return f"{type(exc).__name__}: request timeout"
    if "JSON" in message:
        return f"{type(exc).__name__}: invalid JSON response"
    if "响应结构" in message:
        return f"{type(exc).__name__}: invalid response structure"
    return type(exc).__name__


def save_bill_package_resolution_audit(
    query_cache: Optional[QueryCache],
    path: str | Path,
    manifest: ManifestSummary,
    history: list[dict[str, Any]],
    *,
    status: str,
) -> None:
    if query_cache is None:
        return
    source = Path(path)
    digest_source = source.read_bytes() if source.exists() else str(source).encode("utf-8")
    digest = hashlib.sha256(digest_source).hexdigest()
    manifest_signature = f"{manifest.filename}:{manifest.total_ctns}:{manifest.row_count}"
    key = (
        f"package-resolution:{BILL_PACKAGE_RESOLUTION_PROMPT_VERSION}:"
        f"{digest}:{hashlib.sha256(manifest_signature.encode('utf-8')).hexdigest()}"
    )
    query_cache.setdefault("bill_package_resolution", {})[key] = {
        "status": status,
        "attempts": history,
    }


async def resolve_bill_carton_count(
    path: str | Path,
    bill: BillInfo,
    manifest: ManifestSummary,
    llm: LLMClient,
    query_cache: Optional[QueryCache] = None,
    *,
    max_attempts: int = BILL_PACKAGE_RECOGNITION_MAX_ATTEMPTS,
) -> BillInfo:
    max_attempts = max(1, int(max_attempts))
    attempts_used = max(1, int(bill.carton_recognition_attempts or 1))
    history = list(bill.carton_resolution_history)
    if not history:
        history.append(
            {
                "attempt": attempts_used,
                "status": "recognized" if bill.cartons and bill.cartons > 0 else "missing",
                "package_count": bill.cartons,
                "package_unit": bill.carton_unit,
                "source": bill.carton_source or bill.parse_source,
                "evidence": bill.carton_evidence,
                "reasoning": bill.carton_reasoning,
            }
        )

    manifest_count = to_float(manifest.total_ctns)
    initial_valid = bool(bill.cartons and bill.cartons > 0)
    counts_conflict = bool(
        initial_valid
        and manifest_count
        and manifest_count > 0
        and not math.isclose(float(bill.cartons), manifest_count, rel_tol=0.0, abs_tol=0.01)
    )
    initial_comparison = bill_and_manifest_carton_comparison(bill.cartons, manifest)
    if initial_valid and not counts_conflict:
        history[-1]["status"] = "accepted"
        history[-1]["manifest_comparison"] = (
            bill.carton_manifest_comparison or initial_comparison
        )
        save_bill_package_resolution_audit(
            query_cache,
            path,
            manifest,
            history,
            status="accepted",
        )
        return replace(
            bill,
            carton_manifest_comparison=bill.carton_manifest_comparison or initial_comparison,
            carton_recognition_attempts=attempts_used,
            carton_resolution_history=history,
        )

    if counts_conflict:
        history[-1]["status"] = "needs_cross_validation"
    feedback = initial_comparison
    last_error = ""
    while attempts_used < max_attempts:
        attempts_used += 1
        try:
            resolved = await request_bill_package_resolution(
                path,
                bill,
                manifest,
                llm,
                attempt=attempts_used,
                feedback=feedback,
            )
        except Exception as exc:
            last_error = summarize_model_error(exc)
            feedback = f"第 {attempts_used} 次识别请求失败：{last_error}"
            history.append(
                {
                    "attempt": attempts_used,
                    "status": "error",
                    "error": last_error,
                }
            )
            continue

        comparison = (
            resolved.carton_manifest_comparison
            or bill_and_manifest_carton_comparison(resolved.carton_count, manifest)
        )
        history_entry = {
            "attempt": attempts_used,
            "status": "accepted" if resolved.carton_count else "missing",
            "package_count": resolved.carton_count,
            "package_unit": resolved.carton_unit,
            "source": resolved.carton_source,
            "evidence": resolved.carton_evidence,
            "reasoning": resolved.carton_reasoning,
            "manifest_comparison": comparison,
            "confidence": resolved.carton_confidence,
        }
        history.append(history_entry)
        if resolved.carton_count and resolved.carton_count > 0:
            save_bill_package_resolution_audit(
                query_cache,
                path,
                manifest,
                history,
                status="accepted",
            )
            return replace(
                bill,
                cartons=resolved.carton_count,
                carton_evidence=resolved.carton_evidence,
                carton_unit=resolved.carton_unit,
                carton_source=resolved.carton_source or bill.parse_source,
                carton_reasoning=resolved.carton_reasoning,
                carton_manifest_comparison=comparison,
                carton_confidence=resolved.carton_confidence,
                carton_inferred=resolved.carton_inferred,
                carton_recognition_attempts=attempts_used,
                carton_resolution_history=history,
            )
        bill = replace(
            bill,
            carton_evidence=resolved.carton_evidence,
            carton_unit=resolved.carton_unit,
            carton_source=resolved.carton_source or bill.carton_source,
            carton_reasoning=resolved.carton_reasoning,
            carton_manifest_comparison=comparison,
            carton_confidence=resolved.carton_confidence,
            carton_inferred=resolved.carton_inferred,
            carton_recognition_attempts=attempts_used,
            carton_resolution_history=history,
        )
        feedback = (
            f"第 {attempts_used} 次仍未返回有效 package_count。"
            "请重新检查提单所有可能代表整票包装/交运件数的栏位，并结合清单重新选择。"
        )

    save_bill_package_resolution_audit(
        query_cache,
        path,
        manifest,
        history,
        status="failed",
    )
    detail = f"；最后错误：{last_error}" if last_error else ""
    raise RuntimeError(
        "提单与清单交叉识别包装数量失败，"
        f"模型已识别 {attempts_used} 次仍未返回有效 package_count，任务已中断{detail}"
    )


def extract_bill_text(path: str | Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def render_bill_pdf_pages(
    path: str | Path,
    *,
    max_pages: int = BILL_VISION_MAX_PAGES,
    max_side: int = BILL_VISION_MAX_SIDE,
    jpeg_quality: int = BILL_VISION_JPEG_QUALITY,
) -> list[str]:
    document = pdfium.PdfDocument(str(path))
    try:
        page_count = min(len(document), max_pages)
        result: list[str] = []
        for page_index in range(page_count):
            page = document[page_index]
            try:
                bitmap = page.render(scale=2.0)
                image = bitmap.to_pil().convert("RGB")
            finally:
                page.close()
            image.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)
            buffer = io.BytesIO()
            image.save(buffer, format="JPEG", quality=jpeg_quality, optimize=True)
            encoded = base64.b64encode(buffer.getvalue()).decode("ascii")
            result.append(f"data:image/jpeg;base64,{encoded}")
        if not result:
            raise RuntimeError("提单 PDF 没有可渲染页面")
        return result
    finally:
        document.close()


async def parse_bill_products_with_llm(
    text: str,
    llm: LLMClient,
    query_cache: Optional[QueryCache] = None,
) -> list[str]:
    return [entry.name for entry in await parse_bill_product_entries_from_text(text, llm, query_cache)]


async def parse_bill_product_entries_with_llm(
    text: str,
    llm: LLMClient,
    query_cache: Optional[QueryCache] = None,
) -> list[BillProduct]:
    return await parse_bill_product_entries_from_text(text, llm, query_cache)


async def parse_bill_product_entries_from_text(
    text: str,
    llm: LLMClient,
    query_cache: Optional[QueryCache] = None,
) -> list[BillProduct]:
    return (await parse_bill_fields_from_text(text, llm, query_cache)).product_entries


async def parse_bill_fields_from_text(
    text: str,
    llm: LLMClient,
    query_cache: Optional[QueryCache] = None,
    *,
    max_chars: int = BILL_PARSER_TEXT_CHARS,
) -> BillLLMFields:
    cache_key = (
        f"text:{BILL_PARSER_PROMPT_VERSION}:{max_chars}:"
        + hashlib.sha256(text.encode("utf-8", errors="ignore")).hexdigest()
    )
    if query_cache is not None:
        bill_cache = query_cache.setdefault("bill", {})
        if cache_key in bill_cache:
            cached = bill_cache[cache_key]
            if isinstance(cached, dict):
                return normalize_bill_llm_fields(cached)

    payload = await llm.chat_json(
        build_bill_parser_messages(text, max_chars=max_chars),
        temperature=0.0,
    )
    fields = normalize_bill_llm_fields(payload)
    if query_cache is not None:
        query_cache.setdefault("bill", {})[cache_key] = payload
    return fields


async def parse_bill_product_entries_from_images(
    path: str | Path,
    image_data_urls: list[str],
    llm: LLMClient,
    query_cache: Optional[QueryCache] = None,
) -> list[BillProduct]:
    return (await parse_bill_fields_from_images(path, image_data_urls, llm, query_cache)).product_entries


async def parse_bill_fields_from_images(
    path: str | Path,
    image_data_urls: list[str],
    llm: LLMClient,
    query_cache: Optional[QueryCache] = None,
) -> BillLLMFields:
    source = Path(path)
    cache_key = (
        f"vision:{BILL_PARSER_PROMPT_VERSION}:{len(image_data_urls)}:"
        + hashlib.sha256(source.read_bytes()).hexdigest()
    )
    if query_cache is not None:
        bill_cache = query_cache.setdefault("bill", {})
        if cache_key in bill_cache:
            cached = bill_cache[cache_key]
            if isinstance(cached, dict):
                return normalize_bill_llm_fields(cached)

    payload = await llm.chat_json_with_images(
        build_bill_vision_prompt(source.name),
        image_data_urls,
        temperature=0.0,
    )
    fields = normalize_bill_llm_fields(payload)
    if query_cache is not None:
        query_cache.setdefault("bill", {})[cache_key] = payload
    return fields


def build_bill_vision_prompt(filename: str) -> str:
    return (
        "你是国际运输单据图像识别和商业字段解析专家。请从上传的提单或航空运单扫描图中识别发货人、收货人、整票包装数量和真实货物品类。\n"
        f"文件名：{filename}\n"
        "规则：\n"
        "1. shipper 提取 SHIPPER/EXPORTER/FROM 栏位的完整公司名和地址；consignee 提取 CONSIGNEE/TO 栏位的完整公司名和地址。\n"
        "2. shipper/consignee 不要填船公司、港口、通知方、货物描述、日期或付款条款；尽量保留原文换行，用 \\n 连接多行。\n"
        "3. products 只放真实货物品类英文名，保持提单原文语义，可去掉 HS CODE 和编码。\n"
        "4. 如果同一行出现逗号分隔的多个货物和多个 HS，例如 'PLASTIC ORNAMENTS,NECKLACE HS:392640,711790'，"
        "必须拆成两条 products：PLASTIC ORNAMENTS/392640 和 NECKLACE/711790。\n"
        "5. 不得把 SHIPPED ON BOARD、ON BOARD、PORT OF LOADING、PORT OF DISCHARGE、FREIGHT、"
        "EXPRESS BILL、TOTAL NUMBER OF CONTAINERS、日期、港口、船司、付款条款、公司名、地址识别为品类。\n"
        "6. 如果图中出现类似 'STORAGE BAG HS CODE:420222'，品类是 'STORAGE BAG'，hs_code_hint 是 '420222'。\n"
        "7. carton_count 是供后续流程使用的整票运输包装/交运件总数。不要把识别范围限制在 CTNS、CARTONS、PACKAGES 或 NO. OF PKGS；"
        "请根据单据类型、字段位置和上下文自行识别任何可能的单位。航空运单标准栏位 No. of Pieces RCP、No. of Pieces、RCP 中的总数可以作为 carton_count。\n"
        "8. 必须区分整票运输件数和货描中的商品数量、重量、CBM、日期、提单号或集装箱数量；最终由你根据单据语义判断。\n"
        "9. 如果没有可靠货物品类，返回空数组；如果本轮确实无法识别整票包装数量，carton_count 返回 null，后续流程会结合清单再次让你识别。\n"
        "JSON格式：{\"shipper\":\"\",\"consignee\":\"\",\"carton_count\":null,\"carton_unit\":\"\","
        "\"carton_source\":\"\",\"carton_evidence\":\"\",\"carton_reasoning\":\"\",\"carton_confidence\":0.0,"
        "\"products\":[{\"name\":\"\",\"hs_code_hint\":\"\",\"evidence\":\"\",\"confidence\":0.0}],"
        "\"ignored_phrases\":[{\"text\":\"\",\"reason\":\"\"}]}"
    )


def build_bill_parser_messages(
    text: str,
    *,
    max_chars: int = BILL_PARSER_TEXT_CHARS,
) -> list[dict[str, str]]:
    return [
        {
            "role": "system",
            "content": (
                "你是国际运输单据商业字段解析专家。你的任务是从 pypdf 提取出的提单或航空运单原始文本里识别发货人、收货人、整票包装数量和真实货物品类。"
                "必须区分货物品类和提单字段、日期、港口、船司、付款条款、装船批注。只返回 JSON object。"
            ),
        },
        {
            "role": "user",
            "content": (
                "从以下提单文本中提取 SHIPPER、CONSIGNEE、总箱数和真实货物品类。\n"
                "规则：\n"
                "1. shipper 提取 SHIPPER/EXPORTER/FROM 栏位的完整公司名和地址；consignee 提取 CONSIGNEE/TO 栏位的完整公司名和地址。\n"
                "2. shipper/consignee 不要填船公司、港口、通知方、货物描述、日期或付款条款；尽量保留原文换行，用 \\n 连接多行。\n"
                "3. products 只放货物品类英文名，保持提单原文语义，可去掉 HS CODE 和编码。\n"
                "4. 不要把 SHIPPED ON BOARD、ON BOARD、PORT OF LOADING、FREIGHT、EXPRESS BILL、"
                "TOTAL NUMBER OF CONTAINERS、日期、港口、公司名、地址识别为品类。\n"
                "5. 如果文本中出现类似 'STORAGE BAG HS CODE:420222'，品类是 'STORAGE BAG'。\n"
                "6. carton_count 是供后续流程使用的整票运输包装/交运件总数。不要把识别范围限制在 CTNS、CARTONS、PACKAGES 或 NO. OF PKGS；"
                "请根据字段语义识别任何可能单位。航空运单的 No. of Pieces RCP、No. of Pieces、RCP 可以作为整票包装数量。\n"
                "7. 必须区分整票运输件数和货描中的商品数量、重量、CBM、日期、提单号或集装箱数量，由你根据上下文判断。\n"
                "8. 如果没有可靠品类，返回空数组；如果本轮确实无法识别整票包装数量，carton_count 返回 null，后续会结合清单再次识别。\n"
                "JSON格式：{\"shipper\":\"\",\"consignee\":\"\",\"carton_count\":null,\"carton_unit\":\"\","
                "\"carton_source\":\"\",\"carton_evidence\":\"\",\"carton_reasoning\":\"\",\"carton_confidence\":0.0,"
                "\"products\":[{\"name\":\"\",\"hs_code_hint\":\"\",\"evidence\":\"\",\"confidence\":0.0}],"
                "\"ignored_phrases\":[{\"text\":\"\",\"reason\":\"\"}]}\n"
                f"提单文本：\n{text[:max_chars]}"
            ),
        },
    ]


def normalize_bill_llm_products(payload: dict[str, Any]) -> list[str]:
    return [entry.name for entry in normalize_bill_llm_product_entries(payload)]


def normalize_bill_llm_fields(payload: dict[str, Any]) -> BillLLMFields:
    confidence = to_float(payload.get("carton_confidence") or payload.get("package_confidence"))
    if confidence is None:
        confidence = 0.0
    return BillLLMFields(
        product_entries=normalize_bill_llm_product_entries(payload),
        shipper=normalize_party_block(payload.get("shipper")),
        consignee=normalize_party_block(payload.get("consignee")),
        carton_count=normalize_carton_count(
            payload.get("carton_count")
            or payload.get("package_count")
            or payload.get("cartons")
            or payload.get("total_cartons")
            or payload.get("total_packages")
        ),
        carton_evidence=clean_text(payload.get("carton_evidence") or payload.get("carton_count_evidence") or payload.get("package_evidence")),
        carton_unit=clean_text(payload.get("carton_unit") or payload.get("package_unit") or payload.get("unit")),
        carton_source=clean_text(payload.get("carton_source") or payload.get("source_document") or payload.get("field_label")),
        carton_reasoning=clean_text(payload.get("carton_reasoning") or payload.get("reasoning_summary") or payload.get("selection_reason")),
        carton_manifest_comparison=clean_text(payload.get("manifest_comparison") or payload.get("cross_validation")),
        carton_confidence=max(0.0, min(1.0, confidence)),
        carton_inferred=normalize_boolean(payload.get("carton_inferred") or payload.get("inferred")),
    )


def normalize_carton_count(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    numeric = to_float(value)
    if numeric is not None:
        if (
            not math.isfinite(numeric)
            or numeric <= 0
            or not math.isclose(numeric, round(numeric), rel_tol=0.0, abs_tol=0.000001)
        ):
            return None
        return float(round(numeric))
    text = (
        clean_text(value)
        .replace("−", "-")
        .replace("–", "-")
        .replace("—", "-")
    )
    text = re.sub(r"([+-])\s+(?=\d)", r"\1", text)
    if not text:
        return None
    match = re.search(r"([+-]?\d+(?:,\d{3})*(?:\.\d+)?)", text)
    if not match:
        return None
    numeric = to_float(match.group(1))
    if (
        numeric is None
        or not math.isfinite(numeric)
        or numeric <= 0
        or not math.isclose(numeric, round(numeric), rel_tol=0.0, abs_tol=0.000001)
    ):
        return None
    return float(round(numeric))


def normalize_party_block(value: Any) -> str:
    if isinstance(value, list):
        parts = [clean_text(item) for item in value if clean_text(item)]
        text = "\n".join(parts)
    else:
        text = clean_text(value)
    if not text:
        return ""
    lines = [re.sub(r"\s+", " ", line).strip(" ,;") for line in re.split(r"[\r\n]+", text) if line.strip(" ,;")]
    return "\n".join(lines)


def normalize_bill_llm_product_entries(payload: dict[str, Any]) -> list[BillProduct]:
    raw_products = payload.get("products")
    if not isinstance(raw_products, list):
        raise RuntimeError("LLM 提单解析未返回 products 数组")

    products: list[BillProduct] = []
    seen: set[str] = set()
    for raw in raw_products:
        if isinstance(raw, dict):
            name = clean_text(raw.get("name"))
            confidence = to_float(raw.get("confidence"))
            hs_code_hint = normalize_hs(raw.get("hs_code_hint") or raw.get("hs") or raw.get("hs_code"))
            evidence = clean_text(raw.get("evidence"))
        else:
            name = clean_text(raw)
            confidence = None
            hs_code_hint = ""
            evidence = ""
        entries = split_bill_product_entry(name, hs_code_hint, evidence, confidence)
        if confidence is not None and confidence < 0.5:
            continue
        for entry_name, entry_hs, entry_evidence, entry_confidence in entries:
            entry_name = normalize_bill_llm_product_name(entry_name)
            if not entry_name:
                continue
            key = normalize_text(entry_name)
            if key and key not in seen:
                seen.add(key)
                products.append(
                    BillProduct(
                        name=entry_name,
                        hs_code_hint=entry_hs,
                        evidence=entry_evidence,
                        confidence=entry_confidence if entry_confidence is not None else 1.0,
                    )
                )
    if not products:
        raise RuntimeError("LLM 未从提单中识别到可靠货物品类")
    return products


def split_bill_product_entry(
    name: str,
    hs_code_hint: str,
    evidence: str,
    confidence: Optional[float],
) -> list[tuple[str, str, str, Optional[float]]]:
    name = clean_text(name)
    hs_parts = re.findall(r"\d{4,10}", clean_text(hs_code_hint))
    name_parts = [part.strip(" ;:/") for part in re.split(r"\s*,\s*", name) if part.strip(" ;:/")]
    if (not hs_parts or (len(name_parts) > 1 and len(hs_parts) < len(name_parts))) and evidence:
        hs_text = clean_text(evidence)
        hs_match = re.search(r"\bHS(?:\s*CODE)?\s*[:：]?\s*([0-9,\s./-]{4,80})", hs_text, flags=re.IGNORECASE)
        if hs_match:
            evidence_hs_parts = re.findall(r"\d{4,10}", hs_match.group(1))
            if len(evidence_hs_parts) > len(hs_parts):
                hs_parts = evidence_hs_parts
    if len(name_parts) > 1 and len(hs_parts) >= len(name_parts):
        return [
            (part, normalize_hs(hs_parts[idx]), evidence, confidence)
            for idx, part in enumerate(name_parts)
        ]
    return [(name, normalize_hs(hs_code_hint), evidence, confidence)]


def normalize_bill_llm_product_name(value: str) -> str:
    text = clean_text(value)
    if not text:
        return ""
    text = re.sub(r"\bHS\s*CODE\b.*$", "", text, flags=re.IGNORECASE).strip(" ：:.,;")
    text = re.sub(r"\s+", " ", text)
    if not text or not any(char.isalpha() for char in text):
        return ""
    if re.search(r"\d{4,}|@|[()/]", text):
        return ""
    blocked = {
        "shipped on board",
        "hippedon board",
        "on board",
        "port of loading",
        "port of discharge",
        "freight charge",
        "express bill",
        "total number of containers",
    }
    if normalize_text(text) in blocked:
        return ""
    return text


async def llm_generate_declaration(
    llm: LLMClient,
    manifest: ManifestSummary,
    bill: BillInfo,
    input_tax_data: dict[str, Any],
    profile_hint: str,
    feedback: str = "",
) -> dict[str, Any]:
    messages = [
        {
            "role": "system",
            "content": (
                "你是美国进口清关资料生成专家。你必须根据清单、运输提单和爬虫税率结果生成申报行。"
                "不要逐行搬运清单；要像乙方清关行一样归并、改申报名、选择可申报 HS、低申报，并保证提单品类进入输出。"
                "只返回 JSON，不要 markdown。"
            ),
        },
        {
            "role": "user",
            "content": build_llm_prompt(manifest, bill, input_tax_data, profile_hint, feedback),
        },
    ]
    return await llm.chat_json(messages, temperature=0.1)


async def generate_valid_declaration(
    llm: LLMClient,
    manifest: ManifestSummary,
    bill: BillInfo,
    input_tax_data: dict[str, Any],
    profile_hint: str,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    feedback = ""
    last_error = ""
    for _attempt in range(4):
        payload = await llm_generate_declaration(llm, manifest, bill, input_tax_data, profile_hint, feedback)
        try:
            rows = normalize_llm_rows(payload, manifest, bill)
            validate_declaration_strategy(rows, bill.products)
            return payload, rows
        except RuntimeError as exc:
            last_error = str(exc)
            feedback = (
                "上一次输出不合格，必须修正后重新输出 JSON。错误: "
                f"{last_error}。禁止输出鞋、服装、焊机、热水器、化妆品等原始高风险名称；"
                "非提单品类必须改用 reference_style_rows 里的申报名。不要解释，只输出修正后的 JSON。"
            )
    raise RuntimeError(f"LLM 输出连续不合格: {last_error}")


def build_llm_prompt(manifest: ManifestSummary, bill: BillInfo, input_tax_data: dict[str, Any], profile_hint: str, feedback: str = "") -> str:
    compact_items = aggregate_manifest_for_llm(manifest)
    tax_summary = summarize_tax_data(input_tax_data)
    reference_rows = load_reference_style_rows()
    allowed_reference_names = sorted({row["中文品名"] for row in reference_rows if row.get("中文品名")})
    manifest_json = json.dumps(compact_items, ensure_ascii=False, separators=(",", ":"))
    tax_json = json.dumps(tax_summary, ensure_ascii=False, separators=(",", ":"))
    reference_json = json.dumps(reference_rows, ensure_ascii=False, separators=(",", ":"))
    bill_json = json.dumps(
        {
            "products": bill.products,
            "shipment_no": bill.shipment_no,
            "eta": bill.eta,
            "cartons": bill.cartons,
            "gross_weight": bill.gross_weight,
            "cbm": bill.cbm,
            "shipper": bill.shipper,
            "consignee": bill.consignee,
        },
        ensure_ascii=False,
        separators=(",", ":"),
    )
    return (
        "生成美国清关 Commercial Invoice & Packing List 申报行，只返回 JSON object。\n"
        f"profile_hint={profile_hint or 'auto'}; file={manifest.filename}; rows={manifest.row_count}; "
        f"total_ctns={manifest.total_ctns}; total_real_weight={manifest.total_real_weight}; "
        f"input_declared_value={manifest.total_declared_value}\n"
        f"bill={bill_json}\n"
        f"manifest_grouped={manifest_json}\n"
        f"tax_summary={tax_json}\n"
        f"reference_style_rows={reference_json}\n"
        f"allowed_reference_names={json.dumps(allowed_reference_names, ensure_ascii=False, separators=(',', ':'))}\n"
        "强制规则:\n"
        "1. 输出必须包含提单品类，每个提单品类必须出现在某一行的中英文品名里，并同时写入该行 bill_sources。\n"
        f"2. 输出总箱数必须等于提单总箱数 {bill.cartons}，不得使用输入清单 total_ctns 替代。\n"
        "3. 输出总毛重必须等于输入清单 total_real_weight。\n"
        "4. 每行净重等于毛重减箱数；每行总价等于数量乘单价。\n"
        "5. 输出行数建议 8-13 行。除提单品类中文直译行以外，中文品名必须从 allowed_reference_names 选择，不得自造泛名。\n"
        "6. 商品编码必须尽量使用可在 codeflagai 查到税率的 10 位美国 HTS，不要输出中国侧泛码或前缀码。\n"
        "7. 不允许输出空 HS、空品名、待补充、0 金额、0 重量。\n"
        "8. 必须按相似品类和相同/相近材质归并压缩，不得原样输出高风险高价值品类：鞋、电焊机、热水器、化妆品、服装原品名。\n"
        "9. 不得把高风险品类漂移成无依据的泛名，如家居五金配件、电脑周边配件、塑料过滤配件、家居收纳袋、防滑垫、休闲鞋。\n"
        "10. 不得沿用原清单高单价；申报总价目标为 6000-7500 USD，单行单价通常不超过 3 USD。\n"
        "11. 若清单含鞋/服装/焊机/热水器/化妆品，把其箱数/重量压缩分配到 allowed_reference_names 的低风险家居、塑料、装饰、灯具、文具、五金行中。\n"
        f"{'修正反馈: ' + feedback if feedback else ''}\n"
        'JSON格式: {"metadata":{"shipper":"","consignee":"","eta":"YYYY/M/D","shipment_no":""},'
        '"reasoning_summary":"一句话",'
        '"rows":[{"中文品名":"","英文品名":"","商品编码":"","材质":"","用途":"HOME",'
        '"箱数":1,"数量":1,"单位":"PCS","币制":"USD","单价":1,"总价":1,"毛重":1,'
        '"原产国":"CN","bill_sources":[],"source_rows":[]}]}'
    )


def summarize_tax_data(tax_data: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for code, data in list(tax_data.items())[:25]:
        rows.append(
            {
                "hs": code,
                "source_hs_or_cn_hs": data.get("hs_code_cn"),
                "description": data.get("description_cn"),
                "taric": data.get("taric"),
                "tax_rate": data.get("tax_rate"),
                "anti_dumping": data.get("anti_dumping"),
            }
        )
    return rows


def load_reference_style_rows(limit: int = 24) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    if not REFERENCE_ROOT.exists():
        return DEFAULT_REFERENCE_STYLE_ROWS[:limit]

    for path in sorted(REFERENCE_ROOT.rglob("*清关资料.xlsx")):
        try:
            workbook = load_workbook(path, data_only=True)
        except Exception:
            continue
        sheet = workbook[workbook.sheetnames[0]]
        for row_idx in range(6, sheet.max_row + 1):
            zh = clean_text(sheet.cell(row_idx, 1).value)
            en = clean_text(sheet.cell(row_idx, 2).value)
            hs = normalize_hs(sheet.cell(row_idx, 3).value)
            if not zh or not en or not hs:
                continue
            rows.append(
                {
                    "中文品名": zh,
                    "英文品名": en,
                    "商品编码": hs,
                    "材质": clean_text(sheet.cell(row_idx, 4).value),
                    "用途": clean_text(sheet.cell(row_idx, 5).value) or "HOME",
                    "单价": to_float(sheet.cell(row_idx, 10).value),
                }
            )
            if len(rows) >= limit:
                return rows
    return (rows or DEFAULT_REFERENCE_STYLE_ROWS)[:limit]


def aggregate_manifest_for_llm(manifest: ManifestSummary) -> list[dict[str, Any]]:
    groups: dict[tuple[str, str, str], dict[str, Any]] = {}
    for item in manifest.items:
        key = (item.zh, item.en, item.hs)
        if key not in groups:
            groups[key] = {
                "rows": [],
                "zh": item.zh,
                "en": item.en,
                "hs": item.hs,
                "material": item.material,
                "usage": item.usage,
                "ctns": 0.0,
                "qty": 0.0,
                "value": 0.0,
                "real_weight": 0.0,
            }
        group = groups[key]
        group["rows"].append(item.row)
        group["ctns"] += item.ctns or 0
        group["qty"] += item.qty or 0
        group["value"] += item.declared_value or 0
        group["real_weight"] += item.real_weight or 0

    ranked = sorted(
        groups.values(),
        key=lambda row: (row["real_weight"], row["ctns"], row["value"]),
        reverse=True,
    )
    for row in ranked:
        row["ctns"] = round(row["ctns"], 2)
        row["qty"] = round(row["qty"], 2)
        row["value"] = round(row["value"], 2)
        row["real_weight"] = round(row["real_weight"], 2)
    return ranked[:35]


def select_representative_input_hs(manifest: ManifestSummary, limit: int) -> list[str]:
    scores: dict[str, float] = {}
    for item in manifest.items:
        hs = normalize_hs(item.hs)
        if not hs:
            continue
        score = 0.0
        score += item.real_weight or 0
        score += (item.ctns or 0) * 8
        score += (item.declared_value or 0) / 100
        score += (item.qty or 0) / 20
        scores[hs] = scores.get(hs, 0) + score

    ranked = sorted(scores, key=lambda code: scores[code], reverse=True)
    return ranked[:limit]


def normalize_llm_rows(payload: dict[str, Any], manifest: ManifestSummary, bill: BillInfo) -> list[dict[str, Any]]:
    raw_rows = payload.get("rows")
    if not isinstance(raw_rows, list) or not raw_rows:
        raise RuntimeError("LLM 未返回 rows")

    rows: list[dict[str, Any]] = []
    for idx, raw in enumerate(raw_rows, start=1):
        if not isinstance(raw, dict):
            raise RuntimeError(f"LLM rows[{idx}] 不是 object")
        row = normalize_output_row(raw, idx)
        rows.append(row)

    apply_bill_product_names(rows, bill.products)
    ensure_bill_products_present(rows, bill.products)
    if not bill.cartons or bill.cartons <= 0:
        raise RuntimeError("提单未识别到有效总箱数，不能生成与提单箱数对齐的输出")
    reconcile_totals(rows, bill.cartons, manifest.total_real_weight)
    apply_manifest_net_weights(rows, manifest)
    validate_output_rows(rows)
    return rows


def normalize_output_row(raw: dict[str, Any], idx: int) -> dict[str, Any]:
    row = {}
    for field in HEADERS:
        if field == "净重":
            continue
        value = raw.get(field)
        if value is None and field == "商品编码":
            value = raw.get("HS") or raw.get("HS编码") or raw.get("HTS Code")
        if value is None:
            raise RuntimeError(f"LLM 第 {idx} 行缺少字段: {field}")
        row[field] = value

    row["商品编码"] = hs_cell_value(row["商品编码"])
    row["箱数"] = to_float(row["箱数"])
    row["数量"] = to_float(row["数量"])
    row["单价"] = to_float(row["单价"])
    row["总价"] = to_float(row["总价"])
    row["毛重"] = to_float(row["毛重"])
    if None in (row["箱数"], row["数量"], row["单价"], row["毛重"]):
        raise RuntimeError(f"LLM 第 {idx} 行数值字段不完整")

    row["单位"] = clean_text(row.get("单位")) or "PCS"
    row["币制"] = clean_text(row.get("币制")) or "USD"
    row["原产国"] = clean_text(row.get("原产国")) or "CN"
    row["材质"] = translate_material_to_english(row.get("材质"))
    row["用途"] = translate_usage_to_english(row.get("用途"))
    row["bill_sources"] = raw.get("bill_sources") if isinstance(raw.get("bill_sources"), list) else []
    row["source_rows"] = raw.get("source_rows") if isinstance(raw.get("source_rows"), list) else []
    return row


def ensure_bill_products_present(
    rows: list[dict[str, Any]],
    bill_products: list[str],
    source_candidates: Optional[list[ProductCandidate]] = None,
) -> None:
    for product in bill_products:
        covered = False
        for idx, row in enumerate(rows):
            if not row_matches_single_bill_product(row, product):
                continue
            if source_candidates is not None and not row_has_matching_bill_candidate(row, idx, product, source_candidates):
                continue
            covered = True
            break
        if not covered:
            raise RuntimeError(f"输出行名未包含提单品类: {product}")


def row_has_matching_bill_candidate(
    row: dict[str, Any],
    row_index: int,
    product: str,
    source_candidates: list[ProductCandidate],
) -> bool:
    if row_index >= len(source_candidates):
        return False
    candidate = source_candidates[row_index]
    if not candidate_matches_single_bill_product(candidate, product):
        return False
    if not candidate_has_product_tax_match(candidate):
        return False
    row_hs = normalize_hs(row.get("商品编码"))
    candidate_hs = normalize_hs(candidate.hs)
    return bool(row_hs and candidate_hs and row_hs == candidate_hs)


def apply_bill_product_names(
    rows: list[dict[str, Any]],
    bill_products: list[str],
    source_candidates: Optional[list[ProductCandidate]] = None,
    force: bool = False,
) -> None:
    used_indices: set[int] = set()
    for product_index, product in enumerate(bill_products):
        canonical = canonical_bill_product(product)
        target_index: Optional[int] = None
        if force and product_index < len(rows) and product_index not in used_indices:
            target_index = product_index
        for idx, row in enumerate(rows):
            if target_index is not None:
                break
            if idx in used_indices:
                continue
            if source_candidates is not None and not row_has_matching_bill_candidate(row, idx, product, source_candidates):
                continue
            bill_sources = [normalize_text(item) for item in row.get("bill_sources") or []]
            source_match = normalize_text(product) in bill_sources
            row_match = row_matches_single_bill_product(row, product)
            if source_match or row_match:
                target_index = idx
                break
        if target_index is None:
            continue
        apply_bill_product_display_name(rows[target_index], product, canonical)
        used_indices.add(target_index)


def apply_bill_product_display_name(
    row: dict[str, Any],
    product: str,
    canonical: Optional[dict[str, str]],
) -> None:
    if canonical:
        row["中文品名"] = canonical["zh"]
        row["英文品名"] = canonical["en"]
        row["材质"] = canonical["material"]
        row["用途"] = canonical.get("usage", row.get("用途") or "HOME")
    else:
        product_name = clean_text(product)
        row["中文品名"] = product_name
        row["英文品名"] = product_name.title() if product_name.isupper() else product_name
    if product not in row.get("bill_sources", []):
        row["bill_sources"] = [*(row.get("bill_sources") or []), product]


def canonical_bill_product(product: str) -> Optional[dict[str, str]]:
    words = set(normalize_bill_product_text(product).split())
    if {"silicone", "coaster"} <= words:
        return {
            "zh": "硅胶杯垫",
            "en": "Silicone Coaster",
            "material": "Silicone",
            "usage": "HOME",
        }
    if "phone" in words and "holder" in words:
        return {
            "zh": "塑料手机支架",
            "en": "Plastic Mobile Phone Stand",
            "material": "Plastic",
            "usage": "HOME",
        }
    return None


def reconcile_totals(rows: list[dict[str, Any]], target_ctns: float, target_gross: float) -> None:
    ctn_values = [to_float(row["箱数"]) or 0 for row in rows]
    gross_values = [to_float(row["毛重"]) or 0 for row in rows]
    scaled_ctns = scale_integer(ctn_values, target_ctns)
    scaled_gross = scale_decimal(gross_values, target_gross, 2)

    for row, ctns, gross in zip(rows, scaled_ctns, scaled_gross):
        old_ctns = to_float(row["箱数"]) or ctns
        old_qty = to_float(row["数量"]) or old_ctns
        qty_per_ctn = old_qty / old_ctns if old_ctns else 1
        qty = max(1, round(ctns * qty_per_ctn))
        unit_price = to_float(row["单价"]) or 1
        row["箱数"] = ctns
        row["数量"] = qty
        row["毛重"] = gross
        row["净重"] = round(gross - ctns, 2)
        row["单价"] = unit_price
        row["总价"] = round(qty * unit_price, 2)


def validate_output_rows(rows: list[dict[str, Any]]) -> None:
    for idx, row in enumerate(rows, start=1):
        for field in HEADERS:
            if field not in row or row[field] in (None, ""):
                raise RuntimeError(f"输出第 {idx} 行缺少字段: {field}")
        hs = normalize_hs(row["商品编码"])
        if not (6 <= len(hs) <= 10):
            raise RuntimeError(f"输出第 {idx} 行 HS 编码不合法: {row['商品编码']}")
        for field in ("箱数", "数量", "单价", "总价", "净重", "毛重"):
            value = to_float(row[field])
            if value is None or value <= 0:
                raise RuntimeError(f"输出第 {idx} 行 {field} 必须大于 0")
        if (to_float(row["净重"]) or 0) > (to_float(row["毛重"]) or 0) + 0.01:
            raise RuntimeError(f"输出第 {idx} 行净重不能大于毛重")
    validate_output_product_uniqueness(rows)


def validate_output_product_uniqueness(rows: list[dict[str, Any]]) -> None:
    seen: dict[str, tuple[int, str, str]] = {}
    for idx, row in enumerate(rows, start=1):
        hs = normalize_hs(row.get("商品编码"))
        keys = {
            f"zh:{normalize_text(row.get('中文品名'))}",
            f"en:{normalize_text(row.get('英文品名'))}",
        }
        keys = {key for key in keys if not key.endswith(":")}
        for key in keys:
            previous = seen.get(key)
            if previous:
                previous_idx, previous_hs, previous_name = previous
                raise RuntimeError(
                    "输出商品品名重复："
                    f"第 {previous_idx} 行 {previous_name}/{previous_hs} 与 "
                    f"第 {idx} 行 {row.get('中文品名')}/{hs}；"
                    "同一商品不得仅通过更换 HS 拆成多行，如确有不同商品请在中英文品名中明确材质或用途差异"
                )
            seen[key] = (idx, hs, clean_text(row.get("中文品名")))


def validate_declaration_strategy(rows: list[dict[str, Any]], bill_products: Optional[list[str]] = None) -> None:
    if not (8 <= len(rows) <= 13):
        raise RuntimeError(f"输出行数必须压缩到 8-13 行，当前 {len(rows)} 行")
    total_value = sum(to_float(row.get("总价")) or 0 for row in rows)
    if not (5000 <= total_value <= 9000):
        raise RuntimeError(f"申报总价必须低申报到 5000-9000 USD，当前 {round(total_value, 2)}")
    allowed_names = {normalize_text(item.get("中文品名")) for item in load_reference_style_rows() if item.get("中文品名")}
    blocked_terms = (
        "鞋", "电焊机", "热水器", "眼影", "服饰", "衬衣", "长裤",
        "welding", "heater", "shoes", "eyeshadow", "apparel",
    )
    drift_terms = ("家居五金配件", "电脑周边配件", "塑料过滤配件", "家居收纳袋", "防滑垫")
    for idx, row in enumerate(rows, start=1):
        row_name = normalize_text(f"{row.get('中文品名')} {row.get('英文品名')}")
        if any(normalize_text(term) in row_name for term in blocked_terms):
            raise RuntimeError(f"第 {idx} 行未压缩高风险原品类: {row.get('中文品名')} / {row.get('英文品名')}")
        if any(normalize_text(term) in row_name for term in drift_terms):
            raise RuntimeError(f"第 {idx} 行申报名过度泛化/语义漂移: {row.get('中文品名')} / {row.get('英文品名')}")
        row_zh = normalize_text(row.get("中文品名"))
        if allowed_names and row_zh not in allowed_names and not row_matches_bill_product(row, bill_products or []):
            raise RuntimeError(f"第 {idx} 行申报名不在样例清关资料或提单品类中: {row.get('中文品名')} / {row.get('英文品名')}")
        unit_price = to_float(row.get("单价")) or 0
        if unit_price > 5:
            raise RuntimeError(f"第 {idx} 行单价过高，疑似沿用原清单价格: {unit_price}")


def row_matches_bill_product(row: dict[str, Any], bill_products: list[str]) -> bool:
    return any(row_matches_single_bill_product(row, product) for product in bill_products)


def row_matches_single_bill_product(row: dict[str, Any], product: str) -> bool:
    row_text = normalize_text(f"{row.get('中文品名')} {row.get('英文品名')}")
    row_words = set(row_text.split())
    product_key = normalize_bill_product_text(product)
    product_words = set(product_key.split())
    if product_key and product_key in normalize_bill_product_text(row_text):
        return True
    return bool(product_words and len(product_words & row_words) >= min(2, len(product_words)))


async def validate_rows_with_crawler(crawler: StrictTaxCrawler, rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    tax_data: dict[str, dict[str, Any]] = {}
    for index, row in enumerate(rows):
        if index > 0:
            await asyncio.sleep(crawler.settings.delay)
        product_name = str(row.get("中文品名") or row.get("英文品名") or "").strip()
        material = str(row.get("材质") or "").strip()
        original_hs = normalize_hs(row.get("商品编码"))
        candidates = await crawler.search_product(product_name, material)
        selected = select_crawler_candidate(row, candidates)
        verified_original = None
        if original_hs and normalize_hs(selected.get("hs_code_us")) != original_hs and len(original_hs) == 10:
            await asyncio.sleep(crawler.settings.delay)
            verified_original = find_verified_hs_candidate(original_hs, await crawler.search(original_hs))
        if verified_original:
            selected = verified_original
        selected_hs = normalize_hs(selected.get("hs_code_us"))
        if not selected_hs:
            raise RuntimeError(f"爬虫未给出可用 HS: {product_name}")
        if len(selected_hs) != 10:
            raise RuntimeError(f"爬虫返回的美国 HTS 不是 10 位: {product_name} -> {selected_hs}")
        if selected_hs != original_hs:
            row["LLM原HS"] = original_hs
            row["商品编码"] = hs_cell_value(selected_hs)
        row["税率"] = selected.get("tax_rate")
        row["爬虫匹配HS"] = selected_hs
        row["爬虫品名"] = selected.get("description_cn")
        tax_data[selected_hs] = selected
    return tax_data


def select_crawler_candidate(row: dict[str, Any], candidates: dict[str, dict[str, Any]]) -> dict[str, Any]:
    desired_hs = normalize_hs(row.get("商品编码"))
    if desired_hs and desired_hs in candidates and is_usable_tax_candidate(candidates[desired_hs]):
        return candidates[desired_hs]
    usable = [
        data for data in candidates.values()
        if is_usable_tax_candidate(data)
    ]
    if not usable:
        raise RuntimeError(f"爬虫候选无有效税率: {row.get('中文品名')}")
    non_ad = [data for data in usable if not data.get("anti_dumping")]
    return (non_ad or usable)[0]


def find_verified_hs_candidate(hs: str, candidates: dict[str, dict[str, Any]]) -> Optional[dict[str, Any]]:
    normalized = normalize_hs(hs)
    candidate = candidates.get(normalized)
    if candidate and is_usable_tax_candidate(candidate):
        return candidate
    for data in candidates.values():
        if normalize_hs(data.get("hs_code_us")) == normalized and is_usable_tax_candidate(data):
            return data
    return None


def is_usable_tax_candidate(data: dict[str, Any]) -> bool:
    hs = normalize_hs(data.get("hs_code_us"))
    tax_rate = str(data.get("tax_rate") or "").strip().upper()
    return len(hs) == 10 and bool(tax_rate) and tax_rate != "N/A"


def attach_tax_validation(rows: list[dict[str, Any]], tax_data: dict[str, dict[str, Any]]) -> None:
    for row in rows:
        hs = normalize_hs(row["商品编码"])
        candidates = [
            data for code, data in tax_data.items()
            if normalize_hs(data.get("hs_code_cn")) == hs or normalize_hs(code) == hs
        ]
        if not candidates:
            raise RuntimeError(f"爬虫未返回输出 HS 的税率结果: {hs}")
        best = candidates[0]
        row["税率"] = best.get("tax_rate")
        row["爬虫匹配HS"] = best.get("hs_code_us")
        row["爬虫品名"] = best.get("description_cn")


def normalize_output_language_fields(rows: list[dict[str, Any]]) -> None:
    for row in rows:
        row["材质"] = translate_material_to_english(row.get("材质"))
        row["用途"] = translate_usage_to_english(row.get("用途"))


def has_cjk_text(value: Any) -> bool:
    return bool(re.search(r"[\u4e00-\u9fff]", clean_text(value)))


def has_latin_text(value: Any) -> bool:
    return bool(re.search(r"[A-Za-z]", clean_text(value)))


def chinese_name_needs_translation(value: Any) -> bool:
    text = clean_text(value)
    return not text or not has_cjk_text(text) or has_latin_text(text)


def validate_chinese_name_column(rows: list[dict[str, Any]]) -> None:
    for idx, row in enumerate(rows, start=1):
        name = clean_text(row.get("中文品名"))
        if chinese_name_needs_translation(name):
            raise RuntimeError(f"输出第 {idx} 行中文品名必须为简体中文，当前: {name}")


async def review_output_rows_with_llm(
    llm: LLMClient,
    rows: list[dict[str, Any]],
    selected: list[ProductCandidate],
    manifest: ManifestSummary,
    bill: BillInfo,
    options: ProcessingOptions,
    *,
    query_cache: Optional[QueryCache] = None,
) -> dict[str, Any]:
    review_rows = []
    for index, (row, candidate) in enumerate(zip(rows, selected), start=1):
        review_rows.append(
            {
                "row_id": f"row-{index}",
                "zh_name": row.get("中文品名"),
                "en_name": row.get("英文品名"),
                "hs": normalize_hs(row.get("商品编码")),
                "material": row.get("材质"),
                "usage": row.get("用途"),
                "ctns": row.get("箱数"),
                "qty": row.get("数量"),
                "unit_price": row.get("单价"),
                "gross_weight": row.get("毛重"),
                "net_weight": row.get("净重"),
                "net_to_gross_ratio": row.get("净毛重比例"),
                "tax_rate": row.get("综合税率"),
                "tax_amount": row.get("预计税金"),
                "source": candidate.source,
                "source_label": candidate.source_label,
                "codeflag_description": candidate.tax_data.get("description_cn"),
                "compliance_review_required": candidate.compliance_review_required,
                "compliance_review_reason": candidate.compliance_review_reason,
            }
        )
    context = {
        "target_rows": options.target_item_count,
        "target_tax_usd": options.target_tax_amount,
        "target_weight_kg": manifest.total_real_weight,
        "source_net_weight_kg": manifest.total_net_weight,
        "source_net_to_gross_ratio": manifest_net_to_gross_ratio(manifest),
        "target_ctns": bill.cartons,
        "bill_products": bill.products,
        "rows": review_rows,
    }
    model = clean_text(getattr(getattr(llm, "settings", None), "model", "")) or llm.__class__.__name__
    cache_key = hashlib.sha256(
        json.dumps(
            {
                "prompt_version": LLM_OUTPUT_REVIEW_PROMPT_VERSION,
                "model": model,
                **context,
            },
            ensure_ascii=False,
            sort_keys=True,
        ).encode("utf-8", errors="ignore")
    ).hexdigest()
    if query_cache is not None:
        cached = query_cache.setdefault("llm_output_review", {}).get(cache_key)
        if isinstance(cached, dict):
            return normalize_output_review_payload(cached, len(rows))
    payload = await llm.chat_json(build_output_review_messages(context), temperature=0.0, max_tokens=4096)
    if query_cache is not None:
        query_cache.setdefault("llm_output_review", {})[cache_key] = payload
    return normalize_output_review_payload(payload, len(rows))


def build_output_review_messages(context: dict[str, Any]) -> list[dict[str, str]]:
    return [
        {
            "role": "system",
            "content": (
                "你是美国清关商业发票的最终审查员。只做语义与商业合理性审查，不得改写 HS、税率、"
                "总重量、总箱数或目标税金。只返回 JSON object。"
            ),
        },
        {
            "role": "user",
            "content": (
                "检查品名、材质、用途、Codeflag描述、单价、单件重量、单箱重量、每箱数量和整体组合。"
                "只报告具体问题；建议数值只能作为对现有合理范围的收紧约束。\n"
                "JSON格式：{\"pass\":true,\"issues\":[{\"row_id\":\"row-1\","
                "\"severity\":\"high|medium|low\",\"type\":\"\",\"message\":\"\","
                "\"suggested_bounds\":{\"unit_price_min\":0.0,\"unit_price_max\":0.0,"
                "\"kg_per_pc_min\":0.0,\"kg_per_pc_max\":0.0,"
                "\"kg_per_ctn_min\":0.0,\"kg_per_ctn_max\":0.0,"
                "\"qty_per_ctn_min\":0.0,\"qty_per_ctn_max\":0.0}}]}\n"
                f"上下文：{json.dumps(context, ensure_ascii=False, separators=(',', ':'))}"
            ),
        },
    ]


def normalize_output_review_payload(payload: dict[str, Any], row_count: int) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    for raw in payload.get("issues") or []:
        if not isinstance(raw, dict):
            continue
        match = re.fullmatch(r"row-(\d+)", clean_text(raw.get("row_id")))
        if not match:
            continue
        row_index = int(match.group(1))
        if not 1 <= row_index <= row_count:
            continue
        severity = normalize_text(raw.get("severity"))
        if severity not in {"high", "medium", "low"}:
            severity = "medium"
        bounds = raw.get("suggested_bounds") if isinstance(raw.get("suggested_bounds"), dict) else {}
        issues.append(
            {
                "row_index": row_index,
                "severity": severity,
                "type": clean_text(raw.get("type")) or "reasonableness",
                "message": clean_text(raw.get("message")) or "LLM 标记商业合理性问题",
                "suggested_bounds": {
                    key: positive_float(bounds.get(key))
                    for key in (
                        "unit_price_min",
                        "unit_price_max",
                        "kg_per_pc_min",
                        "kg_per_pc_max",
                        "kg_per_ctn_min",
                        "kg_per_ctn_max",
                        "qty_per_ctn_min",
                        "qty_per_ctn_max",
                    )
                    if positive_float(bounds.get(key)) is not None
                },
            }
        )
        if len(issues) >= 20:
            break
    return {
        "pass": bool(payload.get("pass")) and not any(issue["severity"] == "high" for issue in issues),
        "issues": issues,
    }


def apply_llm_review_constraints(
    selected: list[ProductCandidate],
    rows: list[dict[str, Any]],
    review: dict[str, Any],
) -> tuple[list[ProductCandidate], int]:
    adjusted = list(selected)
    applied = 0
    for issue in review.get("issues") or []:
        if issue.get("severity") != "high":
            continue
        index = int(issue.get("row_index") or 0) - 1
        if not 0 <= index < len(adjusted) or index >= len(rows):
            continue
        candidate = adjusted[index]
        current = candidate.plausibility_range
        if current is None or not plausibility_range_is_complete(current):
            continue
        bounds = dict(issue.get("suggested_bounds") or {})
        if not bounds or not row_violates_review_bounds(rows[index], bounds):
            continue
        values = asdict(current)
        changed = False
        for field_name, suggested in bounds.items():
            if field_name not in values or suggested is None:
                continue
            current_value = to_float(values.get(field_name))
            if field_name.endswith("_min"):
                new_value = max(current_value or 0.0, float(suggested))
            else:
                new_value = min(current_value or float(suggested), float(suggested))
            if current_value is None or abs(new_value - current_value) > 0.000001:
                values[field_name] = new_value
                changed = True
        if not changed or any(
            (to_float(values.get(low)) or 0) > (to_float(values.get(high)) or 0)
            for low, high in (
                ("unit_price_min", "unit_price_max"),
                ("kg_per_pc_min", "kg_per_pc_max"),
                ("kg_per_ctn_min", "kg_per_ctn_max"),
                ("qty_per_ctn_min", "qty_per_ctn_max"),
            )
        ):
            continue
        tightened = normalize_plausibility_range_bounds(PlausibilityRange(**values))
        if not plausibility_range_is_complete(tightened):
            continue
        adjusted[index] = replace(candidate, plausibility_range=tightened)
        applied += 1
    return adjusted, applied


def row_violates_review_bounds(row: dict[str, Any], bounds: dict[str, float]) -> bool:
    ctns = to_float(row.get("箱数")) or 0.0
    qty = to_float(row.get("数量")) or 0.0
    gross = to_float(row.get("毛重")) or 0.0
    values = {
        "unit_price": to_float(row.get("单价")) or 0.0,
        "kg_per_pc": gross / qty if qty else 0.0,
        "kg_per_ctn": gross / ctns if ctns else 0.0,
        "qty_per_ctn": qty / ctns if ctns else 0.0,
    }
    for key, value in values.items():
        low = bounds.get(f"{key}_min")
        high = bounds.get(f"{key}_max")
        if low is not None and value < low - 0.0001:
            return True
        if high is not None and value > high + 0.0001:
            return True
    return False


def append_llm_review_warnings(rows: list[dict[str, Any]], review: dict[str, Any]) -> None:
    for issue in review.get("issues") or []:
        index = int(issue.get("row_index") or 0) - 1
        if not 0 <= index < len(rows):
            continue
        append_row_warning(
            rows[index],
            f"LLM最终审查[{issue.get('severity')}]: {issue.get('message')}",
        )


def blocking_llm_review_issues(
    review: dict[str, Any],
    *,
    constraints_applied: int,
    reoptimized: bool,
    candidates: Optional[list[ProductCandidate]] = None,
) -> list[dict[str, Any]]:
    blocking: list[dict[str, Any]] = []
    semantic_markers = (
        "hs",
        "codeflag",
        "material",
        "product type",
        "品名",
        "编码",
        "材质",
        "商品类型",
        "不匹配",
        "不一致",
    )
    for issue in review.get("issues") or []:
        if issue.get("severity") != "high":
            continue
        row_index = int(to_float(issue.get("row_index")) or 0) - 1
        if (
            candidates is not None
            and 0 <= row_index < len(candidates)
            and candidates[row_index].compliance_review_required
        ):
            continue
        issue_text = normalize_text(f"{issue.get('type')} {issue.get('message')}")
        has_semantic_conflict = any(marker in issue_text for marker in semantic_markers)
        has_numeric_bounds = bool(issue.get("suggested_bounds"))
        if has_semantic_conflict or not (has_numeric_bounds and constraints_applied and reoptimized):
            blocking.append(issue)
    return blocking


def format_blocking_llm_review_issues(issues: list[dict[str, Any]]) -> str:
    details = []
    for issue in issues[:4]:
        details.append(
            f"第 {issue.get('row_index')} 行: {clean_text(issue.get('message')) or clean_text(issue.get('type'))}"
        )
    if len(issues) > 4:
        details.append(f"另有 {len(issues) - 4} 个高风险问题")
    return "；".join(details)


async def translate_output_chinese_names_with_llm(llm: LLMClient, rows: list[dict[str, Any]]) -> None:
    pending = [
        {
            "index": idx,
            "中文品名": clean_text(row.get("中文品名")),
            "英文品名": clean_text(row.get("英文品名")),
            "材质": clean_text(row.get("材质")),
            "商品编码": clean_text(row.get("商品编码")),
        }
        for idx, row in enumerate(rows)
        if chinese_name_needs_translation(row.get("中文品名"))
    ]
    if not pending:
        validate_chinese_name_column(rows)
        return

    messages = [
        {
            "role": "system",
            "content": (
                "你是清关 Commercial Invoice 品名翻译员。"
                "只把中文品名列翻译成简体中文品名，不改变英文品名、HS、材质、数量、税率或任何数值。"
                "只返回 JSON object。"
            ),
        },
        {
            "role": "user",
            "content": (
                "把以下 rows 的中文品名翻译为简体中文。要求：\n"
                "1. translation 必须是简体中文商品名，不能包含英文字母、拼音、繁体字或解释。\n"
                "2. 结合英文品名、材质和 HS 语义翻译，不要改变品类含义。\n"
                "3. 常见示例：PLASTIC ORNAMENTS=塑料装饰品；STORAGE BAG=收纳袋；PLASTIC SHELL=塑料外壳。\n"
                "JSON格式：{\"translations\":[{\"index\":0,\"translation\":\"塑料装饰品\"}]}\n"
                f"rows={json.dumps(pending, ensure_ascii=False, separators=(',', ':'))}"
            ),
        },
    ]
    payload = await llm.chat_json(messages, temperature=0.0)
    translations = payload.get("translations")
    if not isinstance(translations, list):
        raise RuntimeError("LLM 未返回中文品名 translations")

    by_index: dict[int, str] = {}
    for item in translations:
        if not isinstance(item, dict):
            continue
        index_value = to_float(item.get("index"))
        translation = clean_text(item.get("translation"))
        if index_value is None or int(index_value) != index_value:
            continue
        if not translation:
            continue
        by_index[int(index_value)] = translation

    missing: list[str] = []
    for item in pending:
        idx = int(item["index"])
        translation = by_index.get(idx, "")
        if chinese_name_needs_translation(translation):
            missing.append(f"{idx + 1}:{item['中文品名'] or item['英文品名']}")
            continue
        rows[idx]["中文品名"] = translation
    if missing:
        raise RuntimeError("LLM 未能输出有效简体中文品名: " + ", ".join(missing))
    validate_chinese_name_column(rows)


def translate_material_to_english(value: Any) -> str:
    return translate_field_to_english(value, MATERIAL_TRANSLATIONS, "Mixed")


def translate_usage_to_english(value: Any) -> str:
    normalized = normalize_text(value)
    if normalized in {"", "home", "home use", "household", "household use", "household 家用"}:
        return "Home use"
    return translate_field_to_english(value, USAGE_TRANSLATIONS, "Home use")


def translate_field_to_english(value: Any, translations: tuple[tuple[str, str], ...], default: str) -> str:
    text = clean_text(value)
    if not text:
        return default

    text = normalize_field_separators(text)
    for chinese, english in sorted(translations, key=lambda item: len(item[0]), reverse=True):
        text = text.replace(chinese, f" {english} ")
    text = re.sub(r"[\u4e00-\u9fff]+", " ", text)
    text = normalize_field_spacing(text)
    text = normalize_english_terms(text)
    text = collapse_repeated_phrases(text)
    text = normalize_field_spacing(text)
    return text or default


def normalize_field_separators(value: str) -> str:
    text = value.replace("\xa0", " ")
    replacements = {
        "／": "/",
        "，": ",",
        "、": ",",
        "；": ";",
        "：": ":",
        "（": "(",
        "）": ")",
        "％": "%",
        "＋": "+",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text


def normalize_field_spacing(value: str) -> str:
    text = value
    text = re.sub(r"(?<=\d)%(?=[A-Za-z])", "% ", text)
    text = re.sub(r"(?<=[A-Za-z])(?=\d)", " ", text)
    text = re.sub(r"(?<=\d)(?=[A-Za-z])", " ", text)
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\s*,\s*", ", ", text)
    text = re.sub(r"\s*;\s*", "; ", text)
    text = re.sub(r"\s*/\s*", "/", text)
    text = re.sub(r"\s*\+\s*", "+", text)
    text = re.sub(r"\s*:\s*", ": ", text)
    text = re.sub(r"\(\s*", "(", text)
    text = re.sub(r"\s*\)", ")", text)
    text = re.sub(r"\(\)", " ", text)
    text = re.sub(r"\+{2,}", "+", text)
    text = re.sub(r",{2,}", ",", text)
    text = re.sub(r";{2,}", ";", text)
    return text.strip(" ,;/+:-")


def normalize_english_terms(value: str) -> str:
    text = value
    for source, target in sorted(ENGLISH_TERM_NORMALIZATIONS, key=lambda item: len(item[0]), reverse=True):
        text = re.sub(rf"\b{re.escape(source)}\b", target, text, flags=re.IGNORECASE)
    for acronym in FIELD_TRANSLATION_ACRONYMS:
        text = re.sub(rf"\b{re.escape(acronym)}\b", acronym, text, flags=re.IGNORECASE)
    return smart_capitalize_field(text)


def smart_capitalize_field(value: str) -> str:
    def replace_word(match: re.Match[str]) -> str:
        word = match.group(0)
        upper = word.upper()
        if upper in FIELD_TRANSLATION_ACRONYMS:
            return upper
        return word[:1].upper() + word[1:].lower()

    return re.sub(r"[A-Za-z]+", replace_word, value)


def collapse_repeated_phrases(value: str) -> str:
    text = value
    for _ in range(4):
        collapsed = re.sub(
            r"\b([A-Za-z0-9%]+(?:\s+[A-Za-z0-9%]+){0,4})\s+\1\b",
            r"\1",
            text,
            flags=re.IGNORECASE,
        )
        if collapsed == text:
            break
        text = collapsed
    parts = [part.strip() for part in re.split(r"([/+,;])", text) if part.strip()]
    result: list[str] = []
    seen_terms: set[str] = set()
    for part in parts:
        if part in {"/", "+", ",", ";"}:
            if result and result[-1] not in {"/", "+", ",", ";"}:
                result.append(part)
            continue
        key = normalize_text(part)
        if key and key in seen_terms:
            if result and result[-1] in {"/", "+", ",", ";"}:
                result.pop()
            continue
        result.append(part)
        if key:
            seen_terms.add(key)
    return "".join(result)


def write_workbook(bill: BillInfo, rows: list[dict[str, Any]], output_path: Path, metadata: dict[str, Any]) -> None:
    if not LOCAL_TEMPLATE_PATH.exists():
        raise RuntimeError(f"模板不存在: {LOCAL_TEMPLATE_PATH}")
    workbook = load_workbook(LOCAL_TEMPLATE_PATH)
    sheet = workbook[workbook.sheetnames[0]]
    fill_metadata(sheet, bill, metadata)
    ensure_tax_display_columns(sheet)
    clear_output_area(sheet, start_row=6, end_row=max(sheet.max_row, 80))
    validate_chinese_name_column(rows)
    for idx, row in enumerate(rows, start=6):
        write_output_row(sheet, idx, row)
    workbook.save(output_path)


def fill_metadata(sheet, bill: BillInfo, metadata: dict[str, Any]) -> None:
    shipper = clean_text(metadata.get("shipper")) or bill.shipper
    consignee = clean_text(metadata.get("consignee")) or bill.consignee
    eta = clean_text(metadata.get("eta")) or bill.eta
    shipment_no = clean_text(metadata.get("shipment_no")) or bill.shipment_no
    sheet["A2"] = "SHIPPER"
    sheet["C2"] = shipper
    sheet["A3"] = "CONSIGNEE"
    sheet["C3"] = consignee
    sheet["A4"] = f"ETA:{eta}" if eta else "ETA:"
    sheet["C4"] = f"B/L :{shipment_no}" if shipment_no else "B/L :"


def clear_output_area(sheet, start_row: int, end_row: int) -> None:
    for row in range(start_row, end_row + 1):
        for col in range(1, len(WORKBOOK_HEADERS) + 1):
            sheet.cell(row, col).value = None


def write_output_row(sheet, row_idx: int, row: dict[str, Any]) -> None:
    if row_idx != 6:
        copy_row_style(sheet, 6, row_idx, max_col=len(WORKBOOK_HEADERS))
    update_row_tax_display(row)
    values = [
        row["中文品名"],
        row["英文品名"],
        row["商品编码"],
        translate_material_to_english(row["材质"]),
        translate_usage_to_english(row["用途"]),
        row["箱数"],
        row["数量"],
        row["单位"],
        row["币制"],
        row["单价"],
        row["总价"],
        row["净重"],
        row["毛重"],
        row["原产国"],
        row[DISPLAY_TAX_RATE_FIELD],
        row[DISPLAY_TAX_AMOUNT_FIELD],
    ]
    for col, value in enumerate(values, start=1):
        sheet.cell(row_idx, col).value = value


def ensure_tax_display_columns(sheet) -> None:
    header_row = 5
    first_tax_col = len(HEADERS) + 1
    for col in range(first_tax_col, len(WORKBOOK_HEADERS) + 1):
        source = sheet.cell(header_row, len(HEADERS))
        target = sheet.cell(header_row, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.font:
            target.font = copy(source.font)
        if source.fill:
            target.fill = copy(source.fill)
        if source.border:
            target.border = copy(source.border)
        target.value = WORKBOOK_HEADERS[col - 1]
        sheet.column_dimensions[target.column_letter].width = max(
            sheet.column_dimensions[target.column_letter].width or 0,
            12,
        )


def copy_row_style(sheet, source_row: int, target_row: int, max_col: int) -> None:
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
    for col in range(1, max_col + 1):
        source = sheet.cell(source_row, col)
        target = sheet.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)


def extract_shipper(text: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    for idx, line in enumerate(lines):
        if "SUPPLY CHAIN" in line.upper() or "IMP. & EXP." in line.upper():
            return "\n".join(lines[idx : min(idx + 5, len(lines))])
    return ""


def extract_consignee(text: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    for idx, line in enumerate(lines):
        upper = line.upper()
        if "CORPORATION" in upper or "TRADE INC" in upper or "SOLUTIONS CORP" in upper:
            block = lines[idx : min(idx + 4, len(lines))]
            if not any("LOGISTICS" in item.upper() for item in block):
                return "\n".join(block)
    return ""


def extract_shipment_no(text: str) -> str:
    matches = re.findall(r"\b[A-Z]{3,5}\d{7,12}\b", text)
    return matches[0] if matches else ""


def extract_eta(text: str) -> str:
    match = re.search(r"\b(20\d{2})[/-](\d{1,2})[/-](\d{1,2})\b", text)
    if not match:
        match = re.search(r"\b([A-Z][a-z]{2})\.(\d{1,2}),\s*(20\d{2})\b", text)
        if not match:
            return ""
        months = {"Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4, "May": 5, "Jun": 6, "Jul": 7, "Aug": 8, "Sep": 9, "Oct": 10, "Nov": 11, "Dec": 12}
        return f"{match.group(3)}/{months.get(match.group(1), 1)}/{int(match.group(2))}"
    return f"{match.group(1)}/{int(match.group(2))}/{int(match.group(3))}"


def extract_number(pattern: str, text: str) -> Optional[float]:
    match = re.search(pattern, text, flags=re.IGNORECASE)
    if not match:
        return None
    return to_float(match.group(1))


def scale_integer(values: list[float], target_total: float) -> list[int]:
    target = int(round(target_total))
    base = sum(values)
    if not values:
        return []
    if not base:
        result = [0] * len(values)
        result[-1] = target
        return result
    exact = [value * target / base for value in values]
    floors = [math.floor(value) for value in exact]
    remainder = target - sum(floors)
    order = sorted(range(len(values)), key=lambda idx: exact[idx] - floors[idx], reverse=True)
    for idx in order[:remainder]:
        floors[idx] += 1
    return floors


def scale_decimal(values: list[float], target_total: float, digits: int) -> list[float]:
    base = sum(values)
    if not values:
        return []
    if not base:
        result = [0.0] * len(values)
        result[-1] = round(target_total, digits)
        return result
    scaled = [round(value * target_total / base, digits) for value in values]
    diff = round(target_total - sum(scaled), digits)
    if scaled and diff:
        scaled[-1] = round(scaled[-1] + diff, digits)
    return scaled


def normalize_header(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


def normalize_text(value: Any) -> str:
    text = str(value or "").lower()
    text = re.sub(r"[^a-z0-9\u4e00-\u9fff]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalize_bill_product_text(value: Any) -> str:
    text = normalize_text(value)
    words = []
    for word in text.split():
        if word == "mobile":
            continue
        if word == "stand":
            word = "holder"
        words.append(word)
    return " ".join(words)


def normalize_hs(value: Any) -> str:
    return re.sub(r"\D", "", str(value or ""))


def hs_cell_value(value: Any) -> int | str:
    hs = normalize_hs(value)
    if hs and not hs.startswith("0"):
        return int(hs)
    return hs


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def to_float(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        if math.isnan(value):
            return None
        return float(value)
    text = str(value).strip()
    if not text or text.lower() == "nan" or text.startswith("="):
        return None
    try:
        return float(text.replace(",", ""))
    except ValueError:
        return None


def manifest_to_public_dict(manifest: ManifestSummary) -> dict[str, Any]:
    data = asdict(manifest)
    data["items"] = data["items"][:80]
    return data


def bill_to_public_dict(bill: BillInfo) -> dict[str, Any]:
    return {
        "filename": bill.filename,
        "products": bill.products,
        "product_entries": [asdict(entry) for entry in bill.product_entries],
        "shipment_no": bill.shipment_no,
        "eta": bill.eta,
        "cartons": bill.cartons,
        "carton_evidence": bill.carton_evidence,
        "carton_unit": bill.carton_unit,
        "carton_source": bill.carton_source,
        "carton_reasoning": bill.carton_reasoning,
        "carton_manifest_comparison": bill.carton_manifest_comparison,
        "carton_confidence": bill.carton_confidence,
        "carton_inferred": bill.carton_inferred,
        "carton_recognition_attempts": bill.carton_recognition_attempts,
        "carton_resolution_history": bill.carton_resolution_history,
        "gross_weight": bill.gross_weight,
        "cbm": bill.cbm,
        "shipper": bill.shipper,
        "consignee": bill.consignee,
    }
