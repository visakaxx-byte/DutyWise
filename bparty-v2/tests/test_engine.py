from __future__ import annotations

import asyncio
import unittest
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import httpx
from openpyxl import Workbook, load_workbook

from engine import (
    BillInfo,
    BillLLMFields,
    BillProduct,
    blocking_llm_review_issues,
    ManifestHsGroup,
    ManifestItem,
    ManifestWeightInfo,
    ManifestSummary,
    ProductCandidate,
    ProcessingOptions,
    SelectionRules,
    assert_selected_price_fit_resolved,
    apply_manifest_net_weights,
    apply_bill_product_names,
    adjust_price_gap,
    allocate_price_first_cartons,
    allocate_plausible_prices,
    allocate_row_tax_budgets,
    build_clearance,
    build_manual_invoice_candidate_pool,
    build_manifest_candidates,
    build_manifest_hs_groups,
    build_manifest_weight_context,
    build_output_rows,
    write_workbook,
    build_plausible_row_plans,
    certification_filter_reason,
    effective_tax_rate,
    ensure_candidate_plausibility_ranges,
    ensure_manifest_codeflag_candidates,
    ensure_bill_products_present,
    exclude_bill_product_replacements,
    generate_valid_output_rows_with_llm,
    load_plausibility_ranges,
    load_replacement_candidates,
    normalize_manifest_weight_payload,
    normalize_bill_llm_products,
    normalize_bill_llm_fields,
    normalize_bill_llm_product_entries,
    normalize_carton_count,
    output_package_total,
    parse_bill,
    parse_manifest,
    parse_bill_product_entries_from_images,
    parse_bill_product_entries_from_text,
    parse_non_exempt_additional_tax_rate,
    parse_percentage_component_tax_rate,
    parse_tax_rate,
    parse_bill_products_with_llm,
    PlausibilityRange,
    qualify_manual_invoice_replacements,
    qualify_single_candidate,
    qualify_bill_product_candidates,
    select_initial_candidates,
    select_qualified_tax_data,
    tax_filter_reason,
    translate_material_to_english,
    translate_usage_to_english,
    translate_output_chinese_names_with_llm,
    normalize_generated_candidate,
    optimize_selected_candidates_for_price_fit,
    validate_llm_output_rows,
    validate_price_evidence,
    validate_qty_ctn_relationship,
    infer_bill_material_from_entry,
    is_bill_required_candidate,
    is_undetailed_bill_candidate,
    load_default_reference_manual_candidates,
    manual_invoice_search_pool,
    optimize_selected_candidates_for_manual_invoice,
    validate_output_product_uniqueness,
)
from config import CrawlerSettings
from crawler_client import (
    StrictTaxCrawler,
    aes_encrypt,
    is_auth_expired,
    login_body_error_detail,
    parse_classification_results,
)
from price_search import build_price_evidence, extract_price_samples, parse_pack_qty


def tax_result(rate: str = "3.4%", hs: str = "3924104000", certifications: list[str] | None = None) -> dict[str, dict]:
    return {
        hs: {
            "hs_code_us": hs,
            "tax_rate": rate,
            "certification_texts": certifications or [],
            "description_cn": "测试品名",
        }
    }


def manifest_schema_payload(
    *,
    sheet_name: str = "箱单",
    header_row: int = 1,
    summary_rows: list[int] | None = None,
    data_start_row: int = 3,
    data_end_row: int = 4,
    columns: dict[str, int] | None = None,
    header_labels: dict[str, str] | None = None,
    confidence: float = 0.98,
    strategy: str = "explicit_total",
    total_cell: str = "箱单!F2",
    detail_range: str = "箱单!F3:F4",
) -> dict:
    return {
        "sheet_name": sheet_name,
        "header_row": header_row,
        "summary_rows": [2] if summary_rows is None else summary_rows,
        "data_start_row": data_start_row,
        "data_end_row": data_end_row,
        "columns": columns
        or {
            "zh_name": 1,
            "en_name": 2,
            "material": 3,
            "qty": 4,
            "ctns": 5,
            "gross_weight": 6,
            "hs_code": 9,
        },
        "header_labels": header_labels
        or {
            "zh_name": "中文品名",
            "en_name": "品名",
            "material": "材质",
            "qty": "产品数量",
            "ctns": "件数",
            "gross_weight": "重量",
            "hs_code": "海关编码",
        },
        "weight_unit": "kg",
        "confidence": confidence,
        "weight_strategy": strategy,
        "weight_total_cell": total_cell,
        "weight_detail_range": detail_range,
    }


def write_manifest_workbook(path: Path, *, include_summary: bool = True, include_details: bool = True) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "箱单"
    sheet.append(["中文品名", "品名", "材质", "产品数量", "件数", "重量", "立方", "计费重", "海关编码"])
    sheet.append([None, None, None, None, 3, 30 if include_summary else None, None, None, None])
    sheet.append(["水杯", "Water cup", "塑料", 20, 1, 10 if include_details else None, None, None, "3924104000"])
    sheet.append(["枕套", "Pillowcase", "涤纶", 40, 2, 20 if include_details else None, None, None, "6302322020"])
    workbook.save(path)


class FakeCrawler:
    def __init__(self, result: dict[str, dict]):
        self.result = result
        self.settings = SimpleNamespace(delay=0.0)

    async def search_product(self, product_name: str, material: str = "") -> dict[str, dict]:
        return self.result

    async def search(self, hs_code: str) -> dict[str, dict]:
        normalized = "".join(ch for ch in str(hs_code or "") if ch.isdigit())
        if len(normalized) == 10 and len(self.result) == 1:
            data = dict(next(iter(self.result.values())))
            data["hs_code_us"] = normalized
            return {normalized: data}
        return self.result


class RoutedFakeCrawler:
    def __init__(self, product_results: dict[str, dict[str, dict]], hs_results: dict[str, dict[str, dict]] | None = None):
        self.product_results = product_results
        self.hs_results = hs_results or {}
        self.settings = SimpleNamespace(delay=0.0)
        self.product_calls: list[tuple[str, str]] = []
        self.hs_calls: list[str] = []

    async def search_product(self, product_name: str, material: str = "") -> dict[str, dict]:
        self.product_calls.append((product_name, material))
        key = product_name.lower()
        if key not in self.product_results:
            raise RuntimeError(f"missing product result: {product_name}")
        return self.product_results[key]

    async def search(self, hs_code: str) -> dict[str, dict]:
        self.hs_calls.append(hs_code)
        if hs_code not in self.hs_results:
            raise RuntimeError(f"missing hs result: {hs_code}")
        return self.hs_results[hs_code]


class FakeBillParser:
    def __init__(self, payload: dict):
        self.payload = payload
        self.calls = 0
        self.messages = []
        self.temperatures = []

    async def chat_json(self, messages: list[dict[str, str]], *, temperature: float = 0.1, **kwargs) -> dict:
        self.calls += 1
        self.messages.append(messages)
        self.temperatures.append(temperature)
        return self.payload


class FakeVisionBillParser(FakeBillParser):
    def __init__(self, payload: dict):
        super().__init__(payload)
        self.image_calls = 0
        self.image_prompts: list[str] = []
        self.image_data_urls: list[list[str]] = []

    async def chat_json_with_images(
        self,
        text: str,
        image_data_urls: list[str],
        *,
        temperature: float = 0.0,
        model: str | None = None,
        max_tokens: int = 4096,
    ) -> dict:
        self.image_calls += 1
        self.image_prompts.append(text)
        self.image_data_urls.append(image_data_urls)
        self.temperatures.append(temperature)
        return self.payload


class QueueFakeLLM:
    def __init__(self, payloads: list[dict | BaseException]):
        self.payloads = list(payloads)
        self.calls = 0
        self.kwargs = []

    async def chat_json(self, messages: list[dict[str, str]], *, temperature: float = 0.1, **kwargs) -> dict:
        self.calls += 1
        self.kwargs.append({"temperature": temperature, **kwargs})
        if not self.payloads:
            raise RuntimeError("no fake LLM payload left")
        payload = self.payloads.pop(0)
        if isinstance(payload, BaseException):
            raise payload
        return payload


class DictFakeLLM:
    def __init__(self, payload: dict):
        self.payload = payload
        self.calls = 0

    async def chat_json(self, messages: list[dict[str, str]], *, temperature: float = 0.1, **kwargs) -> dict:
        self.calls += 1
        return self.payload


class RoutedFakeParser:
    def __init__(self, bill_payload: dict, manifest_payload: dict):
        self.bill_payload = bill_payload
        self.manifest_payload = manifest_payload

    async def chat_json(self, messages: list[dict[str, str]], *, temperature: float = 0.1, **kwargs) -> dict:
        content = "\n".join(message.get("content", "") for message in messages)
        if "total_weight_kg" in content:
            return self.manifest_payload
        return self.bill_payload


class SequencedAsyncClient:
    responses: list[httpx.Response] = []
    requests: list[dict] = []

    def __init__(self, *args, **kwargs):
        self.args = args
        self.kwargs = kwargs

    async def __aenter__(self):
        return self

    async def __aexit__(self, exc_type, exc, tb):
        return False

    async def post(self, url: str, **kwargs) -> httpx.Response:
        self.__class__.requests.append({"url": url, **kwargs})
        if not self.__class__.responses:
            raise AssertionError(f"no fake response left for {url}")
        return self.__class__.responses.pop(0)


def codeflag_payload(hs: str = "3924104000", tax_rate: str = "3.4%") -> dict:
    return {
        "classificationCodeList": [
            {
                "hsCode": hs,
                "importTariffRate": tax_rate,
            }
        ],
        "classificationResultList": [
            {
                "hsCode": hs,
                "gName": "测试品名",
            }
        ],
    }


class TaxRateTests(unittest.IsolatedAsyncioTestCase):
    def test_login_body_error_detail_decrypts_nested_message(self) -> None:
        encrypted = aes_encrypt('{"code":500,"message":"暂未启用，请联系管理员"}')
        self.assertEqual(login_body_error_detail({"code": 200, "data": encrypted}), "暂未启用，请联系管理员")

    def test_is_auth_expired_matches_codeflag_logout_text(self) -> None:
        self.assertTrue(is_auth_expired("你已下线，请重新登陆"))
        self.assertTrue(is_auth_expired("你的账号在另一地点登录，如果不是本人操作"))
        self.assertTrue(is_auth_expired("token expired"))
        self.assertFalse(is_auth_expired("您的操作过于频繁，请联系管理员"))

    async def test_hs_search_relogs_once_when_codeflag_says_logged_out(self) -> None:
        settings = CrawlerSettings(
            username="user",
            password="password",
            base_url="https://codeflag.test",
            timeout=1,
            delay=0,
            max_retries=1,
            rate_limit_backoff=0,
        )
        SequencedAsyncClient.responses = [
            httpx.Response(200, json={"code": 200}, headers={"CusAuthorization": "token-1"}),
            httpx.Response(200, json={"code": 500, "message": "你已下线，请重新登陆"}),
            httpx.Response(200, json={"code": 200}, headers={"CusAuthorization": "token-2"}),
            httpx.Response(200, json={"code": 200, "data": codeflag_payload()}),
        ]
        SequencedAsyncClient.requests = []

        with patch("crawler_client.httpx.AsyncClient", SequencedAsyncClient):
            crawler = StrictTaxCrawler(settings)
            result = await crawler.search("3924104000")

        self.assertEqual(result["3924104000"]["tax_rate"], "3.4%")
        self.assertEqual(crawler.auth_expired_retry_count, 1)
        search_tokens = [
            request["headers"]["CusAuthorization"]
            for request in SequencedAsyncClient.requests
            if request["url"].endswith("/classification/search")
        ]
        self.assertEqual(search_tokens, ["token-1", "token-2"])

    def test_parse_tax_rate_formats(self) -> None:
        self.assertEqual(parse_tax_rate("N/A"), None)
        self.assertAlmostEqual(parse_tax_rate(0.034), 0.034)
        self.assertAlmostEqual(parse_tax_rate("3.4%"), 0.034)
        self.assertAlmostEqual(parse_tax_rate("25%+10%"), 0.35)
        self.assertAlmostEqual(parse_tax_rate("20%"), 0.2)
        self.assertAlmostEqual(parse_tax_rate("20"), 0.2)
        self.assertAlmostEqual(parse_tax_rate("Free"), 0.0)
        self.assertAlmostEqual(parse_percentage_component_tax_rate("0.8¢ each + 4.6%"), 0.046)

    def test_select_qualified_tax_data_can_prefer_first_and_ignore_certifications(self) -> None:
        rules = SelectionRules()
        results = {
            "3924104000": {
                "hs_code_us": "3924104000",
                "tax_rate": "3.4%",
                "additional_tax_rate": "10%",
                "certification_texts": ["FD1: FDA data MAY BE required"],
            },
            "3926400090": {
                "hs_code_us": "3926400090",
                "tax_rate": "Free",
                "additional_tax_rate": "10%",
                "certification_texts": [],
            },
        }

        selected = select_qualified_tax_data(
            results,
            rules,
            enforce_tax_limit=False,
            prefer_first=True,
            ignore_certifications=True,
        )

        self.assertEqual(selected["hs_code_us"], "3924104000")

    def test_bill_tax_selection_ranks_semantic_match_before_lowest_tax(self) -> None:
        results = {
            "6307908950": {
                "hs_code_us": "6307908950",
                "tax_rate": "7%",
                "additional_tax_rate": "12.5%",
                "certification_texts": [],
                "description_cn": "其他纺织材料制未列名制品",
            },
            "4202923131": {
                "hs_code_us": "4202923131",
                "tax_rate": "17.6%",
                "additional_tax_rate": "25%+12.5%",
                "certification_texts": [],
                "description_cn": "其他化学纤维制货物包装袋",
            },
            "6305390000": {
                "hs_code_us": "6305390000",
                "tax_rate": "8.4%",
                "additional_tax_rate": "7.5%+12.5%",
                "certification_texts": [],
                "description_cn": "其他化学纤维制货物包装袋",
            },
            "6305200000": {
                "hs_code_us": "6305200000",
                "tax_rate": "6.2%",
                "additional_tax_rate": "7.5%+12.5%",
                "certification_texts": [],
                "description_cn": "用于货物包装的袋及包：棉制",
            },
            "6305900000": {
                "hs_code_us": "6305900000",
                "tax_rate": "6.2%",
                "additional_tax_rate": "7.5%+12.5%",
                "certification_texts": [],
                "description_cn": "用于货物包装的袋及包：其他纺织材料制",
            },
        }

        selected = select_qualified_tax_data(
            results,
            SelectionRules(),
            enforce_tax_limit=False,
            ignore_certifications=True,
            semantic_name="Polyester bag",
            semantic_material="Polyester",
            semantic_usage="Storage",
        )

        self.assertEqual(selected["hs_code_us"], "6305390000")

    def test_bill_tax_selection_rejects_explicit_material_conflict(self) -> None:
        results = {
            "6303910010": {
                "hs_code_us": "6303910010",
                "tax_rate": "10.3%",
                "certification_texts": [],
                "description_cn": "棉制非针织非钩编窗帘",
            },
            "6303921000": {
                "hs_code_us": "6303921000",
                "tax_rate": "11.3%",
                "certification_texts": [],
                "description_cn": "合纤制非针织非钩编窗帘",
            },
        }

        selected = select_qualified_tax_data(
            results,
            SelectionRules(),
            enforce_tax_limit=False,
            ignore_certifications=True,
            semantic_name="Polyester curtain",
            semantic_material="Polyester",
        )

        self.assertEqual(selected["hs_code_us"], "6303921000")

    def test_customer_codebook_candidates_use_combined_tax_below_30_percent(self) -> None:
        candidates = [
            candidate
            for candidate in load_replacement_candidates()
            if candidate.tax_match_source == "customer_codebook"
        ]

        self.assertGreater(len(candidates), 1000)
        self.assertTrue(all(candidate.effective_tax_rate < 0.30 for candidate in candidates))
        self.assertTrue(any(candidate.row_warnings for candidate in candidates))

    def test_effective_tax_rate_adds_non_exempt_additional_tax(self) -> None:
        self.assertAlmostEqual(
            effective_tax_rate({"tax_rate": "3%", "additional_tax_rate": "7%"}),
            0.10,
        )
        self.assertAlmostEqual(
            effective_tax_rate({"tax_rate": "3%", "additional_tax_rate": "7%", "applicable_tax_rate": "12%"}),
            0.10,
        )
        self.assertAlmostEqual(
            effective_tax_rate({"additional_tax_rate": "7%", "applicable_tax_rate": "12%"}),
            0.07,
        )
        self.assertAlmostEqual(
            effective_tax_rate({"tax_rate": "5%", "additional_tax_rate": "7.5%+10%", "additional_tax_details": []}),
            0.225,
        )

    def test_additional_tax_rate_treats_conditional_exemption_as_exempt(self) -> None:
        self.assertAlmostEqual(
            parse_non_exempt_additional_tax_rate("无豁免7.5%；条件豁免10%"),
            0.075,
        )
        self.assertAlmostEqual(
            effective_tax_rate({"tax_rate": "3.4%", "additional_tax_rate": "无豁免7.5%；条件豁免10%"}),
            0.109,
        )
        self.assertAlmostEqual(
            effective_tax_rate(
                {
                    "tax_rate": "3.4%",
                    "additional_tax_details": [
                        {"label": "通用税率", "rate": "3.4%"},
                        {"label": "无豁免", "rate": "7.5%"},
                        {"label": "条件豁免", "rate": "10%"},
                    ],
                }
            ),
            0.109,
        )

    def test_parse_codeflagai_additional_tax_and_authentication(self) -> None:
        parsed = parse_classification_results(
            "9017800000",
            {
                "classificationResultList": [{"hsCode": "9017800000", "gName": "其他手用测量长度的器具"}],
                "classificationCodeList": [
                    {
                        "hsCode": "9017800000",
                        "importTariffRate": "5.3%",
                        "applicableTaxRate": "5.3%",
                        "additionalTaxRate": "25%+10%",
                        "authentication": "FD1: FDA data MAY BE required",
                    }
                ],
            },
        )
        item = parsed["9017800000"]
        self.assertEqual(item["additional_tax_rate"], "25%+10%")
        self.assertEqual(item["certification_texts"], ["FD1: FDA data MAY BE required"])

    def test_parse_codeflagai_prefers_description_attached_to_us_hts(self) -> None:
        parsed = parse_classification_results(
            "Polyester bag",
            {
                "classificationResultList": [
                    {"hsCode": "4202920000", "gName": "塑料片或纺织材料作面的其他容器"}
                ],
                "classificationCodeList": [
                    {
                        "hsCode": "6305200000",
                        "importTariffRate": "6.2%",
                        "taricCn": "用于货物包装的袋及包：棉制",
                        "taric": "Sacks and bags used for packing goods: Of cotton",
                    }
                ],
            },
        )

        item = parsed["6305200000"]
        self.assertEqual(item["description_cn"], "用于货物包装的袋及包：棉制")
        self.assertEqual(item["source_description_cn"], "塑料片或纺织材料作面的其他容器")


class ManifestHsGroupTests(unittest.TestCase):
    def test_build_manifest_hs_groups_preserves_distinct_products_under_same_hs(self) -> None:
        manifest = ManifestSummary(
            filename="input.xlsx",
            row_count=3,
            total_ctns=6,
            total_real_weight=15,
            total_declared_value=30,
            categories=[],
            items=[
                ManifestItem(2, "塑料发夹", "Plastic hair clip", "9615900000", "Plastic", "Hair", 2, 100, 0.1, 10, 4, 4),
                ManifestItem(20, "发夹", "Hair clips", "9615900000", "Plastic", "Hair", 3, 200, 0.08, 16, 6, 6),
                ManifestItem(21, "钥匙扣", "Keychain", "7326209000", "Metal", "Home", 1, 50, 0.12, 6, 5, 5),
            ],
        )

        groups = build_manifest_hs_groups(manifest)
        same_hs_groups = [group for group in groups if group.hs == "9615900000"]

        self.assertEqual(len(same_hs_groups), 2)
        self.assertEqual(sorted(group.source_rows[0] for group in same_hs_groups), [2, 20])
        self.assertEqual(sum(group.total_qty for group in same_hs_groups), 300)
        self.assertEqual(sum(group.total_ctns for group in same_hs_groups), 5)
        self.assertEqual(sum(group.total_gross_weight for group in same_hs_groups), 10)


class PriceSearchTests(unittest.TestCase):
    def test_parse_pack_qty_and_price_evidence(self) -> None:
        self.assertEqual(parse_pack_qty("Plastic hair clips 100 pcs $6.99"), 100)
        self.assertEqual(parse_pack_qty("High-Waisted Everyday Cotton Underwear 6-Pack $39.99"), 6)
        samples = extract_price_samples("<html>Plastic hair clips 100 pcs $6.99 Another pack of 50 $5.00</html>")
        evidence = build_price_evidence("plastic hair clip", samples)

        self.assertGreaterEqual(len(evidence.samples), 2)
        self.assertAlmostEqual(evidence.declared_unit_price, evidence.retail_unit_price * 0.3, places=4)

    def test_validate_price_evidence_allows_low_total_value(self) -> None:
        candidate = ProductCandidate(
            source="manifest_group",
            source_label="input",
            zh="塑料发夹",
            en="Plastic hair clip",
            hs="9615900000",
            material="Plastic",
            usage="HOME",
            price_evidence={"declared_unit_price": 0.1, "confidence": 0.5, "basis": "test", "retail_unit_price": 0.33},
        )
        rows = [
            {
                "中文品名": "塑料发夹",
                "英文品名": "Plastic hair clip",
                "商品编码": "9615900000",
                "单价": 0.1,
                "数量": 10,
                "总价": 1.0,
            }
        ]

        validate_price_evidence(rows, [candidate])

        self.assertEqual(rows[0]["价格依据"], "test")
        self.assertEqual(rows[0]["零售参考单价"], 0.33)

    def test_validate_price_evidence_warns_when_tax_price_is_below_reference(self) -> None:
        candidate = ProductCandidate(
            source="manifest_group",
            source_label="input",
            zh="塑料发夹",
            en="Plastic hair clip",
            hs="9615900000",
            material="Plastic",
            usage="HOME",
            price_evidence={"declared_unit_price": 1.0, "confidence": 0.5, "basis": "test", "retail_unit_price": 3.33},
        )
        rows = [
            {
                "中文品名": "塑料发夹",
                "英文品名": "Plastic hair clip",
                "商品编码": "9615900000",
                "单价": 0.05,
                "数量": 10,
                "总价": 0.5,
            }
        ]

        validate_price_evidence(rows, [candidate])

        self.assertEqual(rows[0]["价格依据"], "test")
        self.assertIn("单价超出价格证据范围", rows[0]["价格提示"])
        self.assertIn("单价超出价格证据范围", rows[0]["约束提示"])


class CertificationRuleTests(unittest.TestCase):
    def setUp(self) -> None:
        self.rules = SelectionRules(allowed_certifications=["Lacey Act", "TSCA"])

    def test_allowed_certification_combinations(self) -> None:
        self.assertEqual(certification_filter_reason([], self.rules), "")
        self.assertEqual(certification_filter_reason(["May need Lacey Act"], self.rules), "")
        self.assertEqual(certification_filter_reason(["TSCA"], self.rules), "")
        self.assertEqual(certification_filter_reason(["Lacey Act and TSCA"], self.rules), "")

    def test_unknown_or_blocked_certification_is_filtered(self) -> None:
        self.assertIn("未允许", certification_filter_reason(["FDA"], self.rules))
        blocked = SelectionRules(
            allowed_certifications=["Lacey Act", "TSCA"],
            blocked_certifications=["FDA"],
        )
        self.assertIn("禁用认证", certification_filter_reason(["FDA and TSCA"], blocked))

    def test_tax_filter_combines_rate_and_certification_rules(self) -> None:
        good = {"hs_code_us": "3924104000", "tax_rate": "3.4%", "certification_texts": ["TSCA"]}
        self.assertEqual(tax_filter_reason(good, self.rules), "")
        high = {"hs_code_us": "3924104000", "tax_rate": "20%", "certification_texts": []}
        self.assertIn("不小于 20%", tax_filter_reason(high, self.rules))
        high_effective = {
            "hs_code_us": "7020006000",
            "tax_rate": "5%",
            "additional_tax_rate": "25%+10%",
            "certification_texts": [],
        }
        self.assertIn("综合税率", tax_filter_reason(high_effective, self.rules))
        self.assertEqual(tax_filter_reason(high_effective, self.rules, enforce_tax_limit=False), "")
        unknown = {"hs_code_us": "3924104000", "tax_rate": "3.4%", "certification_texts": ["FDA"]}
        self.assertIn("未允许", tax_filter_reason(unknown, self.rules))
        self.assertIn("未允许", tax_filter_reason(unknown, self.rules, enforce_tax_limit=False))


class ReplacementWorkbookTests(unittest.TestCase):
    def test_load_replacement_candidates_from_known_sheets(self) -> None:
        path = Path(__file__).resolve().parents[1] / "docs/海关编码查找.xlsx"
        candidates = load_replacement_candidates(path)
        self.assertGreater(len(candidates), 50)
        common = next(item for item in candidates if item.zh == "托盘")
        self.assertEqual(common.en, "Pallet")
        self.assertEqual(common.hs, "3924104000")
        self.assertGreater(common.unit_price, 0)
        self.assertTrue(any(item.source_label.endswith("/20260330") for item in candidates))

    def test_load_plausibility_ranges_from_known_sheets(self) -> None:
        path = Path(__file__).resolve().parents[1] / "docs/海关编码查找.xlsx"
        ranges = load_plausibility_ranges(path)

        towel = ranges[("毛巾", "毛巾", "6302932000")]
        self.assertAlmostEqual(towel.unit_price_min, 0.5)
        self.assertAlmostEqual(towel.unit_price_max, 0.8)
        self.assertAlmostEqual(towel.kg_per_pc_min, 0.5)
        self.assertAlmostEqual(towel.kg_per_pc_max, 0.9)
        self.assertAlmostEqual(towel.ctns_min, 40)
        self.assertAlmostEqual(towel.ctns_max, 60)
        self.assertAlmostEqual(towel.qty_per_ctn_min, 6)
        self.assertAlmostEqual(towel.qty_per_ctn_max, 10)

        pallet = ranges[("托盘", "pallet", "3924104000")]
        self.assertAlmostEqual(pallet.kg_per_ctn_min, 5.1746, places=3)
        self.assertAlmostEqual(pallet.kg_per_ctn_max, 32.0)
        self.assertAlmostEqual(pallet.kg_per_pc_min, 0.265, places=3)
        self.assertAlmostEqual(pallet.kg_per_pc_max, 1.88, places=2)


class BillProductParsingTests(unittest.IsolatedAsyncioTestCase):
    async def test_llm_parser_extracts_products_and_ignores_broken_shipped_on_board_phrase(self) -> None:
        text = """
        STORAGE BAG HS CODE:420222
        PLASTIC SHELL HS CODE:392690
        S
        HIPPEDON BOARD:
        MAY.06 2026
        """
        parser = FakeBillParser(
            {
                "products": [
                    {"name": "STORAGE BAG", "hs_code_hint": "420222", "evidence": "STORAGE BAG HS CODE:420222", "confidence": 0.98},
                    {"name": "PLASTIC SHELL", "hs_code_hint": "392690", "evidence": "PLASTIC SHELL HS CODE:392690", "confidence": 0.97},
                ],
                "ignored_phrases": [{"text": "HIPPEDON BOARD", "reason": "broken shipped on board field"}],
            }
        )
        cache = {"product": {}, "hs": {}, "bill": {}}
        self.assertEqual(await parse_bill_products_with_llm(text, parser, cache), ["STORAGE BAG", "PLASTIC SHELL"])
        self.assertEqual(await parse_bill_products_with_llm(text, parser, cache), ["STORAGE BAG", "PLASTIC SHELL"])
        self.assertEqual(parser.calls, 1)
        self.assertEqual(parser.temperatures, [0.0])
        prompt = "\n".join(message.get("content", "") for message in parser.messages[0])
        self.assertIn("ignored_phrases", prompt)

    async def test_text_bill_cache_key_is_namespaced_and_preserves_product_entry_details(self) -> None:
        text = "STORAGE BAG HS CODE:420222"
        parser = FakeBillParser(
            {
                "shipper": "ACME EXPORT LTD\n1 SHIPPER ROAD",
                "consignee": "BUYER INC\n9 CONSIGNEE AVE",
                "products": [
                    {"name": "STORAGE BAG", "hs": "HS CODE: 4202.22", "evidence": text, "confidence": 0.98},
                ]
            }
        )
        cache = {"bill": {}}

        first = await parse_bill_product_entries_from_text(text, parser, cache)
        second = await parse_bill_product_entries_from_text(text, parser, cache)

        self.assertEqual(parser.calls, 1)
        self.assertTrue(list(cache["bill"].keys())[0].startswith("text:v4:"))
        self.assertEqual(second[0].name, first[0].name)
        self.assertEqual(second[0].hs_code_hint, "420222")
        self.assertEqual(second[0].evidence, text)

    async def test_parse_text_bill_uses_llm_shipper_and_consignee(self) -> None:
        parser = FakeBillParser(
            {
                "shipper": "ACME EXPORT LTD\n1 SHIPPER ROAD",
                "consignee": "BUYER INC\n9 CONSIGNEE AVE",
                "carton_count": 930,
                "carton_evidence": "SAY NINE HUNDRED AND THIRTY CARTONS ONLY",
                "products": [
                    {"name": "STORAGE BAG", "hs_code_hint": "420222", "evidence": "text evidence", "confidence": 0.96},
                ],
            }
        )
        text = "STORAGE BAG HS CODE:420222 " * 5
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "text.pdf"
            pdf.write_bytes(b"fake pdf bytes")
            with patch("engine.extract_bill_text", return_value=text):
                bill = await parse_bill(pdf, parser, {"bill": {}})

        self.assertEqual(bill.shipper, "ACME EXPORT LTD\n1 SHIPPER ROAD")
        self.assertEqual(bill.consignee, "BUYER INC\n9 CONSIGNEE AVE")
        self.assertEqual(bill.cartons, 930)

    async def test_image_bill_parser_uses_vision_cache_key_and_image_llm(self) -> None:
        parser = FakeVisionBillParser(
            {
                "shipper": "VISION SHIPPER LTD\n2 EXPORT ST",
                "consignee": "VISION BUYER INC\n3 IMPORT AVE",
                "products": [
                    {"name": "STORAGE BAG", "hs_code_hint": "420222", "evidence": "image evidence", "confidence": 0.96},
                ]
            }
        )
        cache = {"bill": {}}
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "scan.pdf"
            pdf.write_bytes(b"fake pdf bytes")
            first = await parse_bill_product_entries_from_images(pdf, ["data:image/jpeg;base64,abc"], parser, cache)
            second = await parse_bill_product_entries_from_images(pdf, ["data:image/jpeg;base64,abc"], parser, cache)

        self.assertEqual(parser.image_calls, 1)
        self.assertEqual(first[0].name, "STORAGE BAG")
        self.assertEqual(second[0].hs_code_hint, "420222")
        self.assertTrue(next(iter(cache["bill"])).startswith("vision:v4:"))
        self.assertIn("扫描图", parser.image_prompts[0])

    async def test_parse_bill_falls_back_to_vision_when_pdf_text_is_empty(self) -> None:
        parser = FakeVisionBillParser(
            {
                "shipper": "VISION SHIPPER LTD\n2 EXPORT ST",
                "consignee": "VISION BUYER INC\n3 IMPORT AVE",
                "carton_count": 661,
                "carton_evidence": "40'HC/661CARTONS/16536.71KGS/68.00CBM",
                "products": [
                    {"name": "PLASTIC SHELL", "hs_code_hint": "392690", "evidence": "PLASTIC SHELL HS CODE:392690", "confidence": 0.97},
                ]
            }
        )
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "scan.pdf"
            pdf.write_bytes(b"fake pdf bytes")
            with (
                patch("engine.extract_bill_text", return_value=""),
                patch("engine.render_bill_pdf_pages", return_value=["data:image/jpeg;base64,abc"]),
            ):
                bill = await parse_bill(pdf, parser, {"bill": {}})

        self.assertEqual(bill.products, ["PLASTIC SHELL"])
        self.assertEqual(bill.product_entries[0].hs_code_hint, "392690")
        self.assertEqual(bill.parse_source, "vision")
        self.assertEqual(bill.text_chars, 0)
        self.assertEqual(bill.vision_pages, 1)
        self.assertEqual(bill.shipper, "VISION SHIPPER LTD\n2 EXPORT ST")
        self.assertEqual(bill.consignee, "VISION BUYER INC\n3 IMPORT AVE")
        self.assertEqual(bill.cartons, 661)

    async def test_parse_bill_keeps_text_path_when_pdf_text_is_long_enough(self) -> None:
        parser = FakeVisionBillParser(
            {
                "carton_count": 748,
                "products": [
                    {"name": "STORAGE BAG", "hs_code_hint": "420222", "evidence": "text evidence", "confidence": 0.96},
                ]
            }
        )
        text = "STORAGE BAG HS CODE:420222 " * 5
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "text.pdf"
            pdf.write_bytes(b"fake pdf bytes")
            with (
                patch("engine.extract_bill_text", return_value=text),
                patch("engine.render_bill_pdf_pages", side_effect=AssertionError("vision path should not render")),
            ):
                bill = await parse_bill(pdf, parser, {"bill": {}})

        self.assertEqual(bill.products, ["STORAGE BAG"])
        self.assertEqual(bill.parse_source, "text")
        self.assertEqual(parser.image_calls, 0)
        self.assertEqual(bill.cartons, 748)

    async def test_parse_bill_rejects_missing_carton_count_without_manifest_fallback(self) -> None:
        parser = FakeVisionBillParser(
            {
                "products": [
                    {"name": "PLASTIC SHELL", "hs_code_hint": "392690", "evidence": "PLASTIC SHELL HS CODE:392690", "confidence": 0.97},
                ]
            }
        )
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "scan.pdf"
            pdf.write_bytes(b"fake pdf bytes")
            with (
                patch("engine.extract_bill_text", return_value=""),
                patch("engine.render_bill_pdf_pages", return_value=["data:image/jpeg;base64,abc"]),
                self.assertRaisesRegex(RuntimeError, "提单未识别到有效总箱数"),
            ):
                await parse_bill(pdf, parser, {"bill": {}})

    async def test_parse_bill_rejects_fractional_text_carton_fallback(self) -> None:
        parser = FakeBillParser(
            {
                "products": [
                    {
                        "name": "STORAGE BAG",
                        "hs_code_hint": "420222",
                        "evidence": "STORAGE BAG HS CODE:420222",
                        "confidence": 0.97,
                    },
                ]
            }
        )
        text = ("STORAGE BAG HS CODE:420222 " * 4) + "TOTAL 2.6 CARTONS"
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "text.pdf"
            pdf.write_bytes(b"fake pdf bytes")
            with (
                patch("engine.extract_bill_text", return_value=text),
                self.assertRaisesRegex(RuntimeError, "提单未识别到有效总箱数"),
            ):
                await parse_bill(pdf, parser, {"bill": {}})

    async def test_parse_bill_accepts_thousands_separated_text_carton_fallback(self) -> None:
        parser = FakeBillParser(
            {
                "products": [
                    {
                        "name": "STORAGE BAG",
                        "hs_code_hint": "420222",
                        "evidence": "STORAGE BAG HS CODE:420222",
                        "confidence": 0.97,
                    },
                ]
            }
        )
        text = ("STORAGE BAG HS CODE:420222 " * 4) + "TOTAL 1,234 CARTONS"
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "text.pdf"
            pdf.write_bytes(b"fake pdf bytes")
            with patch("engine.extract_bill_text", return_value=text):
                bill = await parse_bill(pdf, parser, {"bill": {}})

        self.assertEqual(bill.cartons, 1234)
        self.assertEqual(bill.carton_source, "bill_text_cartons_regex")
        self.assertEqual(bill.carton_unit, "CARTONS")

    async def test_parse_bill_rejects_negative_text_carton_fallbacks(self) -> None:
        payload = {
            "products": [
                {
                    "name": "STORAGE BAG",
                    "hs_code_hint": "420222",
                    "evidence": "STORAGE BAG HS CODE:420222",
                    "confidence": 0.97,
                },
            ]
        }
        for raw_count in ("- 5 CARTONS", "−5 CARTONS"):
            with self.subTest(raw_count=raw_count), tempfile.TemporaryDirectory() as tmp:
                pdf = Path(tmp) / "text.pdf"
                pdf.write_bytes(b"fake pdf bytes")
                text = ("STORAGE BAG HS CODE:420222 " * 4) + f"TOTAL {raw_count}"
                with (
                    patch("engine.extract_bill_text", return_value=text),
                    self.assertRaisesRegex(RuntimeError, "提单未识别到有效总箱数"),
                ):
                    await parse_bill(pdf, FakeBillParser(payload), {"bill": {}})

    async def test_parse_bill_scanned_empty_vision_result_has_clear_error(self) -> None:
        parser = FakeVisionBillParser({"products": []})
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "scan.pdf"
            pdf.write_bytes(b"fake pdf bytes")
            with (
                patch("engine.extract_bill_text", return_value=""),
                patch("engine.render_bill_pdf_pages", return_value=["data:image/jpeg;base64,abc"]),
                self.assertRaisesRegex(RuntimeError, "提单为扫描件，视觉模型未识别到可靠货物品类"),
            ):
                await parse_bill(pdf, parser, {"bill": {}})

    def test_llm_products_are_validated_by_code(self) -> None:
        payload = {
            "products": [
                {"name": "STORAGE BAG HS CODE:420222", "confidence": 0.99},
                {"name": "HIPPEDON BOARD", "confidence": 0.99},
                {"name": "LOW CONFIDENCE ITEM", "confidence": 0.2},
            ]
        }
        self.assertEqual(normalize_bill_llm_products(payload), ["STORAGE BAG"])

    def test_llm_product_entries_preserve_hs_hint(self) -> None:
        payload = {
            "products": [
                {
                    "name": "STORAGE BAG",
                    "hs_code_hint": "420222",
                    "evidence": "STORAGE BAG HS CODE:420222",
                    "confidence": 0.99,
                }
            ]
        }
        entries = normalize_bill_llm_product_entries(payload)
        self.assertEqual(entries[0].name, "STORAGE BAG")
        self.assertEqual(entries[0].hs_code_hint, "420222")
        self.assertEqual(entries[0].evidence, "STORAGE BAG HS CODE:420222")

    def test_llm_product_entries_deduplicate_and_support_hs_aliases(self) -> None:
        payload = {
            "products": [
                {"name": "STORAGE BAG", "hs_code": "HS CODE: 4202.22", "evidence": "first", "confidence": 0.99},
                {"name": "storage   bag", "hs_code_hint": "999999", "evidence": "duplicate", "confidence": 0.99},
            ]
        }
        entries = normalize_bill_llm_product_entries(payload)
        self.assertEqual(len(entries), 1)
        self.assertEqual(entries[0].hs_code_hint, "420222")
        self.assertEqual(entries[0].evidence, "first")

    def test_llm_product_entries_split_comma_joined_products_and_hs_hints(self) -> None:
        payload = {
            "products": [
                {
                    "name": "PLASTIC ORNAMENTS,NECKLACE",
                    "hs_code_hint": "392640711790",
                    "evidence": "PLASTIC ORNAMENTS,NECKLACE HS:392640,711790",
                    "confidence": 1.0,
                }
            ]
        }
        entries = normalize_bill_llm_product_entries(payload)
        self.assertEqual([entry.name for entry in entries], ["PLASTIC ORNAMENTS", "NECKLACE"])
        self.assertEqual([entry.hs_code_hint for entry in entries], ["392640", "711790"])

    def test_llm_product_entries_reject_bad_structure_and_empty_products(self) -> None:
        with self.assertRaisesRegex(RuntimeError, "LLM 提单解析未返回 products 数组"):
            normalize_bill_llm_product_entries({"items": []})
        with self.assertRaisesRegex(RuntimeError, "LLM 未从提单中识别到可靠货物品类"):
            normalize_bill_llm_product_entries({"products": [{"name": "LOW CONFIDENCE ITEM", "confidence": 0.1}]})

    def test_bill_llm_fields_normalize_party_blocks(self) -> None:
        fields = normalize_bill_llm_fields(
            {
                "shipper": [" ACME EXPORT LTD  ", " 1 SHIPPER ROAD "],
                "consignee": " BUYER INC \n 9 CONSIGNEE AVE ",
                "carton_count": "1,234 CARTONS",
                "carton_evidence": "TOTAL 1,234 CARTONS",
                "products": [{"name": "STORAGE BAG", "confidence": 0.99}],
            }
        )

        self.assertEqual(fields.shipper, "ACME EXPORT LTD\n1 SHIPPER ROAD")
        self.assertEqual(fields.consignee, "BUYER INC\n9 CONSIGNEE AVE")
        self.assertEqual(fields.product_entries[0].name, "STORAGE BAG")
        self.assertEqual(fields.carton_count, 1234)
        self.assertEqual(fields.carton_evidence, "TOTAL 1,234 CARTONS")

    def test_carton_count_rejects_boolean_negative_zero_and_non_finite_values(self) -> None:
        self.assertIsNone(normalize_carton_count(True))
        self.assertIsNone(normalize_carton_count(False))
        self.assertIsNone(normalize_carton_count(-5))
        self.assertIsNone(normalize_carton_count("-5 PIECES"))
        self.assertIsNone(normalize_carton_count("- 5 PIECES"))
        self.assertIsNone(normalize_carton_count("−5 PIECES"))
        self.assertIsNone(normalize_carton_count(0))
        self.assertIsNone(normalize_carton_count(float("inf")))
        self.assertIsNone(normalize_carton_count(2.4))
        self.assertIsNone(normalize_carton_count("2.6 PALLETS"))
        self.assertEqual(normalize_carton_count("153 PIECES RCP"), 153)

    def test_output_package_count_is_never_silently_raised_to_match_row_count(self) -> None:
        bill = BillInfo(
            "bill.pdf",
            "",
            [],
            cartons=2,
            carton_unit="PALLETS",
        )
        with self.assertRaisesRegex(
            RuntimeError,
            "整票包装数量 2 PALLETS 小于输出行数 5",
        ):
            output_package_total(bill, 5)
        with self.assertRaisesRegex(RuntimeError, "未识别到有效包装数量"):
            output_package_total(
                BillInfo("bill.pdf", "", [], cartons=2.6, carton_unit="CARTONS"),
                2,
            )

    def test_bill_material_inference_prefers_modifier_and_does_not_overmatch_pe_pp(self) -> None:
        self.assertEqual(
            "Plastic",
            infer_bill_material_from_entry(
                BillProduct(name="PLASTIC ORNAMENTS", hs_code_hint="392640", evidence="PLASTIC ORNAMENTS HS:392640")
            ),
        )
        self.assertEqual(
            "",
            infer_bill_material_from_entry(BillProduct(name="PIECE", hs_code_hint="392640", evidence="PIECE HS:392640")),
        )
        self.assertEqual(
            "Polyester",
            infer_bill_material_from_entry(
                BillProduct(name="POLYESTER BAG", evidence="POLYESTER BAG")
            ),
        )


class ManifestWeightParsingTests(unittest.IsolatedAsyncioTestCase):
    def test_manifest_weight_payload_is_normalized(self) -> None:
        payload = {
            "total_weight_kg": "3077.7",
            "source": "Sheet1.总毛重KGS",
            "evidence": "总毛重KGS 列求和",
            "confidence": 0.93,
        }
        info = normalize_manifest_weight_payload(payload)
        self.assertEqual(info.total_weight_kg, 3077.7)
        self.assertEqual(info.source, "Sheet1.总毛重KGS")
        self.assertEqual(info.evidence, "总毛重KGS 列求和")
        self.assertEqual(info.confidence, 0.93)

    async def test_manifest_schema_maps_customs_code_and_reconciles_weight_once(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "manifest.xlsx"
            write_manifest_workbook(path)
            parser = FakeBillParser(manifest_schema_payload())
            cache: dict[str, dict] = {}

            first = await parse_manifest(path, parser, cache)
            second = await parse_manifest(path, parser, cache)

        self.assertEqual(parser.calls, 1)
        self.assertEqual(first.schema_columns["hs_code"], 9)
        self.assertEqual([item.hs for item in first.items], ["3924104000", "6302322020"])
        self.assertEqual(first.total_real_weight, 30)
        self.assertEqual(first.weight_explicit_total, 30)
        self.assertEqual(first.weight_detail_sum, 30)
        self.assertTrue(first.weight_reconciled)
        self.assertEqual(second.total_real_weight, 30)
        self.assertTrue(next(iter(cache["manifest_schema"])).startswith("schema:v1:net-weight-v1:"))

    async def test_manifest_schema_parses_complete_net_weight_column(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "manifest-net-weight.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "箱单"
            sheet.append(["中文品名", "英文品名", "箱数", "数量", "净重", "毛重", "HS"])
            sheet.append(["水杯", "Water cup", 1, 20, 9, 10, "3924104000"])
            sheet.append(["枕套", "Pillowcase", 2, 40, 18, 20, "6302322020"])
            workbook.save(path)
            payload = manifest_schema_payload(
                header_row=1,
                summary_rows=[],
                data_start_row=2,
                data_end_row=3,
                columns={
                    "zh_name": 1,
                    "en_name": 2,
                    "ctns": 3,
                    "qty": 4,
                    "net_weight": 5,
                    "gross_weight": 6,
                    "hs_code": 7,
                },
                header_labels={
                    "zh_name": "中文品名",
                    "en_name": "英文品名",
                    "ctns": "箱数",
                    "qty": "数量",
                    "net_weight": "净重",
                    "gross_weight": "毛重",
                    "hs_code": "HS",
                },
                strategy="detail_sum",
                total_cell="",
                detail_range="箱单!E2:E3",
            )

            manifest = await parse_manifest(path, DictFakeLLM(payload), {})

        self.assertEqual(manifest.schema_columns["net_weight"], 5)
        self.assertEqual([item.real_weight for item in manifest.items], [9, 18])
        self.assertEqual(manifest.total_net_weight, 27)
        self.assertEqual(manifest.total_real_weight, 30)
        self.assertEqual(manifest.weight_detail_range, "F2:F3")

    async def test_manifest_schema_supports_non_first_sheet_and_late_header(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "late-header.xlsx"
            workbook = Workbook()
            workbook.active.title = "说明"
            sheet = workbook.create_sheet("数据区")
            sheet.append(["装箱说明"])
            sheet.append([])
            sheet.append(["商品名称", "English", "材质描述", "数量", "箱数", "毛重", "HS"])
            sheet.append([None, None, None, None, 3, 30, None])
            sheet.append(["水杯", "Water cup", "塑料", 20, 1, 10, "3924104000"])
            sheet.append(["枕套", "Pillowcase", "涤纶", 40, 2, 20, "6302322020"])
            workbook.save(path)
            payload = manifest_schema_payload(
                sheet_name="数据区",
                header_row=3,
                summary_rows=[4],
                data_start_row=5,
                data_end_row=6,
                columns={"zh_name": 1, "en_name": 2, "material": 3, "qty": 4, "ctns": 5, "gross_weight": 6, "hs_code": 7},
                header_labels={"zh_name": "商品名称", "en_name": "English", "material": "材质描述", "qty": "数量", "ctns": "箱数", "gross_weight": "毛重", "hs_code": "HS"},
                total_cell="数据区!F4",
                detail_range="数据区!F5:F6",
            )

            manifest = await parse_manifest(path, DictFakeLLM(payload), {})

        self.assertEqual(manifest.schema_sheet, "数据区")
        self.assertEqual(manifest.schema_header_row, 3)
        self.assertEqual(manifest.row_count, 2)
        self.assertEqual(manifest.total_real_weight, 30)

    async def test_manifest_schema_rejects_invalid_mapping_and_low_confidence(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "manifest.xlsx"
            write_manifest_workbook(path)
            invalid_column = manifest_schema_payload()
            invalid_column["columns"] = {**invalid_column["columns"], "hs_code": 99}
            with self.assertRaisesRegex(RuntimeError, "manifest_schema_unresolved.*列号越界"):
                await parse_manifest(path, DictFakeLLM(invalid_column), {})
            with self.assertRaisesRegex(RuntimeError, "manifest_schema_unresolved.*置信度"):
                await parse_manifest(path, DictFakeLLM(manifest_schema_payload(confidence=0.4)), {})

    def test_manifest_product_rows_cannot_skip_codeflag_when_candidates_are_empty(self) -> None:
        manifest = ManifestSummary(
            filename="input.xlsx",
            row_count=1,
            total_ctns=1,
            total_real_weight=1,
            total_declared_value=0,
            categories=[],
            items=[],
        )
        with self.assertRaisesRegex(RuntimeError, "manifest_schema_unresolved.*没有生成 Codeflag 查询候选"):
            ensure_manifest_codeflag_candidates(manifest, [])

    async def test_product_without_hs_uses_name_and_material_codeflag_search(self) -> None:
        manifest = ManifestSummary(
            filename="input.xlsx",
            row_count=1,
            total_ctns=1,
            total_real_weight=10,
            total_declared_value=0,
            categories=["水杯"],
            items=[
                ManifestItem(
                    row=3,
                    zh="水杯",
                    en="Water cup",
                    hs="",
                    material="塑料",
                    usage="",
                    ctns=1,
                    qty=20,
                    unit_price=None,
                    declared_value=None,
                    real_weight=None,
                    gross_weight=10,
                )
            ],
        )
        candidates = build_manifest_candidates(manifest)
        crawler = RoutedFakeCrawler(product_results={"水杯": tax_result("3.4%", hs="3924104000")})

        result = await qualify_single_candidate(
            crawler,
            candidates[0],
            SelectionRules(),
            query_cache={},
            enforce_tax_limit=False,
        )

        self.assertEqual(candidates[0].hs, "")
        self.assertEqual(crawler.hs_calls, [])
        self.assertEqual(crawler.product_calls, [("水杯", "塑料")])
        self.assertEqual(result.hs, "3924104000")

    async def test_manifest_weight_supports_detail_only_and_summary_only(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            detail_path = Path(tmp) / "detail.xlsx"
            write_manifest_workbook(detail_path, include_summary=False)
            detail_payload = manifest_schema_payload(
                summary_rows=[],
                strategy="detail_sum",
                total_cell="",
            )
            detail = await parse_manifest(detail_path, DictFakeLLM(detail_payload), {})

            summary_path = Path(tmp) / "summary.xlsx"
            write_manifest_workbook(summary_path, include_details=False)
            summary_payload = manifest_schema_payload(detail_range="")
            summary = await parse_manifest(summary_path, DictFakeLLM(summary_payload), {})

        self.assertEqual(detail.weight_strategy, "detail_sum")
        self.assertEqual(detail.total_real_weight, 30)
        self.assertIsNone(detail.weight_explicit_total)
        self.assertEqual(summary.weight_strategy, "explicit_total")
        self.assertEqual(summary.total_real_weight, 30)
        self.assertIsNone(summary.weight_detail_sum)

    async def test_manifest_weight_rejects_conflict_and_unrecomputable_total_cell(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            conflict_path = Path(tmp) / "conflict.xlsx"
            write_manifest_workbook(conflict_path)
            workbook = load_workbook(conflict_path)
            workbook["箱单"]["F2"] = 35
            workbook.save(conflict_path)
            with self.assertRaisesRegex(RuntimeError, "manifest_weight_conflict"):
                await parse_manifest(conflict_path, DictFakeLLM(manifest_schema_payload()), {})

            invalid_path = Path(tmp) / "invalid-total.xlsx"
            write_manifest_workbook(invalid_path)
            workbook = load_workbook(invalid_path)
            workbook["箱单"]["F2"] = "TOTAL"
            workbook.save(invalid_path)
            with self.assertRaisesRegex(RuntimeError, "manifest_schema_unresolved.*无法从原表重算"):
                await parse_manifest(invalid_path, DictFakeLLM(manifest_schema_payload()), {})

    async def test_clearance_queries_manifest_before_customer_codebook_fill(self) -> None:
        manifest_candidate = ProductCandidate(
            source="manifest_group",
            source_label="input.xlsx/HS归并",
            zh="水杯",
            en="Water cup",
            hs="3924104000",
            material="Plastic",
            usage="HOME",
            ctns=1,
            qty=20,
            gross_weight=10,
        )
        codebook_candidate = ProductCandidate(
            source="replacement",
            source_label="客户编码库/Sheet1",
            zh="枕套",
            en="Pillowcase",
            hs="6302322020",
            material="Polyester",
            usage="HOME",
            ctns=1,
            qty=20,
            gross_weight=20,
            tax_match_source="customer_codebook",
            effective_tax_rate=0.1,
        )
        manifest = ManifestSummary(
            filename="input.xlsx",
            row_count=1,
            total_ctns=2,
            total_real_weight=30,
            total_declared_value=0,
            categories=["水杯"],
            items=[
                ManifestItem(3, "水杯", "Water cup", "3924104000", "Plastic", "", 1, 20, None, None, None, 10)
            ],
        )
        bill = BillInfo(
            filename="bill.pdf",
            raw_text="",
            products=["WATER CUP"],
            cartons=2,
            product_entries=[BillProduct(name="WATER CUP", evidence="WATER CUP")],
        )
        events: list[str] = []

        async def fake_qualify(crawler, candidates, rules, **kwargs):
            events.append("manifest_codeflag")
            self.assertEqual(candidates, [manifest_candidate])
            qualified = ProductCandidate(
                **{
                    **manifest_candidate.__dict__,
                    "tax_data": tax_result()["3924104000"],
                    "base_tax_rate": 0.1,
                    "effective_tax_rate": 0.1,
                    "tax_match_source": "manifest_group_hs",
                }
            )
            return [qualified], []

        async def fake_plausibility(selected, *args, **kwargs):
            return selected, False

        async def fake_manual_replacements(*args, **kwargs):
            events.append("post_fill_optimization")
            return [], [], 0

        async def fake_translate(*args, **kwargs):
            return None

        def fake_rows(selected, *args, **kwargs):
            events.append("codebook_fill_complete")
            self.assertEqual([candidate.source for candidate in selected], ["manifest_group", "replacement"])
            return [
                {"中文品名": "水杯", "英文品名": "Water cup", "总价": 100, "综合税率": 0.1, "箱数": 1, "毛重": 10, "来源": "manifest_group", "约束提示": "", "商品编码": "3924104000", "爬虫品名": "Water cup"},
                {"中文品名": "枕套", "英文品名": "Pillowcase", "总价": 100, "综合税率": 0.1, "箱数": 1, "毛重": 20, "来源": "replacement", "约束提示": "", "商品编码": "6302322020", "爬虫品名": "Pillowcase"},
            ]

        crawler = SimpleNamespace(settings=SimpleNamespace(delay=0), auth_expired_retry_count=0)
        with tempfile.TemporaryDirectory() as tmp:
            with (
                patch("engine.parse_manifest", return_value=manifest),
                patch("engine.parse_bill", return_value=bill),
                patch("engine.StrictTaxCrawler", return_value=crawler),
                patch("engine.load_selection_rules", return_value=SelectionRules()),
                patch("engine.load_replacement_candidates", return_value=[codebook_candidate]),
                patch("engine.load_plausibility_ranges", return_value={}),
                patch("engine.build_manifest_candidates", return_value=[manifest_candidate]),
                patch("engine.qualify_candidates", side_effect=fake_qualify),
                patch("engine.ensure_candidate_plausibility_ranges", side_effect=fake_plausibility),
                patch("engine.attach_price_evidence_to_candidates", side_effect=lambda selected, **kwargs: selected),
                patch("engine.qualify_manual_invoice_replacements", side_effect=fake_manual_replacements),
                patch("engine.optimize_selected_candidates_for_manual_invoice", side_effect=lambda selected, *args, **kwargs: (selected, {"strategy": "test", "swaps": 0})),
                patch("engine.build_output_rows", side_effect=fake_rows),
                patch("engine.translate_output_chinese_names_with_llm", side_effect=fake_translate),
                patch("engine.validate_output_rows"),
                patch("engine.validate_price_evidence"),
                patch("engine.write_workbook"),
                patch("engine.build_audit_summary", return_value={}),
            ):
                result = await build_clearance(
                    "input.xlsx",
                    "bill.pdf",
                    tmp,
                    target_tax_amount=100,
                    target_item_count=2,
                    llm=DictFakeLLM({}),
                )

        self.assertLess(events.index("manifest_codeflag"), events.index("codebook_fill_complete"))
        self.assertEqual(result["stats"]["codeflag_queried"], 1)
        self.assertEqual(result["stats"]["codebook_fill_needed"], 1)
        self.assertEqual(result["stats"]["codebook_fill_used"], 1)
        self.assertEqual(result["stats"]["constraint_status"], "needs_review")
        self.assertEqual(result["stats"]["bill_required_locked"], 1)
        self.assertEqual(result["stats"]["bill_manifest_detailed_locked"], 1)
        self.assertEqual(result["stats"]["bill_synthetic_locked"], 0)
        self.assertEqual(result["stats"]["solver_status"], "optimized")
        self.assertTrue(result["stats"]["tax_optimization_applied"])
        self.assertEqual(result["stats"]["best_effort_reason"], "")
        bill_stage = next(item for item in result["flow"] if item["stage"] == "bill_products")
        self.assertEqual(bill_stage["manifest_reused"], 1)
        self.assertEqual(bill_stage["queried_separately"], 0)


class OutputOptimizationTests(unittest.TestCase):
    def test_manifest_detailed_bill_candidates_are_not_subject_to_synthetic_caps(self) -> None:
        original_cartons = [28, 2, 18, 19, 20, 15, 18, 5, 17, 14]
        selected = [
            ProductCandidate(
                source="manifest_group",
                source_label="input.xlsx/HS归并",
                zh=f"商品{index}",
                en=f"Item {index}",
                hs=f"3926909{index:03d}",
                material="Polyester",
                usage="HOME",
                ctns=ctns,
                effective_tax_rate=0.2,
                bill_product_name=f"BILL ITEM {index}" if index < 3 else "",
                bill_has_manifest_detail=index < 3,
            )
            for index, ctns in enumerate(original_cartons)
        ]

        allocated_cartons = allocate_price_first_cartons(selected, 153)
        tax_budgets = allocate_row_tax_budgets(selected, 2800)

        self.assertEqual(sum(allocated_cartons), 153)
        self.assertEqual(allocated_cartons[:3], original_cartons[:3])
        self.assertGreater(allocated_cartons[0], 3)
        self.assertGreater(sum(allocated_cartons[:3]), round(153 * 0.06))
        self.assertGreater(sum(tax_budgets[:3]), 2800 * 0.20)
        self.assertTrue(all(is_bill_required_candidate(item) for item in selected[:3]))
        self.assertTrue(all(not is_undetailed_bill_candidate(item) for item in selected[:3]))

    def test_bill_only_synthetic_candidates_keep_carton_and_tax_caps(self) -> None:
        original_cartons = [28, 2, 18, 19, 20, 15, 18, 5, 17, 14]
        selected = [
            ProductCandidate(
                source="bill" if index < 3 else "manifest_group",
                source_label="bill.pdf" if index < 3 else "input.xlsx/HS归并",
                zh=f"商品{index}",
                en=f"Item {index}",
                hs=f"3926909{index:03d}",
                material="Polyester",
                usage="HOME",
                ctns=ctns,
                effective_tax_rate=0.2,
                bill_product_name=f"BILL ITEM {index}" if index < 3 else "",
            )
            for index, ctns in enumerate(original_cartons)
        ]

        allocated_cartons = allocate_price_first_cartons(selected, 153)
        tax_budgets = allocate_row_tax_budgets(selected, 2800)

        self.assertEqual(sum(allocated_cartons), 153)
        self.assertLessEqual(sum(allocated_cartons[:3]), round(153 * 0.06))
        self.assertAlmostEqual(sum(tax_budgets[:3]), 2800 * 0.20, places=2)
        self.assertTrue(all(is_undetailed_bill_candidate(item) for item in selected[:3]))

    def test_candidate_library_is_excluded_when_manifest_already_fills_target(self) -> None:
        manifest_candidates = [
            ProductCandidate(
                source="manifest_group",
                source_label="input.xlsx/HS归并",
                zh=f"清单商品{index}",
                en=f"Manifest item {index}",
                hs=f"3926909{index:03d}",
                material="Plastic",
                usage="Home use",
            )
            for index in range(10)
        ]
        library_candidate = ProductCandidate(
            source="replacement",
            source_label="DEFAULT_REFERENCE_STYLE_ROWS",
            zh="键盘",
            en="Keyboard",
            hs="8471602000",
            material="ABS",
            usage="Home use",
        )

        pool = build_manual_invoice_candidate_pool(
            manifest_candidates,
            manifest_candidates,
            [library_candidate],
            allow_candidate_library=False,
        )

        self.assertEqual(pool, manifest_candidates)
        self.assertNotIn(library_candidate, pool)

    def test_candidate_library_is_available_for_real_row_shortage(self) -> None:
        manifest_candidate = ProductCandidate(
            source="manifest_group",
            source_label="input.xlsx/HS归并",
            zh="清单商品",
            en="Manifest item",
            hs="3926909989",
            material="Plastic",
            usage="Home use",
        )
        library_candidate = ProductCandidate(
            source="replacement",
            source_label="DEFAULT_REFERENCE_STYLE_ROWS",
            zh="键盘",
            en="Keyboard",
            hs="8471602000",
            material="ABS",
            usage="Home use",
        )

        pool = build_manual_invoice_candidate_pool(
            [manifest_candidate],
            [manifest_candidate],
            [library_candidate],
            allow_candidate_library=True,
        )

        self.assertEqual(pool, [manifest_candidate, library_candidate])

    def test_price_gap_increases_low_tax_plan_to_target(self) -> None:
        candidates = [
            ProductCandidate(
                source="manifest_group",
                source_label="input.xlsx/HS归并",
                zh=f"商品{index}",
                en=f"Item {index}",
                hs=f"3926909{index:03d}",
                material="Plastic",
                usage="Home use",
                effective_tax_rate=rate,
            )
            for index, rate in enumerate((0.2, 0.3))
        ]
        unit_prices = [1.0, 1.0]

        adjust_price_gap(
            unit_prices,
            [0.5, 0.5],
            [3.0, 3.0],
            candidates,
            [100, 100],
            100.0,
        )

        estimated_tax = sum(
            price * 100 * candidate.effective_tax_rate
            for price, candidate in zip(unit_prices, candidates)
        )
        self.assertAlmostEqual(estimated_tax, 100.0, places=4)

    def test_formal_price_solver_uses_maximum_safe_price_when_target_is_unreachable(self) -> None:
        candidate = ProductCandidate(
            source="manifest_group",
            source_label="input.xlsx/HS归并",
            zh="塑料商品",
            en="Plastic item",
            hs="3926909989",
            material="Plastic",
            usage="Home use",
            effective_tax_rate=0.1,
        )
        plausibility = PlausibilityRange(
            kg_per_ctn_min=1,
            kg_per_ctn_max=20,
            kg_per_pc_min=0.1,
            kg_per_pc_max=5,
            unit_price_min=0.5,
            unit_price_max=1.0,
            qty_per_ctn_min=1,
            qty_per_ctn_max=100,
            source="test range",
        )

        prices = allocate_plausible_prices(
            [candidate],
            [10],
            [10.0],
            [plausibility],
            100.0,
            [100.0],
        )

        self.assertEqual(prices, [(1.0, 10.0)])

    def test_output_net_weight_inherits_manifest_net_to_gross_ratio(self) -> None:
        rows = [
            {"箱数": 10, "毛重": 100, "净重": 90},
            {"箱数": 20, "毛重": 200, "净重": 180},
        ]
        manifest = ManifestSummary(
            filename="input.xlsx",
            row_count=2,
            total_ctns=30,
            total_real_weight=300,
            total_declared_value=0,
            categories=[],
            total_net_weight=288,
        )

        apply_manifest_net_weights(rows, manifest)

        self.assertEqual(sum(row["净重"] for row in rows), 288)
        self.assertAlmostEqual(
            sum(row["净重"] for row in rows) / sum(row["毛重"] for row in rows),
            0.96,
            places=6,
        )
        self.assertTrue(all("原始清单总净重/总毛重" in row["净重计算依据"] for row in rows))

    def test_material_and_usage_are_translated_to_english(self) -> None:
        material = translate_material_to_english("铁+塑料/Iron + Plastic")
        usage = translate_usage_to_english("测量戒指大小、ring sizer")
        self.assertNotRegex(material, r"[\u4e00-\u9fff]")
        self.assertIn("Iron", material)
        self.assertIn("Plastic", material)
        self.assertNotRegex(usage, r"[\u4e00-\u9fff]")
        self.assertIn("Ring Size Measurement", usage)
        self.assertIn("Ring Sizer", usage)

    def test_output_rows_match_count_weight_and_tax_target(self) -> None:
        candidates = [
            ProductCandidate(
                source="manifest",
                source_label="input.xlsx",
                zh=f"品名{i}",
                en=f"Item {i}",
                hs="3924104000",
                material="Plastic",
                usage="HOME",
                ctns=10,
                qty=100,
                unit_price=2,
                real_weight=100,
                gross_weight=100,
                base_tax_rate=0.034,
                effective_tax_rate=0.234,
                plausibility_range=PlausibilityRange(
                    kg_per_ctn_min=0.5,
                    kg_per_ctn_max=20,
                    kg_per_pc_min=0.05,
                    kg_per_pc_max=2,
                    unit_price_min=0.5,
                    unit_price_max=10,
                    qty_per_ctn_min=5,
                    qty_per_ctn_max=20,
                    source="test range",
                ),
            )
            for i in range(3)
        ]
        candidates[0].material = "塑料 Plastic"
        candidates[0].usage = "装饰/Decoration"
        bill = BillInfo(filename="bill.pdf", raw_text="", products=[], gross_weight=300, cartons=30)
        manifest = ManifestSummary(
            filename="input.xlsx",
            row_count=3,
            total_ctns=30,
            total_real_weight=280,
            total_declared_value=1000,
            categories=[],
        )
        options = ProcessingOptions(target_tax_amount=120, target_item_count=3)

        rows = build_output_rows(candidates, manifest, bill, options)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[0]["材质"], "Plastic")
        self.assertEqual(rows[0]["用途"], "Decoration")
        self.assertAlmostEqual(sum(row["毛重"] for row in rows), 280, places=2)
        total_tax = sum(row["总价"] * row["综合税率"] for row in rows)
        self.assertLessEqual(abs(total_tax - 120), 20)
        self.assertEqual(rows[0]["税率"], "23.4%")
        self.assertAlmostEqual(rows[0]["税金"], round(rows[0]["总价"] * rows[0]["综合税率"], 2))

    def test_workbook_writes_tax_rate_and_tax_amount_as_last_columns(self) -> None:
        row = {
            "中文品名": "塑料杯",
            "英文品名": "Plastic Cup",
            "商品编码": "3924104000",
            "材质": "Plastic",
            "用途": "HOME",
            "箱数": 10,
            "数量": 100,
            "单位": "PCS",
            "币制": "USD",
            "单价": 2.5,
            "总价": 250,
            "净重": 90,
            "毛重": 100,
            "原产国": "CN",
            "综合税率": 0.109,
        }
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "clearance.xlsx"
            write_workbook(BillInfo("bill.pdf", "", []), [row], output, metadata={})
            workbook = load_workbook(output, data_only=True)
            sheet = workbook[workbook.sheetnames[0]]

        self.assertEqual(sheet.cell(5, 15).value, "税率")
        self.assertEqual(sheet.cell(5, 16).value, "税金")
        self.assertEqual(sheet.cell(6, 15).value, "10.9%")
        self.assertEqual(sheet.cell(6, 16).value, 27.25)
        self.assertEqual(row["税率"], "10.9%")
        self.assertEqual(row["税金"], 27.25)

    def test_workbook_writes_bill_shipper_and_consignee(self) -> None:
        row = {
            "中文品名": "塑料杯",
            "英文品名": "Plastic Cup",
            "商品编码": "3924104000",
            "材质": "Plastic",
            "用途": "HOME",
            "箱数": 10,
            "数量": 100,
            "单位": "PCS",
            "币制": "USD",
            "单价": 2.5,
            "总价": 250,
            "净重": 90,
            "毛重": 100,
            "原产国": "CN",
            "综合税率": 0.109,
        }
        bill = BillInfo(
            "bill.pdf",
            "",
            [],
            shipper="ACME EXPORT LTD\n1 SHIPPER ROAD",
            consignee="BUYER INC\n9 CONSIGNEE AVE",
        )
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "clearance.xlsx"
            write_workbook(bill, [dict(row)], output, metadata={})
            workbook = load_workbook(output, data_only=True)
            sheet = workbook[workbook.sheetnames[0]]

        self.assertEqual(sheet["C2"].value, "ACME EXPORT LTD\n1 SHIPPER ROAD")
        self.assertEqual(sheet["C3"].value, "BUYER INC\n9 CONSIGNEE AVE")

    def test_workbook_metadata_overrides_bill_parties(self) -> None:
        row = {
            "中文品名": "塑料杯",
            "英文品名": "Plastic Cup",
            "商品编码": "3924104000",
            "材质": "Plastic",
            "用途": "HOME",
            "箱数": 10,
            "数量": 100,
            "单位": "PCS",
            "币制": "USD",
            "单价": 2.5,
            "总价": 250,
            "净重": 90,
            "毛重": 100,
            "原产国": "CN",
            "综合税率": 0.109,
        }
        bill = BillInfo("bill.pdf", "", [], shipper="BILL SHIPPER", consignee="BILL CONSIGNEE")
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "clearance.xlsx"
            write_workbook(
                bill,
                [dict(row)],
                output,
                metadata={"shipper": "META SHIPPER", "consignee": "META CONSIGNEE"},
            )
            workbook = load_workbook(output, data_only=True)
            sheet = workbook[workbook.sheetnames[0]]

        self.assertEqual(sheet["C2"].value, "META SHIPPER")
        self.assertEqual(sheet["C3"].value, "META CONSIGNEE")

    def test_chinese_name_column_is_translated_to_simplified_chinese(self) -> None:
        rows = [
            {
                "中文品名": "PLASTIC ORNAMENTS",
                "英文品名": "Plastic Ornaments",
                "商品编码": "3926400090",
                "材质": "Plastic",
            }
        ]
        llm = DictFakeLLM({"translations": [{"index": 0, "translation": "塑料装饰品"}]})

        asyncio.run(translate_output_chinese_names_with_llm(llm, rows))

        self.assertEqual(rows[0]["中文品名"], "塑料装饰品")
        self.assertEqual(llm.calls, 1)

    def test_chinese_name_column_rejects_latin_translation(self) -> None:
        rows = [
            {
                "中文品名": "STORAGE BAG",
                "英文品名": "Storage Bag",
                "商品编码": "4202220000",
                "材质": "Textile",
            }
        ]
        llm = DictFakeLLM({"translations": [{"index": 0, "translation": "Storage Bag"}]})

        with self.assertRaisesRegex(RuntimeError, "有效简体中文品名"):
            asyncio.run(translate_output_chinese_names_with_llm(llm, rows))

    def test_output_rejects_duplicate_translated_names_even_when_hs_differs(self) -> None:
        rows = [
            {"中文品名": "窗帘", "英文品名": "Curtain", "商品编码": "6303910010"},
            {"中文品名": "窗帘", "英文品名": "Curtain", "商品编码": "6303921000"},
        ]

        with self.assertRaisesRegex(RuntimeError, "输出商品品名重复"):
            validate_output_product_uniqueness(rows)

    def test_output_allows_explicit_material_qualifiers_for_distinct_products(self) -> None:
        rows = [
            {"中文品名": "棉制窗帘", "英文品名": "Cotton curtain", "商品编码": "6303910010"},
            {"中文品名": "涤纶窗帘", "英文品名": "Polyester curtain", "商品编码": "6303921000"},
        ]

        validate_output_product_uniqueness(rows)

    def test_high_semantic_llm_review_issue_is_blocking_without_numeric_bounds(self) -> None:
        review = {
            "issues": [
                {
                    "row_index": 1,
                    "severity": "high",
                    "type": "material_hs_mismatch",
                    "message": "Polyester material is inconsistent with cotton HS description",
                    "suggested_bounds": {},
                }
            ]
        }

        issues = blocking_llm_review_issues(
            review,
            constraints_applied=0,
            reoptimized=False,
        )

        self.assertEqual(len(issues), 1)

    def test_review_only_candidate_downgrades_high_review_issue_to_warning(self) -> None:
        review = {
            "issues": [
                {
                    "row_index": 1,
                    "severity": "high",
                    "type": "unit_value_hs_mismatch",
                    "message": "Unit value exceeds the HTS per-piece limit",
                    "suggested_bounds": {},
                }
            ]
        }
        candidate = ProductCandidate(
            source="manifest_group",
            source_label="input.xlsx/HS归并",
            zh="不锈钢首饰",
            en="Stainless steel jewelry",
            hs="7117190500",
            material="Stainless steel",
            usage="Decoration",
            compliance_review_required=True,
            compliance_review_reason="HTS value cap mismatch",
        )

        issues = blocking_llm_review_issues(
            review,
            constraints_applied=0,
            reoptimized=False,
            candidates=[candidate],
        )

        self.assertEqual(issues, [])

    def test_manual_invoice_pool_excludes_review_only_candidate_when_clean_pool_is_sufficient(self) -> None:
        clean = [
            ProductCandidate(
                source="manifest_group",
                source_label="input.xlsx/HS归并",
                zh=f"清单商品{i}",
                en=f"Manifest item {i}",
                hs=f"39269099{i:02d}",
                material="Plastic",
                usage="HOME",
                gross_weight=10,
                ctns=1,
                qty=10,
                unit_price=1,
                effective_tax_rate=0.1,
            )
            for i in range(10)
        ]
        review_only = ProductCandidate(
            source="manifest_group",
            source_label="input.xlsx/HS归并",
            zh="不锈钢首饰",
            en="Stainless steel jewelry",
            hs="7117190500",
            material="Stainless steel",
            usage="Decoration",
            compliance_review_required=True,
        )

        pool = manual_invoice_search_pool([review_only, *clean], 10)

        self.assertNotIn(review_only, pool)

    def test_generated_candidates_do_not_inherit_historical_weight_constraints(self) -> None:
        candidate = ProductCandidate(
            source="replacement",
            source_label="海关编码查找.xlsx/常用1",
            zh="托盘",
            en="Pallet",
            hs="3924104000",
            material="Plastic",
            usage="HOME",
            ctns=12,
            qty=360,
            unit_price=0.8,
            real_weight=192,
            gross_weight=192,
                    base_tax_rate=0.1,
                    effective_tax_rate=0.1,
        )
        self.assertEqual(candidate.source, "replacement")
        self.assertEqual(candidate.unit_price, 0.8)
        self.assertEqual(candidate.qty, 360)
        self.assertEqual(candidate.gross_weight, 192)

    def test_output_rows_require_target_item_count(self) -> None:
        candidates = [
            ProductCandidate(
                source="manifest",
                source_label="input.xlsx",
                zh="品名",
                en="Item",
                hs="3924104000",
                material="Plastic",
                usage="HOME",
                ctns=10,
                qty=100,
                real_weight=100,
                gross_weight=100,
                base_tax_rate=0.034,
            )
        ]
        manifest = ManifestSummary("input.xlsx", 1, 10, 100, 100, [])
        bill = BillInfo("bill.pdf", "", [], gross_weight=300, cartons=10)
        options = ProcessingOptions(target_tax_amount=30, target_item_count=2)
        with self.assertRaisesRegex(RuntimeError, "输出行数必须等于 2"):
            build_output_rows(candidates, manifest, bill, options)

    def test_plausible_row_plans_obey_weight_and_price_ranges(self) -> None:
        candidates = [
            ProductCandidate(
                source="manifest",
                source_label="input.xlsx",
                zh="托盘",
                en="Pallet",
                hs="3924104000",
                material="Plastic",
                usage="HOME",
                ctns=12,
                qty=900,
                unit_price=0.8,
                real_weight=192,
                gross_weight=192,
                base_tax_rate=0.1,
                effective_tax_rate=0.1,
            ),
            ProductCandidate(
                source="manifest",
                source_label="input.xlsx",
                zh="垫子",
                en="Mat",
                hs="3924901050",
                material="Plastic",
                usage="HOME",
                ctns=13,
                qty=500,
                unit_price=1.2,
                real_weight=300,
                gross_weight=300,
                base_tax_rate=0.1,
                effective_tax_rate=0.1,
            ),
        ]
        manifest = ManifestSummary("input.xlsx", 2, 25, 492, 912, [])
        bill = BillInfo("bill.pdf", "", [], gross_weight=1200, cartons=25)
        options = ProcessingOptions(target_tax_amount=80, target_item_count=2)
        ranges = load_plausibility_ranges(Path(__file__).resolve().parents[1] / "docs/海关编码查找.xlsx")

        plans = build_plausible_row_plans(
            selected=candidates,
            manifest=manifest,
            bill=bill,
            options=options,
            plausibility_ranges=ranges,
        )

        self.assertAlmostEqual(sum(plan.gross_weight for plan in plans), 492, places=2)
        total_tax = sum(plan.total_value * (plan.candidate.effective_tax_rate or plan.candidate.base_tax_rate) for plan in plans)
        self.assertLessEqual(abs(total_tax - 80), 20)
        for plan in plans:
            self.assertGreaterEqual(plan.total_value * (plan.candidate.effective_tax_rate or plan.candidate.base_tax_rate), 30)
        for plan in plans:
            kg_per_ctn = plan.gross_weight / plan.ctns
            kg_per_pc = plan.gross_weight / plan.qty
            self.assertGreaterEqual(kg_per_ctn, plan.plausibility.kg_per_ctn_min)
            self.assertLessEqual(kg_per_ctn, plan.plausibility.kg_per_ctn_max)
            self.assertGreaterEqual(kg_per_pc, plan.plausibility.kg_per_pc_min)
            self.assertLessEqual(kg_per_pc, plan.plausibility.kg_per_pc_max)
            self.assertGreaterEqual(plan.unit_price, plan.plausibility.unit_price_min)
            self.assertLessEqual(plan.unit_price, plan.plausibility.unit_price_max)

    def test_plausible_row_plans_block_impossible_weight(self) -> None:
        candidates = [
            ProductCandidate(
                source="manifest",
                source_label="input.xlsx",
                zh="塑料花",
                en="Plastic flowers",
                hs="6702104000",
                material="Plastic",
                usage="HOME",
                ctns=1,
                qty=100,
                unit_price=0.15,
                real_weight=1,
                gross_weight=1,
                base_tax_rate=0.034,
            )
        ]
        manifest = ManifestSummary("input.xlsx", 1, 1, 500, 100, [])
        bill = BillInfo("bill.pdf", "", [], gross_weight=500, cartons=1)
        options = ProcessingOptions(target_tax_amount=10, target_item_count=1)
        with self.assertRaisesRegex(RuntimeError, "合理单箱重量上限"):
            build_plausible_row_plans(
                selected=candidates,
                manifest=manifest,
                bill=bill,
                options=options,
                plausibility_ranges={},
            )

    def test_plausible_row_plans_allow_tax_below_target_when_upper_bound_is_safe(self) -> None:
        candidates = [
            ProductCandidate(
                source="manifest",
                source_label="input.xlsx",
                zh="托盘",
                en="Pallet",
                hs="3924104000",
                material="Plastic",
                usage="HOME",
                ctns=12,
                qty=360,
                unit_price=0.8,
                real_weight=192,
                gross_weight=192,
                base_tax_rate=0.034,
            )
        ]
        manifest = ManifestSummary("input.xlsx", 1, 12, 192, 288, [])
        bill = BillInfo("bill.pdf", "", [], gross_weight=192, cartons=12)
        options = ProcessingOptions(target_tax_amount=1000, target_item_count=1)
        ranges = load_plausibility_ranges(Path(__file__).resolve().parents[1] / "docs/海关编码查找.xlsx")
        plans = build_plausible_row_plans(
            selected=candidates,
            manifest=manifest,
            bill=bill,
            options=options,
            plausibility_ranges=ranges,
        )

        total_tax = sum(plan.total_value * (plan.candidate.effective_tax_rate or plan.candidate.base_tax_rate) for plan in plans)
        self.assertLess(total_tax, 1000)
        self.assertLessEqual(total_tax, 1100)

    def test_plausible_row_plans_relax_price_floor_to_meet_tax_upper_bound(self) -> None:
        candidates = [
            ProductCandidate(
                source="manifest_group",
                source_label="input.xlsx",
                zh="品名",
                en="Item",
                hs="3926909985",
                material="Plastic",
                usage="HOME",
                ctns=10,
                qty=100,
                unit_price=90.0,
                real_weight=100,
                gross_weight=100,
                base_tax_rate=0.1,
                effective_tax_rate=0.1,
                plausibility_range=PlausibilityRange(
                    kg_per_ctn_min=5,
                    kg_per_ctn_max=20,
                    kg_per_pc_min=0.5,
                    kg_per_pc_max=2,
                    unit_price_min=90,
                    unit_price_max=91,
                    qty_per_ctn_min=5,
                    qty_per_ctn_max=20,
                    source="test range",
                ),
            )
        ]
        plans = build_plausible_row_plans(
            selected=candidates,
            manifest=ManifestSummary("input.xlsx", 1, 10, 100, 1000, []),
            bill=BillInfo("bill.pdf", "", [], cartons=10),
            options=ProcessingOptions(target_tax_amount=800, target_item_count=1),
            plausibility_ranges={},
        )

        self.assertLess(plans[0].unit_price, plans[0].plausibility.unit_price_min)
        self.assertLessEqual(plans[0].total_value * 0.1, 880)
        self.assertIn("单价低于合理下限", plans[0].warnings[0])

    def test_price_fit_repair_swaps_low_reference_ratio_manifest_candidate(self) -> None:
        bad = ProductCandidate(
            source="manifest_group",
            source_label="input.xlsx",
            zh="高价鞋",
            en="Shoes",
            hs="6402999000",
            material="Textile",
            usage="HOME",
            ctns=10,
            qty=1000,
            unit_price=50,
            gross_weight=100,
            base_tax_rate=0.1,
            effective_tax_rate=0.1,
            price_evidence={"declared_unit_price": 50, "confidence": 0.7},
            plausibility_range=PlausibilityRange(
                kg_per_ctn_min=5,
                kg_per_ctn_max=20,
                kg_per_pc_min=0.1,
                kg_per_pc_max=0.2,
                unit_price_min=30,
                unit_price_max=60,
                qty_per_ctn_min=50,
                qty_per_ctn_max=200,
                source="test range",
            ),
        )
        anchor = ProductCandidate(
            source="manifest_group",
            source_label="input.xlsx",
            zh="塑料桶",
            en="Plastic bucket",
            hs="3924901050",
            material="Plastic",
            usage="HOME",
            ctns=10,
            qty=80,
            unit_price=2,
            gross_weight=100,
            base_tax_rate=0.1,
            effective_tax_rate=0.1,
            price_evidence={"declared_unit_price": 2, "confidence": 0.7},
            plausibility_range=PlausibilityRange(
                kg_per_ctn_min=5,
                kg_per_ctn_max=20,
                kg_per_pc_min=0.5,
                kg_per_pc_max=5,
                unit_price_min=1,
                unit_price_max=10,
                qty_per_ctn_min=1,
                qty_per_ctn_max=50,
                source="test range",
            ),
        )
        better = ProductCandidate(
            source="manifest_group",
            source_label="input.xlsx",
            zh="塑料收纳盒",
            en="Plastic storage box",
            hs="3924905650",
            material="Plastic",
            usage="HOME",
            ctns=10,
            qty=80,
            unit_price=1.5,
            gross_weight=100,
            base_tax_rate=0.1,
            effective_tax_rate=0.1,
            price_evidence={"declared_unit_price": 1.5, "confidence": 0.7},
            plausibility_range=PlausibilityRange(
                kg_per_ctn_min=5,
                kg_per_ctn_max=20,
                kg_per_pc_min=1,
                kg_per_pc_max=5,
                unit_price_min=0.5,
                unit_price_max=10,
                qty_per_ctn_min=1,
                qty_per_ctn_max=50,
                source="test range",
            ),
        )

        selected, summary = optimize_selected_candidates_for_price_fit(
            [bad, anchor],
            [better],
            ManifestSummary("input.xlsx", 2, 20, 200, 1000, []),
            BillInfo("bill.pdf", "", [], cartons=20),
            ProcessingOptions(target_tax_amount=100, target_item_count=2),
        )

        self.assertEqual(summary["swaps"], 1)
        self.assertEqual(selected[0].zh, "塑料收纳盒")
        self.assertNotIn("高价鞋", [candidate.zh for candidate in selected])

    def test_price_fit_resolution_rejects_unresolved_low_reference_price(self) -> None:
        candidate = ProductCandidate(
            source="manifest_group",
            source_label="input.xlsx",
            zh="高价鞋",
            en="Shoes",
            hs="6402999000",
            material="Textile",
            usage="HOME",
            ctns=10,
            qty=1000,
            unit_price=50,
            gross_weight=100,
            base_tax_rate=0.1,
            effective_tax_rate=0.1,
            price_evidence={"declared_unit_price": 50, "confidence": 0.7},
            plausibility_range=PlausibilityRange(
                kg_per_ctn_min=5,
                kg_per_ctn_max=20,
                kg_per_pc_min=0.1,
                kg_per_pc_max=0.2,
                unit_price_min=33.5,
                unit_price_max=75,
                qty_per_ctn_min=50,
                qty_per_ctn_max=200,
                source="test range",
            ),
        )

        with self.assertRaisesRegex(RuntimeError, "优化无解"):
            assert_selected_price_fit_resolved(
                [candidate],
                ManifestSummary("input.xlsx", 1, 10, 100, 1000, []),
                BillInfo("bill.pdf", "", [], cartons=10),
                ProcessingOptions(target_tax_amount=100, target_item_count=1),
                {"swaps": 0, "replacement_attempts": 0},
            )

    def test_manual_reference_candidates_keep_manual_tax_and_plausibility(self) -> None:
        candidates = load_default_reference_manual_candidates()
        self.assertGreaterEqual(len(candidates), 11)

        keyboard = next(candidate for candidate in candidates if candidate.hs == "8471602000" and candidate.zh == "键盘")
        self.assertEqual(keyboard.effective_tax_rate, 0.0)
        self.assertEqual(keyboard.tax_match_source, "manual_reference")
        self.assertEqual(keyboard.ctns, 115)
        self.assertEqual(keyboard.qty, 1380)
        self.assertTrue(keyboard.plausibility_range)
        assert keyboard.plausibility_range is not None
        self.assertGreater(keyboard.plausibility_range.kg_per_ctn_max or 0, keyboard.plausibility_range.kg_per_ctn_min or 0)

    def test_manual_invoice_optimizer_prefers_diverse_reference_rows(self) -> None:
        bill_hairpin = ProductCandidate(
            source="bill",
            source_label="bill.pdf",
            zh="塑料发夹",
            en="Plastic Hairpin",
            hs="9615115000",
            material="Plastic",
            usage="Decoration",
            ctns=10,
            qty=300,
            unit_price=0.32,
            gross_weight=180,
            base_tax_rate=0.1,
            effective_tax_rate=0.1,
            tax_match_source="bill_product",
            plausibility_range=PlausibilityRange(
                kg_per_ctn_min=1,
                kg_per_ctn_max=30,
                kg_per_pc_min=0.01,
                kg_per_pc_max=1,
                unit_price_min=0.05,
                unit_price_max=1,
                qty_per_ctn_min=5,
                qty_per_ctn_max=500,
                source="test bill range",
            ),
        )
        bill_keychain = ProductCandidate(
            source="bill",
            source_label="bill.pdf",
            zh="塑料钥匙扣",
            en="Plastic Keychain",
            hs="3926400090",
            material="Plastic",
            usage="Decoration",
            ctns=10,
            qty=300,
            unit_price=0.3,
            gross_weight=180,
            base_tax_rate=0.153,
            effective_tax_rate=0.153,
            tax_match_source="bill_product",
            plausibility_range=PlausibilityRange(
                kg_per_ctn_min=1,
                kg_per_ctn_max=30,
                kg_per_pc_min=0.01,
                kg_per_pc_max=1,
                unit_price_min=0.05,
                unit_price_max=1,
                qty_per_ctn_min=5,
                qty_per_ctn_max=500,
                source="test bill range",
            ),
        )
        duplicate_keyboard = ProductCandidate(
            source="replacement",
            source_label="DEFAULT_REFERENCE_STYLE_ROWS",
            zh="键鼠套装",
            en="Keyboard and mouse set",
            hs="8471602000",
            material="ABS",
            usage="HOME",
            ctns=115,
            qty=1380,
            unit_price=2.0,
            gross_weight=2465.7,
            base_tax_rate=0.0,
            effective_tax_rate=0.0,
            tax_match_source="manual_reference",
            plausibility_range=next(candidate.plausibility_range for candidate in load_default_reference_manual_candidates() if candidate.hs == "8471602000"),
        )
        pool = [*load_default_reference_manual_candidates(), duplicate_keyboard]

        selected, summary = optimize_selected_candidates_for_manual_invoice(
            [bill_hairpin, bill_keychain, duplicate_keyboard],
            pool,
            ManifestSummary("input.xlsx", 11, 903, 15406.15, 182859.56, []),
            BillInfo("bill.pdf", "", ["PLASTIC HAIRPIN", "PLASTIC KEYCHAIN"], cartons=903),
            ProcessingOptions(target_tax_amount=850, target_item_count=11),
        )

        rows = build_output_rows(
            selected,
            ManifestSummary("input.xlsx", 11, 903, 15406.15, 182859.56, []),
            BillInfo("bill.pdf", "", ["PLASTIC HAIRPIN", "PLASTIC KEYCHAIN"], cartons=903),
            ProcessingOptions(target_tax_amount=850, target_item_count=11),
        )
        manual_rows = [candidate for candidate in selected if candidate.source_label == "DEFAULT_REFERENCE_STYLE_ROWS"]
        keyboard_rows = [candidate for candidate in selected if candidate.hs == "8471602000"]
        self.assertGreaterEqual(len(manual_rows), 7)
        self.assertLessEqual(len(keyboard_rows), 1)
        self.assertLessEqual(sum(row["预计税金"] for row in rows), 935)
        self.assertEqual(summary["strategy"], "manual_invoice")

    def test_manual_invoice_optimizer_preserves_reused_manifest_bill_product(self) -> None:
        broad_range = PlausibilityRange(
            kg_per_ctn_min=1,
            kg_per_ctn_max=30,
            kg_per_pc_min=0.01,
            kg_per_pc_max=5,
            unit_price_min=0.1,
            unit_price_max=20,
            qty_per_ctn_min=1,
            qty_per_ctn_max=200,
            source="test range",
        )

        def candidate(name: str, hs: str, *, bill_product_name: str = "") -> ProductCandidate:
            return ProductCandidate(
                source="manifest_group",
                source_label="input.xlsx/HS归并",
                zh=name,
                en=name,
                hs=hs,
                material="Polyester",
                usage="HOME",
                ctns=10,
                qty=100,
                unit_price=2,
                declared_value=200,
                real_weight=90,
                gross_weight=100,
                base_tax_rate=0.2,
                effective_tax_rate=0.2,
                tax_match_source="manifest_group_hs",
                plausibility_range=broad_range,
                bill_product_name=bill_product_name,
                bill_has_manifest_detail=bool(bill_product_name),
            )

        bill_bag = candidate("Polyester bag", "4202923131", bill_product_name="POLYESTER BAG")
        selected = [
            bill_bag,
            candidate("Curtain", "6303921000"),
            candidate("Scarf", "6214300000"),
        ]

        optimized, _ = optimize_selected_candidates_for_manual_invoice(
            selected,
            [*selected, *load_default_reference_manual_candidates()],
            ManifestSummary("input.xlsx", 3, 30, 300, 1000, []),
            BillInfo("bill.pdf", "", ["POLYESTER BAG"], cartons=30),
            ProcessingOptions(target_tax_amount=100, target_item_count=3),
        )

        self.assertIn(bill_bag, optimized)
        self.assertEqual(
            [item.bill_product_name for item in optimized if item.bill_product_name],
            ["POLYESTER BAG"],
        )


class BillProductCoverageTests(unittest.IsolatedAsyncioTestCase):
    async def test_bill_product_reuses_matching_manifest_hs_candidate_without_query(self) -> None:
        bill = BillInfo(
            filename="bill.pdf",
            raw_text="",
            products=["POLYESTER BAG"],
            cartons=20,
            product_entries=[BillProduct(name="POLYESTER BAG", evidence="POLYESTER BAG")],
        )
        manifest_candidate = ProductCandidate(
            source="manifest_group",
            source_label="input.xlsx/HS归并",
            zh="涤纶包",
            en="Polyester bag",
            hs="4202126000",
            material="Polyester Storage",
            usage="Storage",
            ctns=2,
            qty=74,
            gross_weight=24.6,
            base_tax_rate=0.057,
            effective_tax_rate=0.432,
            tax_match_source="manifest_group_hs",
            tax_data={"hs_code_us": "4202126000", "description_cn": "以纺织材料作面的衣箱"},
        )
        crawler = RoutedFakeCrawler(product_results={}, hs_results={})

        required, filtered = await qualify_bill_product_candidates(
            crawler,
            bill,
            [manifest_candidate],
            [],
            SelectionRules(),
            ProcessingOptions(target_tax_amount=100, target_item_count=10),
            query_cache={},
        )

        self.assertEqual(len(required), 1)
        self.assertEqual(required[0].zh, manifest_candidate.zh)
        self.assertEqual(required[0].hs, manifest_candidate.hs)
        self.assertEqual(required[0].bill_product_name, "POLYESTER BAG")
        self.assertTrue(required[0].bill_has_manifest_detail)
        self.assertTrue(is_bill_required_candidate(required[0]))
        self.assertFalse(is_undetailed_bill_candidate(required[0]))
        self.assertEqual(filtered, [])
        self.assertEqual(crawler.product_calls, [])
        self.assertEqual(crawler.hs_calls, [])

    async def test_bill_product_candidates_deduplicate_repeated_bill_products(self) -> None:
        bill = BillInfo(
            filename="bill.pdf",
            raw_text="",
            products=[
                "PLASTIC HAIRPIN",
                "PLASTIC HAIRPIN",
                "PLASTIC HAIRPIN",
                "PLASTIC KEYCHAIN",
                "PLASTIC KEYCHAIN",
                "PLASTIC KEYCHAIN",
            ],
            gross_weight=150,
            cartons=30,
        )
        crawler = RoutedFakeCrawler(
            product_results={
                "plastic hairpin": tax_result("Free", hs="9615115000"),
                "plastic keychain": tax_result("5.3%", hs="3926400090"),
            }
        )

        bill_required, bill_filtered = await qualify_bill_product_candidates(
            crawler,
            bill,
            [],
            [],
            SelectionRules(allowed_certifications=["Lacey Act", "TSCA"]),
            ProcessingOptions(target_tax_amount=850, target_item_count=10),
            query_cache={},
        )

        self.assertEqual(bill_filtered, [])
        self.assertEqual([item.zh for item in bill_required], ["PLASTIC HAIRPIN", "PLASTIC KEYCHAIN"])
        self.assertEqual(
            crawler.product_calls,
            [("PLASTIC HAIRPIN", ""), ("PLASTIC KEYCHAIN", "")],
        )

    async def test_bill_product_candidates_do_not_use_replacement_pool_fallback(self) -> None:
        bill = BillInfo(
            filename="bill.pdf",
            raw_text="",
            products=["PLASTIC HAIRPIN"],
            gross_weight=150,
            cartons=30,
        )
        replacement = ProductCandidate(
            source="replacement",
            source_label="海关编码查找.xlsx/常用1",
            zh="PLASTIC HAIRPIN",
            en="Plastic Hairpin",
            hs="9615115000",
            material="Plastic",
            usage="HOME",
            ctns=10,
            qty=100,
            unit_price=0.3,
            gross_weight=20,
        )
        crawler = RoutedFakeCrawler(
            product_results={},
            hs_results={"9615115000": tax_result("Free", hs="9615115000")},
        )

        bill_required, bill_filtered = await qualify_bill_product_candidates(
            crawler,
            bill,
            [],
            [replacement],
            SelectionRules(allowed_certifications=["Lacey Act", "TSCA"]),
            ProcessingOptions(target_tax_amount=850, target_item_count=10),
            query_cache={},
        )

        self.assertEqual(bill_required, [])
        self.assertEqual(crawler.product_calls, [("PLASTIC HAIRPIN", ""), ("PLASTIC HAIRPIN", "")])
        self.assertEqual(crawler.hs_calls, [])
        self.assertIn("Codeflag", bill_filtered[-1].filter_reason)

    async def test_bill_product_candidates_accept_first_codeflag_result_even_with_certification(self) -> None:
        bill = BillInfo(
            filename="bill.pdf",
            raw_text="",
            products=["SILICONE COASTER"],
            gross_weight=150,
            cartons=30,
            product_entries=[BillProduct(name="SILICONE COASTER", evidence="SILICONE COASTER")],
        )

        class MaterialSensitiveCrawler(RoutedFakeCrawler):
            async def search_product(self, product_name: str, material: str = "") -> dict[str, dict]:
                self.product_calls.append((product_name, material))
                if material:
                    return tax_result("Free", hs="3924905650", certifications=["FDA"])
                return tax_result("5.3%", hs="3926400090")

        crawler = MaterialSensitiveCrawler(product_results={})

        bill_required, bill_filtered = await qualify_bill_product_candidates(
            crawler,
            bill,
            [],
            [],
            SelectionRules(allowed_certifications=["Lacey Act", "TSCA"]),
            ProcessingOptions(target_tax_amount=850, target_item_count=10),
            query_cache={},
        )

        self.assertEqual([item.zh for item in bill_required], ["SILICONE COASTER"])
        self.assertEqual(bill_required[0].hs, "3924905650")
        self.assertEqual(bill_required[0].tax_match_source, "bill_product")
        self.assertEqual(
            crawler.product_calls,
            [("SILICONE COASTER", "Silicone")],
        )
        self.assertEqual(bill_filtered, [])

    async def test_bill_product_query_retries_transient_failures(self) -> None:
        bill = BillInfo(
            filename="bill.pdf",
            raw_text="",
            products=["PLASTIC MOBILE PHONE STAND"],
            gross_weight=150,
            cartons=30,
            product_entries=[BillProduct(name="PLASTIC MOBILE PHONE STAND", evidence="PLASTIC MOBILE PHONE STAND")],
        )

        class FlakyCrawler(RoutedFakeCrawler):
            async def search_product(self, product_name: str, material: str = "") -> dict[str, dict]:
                self.product_calls.append((product_name, material))
                if len(self.product_calls) == 1:
                    raise RuntimeError("temporary Codeflag timeout")
                return tax_result("5.3%", hs="3926400090")

        crawler = FlakyCrawler(product_results={})

        bill_required, bill_filtered = await qualify_bill_product_candidates(
            crawler,
            bill,
            [],
            [],
            SelectionRules(allowed_certifications=["Lacey Act", "TSCA"]),
            ProcessingOptions(target_tax_amount=850, target_item_count=10),
            query_cache=None,
        )

        self.assertEqual(bill_filtered, [])
        self.assertEqual([item.hs for item in bill_required], ["3926400090"])
        self.assertEqual(
            crawler.product_calls,
            [
                ("PLASTIC MOBILE PHONE STAND", "Plastic"),
                ("PLASTIC MOBILE PHONE STAND", "Plastic"),
            ],
        )

    def test_replacement_pool_excludes_bill_product_names(self) -> None:
        pool = [
            ProductCandidate(
                source="replacement",
                source_label="DEFAULT_REFERENCE_STYLE_ROWS",
                zh="塑料发夹",
                en="Plastic hairpin",
                hs="9615115000",
                material="Plastic",
                usage="Decoration",
            ),
            ProductCandidate(
                source="replacement",
                source_label="DEFAULT_REFERENCE_STYLE_ROWS",
                zh="花瓶",
                en="Vase",
                hs="6913105000",
                material="Ceramic",
                usage="Decoration",
            ),
        ]

        filtered = exclude_bill_product_replacements(pool, ["PLASTIC HAIRPIN"])

        self.assertEqual([candidate.zh for candidate in filtered], ["花瓶"])

    async def test_manual_invoice_stable_references_exclude_bill_product_names(self) -> None:
        qualified, filtered, attempts = await qualify_manual_invoice_replacements(
            FakeCrawler({}),
            [],
            [],
            SelectionRules(allowed_certifications=["Lacey Act", "TSCA"]),
            query_cache={},
            bill_products=["PLASTIC HAIRPIN", "PLASTIC KEYCHAIN"],
        )

        names = {candidate.en.lower() for candidate in qualified}
        self.assertNotIn("plastic hairpin", names)
        self.assertNotIn("plastic keychain", names)
        self.assertEqual(filtered, [])
        self.assertEqual(attempts, 0)

    async def test_replacement_candidate_keeps_original_hs_when_product_query_returns_unrelated_low_tax_hs(self) -> None:
        candidate = ProductCandidate(
            source="replacement",
            source_label="海关编码查找.xlsx/常用1",
            zh="电动钻",
            en="Electric drill",
            hs="8467210010",
            material="Plastic",
            usage="HOME",
            ctns=40,
            qty=120,
            unit_price=2.2,
            gross_weight=285.3,
        )
        crawler = RoutedFakeCrawler(
            product_results={
                "电动钻": tax_result("Free", hs="8508110000"),
            },
            hs_results={
                "8467210010": tax_result("1.7%", hs="8467210010"),
            },
        )

        result = await qualify_single_candidate(
            crawler,
            candidate,
            SelectionRules(allowed_certifications=["Lacey Act", "TSCA"]),
            query_cache={},
        )

        self.assertFalse(result.filter_reason)
        self.assertEqual(result.hs, "8467210010")
        self.assertEqual(result.tax_match_source, "replacement_hs")

    async def test_replacement_candidate_rejects_product_query_when_original_hs_has_no_result(self) -> None:
        candidate = ProductCandidate(
            source="replacement",
            source_label="海关编码查找.xlsx/常用1",
            zh="电动钻",
            en="Electric drill",
            hs="8467210010",
            material="Plastic",
            usage="HOME",
            ctns=40,
            qty=120,
            unit_price=2.2,
            gross_weight=285.3,
        )
        crawler = RoutedFakeCrawler(
            product_results={
                "电动钻": tax_result("Free", hs="8508110000"),
            },
            hs_results={
                "8467210010": {},
            },
        )

        result = await qualify_single_candidate(
            crawler,
            candidate,
            SelectionRules(allowed_certifications=["Lacey Act", "TSCA"]),
            query_cache={},
        )

        self.assertIn("返回 HS 与替换清单原始 HS 不一致", result.filter_reason)
        self.assertEqual(result.hs, "8467210010")

    async def test_manifest_candidate_accepts_product_query_when_original_hs_has_no_result(self) -> None:
        candidate = ProductCandidate(
            source="manifest_group",
            source_label="input.xlsx/HS归并",
            zh="塑料收纳盒",
            en="Plastic storage box",
            hs="3922100000",
            material="Plastic",
            usage="HOME",
            ctns=10,
            qty=300,
            unit_price=1,
            gross_weight=100,
            original_hs="3922100000",
        )
        crawler = RoutedFakeCrawler(
            product_results={
                "塑料收纳盒": tax_result("5.3%", hs="3926909985"),
            },
            hs_results={
                "3922100000": {},
            },
        )

        result = await qualify_single_candidate(
            crawler,
            candidate,
            SelectionRules(allowed_certifications=["Lacey Act", "TSCA"]),
            query_cache={},
            enforce_tax_limit=False,
        )

        self.assertFalse(result.filter_reason)
        self.assertEqual(result.hs, "3926909985")
        self.assertEqual(result.original_hs, "3922100000")
        self.assertEqual(result.tax_match_source, "product")

    async def test_manifest_candidate_accepts_actual_tax_over_twenty_percent(self) -> None:
        candidate = ProductCandidate(
            source="manifest_group",
            source_label="input.xlsx/HS归并",
            zh="电极片",
            en="Electrode pads",
            hs="9033000090",
            material="Mixed",
            usage="HOME",
            ctns=50,
            qty=20000,
            unit_price=0.7,
            gross_weight=400,
            original_hs="9033000090",
        )
        crawler = RoutedFakeCrawler(
            product_results={
                "电极片": tax_result("25%", hs="9033000090"),
            },
            hs_results={
                "9033000090": tax_result("25%", hs="9033000090"),
            },
        )

        result = await qualify_single_candidate(
            crawler,
            candidate,
            SelectionRules(allowed_certifications=["Lacey Act", "TSCA"]),
            query_cache={},
            enforce_tax_limit=False,
        )

        self.assertFalse(result.filter_reason)
        self.assertEqual(result.hs, "9033000090")
        self.assertAlmostEqual(result.effective_tax_rate, 0.25)
        self.assertEqual(result.tax_match_source, "manifest_group_hs")

    async def test_value_capped_hs_is_selected_only_when_clean_candidates_are_insufficient(self) -> None:
        candidate = ProductCandidate(
            source="manifest_group",
            source_label="input.xlsx/HS归并",
            zh="不锈钢首饰",
            en="Stainless steel jewelry",
            hs="7117190500",
            material="Stainless steel",
            usage="Decoration",
            ctns=10,
            qty=800,
            unit_price=3,
            declared_value=2400,
            gross_weight=175.9,
        )
        tax_data = tax_result("Free", hs="7117190500")
        tax_data["7117190500"]["description_cn"] = "玩具首饰，每件价值不超过8美分"
        tax_data["7117190500"]["taric"] = "Toy jewelry valued not over 8 cents per piece"
        crawler = RoutedFakeCrawler(product_results={}, hs_results={"7117190500": tax_data})
        rules = SelectionRules(allowed_certifications=["Lacey Act", "TSCA"])

        review_only = await qualify_single_candidate(
            crawler,
            candidate,
            rules,
            query_cache={},
            enforce_tax_limit=False,
        )

        self.assertFalse(review_only.filter_reason)
        self.assertTrue(review_only.compliance_review_required)
        self.assertIn("0.0800 USD", review_only.compliance_review_reason)
        clean = [
            ProductCandidate(
                source="manifest_group",
                source_label="input.xlsx/HS归并",
                zh=f"清单商品{i}",
                en=f"Manifest item {i}",
                hs=f"39269099{i:02d}",
                material="Plastic",
                usage="HOME",
                ctns=1,
                qty=10,
                unit_price=1,
                gross_weight=10,
                effective_tax_rate=0.1,
            )
            for i in range(10)
        ]

        selected_with_choice = select_initial_candidates(
            [review_only, *clean], [], rules, target_item_count=10
        )
        selected_with_shortage = select_initial_candidates(
            [review_only, *clean[:9]], [], rules, target_item_count=10
        )

        self.assertNotIn(review_only, selected_with_choice)
        self.assertIn(review_only, selected_with_shortage)

    async def test_bill_product_from_manifest_is_selected_before_other_items(self) -> None:
        bill = BillInfo(
            filename="bill.pdf",
            raw_text="",
            products=["PLASTIC MOBILE PHONE STAND"],
            gross_weight=150,
            cartons=15,
        )
        options = ProcessingOptions(target_tax_amount=120, target_item_count=3)
        rules = SelectionRules(allowed_certifications=["Lacey Act", "TSCA"])
        other_candidates = [
            ProductCandidate(
                source="manifest",
                source_label="input.xlsx",
                zh=f"其他品名{i}",
                en=f"Other Item {i}",
                hs="3924104000",
                material="Plastic",
                usage="HOME",
                ctns=50,
                qty=500,
                real_weight=500,
                gross_weight=500,
                unit_price=2,
                base_tax_rate=0.1,
                effective_tax_rate=0.1,
                tax_match_source="product",
                plausibility_range=PlausibilityRange(
                    kg_per_ctn_min=2,
                    kg_per_ctn_max=20,
                    kg_per_pc_min=0.2,
                    kg_per_pc_max=2,
                    unit_price_min=0.5,
                    unit_price_max=10,
                    qty_per_ctn_min=5,
                    qty_per_ctn_max=30,
                    source="test range",
                ),
            )
            for i in range(3)
        ]
        phone_candidate = ProductCandidate(
            source="manifest",
            source_label="input.xlsx",
            zh="手机支架",
            en="Phone holder",
            hs="3926100000",
            material="Plastic",
            usage="HOME",
                ctns=1,
                qty=100,
                unit_price=2,
                real_weight=10,
                gross_weight=50,
            base_tax_rate=0.1,
            effective_tax_rate=0.1,
            tax_match_source="product",
            plausibility_range=PlausibilityRange(
                kg_per_ctn_min=2,
                kg_per_ctn_max=20,
                kg_per_pc_min=0.2,
                kg_per_pc_max=2,
                unit_price_min=0.5,
                unit_price_max=10,
                qty_per_ctn_min=5,
                qty_per_ctn_max=100,
                source="test range",
            ),
        )
        manifest_candidates = [*other_candidates, phone_candidate]

        bill_required, bill_filtered = await qualify_bill_product_candidates(
            FakeCrawler(tax_result("3.4%", hs="3926100000")),
            bill,
            manifest_candidates,
            [],
            rules,
            options,
            query_cache={},
        )
        self.assertEqual(bill_filtered, [])
        self.assertEqual(bill_required[0].zh, "手机支架")
        self.assertEqual(bill_required[0].hs, "3926100000")

        selected = select_initial_candidates(manifest_candidates, bill_required, rules, options.target_item_count)
        self.assertEqual(len(selected), 3)
        self.assertEqual(selected[0].zh, "手机支架")

        manifest = ManifestSummary(
            filename="input.xlsx",
            row_count=3,
            total_ctns=15,
            total_real_weight=150,
            total_declared_value=1000,
            categories=[],
        )
        rows = build_output_rows(selected, manifest, bill, options)
        apply_bill_product_names(rows, bill.products, selected)
        ensure_bill_products_present(rows, bill.products, selected)
        self.assertTrue(any(row["中文品名"] == "塑料手机支架" for row in rows))

    async def test_bill_product_without_matching_hs_candidate_uses_llm_query_terms(self) -> None:
        bill = BillInfo(
            filename="bill.pdf",
            raw_text="",
            products=["PLASTIC SHELL"],
            gross_weight=150,
            cartons=15,
        )
        options = ProcessingOptions(target_tax_amount=90, target_item_count=2)
        crawler = RoutedFakeCrawler(
            product_results={
                "plastic shell": {},
                "plastic enclosure": tax_result("3.4%", hs="3926909985"),
            }
        )
        llm = QueueFakeLLM(
            [
                {
                    "queries": [
                        {"product_name": "Plastic enclosure", "material": "Plastic", "reason": "same product category"},
                    ]
                }
            ]
        )
        rules = SelectionRules(allowed_certifications=["Lacey Act", "TSCA"])

        bill_required, bill_filtered = await qualify_bill_product_candidates(
            crawler,
            bill,
            [],
            [],
            rules,
            options,
            query_cache={},
            llm=llm,
        )

        self.assertEqual(len(bill_required), 1)
        self.assertEqual(bill_required[0].zh, "PLASTIC SHELL")
        self.assertEqual(bill_required[0].hs, "3926909985")
        self.assertEqual(bill_required[0].tax_match_source, "llm_query")
        self.assertGreaterEqual(len(bill_filtered), 1)

    async def test_bill_product_with_hs_hint_is_qualified_from_codeflagai_hs_query(self) -> None:
        bill = BillInfo(
            filename="bill.pdf",
            raw_text="",
            products=["STORAGE BAG"],
            gross_weight=150,
            cartons=15,
            product_entries=[
                BillProduct(
                    name="STORAGE BAG",
                    hs_code_hint="420222",
                    evidence="STORAGE BAG HS CODE:420222",
                    confidence=0.99,
                )
            ],
        )
        options = ProcessingOptions(target_tax_amount=90, target_item_count=3)
        rules = SelectionRules(allowed_certifications=["Lacey Act", "TSCA"])

        bill_required, bill_filtered = await qualify_bill_product_candidates(
            FakeCrawler(tax_result("Free", hs="4202221000")),
            bill,
            [],
            [],
            rules,
            options,
            query_cache={},
        )

        self.assertEqual(bill_filtered, [])
        self.assertEqual(len(bill_required), 1)
        self.assertEqual(bill_required[0].source, "bill")
        self.assertEqual(bill_required[0].zh, "STORAGE BAG")
        self.assertEqual(bill_required[0].hs, "4202221000")
        self.assertEqual(bill_required[0].tax_match_source, "bill_hs")
        self.assertEqual(bill_required[0].base_tax_rate, 0.0)
        self.assertEqual(bill_required[0].tax_data["hs_code_us"], "4202221000")

    async def test_bill_product_with_hs_hint_accepts_actual_tax_over_twenty_percent(self) -> None:
        bill = BillInfo(
            filename="bill.pdf",
            raw_text="",
            products=["POLYESTER PILLOWCASE"],
            gross_weight=150,
            cartons=15,
            product_entries=[
                BillProduct(
                    name="POLYESTER PILLOWCASE",
                    hs_code_hint="630232",
                    evidence="POLYESTER PILLOWCASE HS CODE:630232",
                    confidence=0.99,
                )
            ],
        )
        options = ProcessingOptions(target_tax_amount=90, target_item_count=3)
        rules = SelectionRules(allowed_certifications=["Lacey Act", "TSCA"])
        crawler = FakeCrawler(
            {
                "6302322010": {
                    "hs_code_us": "6302322010",
                    "tax_rate": "11.4%",
                    "additional_tax_rate": "10%",
                    "certification_texts": [],
                    "description_cn": "化纤制床上用品",
                }
            }
        )

        bill_required, bill_filtered = await qualify_bill_product_candidates(
            crawler,
            bill,
            [],
            [],
            rules,
            options,
            query_cache={},
        )

        self.assertEqual(bill_filtered, [])
        self.assertEqual(len(bill_required), 1)
        self.assertEqual(bill_required[0].hs, "6302322010")
        self.assertAlmostEqual(bill_required[0].effective_tax_rate, 0.214)
        self.assertEqual(bill_required[0].tax_match_source, "bill_hs")

    async def test_bill_product_query_accepts_actual_tax_over_twenty_percent(self) -> None:
        bill = BillInfo(
            filename="bill.pdf",
            raw_text="",
            products=["POLYESTER PILLOWCASE"],
            gross_weight=150,
            cartons=15,
        )
        options = ProcessingOptions(target_tax_amount=90, target_item_count=3)
        rules = SelectionRules(allowed_certifications=["Lacey Act", "TSCA"])
        crawler = RoutedFakeCrawler(
            product_results={
                "polyester pillowcase": {
                    "6302322010": {
                        "hs_code_us": "6302322010",
                        "tax_rate": "11.4%",
                        "additional_tax_rate": "10%",
                        "certification_texts": [],
                        "description_cn": "化纤制床上用品",
                    }
                }
            }
        )

        bill_required, bill_filtered = await qualify_bill_product_candidates(
            crawler,
            bill,
            [],
            [],
            rules,
            options,
            query_cache={},
        )

        self.assertEqual(bill_filtered, [])
        self.assertEqual(len(bill_required), 1)
        self.assertEqual(bill_required[0].hs, "6302322010")
        self.assertAlmostEqual(bill_required[0].effective_tax_rate, 0.214)
        self.assertEqual(bill_required[0].tax_match_source, "bill_product")

    async def test_bill_product_with_six_digit_hs_hint_outputs_qualified_ten_digit_hts(self) -> None:
        bill = BillInfo(
            filename="bill.pdf",
            raw_text="",
            products=["PLASTIC ORNAMENTS"],
            gross_weight=220,
            cartons=22,
            product_entries=[
                BillProduct(
                    name="PLASTIC ORNAMENTS",
                    hs_code_hint="392640",
                    evidence="PLASTIC ORNAMENTS HS CODE:392640",
                    confidence=0.99,
                )
            ],
        )
        options = ProcessingOptions(target_tax_amount=58.3, target_item_count=1)
        rules = SelectionRules(allowed_certifications=["Lacey Act", "TSCA"])

        bill_required, bill_filtered = await qualify_bill_product_candidates(
            FakeCrawler(tax_result("5.3%", hs="3926400090")),
            bill,
            [],
            [],
            rules,
            options,
            query_cache={},
        )

        self.assertEqual(bill_filtered, [])
        self.assertEqual(len(bill_required), 1)
        self.assertEqual(bill_required[0].hs, "3926400090")
        self.assertEqual(bill_required[0].tax_data["hs_code_us"], "3926400090")
        self.assertEqual(bill_required[0].tax_match_source, "bill_hs")

    async def test_bill_material_prefers_plastic_modifier(self) -> None:
        bill = BillInfo(
            filename="bill.pdf",
            raw_text="",
            products=["PLASTIC ORNAMENTS"],
            gross_weight=220,
            cartons=22,
            product_entries=[
                BillProduct(
                    name="PLASTIC ORNAMENTS",
                    hs_code_hint="392640",
                    evidence="PLASTIC ORNAMENTS HS CODE:392640",
                    confidence=0.99,
                )
            ],
        )
        options = ProcessingOptions(target_tax_amount=58.3, target_item_count=1)
        rules = SelectionRules(allowed_certifications=["Lacey Act", "TSCA"])

        bill_required, _ = await qualify_bill_product_candidates(
            FakeCrawler(tax_result("5.3%", hs="3926400090")),
            bill,
            [],
            [],
            rules,
            options,
            query_cache={},
        )

        self.assertEqual(bill_required[0].material, "Plastic")
        self.assertEqual(bill_required[0].tax_match_source, "bill_hs")

    async def test_bill_material_can_be_inferred_from_hs_description(self) -> None:
        bill = BillInfo(
            filename="bill.pdf",
            raw_text="",
            products=["ORNAMENT"],
            gross_weight=100,
            cartons=10,
            product_entries=[
                BillProduct(
                    name="ORNAMENT",
                    hs_code_hint="392640",
                    evidence="ORNAMENT HS CODE:392640",
                    confidence=0.99,
                )
            ],
        )
        options = ProcessingOptions(target_tax_amount=90, target_item_count=2)
        rules = SelectionRules(allowed_certifications=["Lacey Act", "TSCA"])

        crawler = RoutedFakeCrawler(
            product_results={},
            hs_results={
                "392640": {
                    "3926400090": {
                        "hs_code_us": "3926400090",
                        "tax_rate": "5.3%",
                        "certification_texts": [],
                        "description_cn": "Other articles of plastics and articles of other materials of headings 3901 to 3914: Statuettes and other ornamental articles",
                    }
                }
            },
        )
        llm = DictFakeLLM({"material": "Plastic"})

        bill_required, _ = await qualify_bill_product_candidates(
            crawler,
            bill,
            [],
            [],
            rules,
            options,
            query_cache={},
            llm=llm,
        )

        self.assertEqual(bill_required[0].material, "Plastic")
        self.assertEqual(llm.calls, 1)

    async def test_bill_product_fallback_uses_inferred_material_without_material_field(self) -> None:
        bill = BillInfo(
            filename="bill.pdf",
            raw_text="",
            products=["PLASTIC ORNAMENTS"],
            gross_weight=100,
            cartons=10,
            product_entries=[
                BillProduct(
                    name="PLASTIC ORNAMENTS",
                    hs_code_hint="",
                    evidence="PLASTIC ORNAMENTS",
                    confidence=0.99,
                )
            ],
        )
        options = ProcessingOptions(target_tax_amount=30, target_item_count=1)
        rules = SelectionRules(allowed_certifications=["Lacey Act", "TSCA"])
        crawler = RoutedFakeCrawler(
            product_results={
                "plastic ornaments": {
                    "3926400090": {
                        "hs_code_us": "3926400090",
                        "tax_rate": "5.3%",
                        "certification_texts": [],
                        "description_cn": "Other articles of plastics and articles of other materials of headings 3901 to 3914: Statuettes and other ornamental articles",
                    }
                }
            },
            hs_results={},
        )

        bill_required, bill_filtered = await qualify_bill_product_candidates(
            crawler,
            bill,
            [],
            [],
            rules,
            options,
            query_cache={},
        )

        self.assertEqual(bill_filtered, [])
        self.assertEqual(len(bill_required), 1)
        self.assertEqual(bill_required[0].material, "Plastic")
        self.assertEqual(bill_required[0].tax_match_source, "bill_product")
        self.assertEqual(bill_required[0].hs, "3926400090")

    async def test_llm_plausibility_range_is_used_when_candidate_lacks_ranges(self) -> None:
        candidate = ProductCandidate(
            source="bill",
            source_label="bill.pdf",
            zh="HANDHELD GARMENT STEAMER",
            en="Handheld Garment Steamer",
            hs="8451300000",
            material="Plastic",
            usage="HOME",
            ctns=1,
            qty=1,
            gross_weight=1,
            base_tax_rate=0.034,
            effective_tax_rate=0.034,
            tax_match_source="bill_product",
        )
        llm = QueueFakeLLM(
            [
                {
                    "unit_price_min": 8,
                    "unit_price_max": 18,
                    "kg_per_pc_min": 0.6,
                    "kg_per_pc_max": 1.5,
                    "qty_per_ctn_min": 4,
                    "qty_per_ctn_max": 12,
                    "kg_per_ctn_min": 4,
                    "kg_per_ctn_max": 18,
                    "confidence": 0.9,
                    "basis": "small appliance",
                }
            ]
        )
        selected, used = await ensure_candidate_plausibility_ranges(
            [candidate],
            ManifestSummary("input.xlsx", 1, 10, 100, 100, []),
            BillInfo("bill.pdf", "", ["HANDHELD GARMENT STEAMER"], cartons=10),
            {},
            llm,
            {},
        )
        self.assertTrue(used)
        self.assertEqual(selected[0].plausibility_range.unit_price_min, 8)
        self.assertIn("LLM合理性估算", selected[0].plausibility_range.source)

    async def test_incomplete_replacement_range_uses_llm_instead_of_placeholder_fields(self) -> None:
        candidate = ProductCandidate(
            source="replacement",
            source_label="海关编码查找.xlsx/20260330",
            zh="加湿器",
            en="Humidifier",
            hs="8509809000",
            material="Plastic",
            usage="HOME",
            ctns=50,
            qty=1000,
            unit_price=1,
            gross_weight=1,
            base_tax_rate=0.034,
            effective_tax_rate=0.034,
            tax_match_source="product",
        )
        known = {
            ("加湿器", "humidifier", "8509809000"): PlausibilityRange(
                kg_per_ctn_min=None,
                kg_per_ctn_max=None,
                kg_per_pc_min=0.5,
                kg_per_pc_max=1.5,
                unit_price_min=8,
                unit_price_max=20,
                qty_per_ctn_min=4,
                qty_per_ctn_max=12,
                source="incomplete test range",
            )
        }
        llm = QueueFakeLLM(
            [
                {
                    "unit_price_min": 8,
                    "unit_price_max": 20,
                    "kg_per_pc_min": 0.5,
                    "kg_per_pc_max": 1.5,
                    "qty_per_ctn_min": 4,
                    "qty_per_ctn_max": 12,
                    "kg_per_ctn_min": 4,
                    "kg_per_ctn_max": 18,
                    "confidence": 0.92,
                    "basis": "small electric appliance",
                }
            ]
        )

        selected, used = await ensure_candidate_plausibility_ranges(
            [candidate],
            ManifestSummary("input.xlsx", 1, 50, 500, 1000, []),
            BillInfo("bill.pdf", "", [], cartons=50),
            known,
            llm,
            {},
        )

        self.assertTrue(used)
        self.assertEqual(selected[0].plausibility_range.kg_per_ctn_min, 4)
        self.assertIn("small electric appliance", selected[0].plausibility_range.source)

    async def test_llm_output_rejects_mechanical_quantity_and_weight_distribution(self) -> None:
        candidates = []
        for idx in range(8):
            candidates.append(
                ProductCandidate(
                    source="replacement",
                    source_label="test",
                    zh=f"品名{idx}",
                    en=f"Item {idx}",
                    hs=f"39269099{idx:02d}",
                    material="Plastic",
                    usage="HOME",
                    ctns=10,
                    qty=100,
                    unit_price=2,
                    gross_weight=50,
                    base_tax_rate=0.1,
                    effective_tax_rate=0.1,
                    tax_match_source="product",
                    plausibility_range=PlausibilityRange(
                        kg_per_ctn_min=2,
                        kg_per_ctn_max=20,
                        kg_per_pc_min=0.2,
                        kg_per_pc_max=2,
                        unit_price_min=1,
                        unit_price_max=5,
                        qty_per_ctn_min=5,
                        qty_per_ctn_max=20,
                        source="test range",
                    ),
                )
            )
        llm = QueueFakeLLM(
            [
                {
                    "rows": [
                        {"candidate_index": idx, "箱数": 10, "数量": 100, "单价": 1.25, "毛重": 47}
                        for idx in range(8)
                    ]
                },
                {
                    "rows": [
                        {"candidate_index": 0, "箱数": 8, "数量": 60, "单价": 5.0, "毛重": 20},
                        {"candidate_index": 1, "箱数": 9, "数量": 60, "单价": 5.0, "毛重": 26},
                        {"candidate_index": 2, "箱数": 10, "数量": 80, "单价": 4.0, "毛重": 32},
                        {"candidate_index": 3, "箱数": 11, "数量": 80, "单价": 4.0, "毛重": 39},
                        {"candidate_index": 4, "箱数": 12, "数量": 90, "单价": 4.0, "毛重": 48},
                        {"candidate_index": 5, "箱数": 13, "数量": 100, "单价": 3.5, "毛重": 58},
                        {"candidate_index": 6, "箱数": 14, "数量": 110, "单价": 3.5, "毛重": 70},
                        {"candidate_index": 7, "箱数": 15, "数量": 120, "单价": 3.0, "毛重": 83},
                    ]
                },
            ]
        )
        manifest = ManifestSummary("input.xlsx", 8, 92, 376, 1000, [])
        bill = BillInfo("bill.pdf", "", [], cartons=92)
        options = ProcessingOptions(target_tax_amount=350, target_item_count=8)

        rows, attempts, feedback = await generate_valid_output_rows_with_llm(llm, candidates, manifest, bill, options)

        self.assertEqual(attempts, 2)
        self.assertIn("数量分布过于机械", feedback[0])
        self.assertEqual(len(rows), 8)
        self.assertLessEqual(sum(row["总价"] * row["综合税率"] for row in rows), 385)
        self.assertTrue(all(row["税金"] >= 0 for row in rows))

    async def test_llm_output_retries_when_llm_draft_is_not_json(self) -> None:
        candidates = [
            ProductCandidate(
                source="replacement",
                source_label="test",
                zh="品名",
                en="Item",
                hs="3926909985",
                material="Plastic",
                usage="HOME",
                ctns=10,
                qty=100,
                unit_price=5,
                gross_weight=100,
                base_tax_rate=0.1,
                effective_tax_rate=0.1,
                tax_match_source="product",
                plausibility_range=PlausibilityRange(
                    kg_per_ctn_min=5,
                    kg_per_ctn_max=20,
                    kg_per_pc_min=0.5,
                    kg_per_pc_max=2,
                    unit_price_min=1,
                    unit_price_max=10,
                    qty_per_ctn_min=5,
                    qty_per_ctn_max=20,
                    source="test range",
                ),
            )
        ]
        llm = QueueFakeLLM(
            [
                RuntimeError("LLM 未返回 JSON object"),
                {"rows": [{"candidate_index": 0, "箱数": 10, "数量": 100, "单价": 5.0, "毛重": 100}]},
            ]
        )

        rows, attempts, feedback = await generate_valid_output_rows_with_llm(
            llm,
            candidates,
            ManifestSummary("input.xlsx", 1, 10, 100, 1000, []),
            BillInfo("bill.pdf", "", [], cartons=10),
            ProcessingOptions(target_tax_amount=50, target_item_count=1),
        )

        self.assertEqual(attempts, 2)
        self.assertIn("LLM 未返回 JSON object", feedback[0])
        self.assertEqual(llm.kwargs[0]["temperature"], 0.0)
        self.assertTrue(llm.kwargs[0]["json_mode"])
        self.assertEqual(llm.kwargs[0]["max_tokens"], 8192)
        self.assertEqual(rows[0]["中文品名"], "品名")

    def test_llm_output_keeps_tax_under_upper_bound_without_forcing_target_gap(self) -> None:
        candidates = [
            ProductCandidate(
                source="replacement",
                source_label="test",
                zh="品名",
                en="Item",
                hs="3926909985",
                material="Plastic",
                usage="HOME",
                ctns=10,
                qty=100,
                unit_price=85.133,
                gross_weight=100,
                base_tax_rate=0.1,
                effective_tax_rate=0.1,
                tax_match_source="product",
                plausibility_range=PlausibilityRange(
                    kg_per_ctn_min=0.5,
                    kg_per_ctn_max=20,
                    kg_per_pc_min=0.05,
                    kg_per_pc_max=2,
                    unit_price_min=90,
                    unit_price_max=91,
                    qty_per_ctn_min=5,
                    qty_per_ctn_max=20,
                    source="test range",
                ),
            )
        ]
        rows = [
            {
                "中文品名": "品名",
                "英文品名": "Item",
                "商品编码": "3926909985",
                "材质": "Plastic",
                "用途": "Home use",
                "箱数": 10,
                "数量": 100,
                "单位": "PCS",
                "币制": "USD",
                "单价": 85.133,
                "总价": 8513.3,
                "净重": 92,
                "毛重": 100,
                "原产国": "CN",
                "综合税率": 0.1,
                "candidate_index": 0,
            }
        ]
        validate_llm_output_rows(
            rows,
            candidates,
            ManifestSummary("input.xlsx", 1, 10, 100, 1000, []),
            BillInfo("bill.pdf", "", [], cartons=10),
            ProcessingOptions(target_tax_amount=900, target_item_count=1),
        )
        self.assertEqual(rows[0]["单价"], 85.133)
        self.assertEqual(rows[0]["税金"], 851.33)
        self.assertLessEqual(sum(row["总价"] * row["综合税率"] for row in rows), 990)

    def test_llm_output_closes_bill_cartons(self) -> None:
        candidates = [
            ProductCandidate(
                source="replacement",
                source_label="test",
                zh="品名A",
                en="Item A",
                hs="3926909985",
                material="Plastic",
                usage="HOME",
                ctns=10,
                qty=4000,
                unit_price=0.2,
                gross_weight=400,
                base_tax_rate=0.1,
                effective_tax_rate=0.1,
                tax_match_source="product",
                plausibility_range=PlausibilityRange(
                    kg_per_ctn_min=0.5,
                    kg_per_ctn_max=20,
                    kg_per_pc_min=0.05,
                    kg_per_pc_max=0.2,
                    unit_price_min=0.1,
                    unit_price_max=1.0,
                    qty_per_ctn_min=5,
                    qty_per_ctn_max=20,
                    source="test range",
                ),
            ),
            ProductCandidate(
                source="replacement",
                source_label="test",
                zh="品名B",
                en="Item B",
                hs="3926909986",
                material="Plastic",
                usage="HOME",
                ctns=10,
                qty=4170,
                unit_price=0.2,
                gross_weight=417,
                base_tax_rate=0.1,
                effective_tax_rate=0.1,
                tax_match_source="product",
                plausibility_range=PlausibilityRange(
                    kg_per_ctn_min=0.5,
                    kg_per_ctn_max=20,
                    kg_per_pc_min=0.05,
                    kg_per_pc_max=0.2,
                    unit_price_min=0.1,
                    unit_price_max=1.0,
                    qty_per_ctn_min=5,
                    qty_per_ctn_max=20,
                    source="test range",
                ),
            ),
        ]
        rows = [
            {
                "中文品名": "品名A",
                "英文品名": "Item A",
                "商品编码": "3926909985",
                "材质": "Plastic",
                "用途": "Home use",
                "箱数": 400,
                "数量": 4000,
                "单位": "PCS",
                "币制": "USD",
                "单价": 0.2,
                "总价": 800.0,
                "净重": 368,
                "毛重": 400,
                "原产国": "CN",
                "综合税率": 0.1,
                "candidate_index": 0,
            },
            {
                "中文品名": "品名B",
                "英文品名": "Item B",
                "商品编码": "3926909986",
                "材质": "Plastic",
                "用途": "Home use",
                "箱数": 417,
                "数量": 4170,
                "单位": "PCS",
                "币制": "USD",
                "单价": 0.2,
                "总价": 834.0,
                "净重": 384.36,
                "毛重": 417,
                "原产国": "CN",
                "综合税率": 0.1,
                "candidate_index": 1,
            },
        ]

        validate_llm_output_rows(
            rows,
            candidates,
            ManifestSummary("input.xlsx", 2, 0, 817, 1000, []),
            BillInfo("bill.pdf", "", [], cartons=817),
            ProcessingOptions(target_tax_amount=160, target_item_count=2),
        )

        self.assertEqual(sum(row["箱数"] for row in rows), 817)

    def test_llm_output_closes_gross_weight_by_scaling_draft(self) -> None:
        candidates = [
            ProductCandidate(
                source="replacement",
                source_label="test",
                zh="品名",
                en="Item",
                hs="3926909985",
                material="Plastic",
                usage="HOME",
                ctns=10,
                qty=100,
                unit_price=85.133,
                gross_weight=100,
                base_tax_rate=0.1,
                effective_tax_rate=0.1,
                tax_match_source="product",
                plausibility_range=PlausibilityRange(
                    kg_per_ctn_min=5,
                    kg_per_ctn_max=20,
                    kg_per_pc_min=0.5,
                    kg_per_pc_max=2,
                    unit_price_min=90,
                    unit_price_max=91,
                    qty_per_ctn_min=5,
                    qty_per_ctn_max=20,
                    source="test range",
                ),
            )
        ]
        rows = [
            {
                "中文品名": "品名",
                "英文品名": "Item",
                "商品编码": "3926909985",
                "材质": "Plastic",
                "用途": "Home use",
                "箱数": 10,
                "数量": 100,
                "单位": "PCS",
                "币制": "USD",
                "单价": 85.133,
                "总价": 8513.3,
                "净重": 91.99,
                "毛重": 100,
                "原产国": "CN",
                "综合税率": 0.1,
                "candidate_index": 0,
            }
        ]

        validate_llm_output_rows(
            rows,
            candidates,
            ManifestSummary("input.xlsx", 1, 10, 120, 1000, []),
            BillInfo("bill.pdf", "", [], cartons=10),
            ProcessingOptions(target_tax_amount=900, target_item_count=1),
        )

        self.assertEqual(rows[0]["毛重"], 120)
        self.assertEqual(rows[0]["LLM草案毛重"], 100)
        self.assertEqual(rows[0]["毛重闭合调整"], 20)
        self.assertIn("毛重按Excel总重量倒推调整", rows[0]["约束提示"])

    def test_llm_output_requires_ten_digit_hs_after_bill_hint_lookup(self) -> None:
        candidates = [
            ProductCandidate(
                source="bill_product",
                source_label="bill.pdf",
                zh="塑料装饰品",
                en="Plastic Ornaments",
                hs="3926400090",
                material="Plastic",
                usage="HOME",
                ctns=10,
                qty=100,
                unit_price=95,
                gross_weight=100,
                base_tax_rate=0.1,
                effective_tax_rate=0.1,
                tax_match_source="bill_hs",
                tax_data={
                    "hs_code_us": "3926400090",
                    "tax_rate": "10%",
                    "description_cn": "Statuettes and other ornamental articles",
                },
                plausibility_range=PlausibilityRange(
                    kg_per_ctn_min=5,
                    kg_per_ctn_max=20,
                    kg_per_pc_min=0.5,
                    kg_per_pc_max=2,
                    unit_price_min=90,
                    unit_price_max=100,
                    qty_per_ctn_min=5,
                    qty_per_ctn_max=20,
                    source="test range",
                ),
            )
        ]
        rows = [
            {
                "中文品名": "塑料装饰品",
                "英文品名": "Plastic Ornaments",
                "商品编码": "3926400090",
                "材质": "Plastic",
                "用途": "Home use",
                "箱数": 10,
                "数量": 100,
                "单位": "PCS",
                "币制": "USD",
                "单价": 95,
                "总价": 9500,
                "净重": 92,
                "毛重": 100,
                "原产国": "CN",
                "综合税率": 0.1,
                "candidate_index": 0,
            }
        ]

        validate_llm_output_rows(
            rows,
            candidates,
            ManifestSummary("input.xlsx", 1, 10, 100, 1000, []),
            BillInfo("bill.pdf", "", ["PLASTIC ORNAMENTS"], cartons=10),
            ProcessingOptions(target_tax_amount=950, target_item_count=1),
        )

        self.assertEqual(str(rows[0]["商品编码"]), "3926400090")
        self.assertEqual(rows[0]["税金"], 950.0)

    def test_bill_product_weight_range_is_advisory_after_bill_closure(self) -> None:
        candidates = [
            ProductCandidate(
                source="bill",
                source_label="bill.pdf",
                zh="PLASTIC HAIRPIN",
                en="Plastic Hairpin",
                hs="9615115000",
                material="Plastic",
                usage="HOME",
                ctns=1,
                qty=100,
                unit_price=5,
                gross_weight=3.57,
                base_tax_rate=0.1,
                effective_tax_rate=0.1,
                tax_match_source="bill_product",
                plausibility_range=PlausibilityRange(
                    kg_per_ctn_min=0.1,
                    kg_per_ctn_max=50,
                    kg_per_pc_min=0.005,
                    kg_per_pc_max=0.03,
                    unit_price_min=1,
                    unit_price_max=10,
                    qty_per_ctn_min=1,
                    qty_per_ctn_max=200,
                    source="test range",
                ),
            )
        ]
        rows = [
            {
                "中文品名": "PLASTIC HAIRPIN",
                "英文品名": "Plastic Hairpin",
                "商品编码": "9615115000",
                "材质": "Plastic",
                "用途": "Home use",
                "箱数": 1,
                "数量": 100,
                "单位": "PCS",
                "币制": "USD",
                "单价": 5,
                "总价": 500,
                "净重": 3.28,
                "毛重": 3.57,
                "原产国": "CN",
                "综合税率": 0.1,
                "candidate_index": 0,
            }
        ]

        validate_llm_output_rows(
            rows,
            candidates,
            ManifestSummary("input.xlsx", 1, 1, 3.57, 1000, []),
            BillInfo("bill.pdf", "", ["PLASTIC HAIRPIN"], cartons=1),
            ProcessingOptions(target_tax_amount=50, target_item_count=1),
        )

        self.assertEqual(rows[0]["单件重量"], 0.0357)
        self.assertIn("单件重量高于合理上限 0.0357 > 0.03", rows[0]["约束提示"])

    def test_non_bill_product_weight_range_still_rejects_outlier(self) -> None:
        candidates = [
            ProductCandidate(
                source="replacement",
                source_label="test",
                zh="PLASTIC HAIRPIN",
                en="Plastic Hairpin",
                hs="9615115000",
                material="Plastic",
                usage="HOME",
                ctns=1,
                qty=100,
                unit_price=5,
                gross_weight=3.57,
                base_tax_rate=0.1,
                effective_tax_rate=0.1,
                tax_match_source="product",
                plausibility_range=PlausibilityRange(
                    kg_per_ctn_min=0.1,
                    kg_per_ctn_max=50,
                    kg_per_pc_min=0.005,
                    kg_per_pc_max=0.03,
                    unit_price_min=1,
                    unit_price_max=10,
                    qty_per_ctn_min=1,
                    qty_per_ctn_max=200,
                    source="test range",
                ),
            )
        ]
        rows = [
            {
                "中文品名": "PLASTIC HAIRPIN",
                "英文品名": "Plastic Hairpin",
                "商品编码": "9615115000",
                "材质": "Plastic",
                "用途": "Home use",
                "箱数": 1,
                "数量": 100,
                "单位": "PCS",
                "币制": "USD",
                "单价": 5,
                "总价": 500,
                "净重": 3.28,
                "毛重": 3.57,
                "原产国": "CN",
                "综合税率": 0.1,
                "candidate_index": 0,
            }
        ]

        with self.assertRaisesRegex(RuntimeError, "单件重量高于合理上限 0.0357 > 0.03"):
            validate_llm_output_rows(
                rows,
                candidates,
                ManifestSummary("input.xlsx", 1, 1, 3.57, 1000, []),
                BillInfo("bill.pdf", "", [], cartons=1),
                ProcessingOptions(target_tax_amount=50, target_item_count=1),
            )

    def test_qty_must_not_be_less_than_ctns_and_must_be_multiple(self) -> None:
        with self.assertRaisesRegex(RuntimeError, "数量不能小于箱数"):
            validate_qty_ctn_relationship([{"箱数": 50, "数量": 25}])
        with self.assertRaisesRegex(RuntimeError, "整数倍"):
            validate_qty_ctn_relationship([{"箱数": 25, "数量": 51}])
        validate_qty_ctn_relationship([{"箱数": 25, "数量": 50}])

    def test_llm_output_allows_low_row_tax_for_manual_invoice_strategy(self) -> None:
        candidates = []
        rows = []
        for idx, (rate, unit_price) in enumerate(((0.1, 10.0), (0.1, 1.0)), start=1):
            candidates.append(
                ProductCandidate(
                    source="replacement",
                    source_label="test",
                    zh=f"品名{idx}",
                    en=f"Item {idx}",
                    hs=f"39269099{idx:02d}",
                    material="Plastic",
                    usage="HOME",
                    ctns=10,
                    qty=10,
                    unit_price=unit_price,
                    gross_weight=100,
                    base_tax_rate=rate,
                    effective_tax_rate=rate,
                    tax_match_source="product",
                    plausibility_range=PlausibilityRange(
                        kg_per_ctn_min=5,
                        kg_per_ctn_max=20,
                        kg_per_pc_min=0.5,
                        kg_per_pc_max=20,
                        unit_price_min=1,
                        unit_price_max=200,
                        qty_per_ctn_min=1,
                        qty_per_ctn_max=1,
                        source="test range",
                    ),
                )
            )
            rows.append(
                {
                    "中文品名": f"品名{idx}",
                    "英文品名": f"Item {idx}",
                    "商品编码": f"39269099{idx:02d}",
                    "材质": "Plastic",
                    "用途": "Home use",
                    "箱数": 10,
                    "数量": 10,
                    "单位": "PCS",
                    "币制": "USD",
                    "单价": unit_price,
                    "总价": unit_price * 10,
                    "净重": 92,
                    "毛重": 100,
                    "原产国": "CN",
                    "综合税率": rate,
                    "candidate_index": idx - 1,
                }
            )

        validate_llm_output_rows(
            rows,
            candidates,
            ManifestSummary("input.xlsx", 2, 20, 200, 1000, []),
            BillInfo("bill.pdf", "", [], cartons=20),
            ProcessingOptions(target_tax_amount=120, target_item_count=2),
        )

        self.assertEqual(rows[0]["税金"], 10.0)
        self.assertEqual(rows[1]["税金"], 1.0)
        self.assertLessEqual(sum(row["税金"] for row in rows), 132)

    def test_llm_output_allows_up_to_two_zero_tax_rows(self) -> None:
        candidates = []
        rows = []
        for idx, rate in enumerate((0.0, 0.0, 0.1), start=1):
            candidates.append(
                ProductCandidate(
                    source="replacement",
                    source_label="test",
                    zh=f"品名{idx}",
                    en=f"Item {idx}",
                    hs=f"39269099{idx:02d}",
                    material="Plastic",
                    usage="HOME",
                    ctns=10,
                    qty=10,
                    unit_price=10,
                    gross_weight=100,
                    base_tax_rate=rate,
                    effective_tax_rate=rate,
                    tax_match_source="product",
                    plausibility_range=PlausibilityRange(
                        kg_per_ctn_min=5,
                        kg_per_ctn_max=20,
                        kg_per_pc_min=0.5,
                        kg_per_pc_max=20,
                        unit_price_min=1,
                        unit_price_max=200,
                        qty_per_ctn_min=1,
                        qty_per_ctn_max=1,
                        source="test range",
                    ),
                )
            )
            rows.append(
                {
                    "中文品名": f"品名{idx}",
                    "英文品名": f"Item {idx}",
                    "商品编码": f"39269099{idx:02d}",
                    "材质": "Plastic",
                    "用途": "Home use",
                    "箱数": 10,
                    "数量": 10,
                    "单位": "PCS",
                    "币制": "USD",
                    "单价": 10,
                    "总价": 100,
                    "净重": 92,
                    "毛重": 100,
                    "原产国": "CN",
                    "综合税率": rate,
                    "candidate_index": idx - 1,
                }
            )

        validate_llm_output_rows(
            rows,
            candidates,
            ManifestSummary("input.xlsx", 3, 30, 300, 1000, []),
            BillInfo("bill.pdf", "", [], cartons=30),
            ProcessingOptions(target_tax_amount=100, target_item_count=3),
        )

        self.assertEqual([row["税金"] for row in rows[:2]], [0.0, 0.0])
        self.assertEqual(rows[2]["税金"], 10.0)

    def test_llm_output_allows_zero_tax_rows_for_manual_invoice_strategy(self) -> None:
        candidates = []
        rows = []
        for idx in range(5):
            candidates.append(
                ProductCandidate(
                    source="replacement",
                    source_label="test",
                    zh=f"免税品名{idx}",
                    en=f"Free Item {idx}",
                    hs=f"39269099{idx:02d}",
                    material="Plastic",
                    usage="HOME",
                    ctns=10,
                    qty=10,
                    unit_price=10,
                    gross_weight=100,
                    base_tax_rate=0.0,
                    effective_tax_rate=0.0,
                    tax_match_source="product",
                    plausibility_range=PlausibilityRange(
                        kg_per_ctn_min=5,
                        kg_per_ctn_max=20,
                        kg_per_pc_min=0.5,
                        kg_per_pc_max=20,
                        unit_price_min=1,
                        unit_price_max=200,
                        qty_per_ctn_min=1,
                        qty_per_ctn_max=1,
                        source="test range",
                    ),
                )
            )
            rows.append(
                {
                    "中文品名": f"免税品名{idx}",
                    "英文品名": f"Free Item {idx}",
                    "商品编码": f"39269099{idx:02d}",
                    "材质": "Plastic",
                    "用途": "Home use",
                    "箱数": 10,
                    "数量": 10,
                    "单位": "PCS",
                    "币制": "USD",
                    "单价": 10,
                    "总价": 100,
                    "净重": 92,
                    "毛重": 100,
                    "原产国": "CN",
                    "综合税率": 0.0,
                    "candidate_index": idx,
                }
            )

        validate_llm_output_rows(
            rows,
            candidates,
            ManifestSummary("input.xlsx", 5, 50, 500, 1000, []),
            BillInfo("bill.pdf", "", [], cartons=50),
            ProcessingOptions(target_tax_amount=1, target_item_count=5),
        )
        self.assertEqual([row["税金"] for row in rows], [0.0, 0.0, 0.0, 0.0, 0.0])

    def test_llm_output_allows_row_below_old_minimum_tax(self) -> None:
        candidates = [
            ProductCandidate(
                source="replacement",
                source_label="test",
                zh="低税金品名",
                en="Low Tax Item",
                hs="3926909985",
                material="Plastic",
                usage="HOME",
                ctns=10,
                qty=10,
                unit_price=1,
                gross_weight=100,
                base_tax_rate=0.1,
                effective_tax_rate=0.1,
                tax_match_source="product",
                plausibility_range=PlausibilityRange(
                    kg_per_ctn_min=5,
                    kg_per_ctn_max=20,
                    kg_per_pc_min=0.5,
                    kg_per_pc_max=20,
                    unit_price_min=1,
                    unit_price_max=2,
                    qty_per_ctn_min=1,
                    qty_per_ctn_max=1,
                    source="test range",
                ),
            )
        ]
        rows = [
            {
                "中文品名": "低税金品名",
                "英文品名": "Low Tax Item",
                "商品编码": "3926909985",
                "材质": "Plastic",
                "用途": "Home use",
                "箱数": 10,
                "数量": 10,
                "单位": "PCS",
                "币制": "USD",
                "单价": 1,
                "总价": 10,
                "净重": 92,
                "毛重": 100,
                "原产国": "CN",
                "综合税率": 0.1,
                "candidate_index": 0,
            }
        ]

        validate_llm_output_rows(
            rows,
            candidates,
            ManifestSummary("input.xlsx", 1, 10, 100, 1000, []),
            BillInfo("bill.pdf", "", [], cartons=10),
            ProcessingOptions(target_tax_amount=30, target_item_count=1),
        )
        self.assertEqual(rows[0]["税金"], 1.0)

    def test_llm_output_relaxes_price_floor_to_keep_tax_under_upper_bound(self) -> None:
        candidates = [
            ProductCandidate(
                source="replacement",
                source_label="test",
                zh="品名",
                en="Item",
                hs="3926909985",
                material="Plastic",
                usage="HOME",
                ctns=10,
                qty=100,
                unit_price=90.1,
                gross_weight=100,
                base_tax_rate=0.1,
                effective_tax_rate=0.1,
                tax_match_source="product",
                plausibility_range=PlausibilityRange(
                    kg_per_ctn_min=5,
                    kg_per_ctn_max=20,
                    kg_per_pc_min=0.5,
                    kg_per_pc_max=2,
                    unit_price_min=90,
                    unit_price_max=91,
                    qty_per_ctn_min=5,
                    qty_per_ctn_max=20,
                    source="test range",
                ),
            )
        ]
        rows = [
            {
                "中文品名": "品名",
                "英文品名": "Item",
                "商品编码": "3926909985",
                "材质": "Plastic",
                "用途": "Home use",
                "箱数": 10,
                "数量": 100,
                "单位": "PCS",
                "币制": "USD",
                "单价": 90.1,
                "总价": 9010,
                "净重": 92,
                "毛重": 100,
                "原产国": "CN",
                "综合税率": 0.1,
                "candidate_index": 0,
            }
        ]
        validate_llm_output_rows(
            rows,
            candidates,
            ManifestSummary("input.xlsx", 1, 10, 100, 1000, []),
            BillInfo("bill.pdf", "", [], cartons=10),
            ProcessingOptions(target_tax_amount=800, target_item_count=1),
        )

        self.assertLess(rows[0]["单价"], 90)
        self.assertLessEqual(rows[0]["税金"], 880)
        self.assertIn("单价低于合理下限", rows[0]["约束提示"])


if __name__ == "__main__":
    unittest.main()
