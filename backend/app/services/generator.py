"""
文件生成模块
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, numbers
from openpyxl.utils import get_column_letter
from datetime import datetime
from typing import List, Dict, Any, Optional
from pathlib import Path
import logging

logger = logging.getLogger(__name__)

# 样式定义
HEADER_FONT = Font(name='微软雅黑', bold=True, size=11, color='FFFFFF')
HEADER_FILL = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
HEADER_ALIGNMENT = Alignment(horizontal='center', vertical='center', wrap_text=True)
DATA_FONT = Font(name='微软雅黑', size=10)
DATA_ALIGNMENT = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_ALIGNMENT = Alignment(horizontal='left', vertical='center', wrap_text=True)
BORDER_THIN = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
HIGHLIGHT_FILL = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')  # 浅黄色高亮
SAVE_FILL = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')  # 绿色表示节省
RISK_FONT = Font(name='微软雅黑', size=10, color='FF0000')


class FileGenerator:
    """文件生成器"""

    def generate(
        self,
        optimized_data: List[Dict[str, Any]],
        template_path: str,
        output_dir: str,
        optimization_logs: Optional[List[Dict[str, Any]]] = None
    ) -> str:
        """
        生成清关文件

        Args:
            optimized_data: 优化后的数据
            template_path: 模板文件路径
            output_dir: 输出目录
            optimization_logs: 优化日志（含 A→B 详情）

        Returns:
            生成的文件路径
        """
        try:
            logger.info(f"开始生成清关文件，共 {len(optimized_data)} 条数据")

            # 加载模板
            wb = openpyxl.load_workbook(template_path)
            ws = wb.active
            ws.title = "清关数据"

            # 移除模板中多余的空白sheet（避免残留 Sheet2/Sheet3）
            for sn in list(wb.sheetnames):
                if sn != "清关数据":
                    sheet = wb[sn]
                    if sheet.max_row <= 1 and sheet.max_column <= 1:
                        wb.remove(sheet)

            # 找到数据起始行
            start_row = self._find_data_start_row(ws)

            # 填充数据
            for idx, item in enumerate(optimized_data):
                row_num = start_row + idx
                self._fill_row(ws, row_num, item)

            # Sheet2: 优化明细
            if optimization_logs:
                self._write_optimization_sheet(wb, optimization_logs, optimized_data)

            # 保存文件
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            output_filename = f"清关_{timestamp}.xlsx"
            output_path = Path(output_dir) / output_filename

            wb.save(output_path)

            logger.info(f"清关文件生成成功: {output_path}")

            return str(output_path)

        except Exception as e:
            logger.error(f"生成清关文件失败: {str(e)}")
            raise

    def _write_optimization_sheet(
        self,
        wb,
        optimization_logs: List[Dict[str, Any]],
        optimized_data: List[Dict[str, Any]]
    ):
        """写入 Sheet2: 优化明细（客户友好格式）"""
        ws2 = wb.create_sheet("优化明细")

        # 表头
        headers = [
            "序号", "中文品名", "原HS编码", "原税率",
            "优化后HS编码", "优化后税率", "税率降幅",
            "品名匹配度", "风险提示"
        ]
        col_widths = [6, 22, 14, 10, 16, 12, 10, 12, 28]

        for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
            cell = ws2.cell(row=1, column=col_idx, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT
            cell.border = BORDER_THIN
            ws2.column_dimensions[get_column_letter(col_idx)].width = width

        # 冻结表头
        ws2.freeze_panes = 'A2'

        # 构建辅助字典：HS编码 → 品名（用于匹配度展示）
        hs_to_name = {}
        for item in optimized_data:
            hs = str(item.get("商品编码", "")).strip()
            name = item.get("中文品名", "")
            if hs and hs not in hs_to_name:
                hs_to_name[hs] = name

        # 填充数据
        for idx, log in enumerate(optimization_logs):
            row_num = idx + 2
            original_rate = self._parse_rate_str(log.get("original_tax_rate", ""))
            new_rate = self._parse_rate_str(log.get("new_tax_rate", ""))

            # 计算税率降幅
            rate_reduction = ""
            if original_rate is not None and new_rate is not None:
                reduction = original_rate - new_rate
                if reduction > 0:
                    rate_reduction = f"↓ {reduction:.1f}%"
                elif reduction == 0:
                    rate_reduction = "—"
                else:
                    rate_reduction = f"↑ {abs(reduction):.1f}%"

            # 品名匹配度
            similarity = log.get("similarity_score", 0)
            similarity_str = f"{similarity*100:.0f}%" if similarity else "—"

            # 风险提示
            risk = log.get("risk_warning", "")

            row_data = [
                idx + 1,
                log.get("original_name", ""),
                log.get("original_hs_code", ""),
                log.get("original_tax_rate", ""),
                log.get("new_hs_code", ""),
                log.get("new_tax_rate", ""),
                rate_reduction,
                similarity_str,
                risk
            ]

            for col_idx, value in enumerate(row_data, start=1):
                cell = ws2.cell(row=row_num, column=col_idx, value=value)
                cell.font = DATA_FONT
                cell.border = BORDER_THIN
                # 居中对齐大部分列，品名列左对齐
                if col_idx == 2:
                    cell.alignment = LEFT_ALIGNMENT
                elif col_idx == 9:  # 风险提示左对齐
                    cell.alignment = LEFT_ALIGNMENT
                    if risk:
                        cell.font = RISK_FONT
                else:
                    cell.alignment = DATA_ALIGNMENT

                # 高亮有实质性优化（税率降低）的行
                if rate_reduction and rate_reduction.startswith("↓"):
                    if col_idx in (5, 6, 7):  # 新HS编码、新税率、降幅 列
                        cell.fill = SAVE_FILL

        # 添加汇总行
        total_row = len(optimization_logs) + 3
        total_optimized = len(optimization_logs)
        total_items = len(optimized_data)

        ws2.cell(row=total_row, column=1, value="汇总").font = Font(name='微软雅黑', bold=True, size=10)
        ws2.cell(row=total_row, column=2, value=f"共 {total_items} 个商品，优化 {total_optimized} 个")
        ws2.cell(row=total_row, column=2).font = Font(name='微软雅黑', bold=True, size=10)

        # 计算平均税率降幅
        reductions = []
        for log in optimization_logs:
            or_val = self._parse_rate_str(log.get("original_tax_rate", ""))
            nr_val = self._parse_rate_str(log.get("new_tax_rate", ""))
            if or_val is not None and nr_val is not None:
                red = or_val - nr_val
                if red > 0:
                    reductions.append(red)

        if reductions:
            avg_reduction = sum(reductions) / len(reductions)
            ws2.cell(row=total_row, column=5, value=f"平均降税: {avg_reduction:.1f}%")
            ws2.cell(row=total_row, column=5).font = Font(name='微软雅黑', bold=True, size=10, color='006100')

        # 优化率
        if total_items > 0:
            opt_rate = total_optimized / total_items * 100
            ws2.cell(row=total_row, column=7, value=f"优化率: {opt_rate:.1f}%")
            ws2.cell(row=total_row, column=7).font = Font(name='微软雅黑', bold=True, size=10)

    def _parse_rate_str(self, rate_str: str) -> Optional[float]:
        """解析税率字符串为数值"""
        if not rate_str:
            return None
        s = str(rate_str).lower().strip()
        if s in ("free", "0", "0%", ""):
            return 0.0
        try:
            return float(s.replace("%", ""))
        except (ValueError, TypeError):
            return None

    def _find_data_start_row(self, ws) -> int:
        """找到数据起始行（表头下一行）"""
        # 简单策略：找到第一行包含"品名"的行，然后返回下一行
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=20), start=1):
            for cell in row:
                if cell.value and "品名" in str(cell.value):
                    return row_idx + 1

        # 默认从第2行开始
        return 2

    def _fill_row(self, ws, row_num: int, item: Dict[str, Any]):
        """填充一行数据"""
        # 字段到列的映射（需与实际模板列顺序一致）
        field_column_mapping = {
            "中文品名": 1,
            "英文品名": 2,
            "商品编码": 3,
            "材质": 4,
            "用途": 5,
            "箱数": 6,
            "数量": 7,
            "单位": 8,
            "币制": 9,
            "单价": 10,
            "总价": 11,
            "净重": 12,
            "毛重": 13,
            "原产国": 14,
        }

        for field, col_idx in field_column_mapping.items():
            value = item.get(field)

            if value is not None:
                ws.cell(row_num, col_idx).value = value
            else:
                # 缺失字段标记
                ws.cell(row_num, col_idx).value = "待补充"
                ws.cell(row_num, col_idx).font = Font(color="FF0000")

    def generate_log_file(
        self,
        logs: Dict[str, Any],
        output_dir: str
    ) -> str:
        """
        生成日志文件

        Args:
            logs: 日志数据
            output_dir: 输出目录

        Returns:
            日志文件路径
        """
        try:
            logger.info("开始生成日志文件")

            wb = openpyxl.Workbook()

            # Sheet1: 替换日志
            ws1 = wb.active
            ws1.title = "替换日志"
            ws1.append(["原品名", "原HS编码", "原税率", "新HS编码", "新税率", "替换原因", "风险提示"])

            for log in logs.get("optimization_logs", []):
                ws1.append([
                    log.get("original_name", ""),
                    log.get("original_hs_code", ""),
                    log.get("original_tax_rate", ""),
                    log.get("new_hs_code", ""),
                    log.get("new_tax_rate", ""),
                    log.get("reason", ""),
                    log.get("risk_warning", "")
                ])

            # Sheet2: 数据冲突日志
            ws2 = wb.create_sheet("数据冲突日志")
            ws2.append(["字段", "文档A的值", "文档B的值", "处理方式"])

            for conflict in logs.get("conflicts", []):
                ws2.append([
                    conflict.get("field", ""),
                    conflict.get("value_a", ""),
                    conflict.get("value_b", ""),
                    "已空着，请手动填写"
                ])

            # Sheet3: 字段映射日志
            ws3 = wb.create_sheet("字段映射日志")
            ws3.append(["输入文档字段", "映射到清关字段", "置信度"])

            for mapping in logs.get("mapping_logs", []):
                confidence = mapping.get("confidence", 0)
                ws3.append([
                    mapping.get("source_field", ""),
                    mapping.get("target_field", ""),
                    f"{confidence*100:.0f}%"
                ])

            # 保存文件
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            output_filename = f"日志_{timestamp}.xlsx"
            output_path = Path(output_dir) / output_filename

            wb.save(output_path)

            logger.info(f"日志文件生成成功: {output_path}")

            return str(output_path)

        except Exception as e:
            logger.error(f"生成日志文件失败: {str(e)}")
            raise
