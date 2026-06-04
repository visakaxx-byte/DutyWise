from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side


ROOT = Path(__file__).resolve().parents[1]
PACKING_FILE = ROOT / "04.03装柜信息(1).xlsx"
BOOKING_FILE = ROOT / "ZIMUSNH21275749--04.21.xls"
TEMPLATE_FILE = ROOT / "清关模板.xlsx"
OUTPUT_DIR = ROOT / "outputs"


def clean_text(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def clean_number(value: Any, default: float = 0.0) -> float:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return default
    try:
        return float(value)
    except Exception:
        return default


def normalize_material(zh: Any, en: Any) -> str:
    zh_text = clean_text(zh).replace("／", "/")
    en_text = clean_text(en).replace("／", "/")

    if zh_text and en_text:
        zh_lower = zh_text.lower()
        en_lower = en_text.lower()
        if zh_lower == en_lower:
            return zh_text
        if en_lower in zh_lower:
            return zh_text
        return f"{zh_text} / {en_text}"

    return zh_text or en_text


def parse_booking_meta(path: Path) -> dict[str, str]:
    df = pd.read_excel(path, sheet_name="Sheet1", header=None)
    meta: dict[str, str] = {}

    for _, row in df.iloc[:15].iterrows():
        key = clean_text(row.iloc[0])
        if not key:
            continue

        values = [clean_text(v) for v in row.iloc[1:].tolist() if clean_text(v)]
        meta[key] = values[0] if values else ""

    return meta


def build_invoice_rows(path: Path) -> list[dict[str, Any]]:
    df = pd.read_excel(path)
    df = df.copy()

    for col in ["运单号码", "申报价值币别"]:
        if col in df.columns:
            df[col] = df[col].ffill()

    records: list[dict[str, Any]] = []

    for _, row in df.iterrows():
        chinese_name = clean_text(row.get("中文品名"))
        english_name = clean_text(row.get("英文品名"))
        hs_code = clean_text(row.get("海关编码"))

        if not chinese_name or not english_name or not hs_code:
            continue

        quantity = clean_number(row.get("总数量"))
        cartons = clean_number(row.get("箱数量"))
        subtotal = clean_number(row.get("invoice总价"))
        unit_value = clean_number(row.get("申报单价"))
        net_weight = clean_number(row.get("净重"))
        gross_weight = clean_number(row.get("毛重"))

        if quantity <= 0 and subtotal <= 0:
            continue

        if unit_value <= 0 and quantity > 0 and subtotal > 0:
            unit_value = subtotal / quantity

        currency = clean_text(row.get("申报价值币别")) or "USD"
        usage = clean_text(row.get("用途"))
        material = normalize_material(row.get("中文材质"), row.get("英文材质"))

        records.append(
            {
                "中文品名": chinese_name,
                "英文品名": english_name,
                "商品编码": hs_code,
                "材质": material,
                "用途": usage,
                "箱数": cartons,
                "数量": quantity,
                "单位": "PCS",
                "币制": currency,
                "单价": unit_value,
                "总价": subtotal,
                "净重": net_weight,
                "毛重": gross_weight,
                "原产国": "CHINA",
            }
        )

    grouped: dict[tuple[Any, ...], dict[str, Any]] = {}

    for item in records:
        key = (
            item["中文品名"],
            item["英文品名"],
            item["商品编码"],
            item["材质"],
            item["用途"],
            item["币制"],
            round(item["单价"], 4),
        )

        if key not in grouped:
            grouped[key] = item.copy()
            continue

        current = grouped[key]
        for field in ["箱数", "数量", "总价", "净重", "毛重"]:
            current[field] += item[field]

    rows = list(grouped.values())
    rows.sort(key=lambda item: (item["中文品名"], item["英文品名"], item["商品编码"]))
    return rows


def write_workbook(template_path: Path, output_path: Path, meta: dict[str, str], rows: list[dict[str, Any]]) -> None:
    wb = load_workbook(template_path)
    ws = wb["Sheet1"]

    if ws.max_row >= 6:
        ws.delete_rows(6, ws.max_row - 5)

    eta = meta.get("预计到港时间", "")
    bl_no = meta.get("提单号", "")
    container_no = meta.get("柜号", "")
    destination = meta.get("目的港", "")
    vessel_voyage = meta.get("船名航次", "")

    ws["A4"] = f"ETA: {eta}" if eta else "ETA:"
    header_bits = [bit for bit in [f"B/L : {bl_no}" if bl_no else "", f"CONTAINER: {container_no}" if container_no else "", f"POD: {destination}" if destination else "", vessel_voyage] if bit]
    ws["C4"] = "    ".join(header_bits) if header_bits else "B/L :"

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    text_alignment = Alignment(vertical="top", wrap_text=True)
    number_alignment = Alignment(horizontal="right", vertical="top")

    widths = {
        "A": 24,
        "B": 28,
        "C": 16,
        "D": 20,
        "E": 24,
        "F": 12,
        "G": 12,
        "H": 10,
        "I": 10,
        "J": 12,
        "K": 12,
        "L": 12,
        "M": 12,
        "N": 14,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    start_row = 6
    for idx, item in enumerate(rows, start=start_row):
        ws.cell(idx, 1, item["中文品名"])
        ws.cell(idx, 2, item["英文品名"])
        ws.cell(idx, 3, item["商品编码"])
        ws.cell(idx, 4, item["材质"])
        ws.cell(idx, 5, item["用途"])
        ws.cell(idx, 6, item["箱数"])
        ws.cell(idx, 7, item["数量"])
        ws.cell(idx, 8, item["单位"])
        ws.cell(idx, 9, item["币制"])
        ws.cell(idx, 10, item["单价"])
        ws.cell(idx, 11, item["总价"])
        ws.cell(idx, 12, item["净重"])
        ws.cell(idx, 13, item["毛重"])
        ws.cell(idx, 14, item["原产国"])

        for col_idx in range(1, 15):
            cell = ws.cell(idx, col_idx)
            cell.border = border
            cell.alignment = text_alignment
            cell.font = Font(name="Arial", size=10)

        for col_idx in [6, 7, 10, 11, 12, 13]:
            ws.cell(idx, col_idx).alignment = number_alignment

        ws.cell(idx, 6).number_format = "0.##"
        ws.cell(idx, 7).number_format = "0.##"
        ws.cell(idx, 10).number_format = "0.00"
        ws.cell(idx, 11).number_format = "0.00"
        ws.cell(idx, 12).number_format = "0.00"
        ws.cell(idx, 13).number_format = "0.00"

    wb.save(output_path)


def main() -> None:
    meta = parse_booking_meta(BOOKING_FILE)
    rows = build_invoice_rows(PACKING_FILE)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = OUTPUT_DIR / f"清关_{timestamp}.xlsx"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    write_workbook(TEMPLATE_FILE, output_file, meta, rows)

    total_qty = sum(item["数量"] for item in rows)
    total_value = sum(item["总价"] for item in rows)
    print(f"output={output_file} rows={len(rows)} total_qty={total_qty:.2f} total_value={total_value:.2f}")


if __name__ == "__main__":
    main()
